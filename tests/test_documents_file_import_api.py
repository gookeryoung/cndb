"""通用文件导入 API 测试（P1）：预览 / 落库 / 文件夹 / 冲突 / 权限 / 边界.

端点：
- ``POST /{ws}/documents/import/files/analyze``
- ``POST /{ws}/documents/import/files``
"""

from __future__ import annotations

import io
import struct
import zlib

import pytest
from docx import Document as DocxDocument
from sqlalchemy import select
from sqlalchemy.orm import Session

from cndb.plugins.accounts.models import User
from cndb.plugins.documents.services import importers
from cndb.plugins.documents.services.importers import reader
from cndb.plugins.files.models import FileOwner
from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

DOCS = "/api/v1/workspaces/{ws}/documents"
ANALYZE = DOCS + "/import/files/analyze"
IMPORT = DOCS + "/import/files"


def _png_bytes() -> bytes:
    def chunk(kind: bytes, payload: bytes) -> bytes:
        head = struct.pack(">I", len(payload)) + kind + payload
        return head + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)

    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", ihdr)
        + chunk(b"IDAT", zlib.compress(b"\x00\x00\x00\x00\x00"))
        + chunk(b"IEND", b"")
    )


PNG_BYTES = _png_bytes()


def _login(client, username: str, password: str = "passw0rd") -> dict[str, str]:
    r = client.post("/api/v1/accounts/auth/login", json={"login": username, "password": password})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _make_user(db: Session, username: str) -> User:
    user = User(username=username, nickname=username)
    user.set_password("passw0rd")
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


@pytest.fixture
def owner(db) -> User:
    return _make_user(db, "fi_owner")


@pytest.fixture
def ws(db, owner) -> Workspace:
    workspace = Workspace(name="文件导入库", created_by_id=owner.id)
    db.add(workspace)
    db.flush()
    db.add(WorkspaceMember(workspace_id=workspace.id, user_id=owner.id, role=WorkspaceRole.OWNER))
    db.commit()
    db.refresh(workspace)
    return workspace


@pytest.fixture
def auth(client, owner) -> dict[str, str]:
    return _login(client, owner.username)


def _files(*entries: tuple[str, bytes, str]) -> list[tuple[str, tuple[str, bytes, str]]]:
    return [("files", (name, data, ctype)) for name, data, ctype in entries]


def _tree(client, ws_id: int, auth) -> list[dict]:
    r = client.get(DOCS.format(ws=ws_id) + "/tree", headers=auth)
    assert r.status_code == 200, r.text
    return r.json()["items"]


def _flat(nodes: list[dict], out: list[dict] | None = None) -> list[dict]:
    out = [] if out is None else out
    for node in nodes:
        out.append(node)
        _flat(node.get("children") or [], out)
    return out


def _content(client, ws_id: int, doc_id: int, auth) -> str:
    r = client.get(f"{DOCS.format(ws=ws_id)}/{doc_id}/content", headers=auth)
    assert r.status_code == 200, r.text
    return r.json()["content"]


def _make_docx(build) -> bytes:
    document = DocxDocument()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


MD_NOTE = "#  我的笔记  \n\n\n\n这是   正文。  \n"
MD_DEEP = "# 子文档\n\n内容\n"
HTML_PAGE = "<html><head><title>网页标题</title></head><body><h1>一级</h1><p>段落   文本</p></body></html>"


class TestAnalyzeFiles:
    def test_analyze_counts_without_writing(self, client, ws, auth):
        r = client.post(
            ANALYZE.format(ws=ws.id),
            files=_files(
                ("笔记.md", MD_NOTE.encode(), "text/markdown"),
                ("目录/笔记/子文档.md", MD_DEEP.encode(), "text/markdown"),
            ),
            data={"paths": ["笔记.md", "目录/笔记/子文档.md"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["total_files"] == 2
        assert data["total_documents"] == 2
        assert data["total_folders"] == 2  # 目录/、目录/笔记/
        assert data["conflicts"] == 0
        assert data["preview_truncated"] is False
        assert {item["title"] for item in data["items"]} == {"我的笔记", "子文档"}
        assert data["items"][0]["warnings"]  # 标题曾从正文剥离 / 取自文件名的提示
        assert _tree(client, ws.id, auth) == []

    def test_analyze_reports_skipped_and_invalid(self, client, ws, auth):
        r = client.post(
            ANALYZE.format(ws=ws.id),
            files=_files(
                ("ok.md", MD_NOTE.encode(), "text/markdown"),
                ("bad.txt", b"nope", "text/plain"),
                ("bad.docx", b"nope", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
            ),
            data={"paths": ["ok.md", "../escape.txt", "/abs/bad.docx"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["total_documents"] == 1
        assert data["total_files"] == 3
        reasons = {entry["path"]: entry["reason"] for entry in data["skipped"]}
        assert "../escape.txt" in reasons and "/abs/bad.docx" in reasons

    def test_analyze_detects_conflicts(self, client, ws, auth):
        client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("笔记.md", MD_NOTE.encode(), "text/markdown")),
            data={"paths": ["笔记.md"]},
            headers=auth,
        )
        r = client.post(
            ANALYZE.format(ws=ws.id),
            files=_files(("笔记.md", MD_NOTE.encode(), "text/markdown")),
            data={"paths": ["笔记.md"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        assert r.json()["conflicts"] == 1

    def test_analyze_empty_batch_rejected(self, client, ws, auth):
        r = client.post(
            ANALYZE.format(ws=ws.id),
            files=_files(("x.txt", b"x", "text/plain")),
            data={"paths": ["../x.txt"]},
            headers=auth,
        )
        assert r.status_code == 400

    def test_analyze_too_many_files(self, client, ws, auth, monkeypatch):
        monkeypatch.setattr(reader, "MAX_FILES", 1)
        r = client.post(
            ANALYZE.format(ws=ws.id),
            files=_files(("a.md", b"# a\n", "text/markdown"), ("b.md", b"# b\n", "text/markdown")),
            data={"paths": ["a.md", "b.md"]},
            headers=auth,
        )
        assert r.status_code == 400


class TestImportFiles:
    def test_single_markdown_file(self, client, ws, auth):
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("我的笔记.md", MD_NOTE.encode(), "text/markdown")),
            data={"paths": ["我的笔记.md"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        stats = r.json()
        assert stats["created"] == 1
        assert stats["folders_created"] == 0
        nodes = _flat(_tree(client, ws.id, auth))
        assert [node["title"] for node in nodes] == ["我的笔记"]
        assert _content(client, ws.id, nodes[0]["id"], auth) == "这是 正文。"

    def test_markdown_title_from_front_matter(self, client, ws, auth):
        payload = "---\ntitle: 元数据标题\n---\n\n正文\n".encode()
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("file.md", payload, "text/markdown")),
            data={"paths": ["file.md"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        nodes = _flat(_tree(client, ws.id, auth))
        assert nodes[0]["title"] == "元数据标题"

    def test_folder_mode_builds_tree(self, client, ws, auth):
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(
                ("目录/笔记/子文档.md", MD_DEEP.encode(), "text/markdown"),
                ("目录/说明.md", "# 说明\n\n正文\n".encode(), "text/markdown"),
            ),
            data={"paths": ["目录/笔记/子文档.md", "目录/说明.md"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        stats = r.json()
        assert stats["folders_created"] == 2
        assert stats["created"] == 2
        folders = [node for node in _flat(_tree(client, ws.id, auth)) if node["is_folder"]]
        assert {folder["title"] for folder in folders} == {"目录", "笔记"}

    def test_html_file(self, client, ws, auth):
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("页面.html", HTML_PAGE.encode(), "text/html")),
            data={"paths": ["页面.html"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        nodes = _flat(_tree(client, ws.id, auth))
        assert nodes[0]["title"] == "网页标题"
        content = _content(client, ws.id, nodes[0]["id"], auth)
        assert "# 一级" in content
        assert "段落 文本" in content

    def test_html_local_image_inlined(self, client, db, ws, auth):
        html = "<html><head><title>网页图片</title></head><body><p><img src='./图片/a.png' alt='图'></p></body></html>"
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(
                ("站点/页面.html", html.encode(), "text/html"),
                ("站点/图片/a.png", PNG_BYTES, "image/png"),
            ),
            data={"paths": ["站点/页面.html", "站点/图片/a.png"], "import_attachments": "true"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        assert r.json()["attachments"] == 1
        nodes = _flat(_tree(client, ws.id, auth))
        page = next(node for node in nodes if node["title"] == "网页图片")
        content = _content(client, ws.id, page["id"], auth)
        assert "file:" in content
        assert "./图片/a.png" not in content
        owners = (
            db.execute(select(FileOwner).where(FileOwner.workspace_id == ws.id, FileOwner.document_id == page["id"]))
            .scalars()
            .all()
        )
        assert len(owners) == 1

    def test_docx_file_with_attachment(self, client, db, ws, auth):
        def build(document):
            document.add_heading("Word 标题", level=1)
            document.add_paragraph("正文段落")
            document.add_picture(io.BytesIO(PNG_BYTES))

        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(
                (
                    "文档.docx",
                    _make_docx(build),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            ),
            data={"paths": ["文档.docx"], "import_attachments": "true"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        assert r.json()["attachments"] == 1
        node = _flat(_tree(client, ws.id, auth))[0]
        assert node["title"] == "Word 标题"
        content = _content(client, ws.id, node["id"], auth)
        assert "正文段落" in content
        assert "file:" in content
        owners = db.execute(select(FileOwner).where(FileOwner.workspace_id == ws.id)).scalars().all()
        assert len(owners) == 1
        assert owners[0].document_id == node["id"]

    def test_docx_without_attachments_keeps_placeholder(self, client, ws, auth):
        def build(document):
            document.add_heading("无图", level=1)
            document.add_picture(io.BytesIO(PNG_BYTES))

        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(
                (
                    "文档.docx",
                    _make_docx(build),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            ),
            data={"paths": ["文档.docx"], "import_attachments": "false"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        stats = r.json()
        assert stats["attachments"] == 0
        assert any("未导入" in entry["message"] for entry in stats["warnings"])

    def test_conflict_skip(self, client, ws, auth):
        payload = MD_NOTE.encode()
        form = {"paths": "笔记.md"}
        assert (
            client.post(
                IMPORT.format(ws=ws.id), files=_files(("笔记.md", payload, "text/markdown")), data=form, headers=auth
            ).status_code
            == 200
        )
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("笔记.md", payload, "text/markdown")),
            data={**form, "conflict": "skip"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        stats = r.json()
        assert stats["created"] == 0
        assert stats["skipped"] == 1
        assert stats["skipped_files"][0]["reason"] == "同名文档已存在"

    def test_conflict_rename(self, client, ws, auth):
        payload = MD_NOTE.encode()
        assert (
            client.post(
                IMPORT.format(ws=ws.id),
                files=_files(("笔记.md", payload, "text/markdown")),
                data={"paths": "笔记.md"},
                headers=auth,
            ).status_code
            == 200
        )
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("笔记.md", payload, "text/markdown")),
            data={"paths": "笔记.md", "conflict": "rename"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        titles = {node["title"] for node in _flat(_tree(client, ws.id, auth))}
        assert titles == {"我的笔记", "我的笔记 (2)"}

    def test_conflict_overwrite_creates_version(self, client, db, ws, auth):
        from cndb.plugins.documents.models import Document, DocumentVersion

        assert (
            client.post(
                IMPORT.format(ws=ws.id),
                files=_files(("笔记.md", MD_NOTE.encode(), "text/markdown")),
                data={"paths": "笔记.md"},
                headers=auth,
            ).status_code
            == 200
        )
        doc = db.execute(select(Document).where(Document.workspace_id == ws.id)).scalars().one()
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("笔记.md", "# 我的笔记\n\n新正文\n".encode(), "text/markdown")),
            data={"paths": "笔记.md", "conflict": "overwrite"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        assert r.json()["updated"] == 1
        assert r.json()["created"] == 0
        versions = db.execute(select(DocumentVersion).where(DocumentVersion.document_id == doc.id)).scalars().all()
        assert len(versions) == 2
        assert versions[-1].change_type == "auto_import"

    def test_import_into_folder(self, client, ws, auth):
        created = client.post(
            DOCS.format(ws=ws.id),
            json={"title": "目标文件夹", "is_folder": True},
            headers=auth,
        )
        assert created.status_code == 201, created.text
        folder_id = created.json()["id"]
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("笔记.md", MD_NOTE.encode(), "text/markdown")),
            data={"paths": "笔记.md", "parent_id": str(folder_id)},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        nodes = {node["title"]: node for node in _flat(_tree(client, ws.id, auth))}
        assert nodes["目标文件夹"]["is_folder"] is True
        assert nodes["我的笔记"]["parent_id"] == folder_id

    def test_parent_must_be_folder(self, client, ws, auth):
        created = client.post(DOCS.format(ws=ws.id), json={"title": "文档不是文件夹", "content": "x"}, headers=auth)
        doc_id = created.json()["id"]
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("笔记.md", MD_NOTE.encode(), "text/markdown")),
            data={"paths": "笔记.md", "parent_id": str(doc_id)},
            headers=auth,
        )
        assert r.status_code == 409

    def test_invalid_conflict_and_status(self, client, ws, auth):
        for key, value in (("conflict", "bogus"), ("status", "bogus")):
            r = client.post(
                IMPORT.format(ws=ws.id),
                files=_files(("笔记.md", MD_NOTE.encode(), "text/markdown")),
                data={"paths": "笔记.md", key: value},
                headers=auth,
            )
            assert r.status_code == 400

    def test_viewer_forbidden(self, client, db, ws):
        viewer = _make_user(db, "fi_viewer2")
        db.add(WorkspaceMember(workspace_id=ws.id, user_id=viewer.id, role=WorkspaceRole.VIEWER))
        db.commit()
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("笔记.md", MD_NOTE.encode(), "text/markdown")),
            data={"paths": "笔记.md"},
            headers=_login(client, viewer.username),
        )
        assert r.status_code == 403

    def test_non_member_forbidden(self, client, db, ws):
        outsider = _make_user(db, "fi_outsider")
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("笔记.md", MD_NOTE.encode(), "text/markdown")),
            data={"paths": "笔记.md"},
            headers=_login(client, outsider.username),
        )
        assert r.status_code == 403

    def test_partial_success_mixed_files(self, client, ws, auth):
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(
                ("good.md", MD_NOTE.encode(), "text/markdown"),
                ("broken.docx", b"not-a-docx", "application/octet-stream"),
                ("note.txt", b"x", "text/plain"),
            ),
            data={"paths": ["good.md", "broken.docx", "note.txt"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        stats = r.json()
        assert stats["created"] == 1
        assert stats["skipped"] == 2
        reasons = {entry["path"]: entry["reason"] for entry in stats["skipped_files"]}
        assert "broken.docx" in reasons
        assert "note.txt" in reasons
        assert len(_flat(_tree(client, ws.id, auth))) == 1

    def test_invalid_paths_reported(self, client, ws, auth):
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("*.md", MD_NOTE.encode(), "text/markdown")),
            data={"paths": "../escape.md"},
            headers=auth,
        )
        assert r.status_code == 400

    def test_alias_extensions(self, client, ws, auth):
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(
                ("笔记.markdown", "# 别名\n\n正文\n".encode(), "text/markdown"), ("页.htm", b"<p>x</p>", "text/html")
            ),
            data={"paths": ["笔记.markdown", "页.htm"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        assert r.json()["created"] == 2

    def test_two_files_same_batch(self, client, ws, auth):
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(
                ("a.md", "# a\n\n正文\n".encode(), "text/markdown"),
                ("b.md", "# b\n\n正文\n".encode(), "text/markdown"),
            ),
            data={"paths": ["a.md", "b.md"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        assert r.json()["created"] == 2

    def test_import_into_existing_folder_reused(self, client, ws, auth):
        created = client.post(DOCS.format(ws=ws.id), json={"title": "已存在目录", "is_folder": True}, headers=auth)
        assert created.status_code == 201, created.text
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("已存在目录/子/文档.md", "# 文档\n\n正文\n".encode(), "text/markdown")),
            data={"paths": ["已存在目录/子/文档.md"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        assert r.json()["folders_created"] == 1
        folders = [node["title"] for node in _flat(_tree(client, ws.id, auth)) if node["is_folder"]]
        assert sorted(folders) == ["子", "已存在目录"]

    def test_analyze_shared_folder_prefix(self, client, ws, auth):
        r = client.post(
            ANALYZE.format(ws=ws.id),
            files=_files(
                ("目录/a.md", "# a\n\n正文\n".encode(), "text/markdown"),
                ("目录/子/b.md", "# b\n\n正文\n".encode(), "text/markdown"),
            ),
            data={"paths": ["目录/a.md", "目录/子/b.md"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["total_folders"] == 2
        assert {folder["title"] for folder in data["folders"]} == {"目录", "子"}

    def test_analyze_reuses_existing_folder(self, client, ws, auth):
        created = client.post(DOCS.format(ws=ws.id), json={"title": "既有目录", "is_folder": True}, headers=auth)
        assert created.status_code == 201, created.text
        r = client.post(
            ANALYZE.format(ws=ws.id),
            files=_files(("既有目录/文档.md", "# 文档\n\n正文\n".encode(), "text/markdown")),
            data={"paths": ["既有目录/文档.md"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        assert r.json()["total_folders"] == 0

    def test_title_matches_existing_folder_renamed(self, client, ws, auth):
        created = client.post(DOCS.format(ws=ws.id), json={"title": "我的笔记", "is_folder": True}, headers=auth)
        assert created.status_code == 201, created.text
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("笔记.md", MD_NOTE.encode(), "text/markdown")),
            data={"paths": "笔记.md"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        titles = {node["title"] for node in _flat(_tree(client, ws.id, auth))}
        assert titles == {"我的笔记", "我的笔记 (2)"}

    def test_analyze_preview_limit_truncates(self, client, ws, auth, monkeypatch):
        monkeypatch.setattr(importers.importer, "PREVIEW_LIMIT", 1)
        r = client.post(
            ANALYZE.format(ws=ws.id),
            files=_files(
                ("a.md", "# a\n\n正文\n".encode(), "text/markdown"),
                ("b.md", "# b\n\n正文\n".encode(), "text/markdown"),
            ),
            data={"paths": ["a.md", "b.md"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert len(data["items"]) == 1
        assert data["preview_truncated"] is True

    def test_import_skipped_detail_limit(self, client, ws, auth, monkeypatch):
        monkeypatch.setattr(importers.importer, "DETAIL_LIMIT", 1)
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(
                ("x.txt", b"x", "text/plain"),
                ("y.txt", b"y", "text/plain"),
            ),
            data={"paths": ["x.txt", "y.txt"]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        stats = r.json()
        assert stats["skipped"] == 2
        assert len(stats["skipped_files"]) == 1

    def test_import_warnings_detail_limit(self, client, ws, auth, monkeypatch):
        monkeypatch.setattr(importers.importer, "DETAIL_LIMIT", 1)

        def build(document):
            document.add_heading("多图", level=1)
            document.add_picture(io.BytesIO(PNG_BYTES))
            document.add_picture(io.BytesIO(PNG_BYTES))

        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(
                (
                    "多图.docx",
                    _make_docx(build),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                )
            ),
            data={"paths": "多图.docx", "import_attachments": "false"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        assert len(r.json()["warnings"]) == 1

    def test_import_conflict_skip_detail_limit(self, client, ws, auth, monkeypatch):
        payload = "---\ntitle: 同名\n---\n\n正文\n".encode()
        assert (
            client.post(
                IMPORT.format(ws=ws.id),
                files=_files(("a.md", payload, "text/markdown")),
                data={"paths": "a.md"},
                headers=auth,
            ).status_code
            == 200
        )
        monkeypatch.setattr(importers.importer, "DETAIL_LIMIT", 1)
        r = client.post(
            IMPORT.format(ws=ws.id),
            files=_files(("b.md", payload, "text/markdown"), ("c.md", payload, "text/markdown")),
            data={"paths": ["b.md", "c.md"], "conflict": "skip"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        stats = r.json()
        assert stats["skipped"] == 2
        assert len(stats["skipped_files"]) == 1
