"""reports 插件集成测试."""

import io
import re

import pytest

from cndb.plugins.accounts.models import User
from cndb.plugins.reports.models import OutputFormat


@pytest.fixture
def auth_headers(client, db):
    """注册并登录，返回 Authorization header."""
    user = User(username="reporter", email="r@b.c", nickname="R")
    user.set_password("pass1234")
    db.add(user)
    db.commit()
    resp = client.post(
        "/api/v1/accounts/auth/login",
        json={"login": "reporter", "password": "pass1234"},
    )
    token = resp.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


class TestReportTemplateCRUD:
    """报告模板 CRUD 测试."""

    def test_create_template(self, client, auth_headers):
        resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "员工名册",
                "description": "列出所有员工姓名",
                "output_format": "docx",
                "template_content": "# 员工名册\n{% for r in records %}{{ r.name }}\n{% endfor %}",
                "parameters": [],
            },
        )
        assert resp.status_code == 201
        data = resp.json()
        assert data["name"] == "员工名册"
        assert data["output_format"] == "docx"
        assert data["id"] > 0

    def test_list_templates(self, client, auth_headers):
        # 创建两个模板
        for i in range(2):
            client.post(
                "/api/v1/reports",
                headers=auth_headers,
                json={
                    "name": f"模板{i}",
                    "output_format": "docx",
                    "template_content": "{{ table_name }}",
                },
            )
        resp = client.get("/api/v1/reports", headers=auth_headers)
        assert resp.status_code == 200
        items = resp.json()
        assert len(items) == 2

    def test_get_template(self, client, auth_headers):
        create_resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "模板A",
                "output_format": "pdf",
                "template_content": "hello {{ params.msg }}",
            },
        )
        tid = create_resp.json()["id"]
        resp = client.get(f"/api/v1/reports/{tid}", headers=auth_headers)
        assert resp.status_code == 200
        assert resp.json()["name"] == "模板A"

    def test_get_template_not_found(self, client, auth_headers):
        resp = client.get("/api/v1/reports/99999", headers=auth_headers)
        assert resp.status_code == 404

    def test_update_template(self, client, auth_headers):
        create_resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "旧名",
                "output_format": "docx",
                "template_content": "old",
            },
        )
        tid = create_resp.json()["id"]
        resp = client.put(
            f"/api/v1/reports/{tid}",
            headers=auth_headers,
            json={"name": "新名", "output_format": "xlsx"},
        )
        assert resp.status_code == 200
        assert resp.json()["name"] == "新名"
        assert resp.json()["output_format"] == "xlsx"

    def test_delete_template(self, client, auth_headers):
        create_resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "待删",
                "output_format": "docx",
                "template_content": "del",
            },
        )
        tid = create_resp.json()["id"]
        resp = client.delete(f"/api/v1/reports/{tid}", headers=auth_headers)
        assert resp.status_code == 204
        resp = client.get(f"/api/v1/reports/{tid}", headers=auth_headers)
        assert resp.status_code == 404


class TestReportTemplateValidation:
    """模板校验测试."""

    def test_invalid_format_create(self, client, auth_headers):
        resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "坏模板",
                "output_format": "zip",
                "template_content": "x",
            },
        )
        assert resp.status_code == 400
        assert "不支持的输出格式" in resp.json()["detail"]

    def test_invalid_format_update(self, client, auth_headers):
        create_resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "模板",
                "output_format": "docx",
                "template_content": "x",
            },
        )
        tid = create_resp.json()["id"]
        resp = client.put(
            f"/api/v1/reports/{tid}",
            headers=auth_headers,
            json={"output_format": "tar"},
        )
        assert resp.status_code == 400

    def test_bad_jinja_syntax_create(self, client, auth_headers):
        resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "坏语法",
                "output_format": "docx",
                "template_content": "{% if %}",
            },
        )
        assert resp.status_code == 400
        assert "模板语法错误" in resp.json()["detail"]

    def test_bad_jinja_syntax_update(self, client, auth_headers):
        create_resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "模板",
                "output_format": "docx",
                "template_content": "ok",
            },
        )
        tid = create_resp.json()["id"]
        resp = client.put(
            f"/api/v1/reports/{tid}",
            headers=auth_headers,
            json={"template_content": "{% xyz %}"},
        )
        assert resp.status_code == 400


class TestReportRender:
    """报告渲染测试."""

    def test_render_docx(self, client, auth_headers, db):
        # 创建 workspace + table + fields
        ws = client.post(
            "/api/v1/workspaces",
            headers=auth_headers,
            json={"name": "测试工作区"},
        )
        wid = ws.json()["id"]
        table_resp = client.post(
            f"/api/v1/workspaces/{wid}/tables",
            headers=auth_headers,
            json={"name": "员工表"},
        )
        tid = table_resp.json()["id"]
        client.post(
            f"/api/v1/workspaces/{wid}/tables/{tid}/fields",
            headers=auth_headers,
            json={"name": "姓名", "field_type": "text"},
        )
        # 添加一条记录
        client.post(
            f"/api/v1/workspaces/{wid}/tables/{tid}/records",
            headers=auth_headers,
            json={"data": {"姓名": "张三"}},
        )
        # 创建模板
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "docx报告",
                "output_format": "docx",
                "template_content": "hello {{ table_name }}\\n{% for r in records %}{{ r.get('姓名', '') }}\\n{% endfor %}",
            },
        )
        tpl_id = tpl.json()["id"]
        # 渲染
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid, "params": {}},
        )
        assert resp.status_code == 200
        assert "application/vnd.openxmlformats" in resp.headers["content-type"]
        assert len(resp.content) > 0

    def test_render_pdf(self, client, auth_headers, db):
        ws = client.post(
            "/api/v1/workspaces",
            headers=auth_headers,
            json={"name": "ws2"},
        )
        wid = ws.json()["id"]
        table_resp = client.post(
            f"/api/v1/workspaces/{wid}/tables",
            headers=auth_headers,
            json={"name": "t2"},
        )
        tid = table_resp.json()["id"]
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "pdf报告",
                "output_format": "pdf",
                "template_content": "PDF 测试 {{ table_name }}",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid, "params": {}},
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/pdf"
        assert len(resp.content) > 0

    def test_render_xlsx(self, client, auth_headers, db):
        ws = client.post(
            "/api/v1/workspaces",
            headers=auth_headers,
            json={"name": "ws3"},
        )
        wid = ws.json()["id"]
        table_resp = client.post(
            f"/api/v1/workspaces/{wid}/tables",
            headers=auth_headers,
            json={"name": "t3"},
        )
        tid = table_resp.json()["id"]
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "xlsx报告",
                "output_format": "xlsx",
                "template_content": "col1\\tcol2\\n{{ table_name }}\\tok",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid, "params": {}},
        )
        assert resp.status_code == 200
        assert "spreadsheetml" in resp.headers["content-type"]

    def test_render_injects_generated_at(self, client, auth_headers, db):
        """渲染上下文应注入 generated_at 生成日期标签（YYYY-MM-DD HH:mm）."""
        ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws-date"})
        wid = ws.json()["id"]
        table_resp = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "日期表"})
        tid = table_resp.json()["id"]
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "日期标签报告",
                "output_format": "docx",
                "template_content": "生成于 {{ generated_at }}",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid, "params": {}},
        )
        assert resp.status_code == 200

        from docx import Document

        doc = Document(io.BytesIO(resp.content))
        text = "\n".join(p.text for p in doc.paragraphs)
        m = re.search(r"\d{4}-\d{2}-\d{2} \d{2}:\d{2}", text)
        assert m is not None, f"正文未包含生成日期，实际：{text!r}"
        assert text.strip() == f"生成于 {m.group(0)}"

    def test_render_filename_with_timestamp(self, client, auth_headers, db):
        """下载文件名应含模板名与生成时间戳，便于区分不同批次生成的报告."""
        ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws-fn"})
        wid = ws.json()["id"]
        table_resp = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "文件名表"})
        tid = table_resp.json()["id"]
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "月度汇总",
                "output_format": "docx",
                "template_content": "ok",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid, "params": {}},
        )
        assert resp.status_code == 200
        disposition = resp.headers["content-disposition"]
        # filename* (RFC 5987) 携带 UTF-8 中文名 + 时间戳 {模板名}_{YYYYMMDD_HHMMSS}.docx
        assert "filename*=UTF-8''" in disposition
        assert re.search(r"%E6%9C%88%E5%BA%A6%E6%B1%87%E6%80%BB_\d{8}_\d{6}\.docx", disposition)

    def test_render_template_not_found(self, client, auth_headers):
        resp = client.post(
            "/api/v1/reports/99999/render",
            headers=auth_headers,
            json={"table_id": 1, "params": {}},
        )
        assert resp.status_code == 404

    def test_render_table_not_found(self, client, auth_headers):
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "模板",
                "output_format": "docx",
                "template_content": "x",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": 99999, "params": {}},
        )
        assert resp.status_code == 404

    def test_render_template_runtime_error(self, client, auth_headers, db):
        ws = client.post(
            "/api/v1/workspaces",
            headers=auth_headers,
            json={"name": "ws4"},
        )
        wid = ws.json()["id"]
        table_resp = client.post(
            f"/api/v1/workspaces/{wid}/tables",
            headers=auth_headers,
            json={"name": "t4"},
        )
        tid = table_resp.json()["id"]
        # 用 StrictUndefined 渲染不存在的变量，应该 400
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "坏模板",
                "output_format": "docx",
                "template_content": "{{ undefined_var }}",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid, "params": {}},
        )
        assert resp.status_code == 400
        assert "模板渲染失败" in resp.json()["detail"]

    def test_output_format_enum_values(self):
        """OutputFormat 枚举应包含三种格式."""
        assert OutputFormat.DOCX == "docx"
        assert OutputFormat.PDF == "pdf"
        assert OutputFormat.XLSX == "xlsx"


class TestReportTemplateOwnership:
    """模板归属 table_id 校验测试."""

    def test_create_template_with_valid_table(self, client, auth_headers):
        """绑定存在的 table_id 应成功."""
        ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws_own"})
        wid = ws.json()["id"]
        tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "t_own"})
        tid = tbl.json()["id"]
        resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "带归属",
                "table_id": tid,
                "output_format": "docx",
                "template_content": "x",
            },
        )
        assert resp.status_code == 201
        assert resp.json()["table_id"] == tid

    def test_create_template_with_invalid_table(self, client, auth_headers):
        """绑定不存在的 table_id 应 404."""
        resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "坏归属",
                "table_id": 99999,
                "output_format": "docx",
                "template_content": "x",
            },
        )
        assert resp.status_code == 404

    def test_update_template_with_invalid_table(self, client, auth_headers):
        """更新时传入不存在的 table_id 应 404."""
        create_resp = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "模板",
                "output_format": "docx",
                "template_content": "x",
            },
        )
        tid = create_resp.json()["id"]
        resp = client.put(
            f"/api/v1/reports/{tid}",
            headers=auth_headers,
            json={"table_id": 99999},
        )
        assert resp.status_code == 404


def test_update_template_parameters(client, auth_headers):
    """更新 parameters 字段应正确生效."""
    create_resp = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={
            "name": "模板",
            "output_format": "docx",
            "template_content": "{{ params.msg }}",
            "parameters": [{"name": "msg", "type": "string", "required": True}],
        },
    )
    tid = create_resp.json()["id"]
    resp = client.put(
        f"/api/v1/reports/{tid}",
        headers=auth_headers,
        json={"parameters": [{"name": "new_param", "type": "number"}]},
    )
    assert resp.status_code == 200
    params = resp.json()["parameters"]
    assert len(params) == 1
    assert params[0]["name"] == "new_param"


def test_update_template_to_null_table_id(client, auth_headers, db):
    """已绑定 table 的模板可以解除绑定（table_id=null）."""
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws_null"})
    wid = ws.json()["id"]
    tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "t_null"})
    tid = tbl.json()["id"]

    create_resp = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={
            "name": "模板",
            "table_id": tid,
            "output_format": "docx",
            "template_content": "x",
        },
    )
    rep_id = create_resp.json()["id"]
    assert create_resp.json()["table_id"] == tid

    resp = client.put(
        f"/api/v1/reports/{rep_id}",
        headers=auth_headers,
        json={"table_id": None},
    )
    assert resp.status_code == 200
    assert resp.json()["table_id"] is None


def test_plugin_register_routes_fallback(monkeypatch):
    """覆盖 ReportsPlugin.register_routes fallback 分支（direct_router 未设置时）."""
    from fastapi import APIRouter

    from cndb.plugins.reports.plugin import ReportsPlugin

    # monkeypatch APIRouter.include_router 避免空 path 校验报错
    monkeypatch.setattr(APIRouter, "include_router", lambda self, r: None)

    plugin = ReportsPlugin()
    router = APIRouter()
    plugin.register_routes(router)
    assert plugin.direct_router is not None


def test_render_unsupported_format(client, auth_headers, db):
    """如果模板格式不被渲染器支持（当前只有 docx/pdf/xlsx），应 400."""
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws_uf"})
    wid = ws.json()["id"]
    tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "t_uf"})
    tid = tbl.json()["id"]
    # 手动在 DB 里把 output_format 改成一个非法值（绕过 create 时的校验）
    from sqlalchemy import update as sa_update

    from cndb.plugins.reports.models import ReportTemplate

    tpl = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={
            "name": "坏模板",
            "output_format": "docx",
            "template_content": "x",
        },
    )
    rep_id = tpl.json()["id"]
    # 直接改 DB
    db.execute(sa_update(ReportTemplate).where(ReportTemplate.id == rep_id).values(output_format="xyz"))
    db.commit()
    resp = client.post(
        f"/api/v1/reports/{rep_id}/render",
        headers=auth_headers,
        json={"table_id": tid, "params": {}},
    )
    assert resp.status_code == 400


def test_render_denied_when_user_no_read_permission(client, auth_headers, db):
    """无工作区成员身份的用户不应能通过 render 端点拿到其他表的数据。"""
    # 用户 A（auth_headers）创建 workspace + table + record + template
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws_secret"})
    wid = ws.json()["id"]
    tbl = client.post(
        f"/api/v1/workspaces/{wid}/tables",
        headers=auth_headers,
        json={"name": "secret_table"},
    )
    tid = tbl.json()["id"]
    client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/fields",
        headers=auth_headers,
        json={"name": "姓名", "field_type": "text"},
    )
    client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/records",
        headers=auth_headers,
        json={"data": {"姓名": "机密数据"}},
    )
    tpl = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={
            "name": "泄漏测试",
            "output_format": "docx",
            "template_content": "{{ records }}",
        },
    )
    rep_id = tpl.json()["id"]

    # 用户 B（全新账号，不在 A 的 workspace 里）
    client.post(
        "/api/v1/accounts/auth/register",
        json={"username": "user_b", "email": "b@b.com", "password": "passw0rd"},
    )
    r = client.post(
        "/api/v1/accounts/auth/login",
        json={"login": "user_b", "password": "passw0rd"},
    )
    token_b = r.json()["access_token"]
    headers_b = {"Authorization": f"Bearer {token_b}"}

    # B 尝试渲染 A 的表 → 应该 403
    resp = client.post(
        f"/api/v1/reports/{rep_id}/render",
        headers=headers_b,
        json={"table_id": tid, "params": {}},
    )
    assert resp.status_code == 403, f"应返回 403 但实际 {resp.status_code}: {resp.text}"


# ── 跨工作区额外引用表（extra_table_ids）──────────────────


def _mk_table(client, headers, ws_name, table_name, field_name="名称", value="v1"):
    """辅助：建工作区 + 表 + 字段 + 一条记录，返回 (wid, tid)."""
    ws = client.post("/api/v1/workspaces", headers=headers, json={"name": ws_name})
    wid = ws.json()["id"]
    tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=headers, json={"name": table_name})
    tid = tbl.json()["id"]
    client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/fields",
        headers=headers,
        json={"name": field_name, "field_type": "text"},
    )
    # records 创建负载契约是 {values: {...}}（FastAPI RowCreate 为准）
    created = client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/records",
        headers=headers,
        json={"values": {field_name: value}},
    )
    assert created.status_code == 201, created.text
    return wid, tid


def test_extra_table_ids_persist_roundtrip(client, auth_headers):
    """extra_table_ids 应随模板持久化：create 写入 → get 读回 → update 修改."""
    _ws, tid_a = _mk_table(client, auth_headers, "ws-per-a", "主表A")
    _ws2, tid_b = _mk_table(client, auth_headers, "ws-per-b", "额外表B")
    _ws3, tid_c = _mk_table(client, auth_headers, "ws-per-c", "额外表C")

    create_resp = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={
            "name": "持久化模板",
            "table_id": tid_a,
            "extra_table_ids": [tid_b],
            "output_format": "docx",
            "template_content": "x",
        },
    )
    assert create_resp.status_code == 201
    assert create_resp.json()["extra_table_ids"] == [tid_b]

    rep_id = create_resp.json()["id"]
    got = client.get(f"/api/v1/reports/{rep_id}", headers=auth_headers)
    assert got.json()["extra_table_ids"] == [tid_b]

    upd = client.put(f"/api/v1/reports/{rep_id}", headers=auth_headers, json={"extra_table_ids": [tid_b, tid_c]})
    assert upd.status_code == 200
    assert upd.json()["extra_table_ids"] == [tid_b, tid_c]

    # 未传 extra_table_ids 的模板默认空列表
    plain = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={"name": "无额外表模板", "output_format": "docx", "template_content": "x"},
    )
    assert plain.json()["extra_table_ids"] == []


def test_extra_table_ids_create_rejects_missing_table(client, auth_headers):
    """create/update 时 extra_table_ids 含不存在的表 id 应 404."""
    _ws, tid_a = _mk_table(client, auth_headers, "ws-ex-a", "表A")
    resp = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={
            "name": "坏引用",
            "table_id": tid_a,
            "extra_table_ids": [999999],
            "output_format": "docx",
            "template_content": "x",
        },
    )
    assert resp.status_code == 404


def test_render_cross_workspace_extra_table(client, auth_headers):
    """渲染时应能引用其他工作区（有 READ 权限）的表数据."""
    _ws1, tid_a = _mk_table(client, auth_headers, "ws-cx-a", "主表甲")
    _ws2, tid_b = _mk_table(client, auth_headers, "ws-cx-b", "跨区表乙", field_name="跨区字段", value="跨区数据")

    tpl = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={
            "name": "跨区渲染",
            "table_id": tid_a,
            "extra_table_ids": [tid_b],
            "output_format": "docx",
            "template_content": "值={{ records_by_table['跨区表乙'][0].跨区字段 }}",
        },
    )
    rep_id = tpl.json()["id"]
    resp = client.post(
        f"/api/v1/reports/{rep_id}/render",
        headers=auth_headers,
        json={"table_id": tid_a, "params": {}, "extra_table_ids": [tid_b]},
    )
    assert resp.status_code == 200

    from docx import Document

    doc = Document(io.BytesIO(resp.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "值=跨区数据" in text


# ── 迭代 3：统计计算函数 ─────────────────────────────


def _mk_table_with_fields(client, auth_headers, wid, name, fields):
    """建表 + 字段（fields 是 list[tuple[str, str]]：(字段名, field_type)），返回 tid."""
    tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": name})
    tid = tbl.json()["id"]
    for fname, ftype in fields:
        client.post(
            f"/api/v1/workspaces/{wid}/tables/{tid}/fields",
            headers=auth_headers,
            json={"name": fname, "field_type": ftype},
        )
    return tid


class TestStatsFunctions:
    """Jinja2 沙箱注册的统计函数测试."""

    def test_render_stats_sum_avg_count(self, client, auth_headers, db):
        """stats(records, '金额') 应返回 {sum, avg, count, min, max}."""
        ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws-stats"}).json()
        wid = ws["id"]
        tid = _mk_table_with_fields(client, auth_headers, wid, "销售表", [("名称", "text"), ("金额", "number")])
        # 三条记录：100 / 200 / 300
        for amt in (100, 200, 300):
            client.post(
                f"/api/v1/workspaces/{wid}/tables/{tid}/records",
                headers=auth_headers,
                json={"values": {"名称": f"商品{amt}", "金额": amt}},
            )
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "统计汇总",
                "output_format": "docx",
                "template_content": (
                    "总金额={{ stats(records, '金额').sum }}\n"
                    "平均={{ stats(records, '金额').avg }}\n"
                    "数量={{ stats(records, '金额').count }}\n"
                    "最小={{ stats(records, '金额').min }}\n"
                    "最大={{ stats(records, '金额').max }}"
                ),
            },
        ).json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl}/render",
            headers=auth_headers,
            json={"table_id": tid, "params": {}},
        )
        assert resp.status_code == 200
        from docx import Document

        text = "\n".join(p.text for p in Document(io.BytesIO(resp.content)).paragraphs)
        # 600.0 或 600 都可以
        assert re.search(r"总金额=600(\.0+)?", text), text
        assert re.search(r"平均=200(\.0+)?", text), text
        assert "数量=3" in text
        assert re.search(r"最小=100(\.0+)?", text), text
        assert re.search(r"最大=300(\.0+)?", text), text

    def test_render_group_stats(self, client, auth_headers, db):
        """group_stats(records, '类别', '金额') 应按类别分组聚合金额."""
        ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws-grp"}).json()
        wid = ws["id"]
        tid = _mk_table_with_fields(client, auth_headers, wid, "分类表", [("类别", "text"), ("金额", "number")])
        for cat, amt in (("A", 100), ("A", 200), ("B", 300), ("B", 400)):
            client.post(
                f"/api/v1/workspaces/{wid}/tables/{tid}/records",
                headers=auth_headers,
                json={"values": {"类别": cat, "金额": amt}},
            )
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "分组汇总",
                "output_format": "docx",
                "template_content": (
                    "{% for g in group_stats(records, '类别', '金额') %}{{ g.key }}={{ g.sum }}\n{% endfor %}"
                ),
            },
        ).json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl}/render",
            headers=auth_headers,
            json={"table_id": tid, "params": {}},
        )
        assert resp.status_code == 200
        from docx import Document

        text = "\n".join(p.text for p in Document(io.BytesIO(resp.content)).paragraphs)
        assert re.search(r"A=300(\.0+)?", text), text
        assert re.search(r"B=700(\.0+)?", text), text

    def test_render_stats_on_extra_table(self, client, auth_headers, db):
        """stats 应能对额外引用表的数据做统计（records_by_table['经费表']）."""
        ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws-stats-extra"}).json()
        wid = ws["id"]
        tid_main = _mk_table_with_fields(client, auth_headers, wid, "项目表", [("项目名", "text")])
        tid_extra = _mk_table_with_fields(client, auth_headers, wid, "经费表", [("项目名", "text"), ("预算", "number")])
        client.post(
            f"/api/v1/workspaces/{wid}/tables/{tid_main}/records",
            headers=auth_headers,
            json={"values": {"项目名": "科研项目X"}},
        )
        for budget in (50000, 80000, 120000):
            client.post(
                f"/api/v1/workspaces/{wid}/tables/{tid_extra}/records",
                headers=auth_headers,
                json={"values": {"项目名": "科研项目X", "预算": budget}},
            )
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "跨表统计",
                "output_format": "docx",
                "template_content": "经费总额={{ stats(records_by_table['经费表'], '预算').sum }}",
                "extra_table_ids": [tid_extra],
            },
        ).json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl}/render",
            headers=auth_headers,
            json={"table_id": tid_main, "params": {}, "extra_table_ids": [tid_extra]},
        )
        assert resp.status_code == 200
        from docx import Document

        text = "\n".join(p.text for p in Document(io.BytesIO(resp.content)).paragraphs)
        assert re.search(r"经费总额=250000(\.0+)?", text), text
