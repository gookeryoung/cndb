"""数据管理健壮性测试 — 语义等价 diff / 预览截断 / 导出特殊字符 / 公式注入防护 / 损坏文件 / 端到端 roundtrip.

聚焦错误场景与边界锁定：
- upsert diff 语义等价（"1.0" vs 1、True vs "true"、文本空白不误报）
- preview 200 行截断（计数不截断、预览截断）
- 导出特殊字符往返一致（逗号/引号/换行/emoji/公式串/参差行）
- CSV 公式注入防护（危险前缀转义 / XLSX 公式单元格回字符串 / 失败行导出）
- 损坏文件解析（垃圾字节 xlsx、仅表头 CSV、重复表头）
- 导入→导出→再导入 roundtrip 一致性
"""

from __future__ import annotations

import csv
import io
import zipfile

import pytest
from openpyxl import load_workbook
from sqlalchemy.orm import Session

from cndb.plugins.tables.services.core import ddl
from cndb.plugins.tables import transfer
from cndb.plugins.tables.services.importing.diff_reporter import _json_safe
from cndb.plugins.tables.services.importing.failed_row_exporter import FailedRowExporter
from cndb.plugins.tables.services.importing.import_tasks import create_import_task, execute_import_task
from cndb.plugins.tables.services.importing.importer import Importer
from cndb.plugins.tables.models import DataField, DataTable
from cndb.plugins.tables.services.importing.row_validator import Issue, ValidationResult
from cndb.plugins.tables.transfer import export_rows_to_csv, export_rows_to_xlsx
from tests.helpers import wait_import_settled

# ── 公共 Fixture 与造数工具 ─────────────────────────


@pytest.fixture
def test_session(db_engine, db):
    """复用 conftest 内存 DB，yield (engine, session)."""
    yield db_engine, db


def _make_table(session: Session, engine, *, workspace_id: int = 1) -> DataTable:
    """创建一个 DataTable（不含字段）."""
    dt = DataTable(workspace_id=workspace_id, name="健壮性测试表")
    dt.ensure_db_name()
    session.add(dt)
    session.commit()
    session.refresh(dt)
    return dt


def _add_field(
    session: Session,
    table: DataTable,
    name: str,
    field_type: str,
    *,
    required: bool = False,
    config: dict | None = None,
    order: int = 0,
) -> DataField:
    """为表添加一个字段（flush 不 commit，由调用方统一提交）."""
    f = DataField(
        table_id=table.id, name=name, field_type=field_type, required=required, config=config or {}, order=order
    )
    f.ensure_db_name()
    session.add(f)
    session.flush()
    return f


def _make_table_with_fields(session: Session, engine, fields: list[tuple[str, str]]) -> DataTable:
    """建表 + 逐个添加字段 + 建物理表，fields 为 (name, field_type) 列表."""
    table = _make_table(session, engine)
    for i, (name, ftype) in enumerate(fields):
        _add_field(session, table, name, ftype, order=i)
    session.commit()
    ddl.create_table(engine, table)
    return table


# ── A: upsert 语义等价 diff ─────────────────────────


class TestUpsertSemanticEquivalence:
    """命中行的语义等价值不应被误报为字段变化."""

    def test_upsert_numeric_string_no_false_diff(self, test_session):
        """number 库内 1 vs 文件 "1.0" —— 数值语义等价，无变化."""
        engine, session = test_session
        from cndb.plugins.tables.services.core import records as rec

        table = _make_table_with_fields(session, engine, [("code", "text"), ("qty", "number")])
        rec.bulk_create(engine, table, [{"code": "A1", "qty": 1}], db=session)

        imp = Importer(engine, session, table)
        rpt = imp.analyze("code,qty\nA1,1.0\n", "csv", match_keys=["code"]).report

        assert rpt["update_count"] == 1
        assert rpt["update_changed_count"] == 0
        assert rpt["update_no_change_count"] == 1
        assert rpt["update_preview"][0]["field_diffs"] == {}

    def test_upsert_boolean_string_no_false_diff(self, test_session):
        """boolean 库内 True vs 文件 "true" —— 布尔语义等价，无变化."""
        engine, session = test_session
        from cndb.plugins.tables.services.core import records as rec

        table = _make_table_with_fields(session, engine, [("code", "text"), ("active", "boolean")])
        rec.bulk_create(engine, table, [{"code": "A1", "active": True}], db=session)

        imp = Importer(engine, session, table)
        rpt = imp.analyze("code,active\nA1,true\n", "csv", match_keys=["code"]).report

        assert rpt["update_no_change_count"] == 1
        assert rpt["update_changed_count"] == 0
        assert rpt["update_preview"][0]["field_diffs"] == {}

    def test_upsert_whitespace_no_false_diff(self, test_session):
        """text "abc" vs " abc " —— 字符串去空白比较，无变化."""
        engine, session = test_session
        from cndb.plugins.tables.services.core import records as rec

        table = _make_table_with_fields(session, engine, [("code", "text"), ("name", "text")])
        rec.bulk_create(engine, table, [{"code": "A1", "name": "abc"}], db=session)

        imp = Importer(engine, session, table)
        rpt = imp.analyze("code,name\nA1, abc \n", "csv", match_keys=["code"]).report

        assert rpt["update_no_change_count"] == 1
        assert rpt["update_changed_count"] == 0
        assert rpt["update_preview"][0]["field_diffs"] == {}


# ── A: 预览截断 ─────────────────────────────────────


class TestPreviewLimit:
    """preview 按 PREVIEW_LIMIT=200 截断，计数保持全量."""

    def test_update_preview_truncated_at_200(self, test_session):
        """205 行全命中已有 key → update_count==205 但 update_preview 仅 200 条."""
        engine, session = test_session
        from cndb.plugins.tables.services.core import records as rec

        table = _make_table_with_fields(session, engine, [("code", "text"), ("name", "text")])
        rec.bulk_create(engine, table, [{"code": f"K{i:04d}", "name": f"旧{i}"} for i in range(205)], db=session)

        csv_content = "code,name\n" + "".join(f"K{i:04d},新{i}\n" for i in range(205))
        imp = Importer(engine, session, table)
        rpt = imp.analyze(csv_content, "csv", match_keys=["code"]).report

        assert rpt["update_count"] == 205
        assert len(rpt["update_preview"]) == 200
        assert rpt["update_changed_count"] == 205

    def test_new_preview_truncated_at_200(self, test_session):
        """205 行全新 key → new_count==205 但 new_preview 仅 200 条."""
        engine, session = test_session
        table = _make_table_with_fields(session, engine, [("code", "text")])

        csv_content = "code\n" + "".join(f"K{i:04d}\n" for i in range(205))
        imp = Importer(engine, session, table)
        rpt = imp.analyze(csv_content, "csv", match_keys=["code"]).report

        assert rpt["new_count"] == 205
        assert len(rpt["new_preview"]) == 200


# ── B: 导出特殊字符往返 ─────────────────────────────


class TestExportSpecialChars:
    """导出函数对特殊字符的转义与参差行的行为锁定."""

    _SPECIAL_ROWS = [
        {"name": "值,带逗号", "note": '含"引号"', "bio": "第一行\n第二行", "cmd": "=cmd()"},
        {"name": "emoji🚀", "note": "中文标点；", "bio": "tab\t制表", "cmd": "+1+1"},
    ]

    def test_csv_roundtrip_special_chars(self):
        """CSV 导出含逗号/引号/换行/emoji → DictReader 解析回逐值相等；公式串按注入防护转义."""
        out = export_rows_to_csv(self._SPECIAL_ROWS)
        reader = csv.DictReader(io.StringIO(out))
        parsed = list(reader)
        assert len(parsed) == 2
        for original, row in zip(self._SPECIAL_ROWS, parsed, strict=True):
            for key, expected in original.items():
                if key == "cmd":
                    continue  # 公式样串走注入防护转义，下方单独断言
                assert row[key] == expected, f"列 {key} 往返不一致: {row[key]!r} != {expected!r}"
        # 公式样串前缀 ' 转义（OWASP CSV 注入防护）
        assert parsed[0]["cmd"] == "'=cmd()"
        assert parsed[1]["cmd"] == "'+1+1"

    def test_xlsx_roundtrip_special_chars(self):
        """XLSX 导出同值域 → load_workbook 读回逐值相等."""
        xlsx_bytes = export_rows_to_xlsx(self._SPECIAL_ROWS)
        wb = load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        assert ws is not None
        rows = list(ws.iter_rows(values_only=True))
        header = list(rows[0])
        assert header == list(self._SPECIAL_ROWS[0].keys())
        for original, row in zip(self._SPECIAL_ROWS, rows[1:], strict=True):
            for key, expected in original.items():
                assert row[header.index(key)] == expected, f"列 {key} 往返不一致: {row!r} != {expected!r}"

    def test_csv_ragged_extra_column_raises(self):
        """第二行含 fieldnames 之外的键 → DictWriter 默认 extrasaction=raise 抛 ValueError."""
        rows = [{"a": "1"}, {"a": "2", "b": "3"}]
        with pytest.raises(ValueError, match="fields not in fieldnames"):
            export_rows_to_csv(rows)

    def test_csv_ragged_missing_column_blank(self):
        """第二行缺键 → CSV 写空串；XLSX 该格为 None."""
        rows = [{"a": "1", "b": "x"}, {"a": "2"}]

        out = export_rows_to_csv(rows)
        parsed = list(csv.DictReader(io.StringIO(out)))
        assert parsed[1]["a"] == "2"
        assert parsed[1]["b"] == ""

        xlsx_bytes = export_rows_to_xlsx(rows)
        wb = load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        assert ws is not None
        data_rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert data_rows[1] == ("2", None)


# ── B2: CSV 公式注入防护 ────────────────────────────


class TestCsvFormulaInjectionProtection:
    """导出链路对 CSV 公式注入的防护（OWASP 缓解：危险前缀加 ' 前缀）."""

    @pytest.mark.parametrize(
        ("payload", "expected"),
        [
            ("=cmd|' /C calc'!A0", "'=cmd|' /C calc'!A0"),
            ("+SUM(A1)", "'+SUM(A1)"),
            ("-2+3", "'-2+3"),
            ("@SUM(1)", "'@SUM(1)"),
            ("\t缩进值", "'\t缩进值"),
            ("\rCR值", "'\rCR值"),
        ],
        ids=["dde_cmd", "plus_formula", "minus_expr", "at_formula", "tab_prefix", "cr_prefix"],
    )
    def test_csv_dangerous_prefix_escaped(self, payload: str, expected: str):
        """以 =/+/-/@/Tab/CR 开头的字符串值导出时前缀 ' 转义，不被 Excel 当公式求值."""
        out = export_rows_to_csv([{"cmd": payload}])
        parsed = list(csv.DictReader(io.StringIO(out)))
        assert parsed[0]["cmd"] == expected

    def test_csv_safe_values_not_escaped(self):
        """普通字符串与非 str 值（int/float/bool/None）不做转义."""
        rows = [{"s": "abc", "paren": "(1,234)", "num": 123, "neg": -1.5, "flag": True, "blank": None}]
        out = export_rows_to_csv(rows)
        parsed = next(iter(csv.DictReader(io.StringIO(out))))
        assert parsed["s"] == "abc"
        assert parsed["paren"] == "(1,234)"
        assert parsed["num"] == "123"
        assert parsed["neg"] == "-1.5"
        assert parsed["flag"] == "True"
        assert parsed["blank"] == ""

    def test_csv_header_dangerous_prefix_escaped(self):
        """危险前缀表头同步转义，数据列仍按转义后列名对齐."""
        out = export_rows_to_csv([{"=bad": "x", "ok": "=SUM(B1)"}])
        parsed = list(csv.DictReader(io.StringIO(out)))
        assert list(parsed[0].keys()) == ["'=bad", "ok"]
        assert parsed[0]["'=bad"] == "x"
        assert parsed[0]["ok"] == "'=SUM(B1)"

    def test_xlsx_formula_cell_forced_to_string(self):
        """= 开头单元格 data_type 强制为 's'，值原样保留，Excel 打开不被求值."""
        xlsx_bytes = export_rows_to_xlsx([{"cmd": "=SUM(A1)", "note": "+1+1"}])
        wb = load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        assert ws is not None
        assert ws["A2"].value == "=SUM(A1)"
        assert ws["A2"].data_type == "s"
        # + 开头在 XLSX 中本就存为字符串类型
        assert ws["B2"].value == "+1+1"
        assert ws["B2"].data_type == "s"

    def test_xlsx_header_formula_cell_forced_to_string(self):
        """表头 = 开头单元格同样强制字符串（值不丢）."""
        xlsx_bytes = export_rows_to_xlsx([{"=bad": "x"}])
        wb = load_workbook(io.BytesIO(xlsx_bytes))
        ws = wb.active
        assert ws is not None
        assert ws["A1"].value == "=bad"
        assert ws["A1"].data_type == "s"

    def test_failed_rows_csv_escaped(self):
        """失败行 CSV 导出对危险前缀值同样转义，_error 列保持可读."""
        result = ValidationResult(
            row_number=3,
            values={"name": "=cmd()", "qty": "abc"},
            status="error",
            issues=[Issue(field="qty", level="error", message="qty 必须为数字")],
        )
        data = FailedRowExporter.export_failed_rows([result], "csv")
        parsed = list(csv.DictReader(io.StringIO(data.decode("utf-8-sig"))))
        assert parsed[0]["name"] == "'=cmd()"
        assert parsed[0]["_error"] == "qty 必须为数字"

    def test_failed_rows_xlsx_formula_cell_forced_to_string(self):
        """失败行 XLSX 导出 = 开头值强制字符串，值原样保留."""
        result = ValidationResult(
            row_number=1,
            values={"name": '=HYPERLINK("http://evil")'},
            status="error",
            issues=[Issue(field="name", level="error", message="必填校验失败")],
        )
        data = FailedRowExporter.export_failed_rows([result], "xlsx")
        wb = load_workbook(io.BytesIO(data))
        ws = wb.active
        assert ws is not None
        cell = ws["A2"]
        assert cell.value == '=HYPERLINK("http://evil")'
        assert cell.data_type == "s"


# ── C: 损坏文件解析 ─────────────────────────────────


class TestCorruptFileParsing:
    """损坏/退化输入的解析错误路径."""

    def test_analyze_garbage_xlsx_raises(self, test_session):
        """垃圾字节标 fmt=xlsx → openpyxl 抛 zipfile.BadZipFile（不被包装吞掉）."""
        engine, session = test_session
        table = _make_table_with_fields(session, engine, [("name", "text")])

        imp = Importer(engine, session, table)
        with pytest.raises(zipfile.BadZipFile, match="not a zip file"):
            imp.analyze(b"garbage-bytes-not-a-zip", "xlsx")

    def test_task_garbage_xlsx_fails_with_message(self, test_session, db):
        """任务层垃圾 xlsx → 状态 failed 且 error_message 非空（兜底不崩溃）."""
        engine, session = test_session
        table = _make_table_with_fields(session, engine, [("name", "text")])

        task = create_import_task(
            session, table_id=table.id, user_id=1, filename="bad.xlsx", fmt="xlsx", content=b"garbage-bytes"
        )
        execute_import_task(session, task.id)
        db.refresh(task)

        assert task.status == "failed"
        assert task.error_message, "失败任务必须携带 error_message"

    def test_analyze_csv_header_only(self, test_session):
        """仅表头无数据行 → total==0，正常返回不异常."""
        engine, session = test_session
        table = _make_table_with_fields(session, engine, [("name", "text"), ("email", "email")])

        imp = Importer(engine, session, table)
        rpt = imp.analyze("name,email\n", "csv").report

        assert rpt["total"] == 0
        assert rpt["valid_count"] == 0
        assert rpt["error_count"] == 0

    def test_parse_csv_duplicate_headers(self):
        """重复表头列名 → file_columns 保留两个 name，行值后列覆盖."""
        rows, file_columns = Importer._parse_csv("name,name\na,b\n")
        assert file_columns == ["name", "name"]
        assert len(rows) == 1
        assert rows[0]["name"] == "b"


# ── E2: transfer 路径 CSV 结构守卫 ────────────────────


class TestCsvStructuralGuard:
    """transfer 建表/导入路径的 CSV 结构守卫：明确报错优于静默丢数据.

    与 Importer._parse_csv（增量导入，按列位置映射，重复表头为既有兼容行为）
    语义不同 —— transfer 路径以表头名建字段/寻址，重复名/空名会静默丢数据，
    多列溢出源于未转义逗号，均属结构性缺陷必须报错。
    """

    def test_analyze_duplicate_headers_raises(self):
        """重复列名 → 报错并列出重复项（修复前后列静默覆盖丢前列）."""
        with pytest.raises(ValueError, match=r"重复列名.*name"):
            transfer.analyze_csv_columns("name,name\na,b\n")

    def test_analyze_blank_header_raises(self):
        """空列名（含 None 列名）→ 报错."""
        with pytest.raises(ValueError, match="空列名"):
            transfer.analyze_csv_columns("name,,email\na,b,c\n")

    def test_analyze_row_overflow_raises(self):
        """行值数超出表头列数（restkey）→ 报错并指明行号."""
        with pytest.raises(ValueError, match=r"第 2 行.*超出表头列数"):
            transfer.analyze_csv_columns("name,email\n张三,a@x.com,多余值\n")

    def test_analyze_short_row_not_error(self):
        """少列短行（尾逗号截断常态）不报错，缺列按空值统计."""
        cols, n = transfer.analyze_csv_columns("name,email,age\n张三,a@x.com,\n李四\n")
        assert n == 2
        by_name = {c["name"]: c for c in cols}
        assert by_name["name"]["field_type"] in ("text", "select")
        assert by_name["age"]["null_ratio"] == 1.0

    def test_import_rows_duplicate_headers_raises(self, test_session):
        """import_rows_from_csv 独立导入路径同样触发表头守卫."""
        engine, session = test_session
        table = _make_table_with_fields(session, engine, [("name", "text")])

        with pytest.raises(ValueError, match="重复列名"):
            transfer.import_rows_from_csv(engine, table, "name,name\na,b\n", db=session)

    def test_import_rows_overflow_raises(self, test_session):
        """import_rows_from_csv 行值溢出同样报错."""
        engine, session = test_session
        table = _make_table_with_fields(session, engine, [("name", "text")])

        with pytest.raises(ValueError, match="超出表头列数"):
            transfer.import_rows_from_csv(engine, table, "name\n张三,多余\n", db=session)

    def test_create_table_from_csv_blank_header_raises(self, test_session):
        """create_table_from_csv 链路经 analyze 天然继承守卫，空列名报错且无残留."""
        engine, session = test_session
        from cndb.plugins.accounts.models import User
        from cndb.plugins.workspaces.models import Workspace

        u = User(username="guard_user")
        u.set_password("pass")
        session.add(u)
        session.flush()
        ws = Workspace(name="GUARD_WS", created_by_id=u.id)
        session.add(ws)
        session.commit()

        with pytest.raises(ValueError, match="空列名"):
            transfer.create_table_from_csv(engine, session, ws.id, "守卫表", "name,,email\na,b,c\n")
        leftovers = session.query(DataTable).filter(DataTable.workspace_id == ws.id).all()
        assert leftovers == [], "守卫报错后不得残留半残表"


# ── E3: XLSX 结构守卫（与 CSV 守卫对齐）────────────────


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    """构建最小 XLSX 字节串（iter_rows values_only 读出与写入一致）."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    assert ws is not None
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class TestXlsxStructuralGuard:
    """XLSX 三条解析路径的结构守卫：空列名/重复列名/行溢出明确报错.

    与 CSV 守卫同款决策：重复列名按名建 dict 时后列覆盖前列、空列名无法
    按名寻址 —— 均属结构性缺陷必须报错；短行保持补 None 空值语义不报错
    （对齐 CSV 尾逗号截断常态）。取代旧的 col_{i} 静默兜底。

    行溢出守卫为防御性兜底（openpyxl 普通 iter_rows 为矩形读取，短行自动
    补 None 到 max_col，实际读出的行不会超宽，见 data_beyond_header 用例：
    真实症状是表头右侧出现空列名）。
    """

    def test_parse_xlsx_bytes_duplicate_headers_raises(self):
        with pytest.raises(ValueError, match=r"重复列名.*name"):
            transfer._parse_xlsx_bytes(_xlsx_bytes([["name", "name"], ["a", "b"]]))

    def test_parse_xlsx_bytes_blank_header_raises(self):
        with pytest.raises(ValueError, match="空列名"):
            transfer._parse_xlsx_bytes(_xlsx_bytes([["name", None, "email"], ["a", "b", "c"]]))

    def test_parse_xlsx_bytes_row_overflow_raises(self):
        """行溢出守卫：值数超出表头列数报错并指明行号（直接单测守卫函数）."""
        with pytest.raises(ValueError, match=r"第 3 行.*超出表头列数"):
            transfer._check_xlsx_row_overflow(("a", "b", "多余值"), 3, 2)

    def test_parse_xlsx_bytes_data_beyond_header_raises_blank(self):
        """数据超出表头宽度：openpyxl 矩形读取把表头补 None → 命中空列名守卫.

        真实世界"行比表头长"的症状即此 —— 表头行右侧实际存在空列名，
        报错优于旧版静默生成 col_{i}/"None" 列。
        """
        with pytest.raises(ValueError, match="空列名"):
            transfer._parse_xlsx_bytes(
                _xlsx_bytes([["name", "email"], ["张三", "a@x.com"], ["李四", "b@x.com", "多余值"]])
            )

    def test_parse_xlsx_bytes_short_row_not_error(self):
        """少列短行不报错，缺列按 None 空值补齐."""
        rows, cols = transfer._parse_xlsx_bytes(_xlsx_bytes([["name", "email"], ["张三", None], ["李四"]]))
        assert cols == ["name", "email"]
        assert rows[0] == {"name": "张三", "email": None}
        assert rows[1] == {"name": "李四", "email": None}

    def test_import_rows_from_xlsx_duplicate_headers_raises(self, test_session):
        """import_rows_from_xlsx 独立导入路径同样触发表头守卫."""
        engine, session = test_session
        table = _make_table_with_fields(session, engine, [("name", "text")])

        with pytest.raises(ValueError, match="重复列名"):
            transfer.import_rows_from_xlsx(engine, table, _xlsx_bytes([["name", "name"], ["a", "b"]]), db=session)

    def test_import_rows_from_xlsx_data_beyond_header_raises(self, test_session):
        """import_rows_from_xlsx 数据超出表头宽度 → 空列名守卫报错."""
        engine, session = test_session
        table = _make_table_with_fields(session, engine, [("name", "text")])

        with pytest.raises(ValueError, match="空列名"):
            transfer.import_rows_from_xlsx(engine, table, _xlsx_bytes([["name"], ["张三", "多余值"]]), db=session)

    def test_importer_parse_xlsx_duplicate_headers_raises(self):
        """Importer._parse_xlsx 增量导入路径同样触发守卫."""
        with pytest.raises(ValueError, match="重复列名"):
            Importer._parse_xlsx(_xlsx_bytes([["name", "name"], ["a", "b"]]))

    def test_importer_parse_xlsx_blank_header_raises(self):
        with pytest.raises(ValueError, match="空列名"):
            Importer._parse_xlsx(_xlsx_bytes([[None, "age"], [1, 2]]))

    def test_task_structural_bad_xlsx_fails_with_message(self, test_session, db):
        """异步任务：重复表头 xlsx → status=failed 且 error_message 含守卫文案."""
        engine, session = test_session
        table = _make_table_with_fields(session, engine, [("name", "text")])

        task = create_import_task(
            session,
            table_id=table.id,
            user_id=1,
            filename="dup.xlsx",
            fmt="xlsx",
            content=_xlsx_bytes([["name", "name"], ["a", "b"]]),
        )
        execute_import_task(session, task.id)
        db.refresh(task)

        assert task.status == "failed"
        assert "重复列名" in (task.error_message or "")

    def test_sync_import_endpoint_bad_xlsx_maps_400(self, client, db, auth_headers):
        """bulk /import 同步端点：结构性缺陷 xlsx → ValueError 映射 400 中文错误."""
        ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "xlsx_guard_ws"})
        assert ws.status_code in (200, 201), f"创建 workspace 失败: {ws.status_code} {ws.text[:100]}"
        wid = ws.json()["id"]
        tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "XlsxGuard"})
        assert tbl.status_code in (200, 201), f"创建 table 失败: {tbl.status_code} {tbl.text[:100]}"
        tid = tbl.json()["id"]

        resp = client.post(
            f"/api/v1/workspaces/{wid}/tables/{tid}/import",
            headers=auth_headers,
            files={
                "file": (
                    "dup.xlsx",
                    io.BytesIO(_xlsx_bytes([["name", "name"], ["a", "b"]])),
                    "application/octet-stream",
                )
            },
        )
        assert resp.status_code == 400, f"期望 400，实际 {resp.status_code} {resp.text[:200]}"
        assert "重复列名" in resp.json()["detail"]


# ── F: 端到端 roundtrip ─────────────────────────────


class TestImportExportRoundtripE2E:
    """导入 → 导出 → 再导入的一致性验证（走真实 API 路由）."""

    def _create_ws_table_fields(self, client, auth_headers, *, ws_name, tbl_name, fields):
        """创建 workspace + table + fields（全部走 API，确保物理建列），返回 (wid, tid)."""
        ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": ws_name})
        assert ws.status_code in (200, 201), f"创建 workspace 失败: {ws.status_code} {ws.text[:100]}"
        wid = ws.json()["id"]

        tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": tbl_name})
        assert tbl.status_code in (200, 201), f"创建 table 失败: {tbl.status_code} {tbl.text[:100]}"
        tid = tbl.json()["id"]

        # 字段必须走 fields API —— 直接 db.add(DataField) 不会触发物理建列
        for f_def in fields:
            r = client.post(f"/api/v1/workspaces/{wid}/tables/{tid}/fields", headers=auth_headers, json=f_def)
            assert r.status_code in (200, 201), f"创建字段 {f_def['name']} 失败: {r.status_code} {r.text[:200]}"
        return wid, tid

    def _import_via_two_phase(self, client, auth_headers, wid, tid, csv_content):
        """两阶段导入：analyze → 等 pending_confirm → confirm → 等 done."""
        resp = client.post(
            f"/api/v1/workspaces/{wid}/tables/{tid}/import/analyze",
            headers=auth_headers,
            files={"file": ("data.csv", io.BytesIO(csv_content.encode("utf-8")), "text/csv")},
        )
        assert resp.status_code == 200, f"analyze 失败: {resp.status_code} {resp.text[:200]}"
        task_id = resp.json()["task_id"]

        async_url = f"/api/v1/workspaces/{wid}/tables/{tid}/import/async/{task_id}"
        settled = wait_import_settled(client, async_url, auth_headers)
        assert settled["status"] == "pending_confirm", f"analyze 后应处于 pending_confirm: {settled['status']}"

        confirm = client.post(f"/api/v1/workspaces/{wid}/tables/{tid}/import/{task_id}/confirm", headers=auth_headers)
        assert confirm.status_code == 200, f"confirm 失败: {confirm.status_code} {confirm.text[:200]}"

        final = wait_import_settled(client, async_url, auth_headers)
        assert final["status"] == "done", f"导入未完成: {final.get('status')} err={final.get('error_message')}"

    def _export_rows(self, client, auth_headers, wid, tid, fmt):
        """导出并返回解析结果（csv → 字符串；json → list[dict]）."""
        resp = client.get(
            f"/api/v1/workspaces/{wid}/tables/{tid}/export?format={fmt}",
            headers=auth_headers,
        )
        assert resp.status_code == 200, f"导出失败: {resp.status_code} {resp.text[:200]}"
        return resp.text if fmt == "csv" else resp.json()

    def test_roundtrip_import_export_reimport(self, client, auth_headers):
        """导入(2好1坏) → 导出 → 再导入新表 → 行数与值一致."""
        fields = [
            {"name": "name", "field_type": "text", "required": True, "order": 0},
            {"name": "email", "field_type": "email", "order": 1},
        ]
        wid, tid = self._create_ws_table_fields(
            client, auth_headers, ws_name="roundtrip_ws", tbl_name="RoundtripA", fields=fields
        )

        csv_content = "name,email\n张三,zhang@x.com\n李四,li@x.com\n王五,bad-email\n"
        self._import_via_two_phase(client, auth_headers, wid, tid, csv_content)

        # 导出表 A 的 JSON 行数据
        rows_a = self._export_rows(client, auth_headers, wid, tid, "json")
        assert len(rows_a) == 2
        values_a = sorted((r["name"], r["email"]) for r in rows_a)

        # 再导入到结构相同的表 B（独立工作区，走同步 bulk 导入端点）
        wid2, tid2 = self._create_ws_table_fields(
            client, auth_headers, ws_name="roundtrip_ws_b", tbl_name="RoundtripB", fields=fields
        )
        csv_export = self._export_rows(client, auth_headers, wid, tid, "csv")
        re_import = client.post(
            f"/api/v1/workspaces/{wid2}/tables/{tid2}/import",
            headers=auth_headers,
            files={"file": ("export.csv", io.BytesIO(csv_export.encode("utf-8")), "text/csv")},
        )
        assert re_import.status_code == 200, f"再导入失败: {re_import.status_code} {re_import.text[:200]}"
        assert re_import.json()["imported"] == 2

        rows_b = self._export_rows(client, auth_headers, wid2, tid2, "json")
        assert len(rows_b) == 2
        assert sorted((r["name"], r["email"]) for r in rows_b) == values_a

    def test_roundtrip_preserves_quoted_values(self, client, auth_headers):
        """含逗号/双引号的值经 导入→导出 往返不变形."""
        wid, tid = self._create_ws_table_fields(
            client,
            auth_headers,
            ws_name="quoted_ws",
            tbl_name="QuotedValues",
            fields=[{"name": "name", "field_type": "text", "order": 0}],
        )

        # CSV 引号转义：第一行含逗号，第二行含双引号
        csv_content = 'name\n"张三, Jr."\n"含""引号"""'
        self._import_via_two_phase(client, auth_headers, wid, tid, csv_content)

        csv_export = self._export_rows(client, auth_headers, wid, tid, "csv")
        parsed = list(csv.DictReader(io.StringIO(csv_export)))
        values = sorted(r["name"] for r in parsed)
        assert values == sorted(["张三, Jr.", '含"引号"'])


# ── 附带: _json_safe bytes 有损解码 ─────────────────


def test_json_safe_lossy_bytes():
    """非 UTF-8 bytes → 替换符 U+FFFD 有损解码，不抛异常."""
    assert _json_safe(b"\xff") == "\ufffd"
