"""用户批量导入 / 导出 API 测试.

覆盖：
- 导出：CSV（BOM + 表头 + 无密码列）、XLSX、筛选口径、非法参数、权限
- 导入：CSV / XLSX、自动生成密码、冲突跳过、校验失败、dry-run、文件内重复、权限
- 审计：导入 / 导出各写一条操作日志
"""

from __future__ import annotations

import io

import pytest

from cndb.plugins.accounts.models import User, UserAuditLog
from cndb.plugins.accounts.services.user_bulk import (
    MAX_IMPORT_BYTES,
    USER_EXPORT_FIELDS,
    import_user_rows,
    parse_user_file,
)

USERS_URL = "/api/v1/accounts/users"
REGISTER_URL = "/api/v1/accounts/auth/register"
LOGIN_URL = "/api/v1/accounts/auth/login"
ADMIN_REGISTER_URL = "/api/v1/accounts/auth/admin-register"


def _auth(token: str) -> dict[str, str]:
    """构造 Bearer 请求头."""
    return {"Authorization": f"Bearer {token}"}


def _register_and_login(client, username: str, password: str = "pw1234") -> str:
    """注册并登录，返回 JWT."""
    client.post(REGISTER_URL, json={"username": username, "password": password})
    r = client.post(LOGIN_URL, json={"login": username, "password": password})
    return r.json()["access_token"]


def _create_user_by_admin(client, token: str, username: str, role: str = "user", email: str | None = None) -> dict:
    """管理员创建用户并返回响应体."""
    payload: dict = {"username": username, "password": "pw1234", "role": role}
    if email:
        payload["email"] = email
    r = client.post(ADMIN_REGISTER_URL, json=payload, headers=_auth(token))
    assert r.status_code == 201, r.text
    return r.json()


def _upload(client, token: str, content: bytes, filename: str = "users.csv", dry_run: bool = False):
    """上传导入文件."""
    return client.post(
        f"{USERS_URL}/import",
        files={"file": (filename, io.BytesIO(content))},
        data={"dry_run": "true" if dry_run else "false"},
        headers=_auth(token),
    )


def _csv(*lines: str) -> bytes:
    """拼装 CSV 字节内容."""
    return ("\n".join(lines) + "\n").encode("utf-8")


@pytest.fixture
def admin_token(client, db):
    """创建超级管理员并登录，返回 JWT."""
    u = User(username="bulk_admin", email="bulk_admin@example.com")
    u.set_password("pw1234")
    u.is_superuser = True
    db.add(u)
    db.commit()
    return _register_and_login(client, "bulk_admin", "pw1234")


@pytest.fixture
def plain_token(client):
    """普通用户 JWT（用于权限断言）."""
    return _register_and_login(client, "bulk_plain")


class TestUserExport:
    """GET /users/export 导出."""

    def test_export_csv_has_bom_header_and_no_password(self, client, admin_token):
        _create_user_by_admin(client, admin_token, "exp_a", role="audit_admin", email="exp_a@example.com")
        r = client.get(f"{USERS_URL}/export", params={"format": "csv"}, headers=_auth(admin_token))
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("text/csv")
        assert "attachment" in r.headers["content-disposition"]

        text = r.content.decode("utf-8-sig")
        header = text.splitlines()[0]
        assert header == "username,email,phone,nickname,role,is_active,is_superuser,created_at"
        assert "exp_a@example.com" in text
        assert "password" not in header
        assert "pw1234" not in text

    def test_export_respects_filters(self, client, admin_token):
        _create_user_by_admin(client, admin_token, "exp_filter_1", role="audit_admin")
        _create_user_by_admin(client, admin_token, "exp_filter_2", role="user")
        r = client.get(
            f"{USERS_URL}/export",
            params={"format": "csv", "role": "audit_admin"},
            headers=_auth(admin_token),
        )
        text = r.content.decode("utf-8-sig")
        assert "exp_filter_1" in text
        assert "exp_filter_2" not in text

    def test_export_keyword_filter(self, client, admin_token):
        _create_user_by_admin(client, admin_token, "kw_alpha", role="user")
        _create_user_by_admin(client, admin_token, "kw_beta", role="user")
        r = client.get(
            f"{USERS_URL}/export",
            params={"format": "csv", "keyword": "kw_alpha"},
            headers=_auth(admin_token),
        )
        text = r.content.decode("utf-8-sig")
        assert "kw_alpha" in text
        assert "kw_beta" not in text

    def test_export_xlsx(self, client, admin_token):
        _create_user_by_admin(client, admin_token, "exp_xlsx", role="user")
        r = client.get(f"{USERS_URL}/export", params={"format": "xlsx"}, headers=_auth(admin_token))
        assert r.status_code == 200, r.text
        assert r.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert len(r.content) > 0

    def test_export_empty_still_has_header(self, client, admin_token, db):
        # 只保留 admin 本身时按不存在的 keyword 过滤，验证空结果也有表头
        r = client.get(
            f"{USERS_URL}/export",
            params={"format": "csv", "keyword": "no_such_user_xyz"},
            headers=_auth(admin_token),
        )
        assert r.status_code == 200
        assert r.content.decode("utf-8-sig").startswith("username,email")

    def test_export_invalid_format_422(self, client, admin_token):
        r = client.get(f"{USERS_URL}/export", params={"format": "pdf"}, headers=_auth(admin_token))
        assert r.status_code == 422

    def test_export_invalid_role_400(self, client, admin_token):
        r = client.get(f"{USERS_URL}/export", params={"role": "root"}, headers=_auth(admin_token))
        assert r.status_code == 400

    def test_export_requires_superuser(self, client, plain_token):
        r = client.get(f"{USERS_URL}/export", headers=_auth(plain_token))
        assert r.status_code == 403

    def test_export_writes_audit_log(self, client, admin_token, db):
        client.get(f"{USERS_URL}/export", params={"format": "csv"}, headers=_auth(admin_token))
        log = db.query(UserAuditLog).filter(UserAuditLog.action == "export").first()
        assert log is not None
        assert log.detail["format"] == "csv"


class TestUserImport:
    """POST /users/import 导入."""

    def test_import_csv_creates_users_with_generated_password(self, client, admin_token, db):
        content = _csv(
            "用户名,邮箱,角色,密码",
            "imp_a,imp_a@example.com,audit_admin,",
            "imp_b,imp_b@example.com,user,secret9",
        )
        r = _upload(client, admin_token, content)
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["total"] == 2 and body["created"] == 2 and body["failed"] == 0
        assert body["dry_run"] is False

        by_name = {item["username"]: item for item in body["results"]}
        assert by_name["imp_a"]["generated_password"]
        assert by_name["imp_b"]["generated_password"] is None
        assert by_name["imp_a"]["status"] == "ok"

        db.expire_all()
        user = db.query(User).filter(User.username == "imp_a").first()
        assert user is not None
        assert user.role == "audit_admin"
        # 自动生成的密码确实可用于登录
        login = client.post(LOGIN_URL, json={"login": "imp_a", "password": by_name["imp_a"]["generated_password"]})
        assert login.status_code == 200

    def test_import_skips_existing_username_and_email(self, client, admin_token):
        _create_user_by_admin(client, admin_token, "imp_exists", email="imp_exists@example.com")
        content = _csv(
            "username,email",
            "imp_exists,new@example.com",
            "imp_email_conflict,imp_exists@example.com",
            "imp_new,imp_new@example.com",
        )
        r = _upload(client, admin_token, content)
        body = r.json()
        assert body["created"] == 1 and body["skipped"] == 2 and body["failed"] == 0
        statuses = {item["username"]: item["status"] for item in body["results"]}
        assert statuses["imp_exists"] == "skipped"
        assert statuses["imp_email_conflict"] == "skipped"

    def test_import_failed_rows_do_not_block_others(self, client, admin_token):
        content = _csv(
            "username,role",
            ",user",
            "imp_badrole,root",
            "imp_ok,user",
        )
        r = _upload(client, admin_token, content)
        body = r.json()
        assert body["created"] == 1 and body["failed"] == 2
        messages = " / ".join(item["message"] for item in body["results"] if item["status"] == "failed")
        assert "缺少 username" in messages
        assert "无效角色" in messages

    def test_import_file_duplicate_username_skipped(self, client, admin_token):
        content = _csv("username", "imp_dup", "imp_dup")
        body = _upload(client, admin_token, content).json()
        assert body["created"] == 1 and body["skipped"] == 1
        assert body["results"][1]["message"] == "文件内 username 重复"

    def test_import_dry_run_does_not_write(self, client, admin_token, db):
        content = _csv("username,email", "imp_dry@example.com,imp_dry@example.com")
        body = _upload(client, admin_token, content, dry_run=True).json()
        assert body["dry_run"] is True
        assert body["created"] == 1
        assert body["results"][0]["user_id"] is None
        db.expire_all()
        assert db.query(User).filter(User.username == "imp_dry@example.com").first() is None

    def test_import_xlsx(self, client, admin_token):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["用户名", "邮箱", "角色"])
        ws.append(["imp_xlsx", "imp_xlsx@example.com", "security_admin"])
        buf = io.BytesIO()
        wb.save(buf)

        r = _upload(client, admin_token, buf.getvalue(), filename="users.xlsx")
        assert r.status_code == 200, r.text
        body = r.json()
        assert body["created"] == 1
        assert body["results"][0]["username"] == "imp_xlsx"

    def test_import_unsupported_format_400(self, client, admin_token):
        r = _upload(client, admin_token, b"x", filename="users.txt")
        assert r.status_code == 400
        assert "不支持的文件格式" in r.json()["detail"]

    def test_import_missing_username_column_400(self, client, admin_token):
        r = _upload(client, admin_token, _csv("email,nickname", "a@example.com,甲"))
        assert r.status_code == 400
        assert "缺少 username 列" in r.json()["detail"]

    def test_import_empty_file_400(self, client, admin_token):
        r = _upload(client, admin_token, b"")
        assert r.status_code == 400
        assert "上传文件为空" in r.json()["detail"]

    def test_import_oversize_file_400(self, client, admin_token):
        r = _upload(client, admin_token, b"x" * (MAX_IMPORT_BYTES + 1))
        assert r.status_code == 400
        assert "文件过大" in r.json()["detail"]

    def test_import_inactive_state_parsed(self, client, admin_token, db):
        content = _csv("username,is_active", "imp_off,禁用", "imp_on,启用")
        body = _upload(client, admin_token, content).json()
        assert body["created"] == 2
        db.expire_all()
        assert db.query(User).filter(User.username == "imp_off").first().is_active is False
        assert db.query(User).filter(User.username == "imp_on").first().is_active is True

    def test_import_invalid_email_failed(self, client, admin_token):
        content = _csv("username,email", "imp_bad_mail,not-an-email")
        body = _upload(client, admin_token, content).json()
        assert body["failed"] == 1
        assert "邮箱格式不正确" in body["results"][0]["message"]

    def test_import_writes_audit_log(self, client, admin_token, db):
        _upload(client, admin_token, _csv("username", "imp_audit"))
        log = db.query(UserAuditLog).filter(UserAuditLog.action == "import").first()
        assert log is not None
        assert log.detail["created"] == 1
        assert "imp_audit" in log.detail["usernames"]

    def test_import_requires_superuser(self, client, plain_token):
        r = _upload(client, plain_token, _csv("username", "imp_denied"))
        assert r.status_code == 403


class TestUserBulkEdgeCases:
    """解析 / 校验 / 导出的边界分支."""

    def test_import_ignores_unknown_columns(self, client, admin_token, db):
        content = _csv("username,备注", "imp_extra,随意填写")
        body = _upload(client, admin_token, content).json()
        assert body["created"] == 1
        db.expire_all()
        assert db.query(User).filter(User.username == "imp_extra").first() is not None

    def test_parse_empty_csv_returns_no_rows(self):
        assert parse_user_file(b"", "users.csv") == []

    def test_parse_empty_xlsx_returns_no_rows(self):
        from openpyxl import Workbook

        buf = io.BytesIO()
        Workbook().save(buf)
        assert parse_user_file(buf.getvalue(), "users.xlsx") == []

    def test_parse_xlsx_skips_blank_rows(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.append(["用户名"])
        ws.append([None])
        ws.append(["imp_xlsx_blank"])
        buf = io.BytesIO()
        wb.save(buf)

        rows = parse_user_file(buf.getvalue(), "users.xlsx")
        assert [row["username"] for row in rows] == ["imp_xlsx_blank"]

    def test_import_too_many_rows_rejected(self, client, admin_token):
        content = _csv("username", *[f"imp_many_{index}" for index in range(1001)])
        r = _upload(client, admin_token, content)
        assert r.status_code == 400
        assert "单次最多导入" in r.json()["detail"]

    def test_export_filters_by_active_state(self, client, admin_token, db):
        user = User(username="exp_inactive", email="exp_inactive@example.com", is_active=False)
        user.set_password("pw1234")
        db.add(user)
        db.commit()

        r = client.get(
            f"{USERS_URL}/export",
            params={"format": "csv", "is_active": "false"},
            headers=_auth(admin_token),
        )
        text = r.content.decode("utf-8-sig")
        assert "exp_inactive" in text
        assert "bulk_admin," not in text

    def test_export_empty_xlsx_still_has_header(self, client, admin_token):
        from openpyxl import load_workbook

        r = client.get(
            f"{USERS_URL}/export",
            params={"format": "xlsx", "keyword": "no_such_user_xyz"},
            headers=_auth(admin_token),
        )
        assert r.status_code == 200, r.text
        wb = load_workbook(io.BytesIO(r.content), read_only=True)
        first_row = next(wb.active.iter_rows(values_only=True))
        assert first_row == USER_EXPORT_FIELDS

    def test_import_row_validation_errors(self, client, admin_token):
        content = _csv(
            "username,password,phone,is_active",
            "a,,,",
            f"{'x' * 151},,,",
            "imp_short_pw,123,,",
            "imp_long_phone,,123456789012345678901234567890123,",
            "imp_bad_state,,,maybe",
        )
        body = _upload(client, admin_token, content).json()
        assert body["failed"] == 5
        messages = " / ".join(item["message"] for item in body["results"])
        for fragment in ("至少 2 个字符", "最多 150 个字符", "6-128", "phone 最多 32", "无法识别的启用状态"):
            assert fragment in messages

    def test_import_duplicate_email_in_file_skipped(self, client, admin_token):
        content = _csv("username,email", "imp_mail_a,dup@example.com", "imp_mail_b,dup@example.com")
        body = _upload(client, admin_token, content).json()
        assert body["created"] == 1 and body["skipped"] == 1
        assert body["results"][1]["message"] == "文件内邮箱重复"

    def test_import_write_failure_is_reported(self, client, admin_token, db):
        from sqlalchemy.exc import OperationalError

        class _FailingSession:
            """commit 必失败，其余方法委托给真实 Session."""

            def __init__(self, session):
                self._session = session

            def __getattr__(self, name):
                return getattr(self._session, name)

            def commit(self):
                raise OperationalError("INSERT", {}, Exception("disk full"))

        actor = db.query(User).filter(User.username == "bulk_admin").one()
        response = import_user_rows(_FailingSession(db), actor=actor, rows=[{"username": "imp_write_fail"}])
        assert response.failed == 1
        assert "写入失败" in response.results[0].message
        db.rollback()
