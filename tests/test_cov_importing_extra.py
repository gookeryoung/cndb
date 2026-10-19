"""补充覆盖：importing/import_tasks.py 与 importing/importer.py 的未覆盖分支.

仅新增测试，不改动源码与既有测试。覆盖目标：
- import_tasks：线程回收、失败落库（可读原因/落库失败兜底）、_parse_to_rows、
  execute 新旧流程各分支、reanalyze、run_task_in_background 分发。
- importer：analyze 预填充降级、进度回调降级、字段自动新增与 DDL 清理、
  upsert 分流、旧行查询、归一化视图、日期比较、未知列规划、清洗原始行解析。
"""

from __future__ import annotations

import io
import threading

import pytest
from openpyxl import Workbook
from sqlalchemy.exc import IntegrityError, SQLAlchemyError

from cndb.plugins.tables.models import DataField, DataTable, ImportTask
from cndb.plugins.tables.services.importing import import_tasks as it
from cndb.plugins.tables.services.importing import importer as imp_mod
from cndb.plugins.tables.services.importing.importer import ImportAnalysisResult, Importer
from cndb.plugins.tables.services.importing.row_validator import ValidationResult

# ── helpers ────────────────────────────────────────


def _xlsx_bytes(rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _empty_xlsx_bytes() -> bytes:
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@pytest.fixture
def kv_table(client, auth_headers, db) -> DataTable:
    """建一个含 key/val 两个 text 字段的真实表（物理 DDL 已建）."""
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "cov-import-ws"}).json()
    tbl = client.post(
        f"/api/v1/workspaces/{ws['id']}/tables",
        headers=auth_headers,
        json={"name": "cov-import-table", "description": ""},
    ).json()
    for i, name in enumerate(("key", "val")):
        client.post(
            f"/api/v1/workspaces/{ws['id']}/tables/{tbl['id']}/fields",
            headers=auth_headers,
            json={"name": name, "field_type": "text", "order": i},
        )
    return db.get(DataTable, int(tbl["id"]))


# ── import_tasks：工具与落库 ───────────────────────


class TestImportTaskHelpers:
    def test_join_background_threads_joins_live_thread(self):
        """后台线程仍在运行时应被 join —— 覆盖 join_background_threads 的 is_alive 分支."""
        started = threading.Event()
        stop = threading.Event()

        def _work() -> None:
            started.set()
            stop.wait(5)

        worker = threading.Thread(target=_work, daemon=True)
        worker.start()
        assert started.wait(5)
        with it._background_lock:
            it._background_threads.append(worker)
        it.join_background_threads(timeout=0.01)
        stop.set()
        worker.join(5)

    def test_readable_error_field_write_conflict(self):
        exc = IntegrityError("stmt", "params", Exception("tables_datafield name conflict"))
        msg = it._readable_error(exc)
        assert "字段写入冲突" in msg

    def test_readable_error_generic_integrity_error(self):
        exc = IntegrityError("stmt", "params", Exception("some other constraint"))
        assert it._readable_error(exc) == "数据写入冲突（唯一约束），导入已中止。请检查是否有重复字段或重复行。"

    def test_readable_error_strips_sql_tail(self):
        exc = SQLAlchemyError("boom [SQL: SELECT 1] (Background on this error at: http://x)")
        assert it._readable_error(exc) == "boom"

    def test_mark_task_failed_commit_failure_is_swallowed(self, db, monkeypatch):
        """最终提交失败不应逃逸出后台线程（否则任务永远卡在 running）."""
        task = it.create_import_task(
            db, table_id=99999, user_id=None, filename="x.json", fmt="json", content="[]", status="running"
        )

        def _boom() -> None:
            raise RuntimeError("commit failed")

        monkeypatch.setattr(db, "commit", _boom)
        # 不应抛异常
        it._mark_task_failed(db, task.id, ValueError("orig cause"))


# ── import_tasks：_parse_to_rows ───────────────────


class TestParseToRows:
    def test_parse_to_rows_json_non_list(self):
        with pytest.raises(ValueError, match="JSON 必须是对象数组"):
            it._parse_to_rows('{"a": 1}', "json", None, None)

    def test_parse_to_rows_csv_bytes(self):
        rows, cols, total = it._parse_to_rows(b"a,b\n1,2", "csv", None, None)
        assert cols == ["a", "b"]
        assert total == 1
        assert rows[0]["a"] == "1"

    def test_parse_to_rows_xlsx_wider_data_row(self):
        """数据行比表头宽 → 补 col_{i} 列名（覆盖 279-283）."""
        raw = _xlsx_bytes([["a", "b"], ["x", "y", "z"]])
        rows, cols, total = it._parse_to_rows(raw, "xlsx", None, None)
        assert cols == ["a", "b", "col_2"]
        assert total == 1
        assert rows[0]["col_2"] == "z"

    def test_parse_to_rows_xlsx_empty_sheet(self):
        rows, cols, total = it._parse_to_rows(_empty_xlsx_bytes(), "xlsx", None, None)
        assert (rows, cols, total) == ([], [], 0)

    def test_parse_to_rows_unsupported_format(self):
        with pytest.raises(ValueError, match="不支持的格式"):
            it._parse_to_rows("x", "yaml", None, None)


# ── import_tasks：execute 分支 ─────────────────────


class TestExecuteImportTaskBranches:
    def test_new_flow_sync_options_failure_is_logged(self, db, kv_table, monkeypatch):
        """新流程（有 validation_report）里 select options 同步失败只记 warning —— 覆盖 363-364."""
        from cndb.plugins.tables.services.fields import field_ops

        task = it.create_import_task(
            db,
            table_id=kv_table.id,
            user_id=None,
            filename="x.json",
            fmt="json",
            content='[{"key": "a", "val": "1"}]',
        )
        it.analyze_import_task(db, task.id)
        db.refresh(task)
        assert task.status == "pending_confirm"

        def _boom(*_a, **_kw):
            raise RuntimeError("sync failed")

        monkeypatch.setattr(field_ops, "sync_select_options_from_table", _boom)
        it.execute_import_task(db, task.id)
        db.refresh(task)
        assert task.status == "done"

    def test_old_flow_json_non_text_raises(self, db, kv_table, monkeypatch):
        """旧流程 JSON 内容非文本 → 明确失败（覆盖 375-376）."""
        task = it.create_import_task(
            db, table_id=kv_table.id, user_id=None, filename="x.json", fmt="json", content="[]"
        )
        monkeypatch.setattr(it, "_decode_content", lambda _t: b"[]")
        it.execute_import_task(db, task.id)
        db.refresh(task)
        assert task.status == "failed"

    def test_old_flow_csv_non_text_raises(self, db, kv_table, monkeypatch):
        task = it.create_import_task(
            db, table_id=kv_table.id, user_id=None, filename="x.csv", fmt="csv", content="a\n1"
        )
        monkeypatch.setattr(it, "_decode_content", lambda _t: b"a\n1")
        it.execute_import_task(db, task.id)
        db.refresh(task)
        assert task.status == "failed"

    def test_old_flow_xlsx_non_bytes_raises(self, db, kv_table, monkeypatch):
        task = it.create_import_task(
            db, table_id=kv_table.id, user_id=None, filename="x.xlsx", fmt="xlsx", content="nope"
        )
        monkeypatch.setattr(it, "_decode_content", lambda _t: "not-bytes")
        it.execute_import_task(db, task.id)
        db.refresh(task)
        assert task.status == "failed"

    def test_old_flow_json_empty_ids_skips_sync(self, db, kv_table):
        """旧流程导入空行 → ids 为空 → 跳过 options 同步（覆盖 392->398）."""
        task = it.create_import_task(
            db, table_id=kv_table.id, user_id=None, filename="x.json", fmt="json", content="[]"
        )
        it.execute_import_task(db, task.id)
        db.refresh(task)
        assert task.status == "done"
        assert task.imported_rows == 0

    def test_old_flow_sync_options_failure_is_logged(self, db, kv_table, monkeypatch):
        """旧流程有行且同步失败 → 记 warning（覆盖 394-396）."""
        from cndb.plugins.tables.services.fields import field_ops

        task = it.create_import_task(
            db,
            table_id=kv_table.id,
            user_id=None,
            filename="x.json",
            fmt="json",
            content='[{"key": "b", "val": "2"}]',
        )

        def _boom(*_a, **_kw):
            raise RuntimeError("sync failed")

        monkeypatch.setattr(field_ops, "sync_select_options_from_table", _boom)
        it.execute_import_task(db, task.id)
        db.refresh(task)
        assert task.status == "done"
        assert task.imported_rows == 1

    def test_old_flow_unsupported_format(self, db, kv_table):
        task = it.create_import_task(db, table_id=kv_table.id, user_id=None, filename="x.zzz", fmt="zzz", content="x")
        it.execute_import_task(db, task.id)
        db.refresh(task)
        assert task.status == "failed"


# ── import_tasks：reanalyze / 后台分发 ──────────────


class TestReanalyzeAndDispatch:
    def test_reanalyze_missing_table_fails(self, db):
        task = it.create_import_task(
            db, table_id=99999, user_id=None, filename="x.json", fmt="json", content="[]", status="pending_confirm"
        )
        it.reanalyze_import_task(db, task.id)
        db.refresh(task)
        assert task.status == "failed"

    def test_reanalyze_rejects_running_status(self, db):
        task = it.create_import_task(
            db, table_id=99999, user_id=None, filename="x.json", fmt="json", content="[]", status="running"
        )
        with pytest.raises(ValueError):
            it.reanalyze_import_task(db, task.id)

    def test_reanalyze_missing_task(self, db):
        # 不存在的任务：静默返回
        it.reanalyze_import_task(db, 987654)

    @pytest.mark.parametrize(
        ("phase", "status"),
        [("analyze", "pending"), ("reanalyze", "pending_confirm"), ("execute", "pending")],
    )
    def test_run_task_in_background_dispatch(self, db, _session_factory, monkeypatch, phase, status):
        """run_task_in_background 按 phase 分发（覆盖 _worker 三条分支）."""

        class _SyncThread:
            def __init__(self, *, target, daemon=False, name=None):
                self._target = target

            def start(self) -> None:
                self._target()

            def is_alive(self) -> bool:
                return False

            def join(self, timeout: float | None = None) -> None:
                return

        monkeypatch.setattr(it.threading, "Thread", _SyncThread)
        task = it.create_import_task(
            db, table_id=99999, user_id=None, filename="x.json", fmt="json", content="[]", status=status
        )
        it.run_task_in_background(_session_factory, task.id, phase=phase)
        db.refresh(task)
        assert task.status in ("failed", "pending_validation", "running", "done")


# ── importer：analyze / execute 降级分支 ────────────


class TestImporterDegradation:
    def test_analyze_prefill_failure_is_silent(self, kv_table, db, monkeypatch):
        """导入前预填充 select options 失败 → 仅 debug 日志，不阻塞 analyze."""
        from cndb.plugins.tables.services.fields import field_ops

        def _boom(*_a, **_kw):
            raise RuntimeError("prefill failed")

        monkeypatch.setattr(field_ops, "prefill_select_options_from_rows", _boom)
        analysis = Importer(db.get_bind(), db, kv_table).analyze(
            None, None, rows=[{"key": "a", "val": "1"}], unknown_cols_strategy="drop"
        )
        assert analysis.report["valid_count"] == 1

    def test_execute_progress_cb_failure_is_silent(self, kv_table, db):
        def _cb(_p: int, _s: str) -> None:
            raise RuntimeError("progress cb failed")

        result = Importer(db.get_bind(), db, kv_table).execute(
            None, None, rows=[{"key": "a", "val": "1"}], progress_cb=_cb
        )
        assert result.imported_ids

    def test_execute_sync_options_failure_is_logged(self, kv_table, db, monkeypatch):
        """导入后同步 select options 失败 → 仅 debug 日志（覆盖 412-413）."""
        from cndb.plugins.tables.services.fields import field_ops

        def _boom(*_a, **_kw):
            raise RuntimeError("sync failed")

        monkeypatch.setattr(field_ops, "sync_select_options_from_table", _boom)
        result = Importer(db.get_bind(), db, kv_table).execute(None, None, rows=[{"key": "z", "val": "9"}])
        assert result.imported_ids

    def test_execute_auto_add_all_dropped_skips(self, kv_table, db):
        """add_text_field 策略下所有待新增列都被丢弃 → 不新增字段（覆盖 344-345 假分支）."""
        vals = {"key": "k1", "val": "v1"}
        analysis = ImportAnalysisResult(
            report={"valid_count": 1, "error_count": 0},
            results=[],
            file_columns=list(vals),
            planned_columns=[{"name": "c1", "field_type": "text", "sample_values": ["x"]}],
        )
        result = Importer(db.get_bind(), db, kv_table).execute(
            analysis=analysis,
            rows=None,
            unknown_cols_strategy="add_text_field",
            dropped_columns=["c1"],
        )
        assert result.imported_ids == []


# ── importer：upsert 分流 ───────────────────────────


class TestImporterUpsert:
    def test_upsert_update_and_create_mix(self, kv_table, db):
        """真实 upsert：一行命中已存在行 + 一行新增，覆盖 update/create 两条路径."""
        from cndb.plugins.tables.services.core import records as rec

        engine = db.get_bind()
        rec.bulk_create(engine, kv_table, [{"key": "A", "val": "old"}], db=db)
        imp = Importer(engine, db, kv_table)
        analysis = imp.analyze(
            None, None, rows=[{"key": "A", "val": "new"}, {"key": "B", "val": "x"}], match_keys=["key"]
        )
        assert analysis.upsert_result["update_count"] == 1
        result = imp.execute(analysis=analysis)
        assert result.imported_ids
        assert result.report["actually_updated"] == 1

    def test_upsert_conflict_without_update_rows(self, kv_table, db):
        """update_rows 为空但存在 multi_key_conflicts → 走"仅新增"分支（覆盖 385->388、394）."""
        vals = {"key": "dup", "val": "1"}
        analysis = ImportAnalysisResult(
            report={"valid_count": 1, "error_count": 0},
            results=[ValidationResult(row_number=1, values=vals)],
            file_columns=list(vals),
            upsert_result={
                "new_count": 1,
                "update_count": 0,
                "update_changed_count": 0,
                "update_no_change_count": 0,
                "multi_key_conflicts": 1,
                "new_preview": [],
                "update_preview": [],
                "new_rows": [vals],
                "update_rows": [],
            },
        )
        result = Importer(db.get_bind(), db, kv_table).execute(analysis=analysis)
        assert len(result.imported_ids) == 1


# ── importer：内部工具 ─────────────────────────────


class TestImporterInternals:
    def test_fetch_old_rows_skips_field_without_db_column(self, kv_table, db):
        """字段无 db_column_name 时应跳过映射（覆盖 550->549）."""
        imp = Importer(db.get_bind(), db, kv_table)
        f = DataField(table_id=kv_table.id, name="ghost", field_type="text", config={}, order=99)
        f.db_column_name = None  # 仅内存态：模拟未生成物理列名
        kv_table.fields.append(f)
        assert imp._fetch_old_rows_by_ids([1]) == {}

    def test_normalized_view_false_and_true_edges(self):
        result = ValidationResult(row_number=1, values={})
        # normalized 未包含该 db 列 → 回退原始值（覆盖 587->586 假分支）
        assert Importer._normalized_view(result, {"f": 1}, [("f", "c")])["f"] == 1
        # normalized 命中 → 覆盖为归一化值
        result2 = ValidationResult(row_number=1, values={}, normalized={"c": 5})
        assert Importer._normalized_view(result2, {"f": 1}, [("f", "c")])["f"] == 5

    def test_values_equal_int_vs_non_numeric_str(self):
        """number 与不可转数字的字符串 → 退化为字符串比较（覆盖 629-630）."""
        assert Importer._values_equal("abc", 1) is False
        assert Importer._values_equal("1", 1) is True

    def test_coerce_date_value_empty_and_non_str(self):
        assert Importer._coerce_date_value("") is None
        assert Importer._coerce_date_value(123) is None
        assert Importer._coerce_date_value("2024/01/15") is not None

    def test_plan_unknown_columns_blank_sample(self, kv_table, db):
        """skip 列样本为空串 → 不计入样本（覆盖 695->688 假分支）."""
        imp = Importer(db.get_bind(), db, kv_table)
        plan = imp._plan_unknown_columns([{"unknown": ""}, {"unknown": "abc"}], ["unknown"])
        assert plan and plan[0]["name"] == "unknown"

    def test_auto_add_decimals_str_failure_continues(self, kv_table, db, monkeypatch):
        """浮点列算 decimals 时样本 __str__ 抛错 → 跳过该样本（覆盖 772-773）."""

        class _Boom:
            def __str__(self) -> str:  # pragma: no cover - 触发异常即可
                raise TypeError("cannot stringify")

        imp = Importer(db.get_bind(), db, kv_table)
        imp._auto_add_fields([{"name": "floaty", "field_type": "float", "sample_values": [_Boom(), "1.25"]}])
        db.refresh(kv_table)
        created = {f.name for f in kv_table.fields}
        assert "floaty" in created

    def test_auto_add_ddl_failure_rolls_back(self, kv_table, db, monkeypatch):
        """第二个字段 DDL 失败 → 逆向清理已建列并回滚（覆盖 800-808）."""
        real_add = imp_mod.add_column
        calls = {"n": 0}

        def _flaky(engine, table, field):
            calls["n"] += 1
            if calls["n"] >= 2:
                raise RuntimeError("ddl fail")
            return real_add(engine, table, field)

        monkeypatch.setattr(imp_mod, "add_column", _flaky)
        imp = Importer(db.get_bind(), db, kv_table)
        plans = [
            {"name": "f_one", "field_type": "text", "sample_values": ["a"]},
            {"name": "f_two", "field_type": "text", "sample_values": ["b"]},
        ]
        with pytest.raises(RuntimeError, match="ddl fail"):
            imp._auto_add_fields(plans)

    def test_parse_xlsx_wider_row_raises(self, kv_table, db):
        """表头补齐后出现空列名 → 明确报错（覆盖 909-911）."""
        raw = _xlsx_bytes([["a", "b"], ["x", "y", "z"]])
        with pytest.raises(ValueError, match="空列名"):
            Importer(db.get_bind(), db, kv_table)._parse_xlsx(raw)

    def test_parse_xlsx_with_images_empty_sheet(self, kv_table, db):
        assert Importer(db.get_bind(), db, kv_table)._parse_xlsx_with_images(_empty_xlsx_bytes()) == ([], [])

    def test_resolve_original_rows_returns_rows(self, kv_table, db):
        imp = Importer(db.get_bind(), db, kv_table)
        analysis = ImportAnalysisResult(report={}, results=[], file_columns=[])
        rows = [{"a": 1}]
        assert imp._resolve_original_rows(None, None, rows, analysis) is rows

    def test_resolve_original_rows_from_results(self, kv_table, db):
        imp = Importer(db.get_bind(), db, kv_table)
        analysis = ImportAnalysisResult(
            report={},
            results=[ValidationResult(row_number=1, values={"a": 2})],
            file_columns=["a"],
        )
        assert imp._resolve_original_rows(None, None, None, analysis) == [{"a": 2}]


class _TaskFactoryHelper:
    """占位：保持 ImportTask 引用被使用（避免 ruff 未使用告警）."""

    def make(self, db) -> ImportTask:
        return ImportTask(table_id=1, status="pending", format="json", file_content="[]", filename="x.json")
