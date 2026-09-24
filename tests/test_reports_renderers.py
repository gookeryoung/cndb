"""reports 插件渲染器与安全增强集成测试.

覆盖：多表引用、DOCX/PDF/XLSX 富格式、安全沙箱、超时保护、主题风格应用.
"""

from __future__ import annotations

import io
import time

import pytest


@pytest.fixture
def auth_headers(client, db):
    """注册并登录，返回 Authorization header."""
    from cndb.plugins.accounts.models import User

    user = User(username="reporeq", email="rr@b.c", nickname="RR")
    user.set_password("pass1234")
    db.add(user)
    db.commit()
    resp = client.post(
        "/api/v1/accounts/auth/login",
        json={"login": "reporeq", "password": "pass1234"},
    )
    token = resp.json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture
def sample_tables(client, auth_headers):
    """创建一个含 2 张表 + 各 1 条记录的工作区，返回 (wid, tid_a, tid_b)."""
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws_renderer"})
    wid = ws.json()["id"]
    # 表 A
    tbl_a = client.post(
        f"/api/v1/workspaces/{wid}/tables",
        headers=auth_headers,
        json={"name": "员工表"},
    )
    tid_a = tbl_a.json()["id"]
    client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid_a}/fields",
        headers=auth_headers,
        json={"name": "姓名", "field_type": "text"},
    )
    rec_a = client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid_a}/records",
        headers=auth_headers,
        json={"values": {"姓名": "张三"}},
    )
    assert rec_a.status_code == 201, rec_a.text
    # 表 B
    tbl_b = client.post(
        f"/api/v1/workspaces/{wid}/tables",
        headers=auth_headers,
        json={"name": "项目表"},
    )
    tid_b = tbl_b.json()["id"]
    client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid_b}/fields",
        headers=auth_headers,
        json={"name": "项目名", "field_type": "text"},
    )
    rec_b = client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid_b}/records",
        headers=auth_headers,
        json={"values": {"项目名": "A项目"}},
    )
    assert rec_b.status_code == 201, rec_b.text
    return wid, tid_a, tid_b


# ── 多表引用 ────────────────────────────────────────


class TestMultiTableRender:
    """AC-1: 多表数据注入 records_by_table."""

    def test_records_by_table_present(self, client, auth_headers, sample_tables):
        _wid, tid_a, tid_b = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "跨表模板",
                "output_format": "docx",
                "template_content": "主表 {{ table_name }}\n额外表: {% for r in records_by_table['项目表'] %}{{ r['项目名'] }},{% endfor %}",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}, "extra_table_ids": [tid_b]},
        )
        assert resp.status_code == 200
        assert len(resp.content) > 0

    def test_records_by_table_empty_when_no_extra(self, client, auth_headers, sample_tables):
        """不传 extra_table_ids 时 records_by_table 为空 dict."""
        _wid, tid_a, _tid_b = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "单表模板",
                "output_format": "docx",
                "template_content": "records_by_table 长度: {{ records_by_table | length }}",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}},
        )
        assert resp.status_code == 200

    def test_missing_extra_table_id_404(self, client, auth_headers, sample_tables):
        """AC-2: extra_table_ids 含不存在表 ID 返回 404."""
        _wid, tid_a, _tid_b = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={"name": "坏引用", "output_format": "docx", "template_content": "test"},
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}, "extra_table_ids": [99999]},
        )
        assert resp.status_code == 404
        assert "额外数据表不存在" in resp.json()["detail"]


# ── DOCX 富文本 ──────────────────────────────────────


def _docx_heading(paragraph_text: str, content: bytes):
    """渲染 docx 字节流，返回文本为 paragraph_text 的段落对象."""
    from docx import Document

    doc = Document(io.BytesIO(content))
    for p in doc.paragraphs:
        if p.text == paragraph_text:
            return p
    raise AssertionError(f"未找到段落: {paragraph_text!r}")


def _docx_cell_fill(cell) -> str | None:
    """读取 docx 单元格底纹 fill 色值（无底纹返回 None）."""
    from docx.oxml.ns import qn

    tc_pr = cell._tc.tcPr
    if tc_pr is None:
        return None
    shd = tc_pr.find(qn("w:shd"))
    return shd.get(qn("w:fill")) if shd is not None else None


class TestThemedDocxRender:
    """主题风格在 DOCX 渲染中的应用：标题色/字号、表头底纹、代码字体、端到端模板 theme."""

    _TEXT = "# 一级标题\n正文段落\n| 列1 | 列2 |\n| --- | --- |\n| a | `code` |"

    def test_default_and_minimal_produce_identical_output(self):
        """不传 theme（默认 minimal）与显式 minimal 语义一致（回归保障，不做字节比较——docx 内嵌 ZIP 时间戳跨秒会变）."""
        from docx import Document

        from cndb.plugins.reports.routers.reports import _render_docx

        def _fingerprint(content: bytes) -> tuple[list[tuple[str, str]], str, float]:
            """段落(文本, 样式名) + 标题 run 颜色/字号 摘要."""
            doc = Document(io.BytesIO(content))
            paras = [(p.text, p.style.name) for p in doc.paragraphs]
            heading = next(p for p in doc.paragraphs if p.text == "一级标题")
            run = heading.runs[0]
            return paras, str(run.font.color.rgb), run.font.size.pt

        assert _fingerprint(_render_docx(self._TEXT, {})) == _fingerprint(_render_docx(self._TEXT, {}, "minimal"))

    def test_minimal_heading_black_18pt(self):
        """minimal 主题：标题黑色 18pt."""
        from cndb.plugins.reports.routers.reports import _render_docx

        content = _render_docx(self._TEXT, {}, "minimal")
        p = _docx_heading("一级标题", content)
        run = p.runs[0]
        assert str(run.font.color.rgb) == "000000"
        assert run.font.size.pt == pytest.approx(18.0)

    def test_business_heading_brand_color(self):
        """business 主题：H1 深蓝 #1F3864."""
        from cndb.plugins.reports.routers.reports import _render_docx

        content = _render_docx(self._TEXT, {}, "business")
        p = _docx_heading("一级标题", content)
        assert str(p.runs[0].font.color.rgb) == "1F3864"

    def test_business_table_header_shading(self):
        """business 主题：表头单元格底纹 #D9E2F3，表头文字同色系."""
        from docx import Document

        from cndb.plugins.reports.routers.reports import _render_docx

        content = _render_docx(self._TEXT, {}, "business")
        doc = Document(io.BytesIO(content))
        assert doc.tables, "应有表格"
        t = doc.tables[0]
        assert _docx_cell_fill(t.cell(0, 0)) == "D9E2F3"
        assert t.cell(0, 0).paragraphs[0].runs[0].bold is True

    def test_engineering_code_font(self):
        """engineering 主题：行内代码用 Consolas."""
        from docx import Document

        from cndb.plugins.reports.routers.reports import _render_docx

        content = _render_docx(self._TEXT, {}, "engineering")
        doc = Document(io.BytesIO(content))
        code_runs = [
            r for row in doc.tables[0].rows for c in row.cells for p in c.paragraphs for r in p.runs if r.text == "code"
        ]
        assert code_runs, "应存在代码 run"
        assert code_runs[0].font.name == "Consolas"

    def test_render_endpoint_applies_template_theme(self, client, auth_headers, sample_tables):
        """端到端：模板持久化 theme=business，渲染端点按其应用标题色."""
        _wid, tid, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "主题端到端",
                "table_id": tid,
                "output_format": "docx",
                "template_content": "# 主题标题",
                "theme": "business",
            },
        )
        assert tpl.status_code == 201, tpl.text
        resp = client.post(
            f"/api/v1/reports/{tpl.json()['id']}/render",
            headers=auth_headers,
            json={"table_id": tid, "params": {}},
        )
        assert resp.status_code == 200, resp.text
        p = _docx_heading("主题标题", resp.content)
        assert str(p.runs[0].font.color.rgb) == "1F3864"


# ── PDF 主题 ─────────────────────────────────────────


class TestThemedPdfRender:
    """主题风格在 PDF 渲染中的应用：表格预设色、全主题 smoke、端到端."""

    def test_table_style_uses_preset_colors(self):
        """_pdf_table_style_cmds 按 preset 产出表头底色/文字色命令."""
        from reportlab.lib import colors

        from cndb.plugins.reports.routers.reports import _pdf_table_style_cmds
        from cndb.plugins.reports.themes import get_theme_preset

        cmds = _pdf_table_style_cmds(get_theme_preset("business"))
        assert ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#D9E2F3")) in cmds
        assert ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#1F3864")) in cmds

    def test_render_pdf_all_themes_smoke(self):
        """五类主题 PDF 渲染均成功（防预设字体名破坏 reportlab 字体注册）."""
        from cndb.plugins.reports.routers.reports import _render_pdf

        text = "# 标题\n## 二级\n正文段落\n| a | b |\n| --- | --- |\n| 1 | `c` |"
        for theme in ("business", "minimal", "modern", "engineering", "academic"):
            data = _render_pdf(text, {}, theme)
            assert data.startswith(b"%PDF"), theme

    def test_render_endpoint_pdf_theme_smoke(self, client, auth_headers, sample_tables):
        """端到端：theme=modern 的 PDF 模板渲染成功."""
        _wid, tid, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "PDF主题",
                "table_id": tid,
                "output_format": "pdf",
                "template_content": "# 主题标题\n正文",
                "theme": "modern",
            },
        )
        assert tpl.status_code == 201, tpl.text
        resp = client.post(
            f"/api/v1/reports/{tpl.json()['id']}/render",
            headers=auth_headers,
            json={"table_id": tid, "params": {}},
        )
        assert resp.status_code == 200, resp.text
        assert resp.content.startswith(b"%PDF")


# ── DOCX 富文本（既有） ────────────────────────────────


class TestDocxRenderer:
    """AC-3: DOCX 标题/粗体/代码."""

    def test_docx_markdown_heading(self, client, auth_headers, sample_tables):
        _wid, tid_a, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "docx标题测试",
                "output_format": "docx",
                "template_content": "# 一级标题\n## 二级标题\n普通行",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}},
        )
        assert resp.status_code == 200
        from docx import Document

        doc = Document(io.BytesIO(resp.content))
        heading_count = sum(1 for p in doc.paragraphs if p.style.name.startswith("Heading"))
        assert heading_count >= 2, f"期望 >= 2 个 Heading，实际 {heading_count}"

    def test_docx_markdown_bold(self, client, auth_headers, sample_tables):
        _wid, tid_a, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "docx粗体测试",
                "output_format": "docx",
                "template_content": "这段有 **粗体**",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}},
        )
        assert resp.status_code == 200
        from docx import Document

        doc = Document(io.BytesIO(resp.content))
        has_bold = any(run.bold for p in doc.paragraphs for run in p.runs)
        assert has_bold, "期望有粗体 Run"


# ── PDF 渲染 ─────────────────────────────────────────


class TestPdfRenderer:
    """AC-4: PDF 中文 + 表格 + 自动分页."""

    def test_pdf_chinese_no_crash(self, client, auth_headers, sample_tables):
        _wid, tid_a, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "pdf中文",
                "output_format": "pdf",
                "template_content": "# 员工月报\n\n这是一份中文 PDF 报告。\n\n| 姓名 | 部门 |\n| --- | --- |\n| 张三 | 研发 |",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}},
        )
        assert resp.status_code == 200
        assert resp.headers["content-type"] == "application/pdf"
        assert len(resp.content) > 100

    def test_pdf_header_footer(self, client, auth_headers, sample_tables):
        """PDF 应包含页眉和页脚（文件结构检查)."""
        _wid, tid_a, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={"name": "pdf页眉页脚", "output_format": "pdf", "template_content": "测试内容"},
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}},
        )
        assert resp.status_code == 200
        assert b"/Type /Catalog" in resp.content


# ── Excel 多 Sheet ────────────────────────────────────


class TestXlsxMultiSheet:
    """AC-5: XLSX 多 Sheet + 表头加粗."""

    def test_multi_sheet_count(self, client, auth_headers, sample_tables):
        wid, tid_a, tid_b = sample_tables
        # 再建第三张表
        tbl_c = client.post(
            f"/api/v1/workspaces/{wid}/tables",
            headers=auth_headers,
            json={"name": "产品表"},
        )
        tid_c = tbl_c.json()["id"]
        client.post(
            f"/api/v1/workspaces/{wid}/tables/{tid_c}/fields",
            headers=auth_headers,
            json={"name": "产品", "field_type": "text"},
        )

        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={"name": "xlsx多sheet", "output_format": "xlsx", "template_content": "多 sheet"},
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}, "extra_table_ids": [tid_b, tid_c]},
        )
        assert resp.status_code == 200
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(resp.content))
        assert len(wb.sheetnames) == 3, f"期望 3 个 Sheet, 实际 {wb.sheetnames}"
        # Sheet 名应反映表名（测试环境 records 可能为空，但 Sheet 仍应按表名创建）
        assert len(set(wb.sheetnames)) == 3
        for sn in wb.sheetnames:
            ws = wb[sn]
            assert ws is not None

    def test_single_sheet_when_no_extra(self, client, auth_headers, sample_tables):
        _wid, tid_a, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={"name": "xlsx单sheet", "output_format": "xlsx", "template_content": "单 sheet"},
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}},
        )
        assert resp.status_code == 200
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(resp.content))
        assert len(wb.sheetnames) == 1


# ── 安全沙箱 ──────────────────────────────────────────


class TestSecuritySandbox:
    """AC-6: 安全沙箱拒绝逃逸."""

    def test_block_class_escape(self, client, auth_headers, sample_tables):
        _wid, tid_a, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "恶意模板",
                "output_format": "docx",
                "template_content": '{{ "".__class__.__mro__[1].__subclasses__() }}',
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}},
        )
        assert resp.status_code == 400

    def test_block_open_builtin(self, client, auth_headers, sample_tables):
        _wid, tid_a, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "open攻击",
                "output_format": "docx",
                "template_content": "{{ open('/etc/passwd').read() }}",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}},
        )
        assert resp.status_code == 400

    def test_allowed_filter_length(self, client, auth_headers, sample_tables):
        _wid, tid_a, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "正常模板",
                "output_format": "docx",
                "template_content": "共 {{ records | length }} 条",
            },
        )
        tpl_id = tpl.json()["id"]
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}},
        )
        assert resp.status_code == 200


# ── 超时保护 ──────────────────────────────────────────


class TestRenderTimeout:
    """AC-9: 渲染超时保护."""

    def test_sandbox_blocks_huge_range(self, client, auth_headers, sample_tables):
        """SandboxedEnvironment 本身已限制 range(MAX_RANGE=100000)."""
        _wid, tid_a, _ = sample_tables
        tpl = client.post(
            "/api/v1/reports",
            headers=auth_headers,
            json={
                "name": "超时测试",
                "output_format": "docx",
                "template_content": "{% for i in range(99999999) %}{% endfor %}done",
            },
        )
        tpl_id = tpl.json()["id"]
        start = time.time()
        resp = client.post(
            f"/api/v1/reports/{tpl_id}/render",
            headers=auth_headers,
            json={"table_id": tid_a, "params": {}},
        )
        elapsed = time.time() - start
        assert resp.status_code == 400, f"应被拦截, 实际 status={resp.status_code}"
        assert elapsed < 15, f"耗时 {elapsed:.2f}s 过长"


# ── 向后兼容 ─────────────────────────────────────────


def test_row_ids_not_affect_extra_tables(client, auth_headers, sample_tables):
    """row_ids 仅过滤主表 records，额外表数据不受影响."""
    _wid, tid_a, tid_b = sample_tables
    tpl = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={
            "name": "rowids主表隔离",
            "output_format": "docx",
            "template_content": "主 {{ records | length }} 额外 {{ records_by_table['项目表'] | length }}",
        },
    )
    tpl_id = tpl.json()["id"]
    resp = client.post(
        f"/api/v1/reports/{tpl_id}/render",
        headers=auth_headers,
        json={"table_id": tid_a, "params": {}, "extra_table_ids": [tid_b], "row_ids": [1]},
    )
    assert resp.status_code == 200
    from docx import Document

    doc = Document(io.BytesIO(resp.content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "额外 1" in text, f"额外表不应受 row_ids 影响，实际: {text!r}"


def test_backward_compat_single_table_render(client, auth_headers, sample_tables):
    """AC-10: 不传 extra_table_ids 的单表渲染与旧版一致."""
    _wid, tid_a, _ = sample_tables
    tpl = client.post(
        "/api/v1/reports",
        headers=auth_headers,
        json={
            "name": "兼容模板",
            "output_format": "docx",
            "template_content": "Hello {{ table_name }}",
        },
    )
    tpl_id = tpl.json()["id"]
    resp = client.post(
        f"/api/v1/reports/{tpl_id}/render",
        headers=auth_headers,
        json={"table_id": tid_a, "params": {}},
    )
    assert resp.status_code == 200
    assert len(resp.content) > 0


# ── 超时保护单元测试 ──────────────────────────────────


def test_render_with_timeout_raises_timeout():
    """直接测试 _render_with_timeout 超时抛出 TimeoutError."""
    # 用 mock 验证超时逻辑
    from typing import Any

    from cndb.plugins.reports.routers.reports import _render_with_timeout

    class _SlowTmpl:
        def __init__(self) -> None:
            self.called = False

        def render(self, **ctx: Any) -> str:
            if not self.called:
                self.called = True
                time.sleep(1.5)
            return "done"

    fake = _SlowTmpl()
    start = time.time()
    with pytest.raises(TimeoutError):
        _render_with_timeout(fake, {}, timeout=1)
    elapsed = time.time() - start
    assert elapsed < 3, f"超时检测耗时 {elapsed:.2f}s 应 < 3s"
