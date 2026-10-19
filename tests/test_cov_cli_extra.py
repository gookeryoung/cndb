"""补充覆盖：cli 包（service / backup / views / docs / restore / users）的遗漏行与分支.

仅新增测试，不改动源码与既有测试。CLI 窄函数用内存替身替换，进程探测与
注册表访问均被 monkeypatch，不真正派生服务或写注册表；文件库落在 tmp_path。
"""

from __future__ import annotations

import argparse
import datetime as dt
import sys
import tarfile
from decimal import Decimal
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest

from cndb.core.config import settings

# ══════════════════════════════════════════════════════════════
# cli/service.py
# ══════════════════════════════════════════════════════════════


class _FakeWinReg:
    """winreg 的极简内存替身，覆盖注册表窄封装函数的真实实现."""

    HKEY_CURRENT_USER = 0x80000001
    KEY_SET_VALUE = 0x0002
    REG_SZ = 1

    def __init__(self) -> None:
        self.values: dict[str, str] = {}
        self.raise_on_open = False

    class _Key:
        def __init__(self, reg: _FakeWinReg) -> None:
            self._reg = reg

        def __enter__(self) -> _FakeWinReg._Key:
            return self

        def __exit__(self, *_a: Any) -> bool:
            return False

    def OpenKey(self, _root: Any, _path: str, *_a: Any) -> _FakeWinReg._Key:
        if self.raise_on_open:
            raise FileNotFoundError(_path)
        return _FakeWinReg._Key(self)

    def QueryValueEx(self, key: _FakeWinReg._Key, name: str) -> tuple[str, int]:
        if name not in key._reg.values:
            raise FileNotFoundError(name)
        return key._reg.values[name], _FakeWinReg.REG_SZ

    def CreateKeyEx(self, _root: Any, _path: str, *_a: Any) -> _FakeWinReg._Key:
        return _FakeWinReg._Key(self)

    def SetValueEx(self, key: _FakeWinReg._Key, name: str, _rez: int, _typ: int, value: str) -> None:
        key._reg.values[name] = value

    def DeleteValue(self, key: _FakeWinReg._Key, name: str) -> None:
        key._reg.values.pop(name, None)


@pytest.fixture
def fake_winreg(monkeypatch: pytest.MonkeyPatch) -> _FakeWinReg:
    reg = _FakeWinReg()
    monkeypatch.setitem(sys.modules, "winreg", reg)
    return reg


class TestServiceRegistryHelpers:
    def test_read_run_value_registered(self, fake_winreg: _FakeWinReg, monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.cli import service

        fake_winreg.values[service.RUN_VALUE_NAME] = r'"cndb.exe" service run'
        assert service._read_run_value() == r'"cndb.exe" service run'

    def test_read_run_value_missing_key(self, fake_winreg: _FakeWinReg) -> None:
        from cndb.cli import service

        fake_winreg.raise_on_open = True
        assert service._read_run_value() is None

    def test_write_and_delete_run_value(self, fake_winreg: _FakeWinReg) -> None:
        from cndb.cli import service

        service._write_run_value("cmd-1")  # 覆盖真实写入分支
        assert fake_winreg.values[service.RUN_VALUE_NAME] == "cmd-1"
        service._delete_run_value()  # 覆盖真实删除分支
        assert service.RUN_VALUE_NAME not in fake_winreg.values

    def test_delete_run_value_missing_is_silent(self, fake_winreg: _FakeWinReg) -> None:
        from cndb.cli import service

        fake_winreg.raise_on_open = True
        service._delete_run_value()  # FileNotFoundError 被吞掉


class TestServiceProcessHelpers:
    def test_run_cmd_success(self) -> None:
        from cndb.cli import service

        rc, out = service._run_cmd([sys.executable, "-c", "print('ok')"])
        assert rc == 0
        assert "ok" in out

    def test_run_cmd_oserror(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.cli import service

        def _boom(*_a: Any, **_k: Any) -> Any:
            raise OSError("exec failed")

        monkeypatch.setattr(service.subprocess, "run", _boom)
        assert service._run_cmd(["nope"]) == (-1, "")

    def test_pid_alive_delegates(self) -> None:
        import os

        from cndb.cli import service

        # 当前进程必然存活
        assert service._pid_alive(os.getpid()) is True


class _FakePopen:
    pid = 5150
    calls: dict[str, Any] = {}

    def __init__(self, cmd: list[str], **kwargs: Any) -> None:
        type(self).calls = {"cmd": cmd, **kwargs}


class TestServiceRunNonWindows:
    def test_run_spawns_posix_with_local(
        self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path, capsys: pytest.CaptureFixture[str]
    ) -> None:
        from cndb.cli import service

        monkeypatch.setattr(service, "PID_FILE", tmp_path / "service.pid")
        monkeypatch.setattr(service, "SERVICE_LOG", tmp_path / "service.log")
        monkeypatch.setattr(service, "_read_run_value", lambda: None)
        monkeypatch.setattr(service.sys, "platform", "linux")
        monkeypatch.setattr("cndb.gui.checks.port_status", lambda port, host="127.0.0.1": SimpleNamespace(used=False))
        monkeypatch.setattr(service.subprocess, "Popen", _FakePopen)
        monkeypatch.setenv("CNDB_LOCAL_MODE", "0")

        rc = service.run("127.0.0.1", 8772, local=True)

        assert rc == 0
        cmd = _FakePopen.calls["cmd"]
        assert cmd[-5:] == ["--host", "127.0.0.1", "--port", "8772", "--local"]
        assert "--local" in cmd
        assert "creationflags" not in _FakePopen.calls  # 非 win32 不加 creationflags
        assert service.PID_FILE.read_text(encoding="utf-8") == "5150"
        assert "已启动" in capsys.readouterr().out


class TestServiceStopPosix:
    def _prep(self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> Any:
        from cndb.cli import service

        monkeypatch.setattr(service, "PID_FILE", tmp_path / "service.pid")
        service.PID_FILE.write_text("5150", encoding="utf-8")
        monkeypatch.setattr(service, "_pid_alive", lambda pid: True)
        monkeypatch.setattr(service.sys, "platform", "linux")
        return service

    def test_stop_posix_kill_success(self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
        service = self._prep(monkeypatch, tmp_path)
        killed: list[tuple[int, int]] = []
        monkeypatch.setattr(service.os, "kill", lambda pid, sig: killed.append((pid, sig)))

        assert service.stop() is True
        assert killed == [(5150, 15)]
        assert not service.PID_FILE.exists()

    def test_stop_posix_permission_error(self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
        service = self._prep(monkeypatch, tmp_path)

        def _deny(pid: int, sig: int) -> None:
            raise PermissionError("no perm")

        monkeypatch.setattr(service.os, "kill", _deny)

        assert service.stop() is True  # 权限不足保守按已终止处理
        assert not service.PID_FILE.exists()

    def test_stop_posix_process_lookup(self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
        service = self._prep(monkeypatch, tmp_path)

        def _gone(pid: int, sig: int) -> None:
            raise ProcessLookupError

        monkeypatch.setattr(service.os, "kill", _gone)

        assert service.stop() is False


class TestServiceCommandDispatch:
    def _ns(self, service_cmd: str | None) -> argparse.Namespace:
        return argparse.Namespace(command="service", service_cmd=service_cmd)

    def test_dispatch_enable(self, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> None:
        from cndb.cli import service

        monkeypatch.setattr(service.sys, "platform", "win32")
        monkeypatch.setattr(service, "enable", lambda host, port, local=False: "CMD")
        service.service_command(argparse.Namespace(command="service", service_cmd="enable", host="0.0.0.0", port=8000))
        assert "已启用开机自启" in capsys.readouterr().out

    def test_dispatch_disable(self, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]) -> None:
        from cndb.cli import service

        monkeypatch.setattr(service.sys, "platform", "win32")
        monkeypatch.setattr(service, "disable", lambda: None)
        service.service_command(self._ns("disable"))
        assert "已取消开机自启" in capsys.readouterr().out

    def test_dispatch_status_with_port(
        self, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
    ) -> None:
        from cndb.cli import service

        monkeypatch.setattr(service.sys, "platform", "win32")
        monkeypatch.setattr(
            service,
            "collect_status",
            lambda: service.ServiceStatus(autostart=True, running=True, pid=42, port=8000),
        )
        service.service_command(self._ns("status"))
        out = capsys.readouterr().out
        assert "已启用" in out and "运行中" in out and "固化端口: 8000" in out

    def test_dispatch_stop_true_and_false(
        self, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
    ) -> None:
        from cndb.cli import service

        monkeypatch.setattr(service.sys, "platform", "win32")
        monkeypatch.setattr(service, "stop", lambda: True)
        service.service_command(self._ns("stop"))
        assert "已停止" in capsys.readouterr().out

        monkeypatch.setattr(service, "stop", lambda: False)
        service.service_command(self._ns("stop"))
        assert "未在运行" in capsys.readouterr().out

    def test_dispatch_run_exits(self, monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.cli import service

        monkeypatch.setattr(service.sys, "platform", "win32")
        monkeypatch.setattr(service, "run", lambda host, port, local=False: 0)
        with pytest.raises(SystemExit) as exc:
            service.service_command(argparse.Namespace(command="service", service_cmd="run", host="0.0.0.0", port=8000))
        assert exc.value.code == 0


# ══════════════════════════════════════════════════════════════
# cli/backup.py
# ══════════════════════════════════════════════════════════════


def _setup_sqlite(tmp_path: Path) -> Path:
    import sqlite3

    db = tmp_path / "b.db"
    conn = sqlite3.connect(str(db))
    conn.executescript("CREATE TABLE t (id INTEGER PRIMARY KEY, v TEXT); INSERT INTO t (v) VALUES ('a');")
    conn.commit()
    conn.close()
    return db


class TestBackupHelpers:
    def test_resolve_sqlite_path_no_leading_slash(self) -> None:
        from cndb.cli.backup import _resolve_sqlite_path

        p = _resolve_sqlite_path("sqlite:relative.db")
        assert p.name == "relative.db"

    def test_detect_format_directory(self, tmp_path: Path) -> None:
        from cndb.cli.backup import _detect_format

        assert _detect_format(tmp_path / "out") == "directory"
        assert _detect_format(tmp_path / "out.tgz") == "archive"

    def test_total_size_file_and_dir(self, tmp_path: Path) -> None:
        from cndb.cli.backup import _total_size

        f = tmp_path / "f.txt"
        f.write_text("12345", encoding="utf-8")
        assert _total_size(f) == 5

        d = tmp_path / "d"
        d.mkdir()
        (d / "a.txt").write_text("123", encoding="utf-8")
        (d / "sub").mkdir()
        (d / "sub" / "b.txt").write_text("12", encoding="utf-8")
        assert _total_size(d) == 5

    def test_copy_to_directory_new_and_empty(self, tmp_path: Path) -> None:
        from cndb.cli.backup import _copy_to_directory

        src = tmp_path / "src"
        (src / "sub").mkdir(parents=True)
        (src / "top.txt").write_text("x", encoding="utf-8")
        (src / "sub" / "inner.txt").write_text("y", encoding="utf-8")

        out_new = tmp_path / "out_new"
        _copy_to_directory(src, out_new)
        assert (out_new / "top.txt").is_file()
        assert (out_new / "sub" / "inner.txt").is_file()

        out_empty = tmp_path / "out_empty"
        out_empty.mkdir()
        _copy_to_directory(src, out_empty)
        assert (out_empty / "top.txt").is_file()

    def test_copy_to_directory_rejects_nonempty(self, tmp_path: Path) -> None:
        from cndb.cli.backup import BackupError, _copy_to_directory

        src = tmp_path / "src"
        src.mkdir()
        (src / "a.txt").write_text("x", encoding="utf-8")

        busy = tmp_path / "busy"
        busy.mkdir()
        (busy / "existing.txt").write_text("x", encoding="utf-8")
        with pytest.raises(BackupError, match="已包含文件"):
            _copy_to_directory(src, busy)

    def test_copy_to_directory_rejects_file_target(self, tmp_path: Path) -> None:
        from cndb.cli.backup import BackupError, _copy_to_directory

        src = tmp_path / "src"
        src.mkdir()
        notdir = tmp_path / "notdir.txt"
        notdir.write_text("x", encoding="utf-8")
        with pytest.raises(BackupError, match="不是目录"):
            _copy_to_directory(src, notdir)


class TestCoerceRawValueEdges:
    def test_datetime_bytes_fallthrough(self) -> None:
        from sqlalchemy.types import Date, DateTime, Time

        from cndb.cli.backup import _coerce_raw_value

        assert _coerce_raw_value(b"raw", DateTime()) == b"raw"
        assert _coerce_raw_value(b"raw", Date()) == b"raw"
        assert _coerce_raw_value(b"raw", Time()) == b"raw"

    def test_date_instances_and_invalid(self) -> None:
        from sqlalchemy.types import Date

        from cndb.cli.backup import _coerce_raw_value

        assert _coerce_raw_value(dt.date(2024, 1, 2), Date()) == dt.date(2024, 1, 2)
        # Unix 时间戳 → date（UTC）
        assert _coerce_raw_value(1704067200, Date()) == dt.date(2024, 1, 1)
        # 非法字符串原样保留
        assert _coerce_raw_value("garbage", Date()) == "garbage"

    def test_time_instances_and_invalid(self) -> None:
        from sqlalchemy.types import Time

        from cndb.cli.backup import _coerce_raw_value

        assert _coerce_raw_value(dt.time(1, 2, 3), Time()) == dt.time(1, 2, 3)
        assert _coerce_raw_value("garbage", Time()) == "garbage"


class TestBackupDirectoryMode:
    def test_create_backup_directory_format(self, tmp_path: Path) -> None:
        from cndb.cli.backup import create_backup

        db = _setup_sqlite(tmp_path)
        out_dir = tmp_path / "exported"
        result = create_backup(
            output=out_dir,
            mode="native",
            include_uploads=False,
            database_url=f"sqlite:///{db}",
            fmt="directory",
        )
        assert result == out_dir.resolve()
        assert (result / "manifest.json").is_file()
        assert (result / "data" / "cndb.db").is_file()

    def test_backup_command_dir_flag(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.cli.backup import backup_command

        db = _setup_sqlite(tmp_path)
        monkeypatch.setattr(settings, "DATABASE_URL", f"sqlite:///{db}")
        out_dir = tmp_path / "cmd_dir"
        args = argparse.Namespace(
            output=str(out_dir),
            mode="native",
            no_uploads=True,
            format=None,
            dir=True,
            no_fallback=True,
        )
        backup_command(args)
        assert (out_dir / "manifest.json").is_file()


# ══════════════════════════════════════════════════════════════
# cli/views.py
# ══════════════════════════════════════════════════════════════


@pytest.fixture
def cli_views_env(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    from sqlalchemy import create_engine

    import cndb.plugins.accounts.models
    import cndb.plugins.tables.models
    import cndb.plugins.workspaces.models  # noqa: F401
    from cndb.models.base import Base

    db_path = tmp_path / "views_extra.db"
    monkeypatch.setattr(settings, "DATABASE_URL", f"sqlite:///{db_path}")
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    engine.dispose()
    return db_path


def _session_for(db_path: Path) -> Any:
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    return sessionmaker(bind=engine)()


class TestViewsResolvers:
    def test_get_session_creates_missing_tables(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        """全新空库 → inspect 发现缺表 → create_all（覆盖建表兜底）."""
        from cndb.cli.views import _get_session

        monkeypatch.setattr(settings, "DATABASE_URL", f"sqlite:///{tmp_path / 'fresh.db'}")
        db = _get_session()
        try:
            from sqlalchemy import inspect

            assert inspect(db.get_bind()).has_table("tables_dataview")
        finally:
            db.close()

    def test_resolve_workspace_name_not_unique(self, cli_views_env: Path) -> None:
        from cndb.cli.views import _resolve_workspace
        from cndb.plugins.workspaces.models import Workspace

        db = _session_for(cli_views_env)
        try:
            db.add(Workspace(name="dup_ws"))
            db.add(Workspace(name="dup_ws"))
            db.commit()
            with pytest.raises(SystemExit, match="不唯一"):
                _resolve_workspace(db, "dup_ws")
        finally:
            db.close()

    def test_resolve_table_by_id_and_errors(self, cli_views_env: Path) -> None:
        from cndb.cli.views import _resolve_table

        db = _session_for(cli_views_env)
        try:
            # 不存在的 id → SystemExit
            with pytest.raises(SystemExit, match="不存在数据表 id"):
                _resolve_table(db, 1, "999")
            with pytest.raises(SystemExit, match="不存在数据表名称"):
                _resolve_table(db, 1, "no_such_table")
        finally:
            db.close()

    def test_resolve_table_name_not_unique(self, cli_views_env: Path) -> None:
        from cndb.cli.views import _resolve_table
        from cndb.plugins.tables.models import DataTable

        db = _session_for(cli_views_env)
        try:
            for _ in range(2):
                t = DataTable(workspace_id=1, name="dup_table")
                t.ensure_db_name()
                db.add(t)
            db.commit()
            with pytest.raises(SystemExit, match="不唯一"):
                _resolve_table(db, 1, "dup_table")
        finally:
            db.close()


class TestViewsPrintingAndDispatch:
    def test_print_prune_removed_branch(self, capsys: pytest.CaptureFixture[str]) -> None:
        from cndb.cli.views import _print_prune

        _print_prune(
            {
                "workspace": {"id": 1, "name": "W"},
                "table": None,
                "dry_run": False,
                "orphans": [
                    {
                        "id": 1,
                        "name": "v",
                        "table_id": 1,
                        "table_name": "T",
                        "owner_id": None,
                        "scope_key": 0,
                        "reason": "创建者已注销",
                        "created_at": "",
                        "updated_at": "",
                    }
                ],
                "removed": [1],
            }
        )
        assert "已删除 1 个孤儿视图" in capsys.readouterr().out

    def test_unknown_subcommand_exits(self, capsys: pytest.CaptureFixture[str]) -> None:
        from cndb.cli.views import views_command

        with pytest.raises(SystemExit) as exc:
            views_command(argparse.Namespace(views_cmd="weird"))
        assert exc.value.code == 2


# ══════════════════════════════════════════════════════════════
# cli/docs.py
# ══════════════════════════════════════════════════════════════


class TestDocsEdges:
    def test_get_session_creates_missing_tables(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.cli.docs import cmd_gc

        monkeypatch.setattr(settings, "DATABASE_URL", f"sqlite:///{tmp_path / 'fresh_docs.db'}")
        from cndb.plugins.documents.services import content as content_service

        monkeypatch.setattr(content_service, "_DOCS_DIR", tmp_path / "docs")
        # 全新空库 → _get_session 兜底建表（覆盖 create_all 分支）
        payload = cmd_gc(argparse.Namespace(apply=False))
        assert payload["dry_run"] is True

    def test_print_gc_dry_run_no_orphans(self, capsys: pytest.CaptureFixture[str]) -> None:
        from cndb.cli.docs import _print_gc

        _print_gc(
            {
                "root": "/tmp/docs",
                "dry_run": True,
                "referenced_files": 0,
                "missing_files": [],
                "orphan_files": [],
                "orphan_bytes": 0,
                "removed_files": 0,
                "removed_dirs": 0,
            }
        )
        out = capsys.readouterr().out
        assert "--apply" not in out  # dry_run 且无孤儿 → 不打印删除提示

    def test_unknown_subcommand_exits(self, capsys: pytest.CaptureFixture[str]) -> None:
        from cndb.cli.docs import docs_command

        with pytest.raises(SystemExit) as exc:
            docs_command(argparse.Namespace(docs_cmd="weird"))
        assert exc.value.code == 2


# ══════════════════════════════════════════════════════════════
# cli/restore/inspect.py
# ══════════════════════════════════════════════════════════════


class TestRestoreInspectEdges:
    def test_resolve_sqlite_path_no_leading_slash(self) -> None:
        from cndb.cli.restore.inspect import _resolve_sqlite_path

        assert _resolve_sqlite_path("sqlite:relative.db").name == "relative.db"

    def test_detect_backup_format_unknown_suffix_file(self, tmp_path: Path) -> None:
        from cndb.cli.restore.inspect import _detect_backup_format

        assert _detect_backup_format(tmp_path / "backup.bin") == "archive"

    def test_directory_backup_missing_manifest(self, tmp_path: Path) -> None:
        from cndb.cli.restore.inspect import inspect_backup
        from cndb.cli.restore.types import RestoreError

        d = tmp_path / "bkdir"
        d.mkdir()
        with pytest.raises(RestoreError, match="缺少 manifest\\.json"):
            inspect_backup(d)

    def test_directory_backup_invalid_json(self, tmp_path: Path) -> None:
        from cndb.cli.restore.inspect import inspect_backup
        from cndb.cli.restore.types import RestoreError

        d = tmp_path / "bkdir2"
        d.mkdir()
        (d / "manifest.json").write_text("not json", encoding="utf-8")
        with pytest.raises(RestoreError, match="不是合法 JSON"):
            inspect_backup(d)

    def test_archive_manifest_is_directory(self, tmp_path: Path) -> None:
        from cndb.cli.restore.inspect import inspect_backup
        from cndb.cli.restore.types import RestoreError

        archive = tmp_path / "weird.tar.gz"
        with tarfile.open(archive, "w:gz") as tar:
            info = tarfile.TarInfo("backup/manifest.json")
            info.type = tarfile.DIRTYPE
            tar.addfile(info)
        with pytest.raises(RestoreError, match="无法读取"):
            inspect_backup(archive)


# ══════════════════════════════════════════════════════════════
# cli/restore/apply.py
# ══════════════════════════════════════════════════════════════


class TestRestoreApplyEdges:
    def test_extracted_db_dir_fallback(self, tmp_path: Path) -> None:
        from cndb.cli.restore.apply import _extracted_db_dir

        assert _extracted_db_dir(tmp_path) == tmp_path / "data"

    def test_from_json_safe_branches(self) -> None:
        import base64

        from cndb.cli.restore.apply import _from_json_safe

        assert _from_json_safe({"__base64__": base64.b64encode(b"hi").decode()}) == b"hi"
        assert _from_json_safe({"__decimal__": "1.5"}) == Decimal("1.5")
        assert _from_json_safe({"other": "x"}) == {"other": "x"}
        assert _from_json_safe("2024-06-15") == dt.date(2024, 6, 15)
        assert _from_json_safe("12:30:00") == dt.time(12, 30, 0)
        assert _from_json_safe("2024-06-15T10:30:00") == dt.datetime(2024, 6, 15, 10, 30, 0)
        assert _from_json_safe(None) is None
        assert _from_json_safe(42) == 42

    def test_migrate_after_restore_stamp_head_failure(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        from sqlalchemy.exc import SQLAlchemyError

        from cndb.cli.restore.apply import _migrate_after_restore
        from cndb.cli.restore.types import RestoreError
        from cndb.core import migrations

        db_path = tmp_path / "no_alembic.db"
        db_path.touch()

        def _boom(_url: str) -> None:
            raise SQLAlchemyError("stamp failed")

        monkeypatch.setattr(migrations, "stamp_head", _boom)
        with pytest.raises(RestoreError, match="补写迁移版本标记失败"):
            _migrate_after_restore(f"sqlite:///{db_path}", "")

    def test_migrate_after_restore_create_all_failure(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        import sqlite3

        from sqlalchemy.exc import SQLAlchemyError

        from cndb.cli.restore.apply import _migrate_after_restore
        from cndb.cli.restore.types import RestoreError
        from cndb.core import migrations, plugin_registry

        db_path = tmp_path / "with_alembic.db"
        conn = sqlite3.connect(str(db_path))
        conn.executescript("CREATE TABLE alembic_version (version_num VARCHAR(32));")
        conn.commit()
        conn.close()

        monkeypatch.setattr(migrations, "upgrade_to_head", lambda _url: None)

        def _boom() -> None:
            raise SQLAlchemyError("discover failed")

        monkeypatch.setattr(plugin_registry.plugin_registry, "discover_and_load", _boom)
        with pytest.raises(RestoreError, match="补建缺失表失败"):
            _migrate_after_restore(f"sqlite:///{db_path}", "rev1")

    def test_migrate_after_restore_heal_failure(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        import sqlite3

        from sqlalchemy.exc import SQLAlchemyError

        from cndb.cli.restore.apply import _migrate_after_restore
        from cndb.cli.restore.types import RestoreError
        from cndb.core import migrations, plugin_registry

        db_path = tmp_path / "heal.db"
        conn = sqlite3.connect(str(db_path))
        conn.executescript("CREATE TABLE alembic_version (version_num VARCHAR(32));")
        conn.commit()
        conn.close()

        monkeypatch.setattr(migrations, "upgrade_to_head", lambda _url: None)
        monkeypatch.setattr(plugin_registry.plugin_registry, "discover_and_load", lambda: None)

        def _boom(_url: str) -> list[str]:
            raise SQLAlchemyError("heal failed")

        monkeypatch.setattr(migrations, "heal_schema_drift", _boom)
        with pytest.raises(RestoreError, match="schema 自愈失败"):
            _migrate_after_restore(f"sqlite:///{db_path}", "rev1")

    def test_restore_uploads_staging_residue_and_bak(self, tmp_path: Path) -> None:
        from cndb.cli.restore.apply import _restore_uploads

        extracted = tmp_path / "extracted"
        (extracted / "uploads" / "1").mkdir(parents=True)
        (extracted / "uploads" / "1" / "a.txt").write_text("hi", encoding="utf-8")

        target = tmp_path / "uploads"
        target.mkdir()
        (target / "old.txt").write_text("old", encoding="utf-8")
        # 残留 staging + 残留 bak，覆盖清理分支
        (tmp_path / "uploads.new").mkdir()
        (tmp_path / "uploads.bak").mkdir()
        (tmp_path / "uploads.bak" / "stale.txt").write_text("x", encoding="utf-8")

        count = _restore_uploads(extracted, target, included=True)
        assert count == 1
        assert (target / "1" / "a.txt").is_file()
        assert not (tmp_path / "uploads.bak").exists()

    def test_restore_uploads_rollback_on_rename_failure(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.cli.restore.apply import _restore_uploads

        extracted = tmp_path / "extracted2"
        (extracted / "uploads").mkdir(parents=True)
        (extracted / "uploads" / "a.txt").write_text("hi", encoding="utf-8")

        target = tmp_path / "uploads2"
        target.mkdir()
        (target / "old.txt").write_text("old", encoding="utf-8")
        (tmp_path / "uploads2.bak").mkdir()
        (tmp_path / "uploads2.bak" / "stale.txt").write_text("x", encoding="utf-8")

        orig_rename = Path.rename

        def _rename(self: Path, dest: Any, **kw: Any) -> Any:
            if self.name.endswith(".new"):
                raise OSError("simulated rename failure")
            return orig_rename(self, dest, **kw)

        monkeypatch.setattr(Path, "rename", _rename)
        with pytest.raises(OSError, match="simulated rename failure"):
            _restore_uploads(extracted, target, included=True)
        # 回滚后旧目录恢复、staging 被清理
        assert target.is_dir()
        assert not (tmp_path / "uploads2.new").exists()

    def test_restore_backup_non_sqlite_sqlalchemy(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.cli.restore import apply as apply_mod
        from cndb.cli.restore.types import RestoreLossReport

        bkdir = tmp_path / "dirbk"
        bkdir.mkdir()
        (bkdir / "manifest.json").write_text("{}", encoding="utf-8")

        monkeypatch.setattr(
            apply_mod,
            "inspect_backup",
            lambda _p: SimpleNamespace(
                manifest={
                    "version": "1",
                    "app_version": "0",
                    "created_at": "now",
                    "database": {"backup_mode": "sqlalchemy", "schema_version": "v"},
                    "uploads": {"included": False},
                }
            ),
        )
        monkeypatch.setattr(apply_mod, "_is_sqlite_url", lambda _u: False)
        monkeypatch.setattr(apply_mod, "_check_target_safe", lambda _u, _f: None)
        monkeypatch.setattr(apply_mod, "_dispose_global_engine", lambda: None)
        monkeypatch.setattr(apply_mod, "_restore_sqlalchemy_json", lambda _e, _u: RestoreLossReport())
        monkeypatch.setattr("cndb.core.migrations.stamp_head", lambda _u: None)

        report = apply_mod.restore_backup(bkdir, database_url="postgresql://x/y")
        assert report is not None

    def test_restore_command_unexpected_error(
        self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture[str]
    ) -> None:
        from cndb.cli.restore import apply as apply_mod

        def _boom(*_a: Any, **_k: Any) -> Any:
            raise RuntimeError("unexpected")

        monkeypatch.setattr(apply_mod, "restore_backup", _boom)
        args = argparse.Namespace(archive=str(tmp_path / "x.tar.gz"), dry_run=False, force=False, mode=None)
        with pytest.raises(SystemExit) as exc:
            apply_mod.restore_command(args)
        assert exc.value.code == 1
        assert "恢复失败" in capsys.readouterr().err


# ══════════════════════════════════════════════════════════════
# cli/users  （commands / io / db）
# ══════════════════════════════════════════════════════════════


@pytest.fixture
def cli_users_db(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    from sqlalchemy import create_engine

    import cndb.plugins.accounts.models
    import cndb.plugins.tables.models
    import cndb.plugins.workspaces.models  # noqa: F401
    from cndb.models.base import Base

    db_path = tmp_path / "cli_users_extra.db"
    monkeypatch.setattr(settings, "DATABASE_URL", f"sqlite:///{db_path}")
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    engine.dispose()
    return db_path


def _ns(**kwargs: Any) -> argparse.Namespace:
    ns = argparse.Namespace()
    for k, v in kwargs.items():
        setattr(ns, k, v)
    return ns


def _edit_ns(target: str, **kwargs: Any) -> argparse.Namespace:
    base = {
        "username": None,
        "email": None,
        "phone": None,
        "nickname": None,
        "role": None,
        "password": None,
        "activate": False,
        "deactivate": False,
        "is_superuser": None,
    }
    base.update(kwargs)
    base["target"] = target
    return _ns(**base)


class TestUsersCommandsEdges:
    def test_is_last_active_superuser_normal_user(self, cli_users_db: Path) -> None:
        from cndb.cli.users import _get_session, cmd_create
        from cndb.cli.users.commands import _is_last_active_superuser

        cmd_create(_ns(username="norm", password="pw1234", email=None, nickname=None, role="user", is_superuser=False))
        db = _get_session()
        try:
            from cndb.plugins.accounts.models import User

            u = db.query(User).filter(User.username == "norm").first()
            assert u is not None
            assert _is_last_active_superuser(db, u) is False
        finally:
            db.close()

    def test_edit_email_change_and_conflict(self, cli_users_db: Path) -> None:
        from cndb.cli.users import cmd_create, cmd_edit

        cmd_create(
            _ns(username="e_a", password="pw1234", email="a@x.com", nickname=None, role="user", is_superuser=False)
        )
        cmd_create(
            _ns(username="e_b", password="pw1234", email="b@x.com", nickname=None, role="user", is_superuser=False)
        )

        result = cmd_edit(_edit_ns("e_a", email="a2@x.com"))
        assert any("email" in c for c in result.changes)

        with pytest.raises(SystemExit, match="已被使用"):
            cmd_edit(_edit_ns("e_a", email="b@x.com"))

        # 空串 → 清空
        result2 = cmd_edit(_edit_ns("e_a", email=""))
        assert any("email" in c for c in result2.changes)

    def test_edit_phone_unchanged_and_role_unchanged(self, cli_users_db: Path) -> None:
        from cndb.cli.users import cmd_create, cmd_edit

        cmd_create(_ns(username="e_c", password="pw1234", email=None, nickname=None, role="user", is_superuser=False))
        # 先设置 phone
        cmd_edit(_edit_ns("e_c", phone="13900000000"))
        # 传相同 phone → 不产生变更（覆盖 phone 相等分支）
        result = cmd_edit(_edit_ns("e_c", phone="13900000000"))
        assert not any("phone" in c for c in result.changes)
        # 传相同 role → 不产生变更
        result2 = cmd_edit(_edit_ns("e_c", role="user"))
        assert not any("role" in c for c in result2.changes)

    def test_edit_promote_to_superuser(self, cli_users_db: Path) -> None:
        from cndb.cli.users import _get_session, cmd_create, cmd_edit
        from cndb.plugins.accounts.models import User

        cmd_create(_ns(username="e_p", password="pw1234", email=None, nickname=None, role="user", is_superuser=False))
        result = cmd_edit(_edit_ns("e_p", is_superuser=True))
        assert any("is_superuser" in c for c in result.changes)

        db = _get_session()
        try:
            u = db.query(User).filter(User.username == "e_p").first()
            assert u is not None
            assert u.is_superuser is True
        finally:
            db.close()

    def test_delete_interactive_confirm_yes(self, cli_users_db: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.cli.users import _get_session, cmd_create, cmd_delete
        from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

        cmd_create(_ns(username="own_y", password="pw1234", email=None, nickname=None, role="user", is_superuser=False))
        db = _get_session()
        try:
            from cndb.plugins.accounts.models import User

            u = db.query(User).filter(User.username == "own_y").first()
            ws = Workspace(name="ws_y")
            db.add(ws)
            db.flush()
            db.add(WorkspaceMember(workspace_id=ws.id, user_id=u.id, role=WorkspaceRole.OWNER))
            db.commit()
        finally:
            db.close()

        monkeypatch.setattr("builtins.input", lambda _prompt="": "y")
        result = cmd_delete(_ns(target="own_y", cascade=True, yes=False, check_only=False))
        assert result["deleted_workspace_count"] == 1

    def test_delete_interactive_confirm_no(self, cli_users_db: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.cli.users import _get_session, cmd_create, cmd_delete
        from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

        cmd_create(_ns(username="own_n", password="pw1234", email=None, nickname=None, role="user", is_superuser=False))
        db = _get_session()
        try:
            from cndb.plugins.accounts.models import User

            u = db.query(User).filter(User.username == "own_n").first()
            ws = Workspace(name="ws_n")
            db.add(ws)
            db.flush()
            db.add(WorkspaceMember(workspace_id=ws.id, user_id=u.id, role=WorkspaceRole.OWNER))
            db.commit()
        finally:
            db.close()

        monkeypatch.setattr("builtins.input", lambda _prompt="": "n")
        with pytest.raises(SystemExit) as exc:
            cmd_delete(_ns(target="own_n", cascade=True, yes=False, check_only=False))
        assert exc.value.code == 1


class TestUsersIoEdges:
    def test_read_xlsx_import_error(self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
        from cndb.cli.users.io import _read_xlsx

        monkeypatch.setitem(sys.modules, "openpyxl", None)
        with pytest.raises(RuntimeError, match="openpyxl 未安装"):
            _read_xlsx(tmp_path / "any.xlsx")

    def test_read_xlsx_active_none(self, monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
        import openpyxl

        from cndb.cli.users.io import _read_xlsx

        monkeypatch.setattr(openpyxl, "load_workbook", lambda *_a, **_k: SimpleNamespace(active=None))
        assert _read_xlsx(tmp_path / "x.xlsx") == []

    def test_read_xlsx_empty_sheet(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        from cndb.cli.users.io import _read_xlsx

        wb = Workbook()
        path = tmp_path / "empty.xlsx"
        wb.save(path)
        assert _read_xlsx(path) == []

    def test_read_xlsx_skips_blank_rows(self, tmp_path: Path) -> None:
        from openpyxl import Workbook

        from cndb.cli.users.io import _read_xlsx

        wb = Workbook()
        ws = wb.active
        ws.append(["username", "password"])
        ws.append([None, None])  # 空行 → continue
        ws.append(["u1", "pw1234"])
        path = tmp_path / "rows.xlsx"
        wb.save(path)
        rows = _read_xlsx(path)
        assert len(rows) == 1
        assert rows[0]["username"] == "u1"


class TestUsersDbEdges:
    def test_get_session_builds_schema(self, tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.cli.users.db import _get_session

        monkeypatch.setattr(settings, "DATABASE_URL", f"sqlite:///{tmp_path / 'fresh_users.db'}")
        db = _get_session()
        try:
            from sqlalchemy import inspect

            assert inspect(db.get_bind()).has_table("accounts_user")
        finally:
            db.close()

    def test_resolve_owned_and_member_counts(self, cli_users_db: Path) -> None:
        from cndb.cli.users.db import _get_session, _list_member_count, _list_owned_workspaces
        from cndb.plugins.accounts.models import User
        from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

        db = _get_session()
        try:
            u = User(username="db_u")
            u.set_password("pw1234")
            db.add(u)
            db.flush()
            ws = Workspace(name="db_ws")
            db.add(ws)
            db.flush()
            db.add(WorkspaceMember(workspace_id=ws.id, user_id=u.id, role=WorkspaceRole.OWNER))
            db.commit()
            owned = _list_owned_workspaces(db, u.id)
            assert len(owned) == 1
            assert _list_member_count(db, u.id) == 1
        finally:
            db.close()
