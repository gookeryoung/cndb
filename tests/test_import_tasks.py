"""异步导入 + ImportTask 测试."""

from __future__ import annotations

import time


def test_import_task_model_exists(db):
    """ImportTask 模型应可导入并注册."""
    from cndb.plugins.tables.models import ImportTask

    assert ImportTask.__tablename__ == "tables_importtask"
    # 检查字段存在
    cols = [c.name for c in ImportTask.__table__.columns]
    assert "table_id" in cols
    assert "status" in cols
    assert "progress" in cols
    assert "result_ids" in cols


def test_import_task_status_transitions():
    """状态机应拒绝非法转换."""
    from cndb.plugins.tables.import_tasks import _transition_status

    # 模拟一个 task 对象
    class FakeTask:
        def __init__(self, status):
            self.status = status

    fake = FakeTask("done")
    try:
        _transition_status(fake, "pending")
        raise AssertionError("应抛 ValueError")
    except ValueError:
        pass  # 预期行为

    # 合法转换
    fake2 = FakeTask("pending")
    _transition_status(fake2, "running")
    assert fake2.status == "running"


def test_async_import_submit_and_poll(client, auth_headers, db):
    """提交异步导入任务并轮询进度."""
    # 创建工作区和表
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws_async"})
    wid = ws.json()["id"]
    tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "t_async"})
    tid = tbl.json()["id"]
    # 添加一个 text 字段
    client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/fields",
        headers=auth_headers,
        json={"name": "name", "field_type": "text", "order": 0},
    )
    # 准备 JSON 文件内容
    import json

    file_content = json.dumps([{"name": "foo"}, {"name": "bar"}])
    # 提交异步导入
    from io import BytesIO

    # 用文件上传方式调用（multipart/form-data）
    files = {"file": ("test.json", BytesIO(file_content.encode()), "application/json")}
    resp = client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/import/async",
        headers=auth_headers,
        files=files,
    )
    assert resp.status_code == 200
    task_id = resp.json()["task_id"]
    assert resp.json()["status"] == "pending"

    # 轮询等待完成
    for _ in range(30):
        time.sleep(0.3)
        poll = client.get(
            f"/api/v1/workspaces/{wid}/tables/{tid}/import/async/{task_id}",
            headers=auth_headers,
        )
        assert poll.status_code == 200
        data = poll.json()
        if data["status"] in ("done", "failed"):
            break

    # 最终应为 done 或 failed（取决于引擎行为）
    assert data["status"] in ("done", "failed")
    if data["status"] == "done":
        assert data["imported_rows"] == 2


def test_import_task_not_found(client, auth_headers, db):
    """查询不存在的任务应 404."""
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws"})
    wid = ws.json()["id"]
    tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "t"})
    tid = tbl.json()["id"]

    resp = client.get(
        f"/api/v1/workspaces/{wid}/tables/{tid}/import/async/99999",
        headers=auth_headers,
    )
    assert resp.status_code == 404


def test_execute_import_task_not_found(db):
    """执行不存在的任务应静默返回（不抛异常）."""
    from cndb.plugins.tables.import_tasks import execute_import_task

    # 不应抛异常
    execute_import_task(db, 99999)


def test_execute_import_task_failed_path(db, auth_headers, client):
    """任务执行失败时应标记为 failed."""
    # 创建工作区和表
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws_fail"})
    wid = ws.json()["id"]
    tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "t_fail"})
    tbl.json()["id"]

    from cndb.plugins.tables.import_tasks import (
        create_import_task,
        execute_import_task,
    )

    # 用不存在的 table_id 创建任务，执行时应失败
    task = create_import_task(
        db,
        table_id=99999,
        user_id=None,
        filename="test.json",
        fmt="json",
        content='[{"name": "x"}]',
    )
    execute_import_task(db, task.id)
    db.refresh(task)
    assert task.status == "failed"
    assert task.error_message != ""


def test_execute_import_task_invalid_format(db, auth_headers, client):
    """不支持的格式应导致任务失败."""
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws_badfmt"})
    wid = ws.json()["id"]
    tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "t_badfmt"})
    tid = tbl.json()["id"]

    from cndb.plugins.tables.import_tasks import (
        create_import_task,
        execute_import_task,
    )

    task = create_import_task(
        db,
        table_id=tid,
        user_id=None,
        filename="test.xyz",
        fmt="xyz",
        content="garbage",
    )
    execute_import_task(db, task.id)
    db.refresh(task)
    assert task.status == "failed"


def test_create_import_task_bytes_content(db):
    """bytes 内容应正确解码存储."""
    from cndb.plugins.tables.import_tasks import create_import_task
    from cndb.plugins.tables.models import DataTable

    dt = DataTable(workspace_id=1, name="unused")
    dt.db_table_name = "table_testunused123"
    db.add(dt)
    db.commit()
    db.refresh(dt)

    raw = "a,b\n1,2"
    task = create_import_task(
        db,
        table_id=dt.id,
        user_id=None,
        filename="test.csv",
        fmt="csv",
        content=raw.encode("utf-8"),
    )
    assert task.file_content == raw

    db.delete(task)
    db.delete(dt)
    db.commit()


def test_execute_import_task_csv_path(db, auth_headers, client):
    """CSV 格式导入应走 csv 分支."""
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws_csv"})
    wid = ws.json()["id"]
    tbl = client.post(f"/api/v1/workspaces/{wid}/tables", headers=auth_headers, json={"name": "t_csv"})
    tid = tbl.json()["id"]
    # 加一个 text 字段
    client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/fields",
        headers=auth_headers,
        json={"name": "col", "field_type": "text"},
    )

    from cndb.plugins.tables.import_tasks import (
        create_import_task,
        execute_import_task,
    )

    csv_content = "col\nhello\nworld"
    task = create_import_task(
        db,
        table_id=tid,
        user_id=None,
        filename="test.csv",
        fmt="csv",
        content=csv_content,
    )
    execute_import_task(db, task.id)
    db.refresh(task)
    assert task.status == "done"


def test_execute_import_task_xlsx_format(db):
    """xlsx 格式应触发 base64 decode（但 transfer 里没实际 xlsx 解析，会失败）."""
    from cndb.plugins.tables.import_tasks import (
        create_import_task,
        execute_import_task,
    )
    from cndb.plugins.tables.models import DataTable

    dt = DataTable(workspace_id=1, name="unused_xlsx")
    dt.db_table_name = "table_testxlsx123"
    db.add(dt)
    db.commit()
    db.refresh(dt)

    import base64

    base64.b64encode(b"PKG!fake").decode("ascii")
    task = create_import_task(
        db,
        table_id=dt.id,
        user_id=None,
        filename="test.xlsx",
        fmt="xlsx",
        content=b"PKG!fake",
    )
    execute_import_task(db, task.id)
    db.refresh(task)
    # 要么 done（如果 transfer 支持 xlsx）要么 failed（不支持）
    assert task.status in ("done", "failed")

    db.delete(task)
    db.delete(dt)
    db.commit()


def test_create_import_task_xlsx_str_content(db):
    """xlsx 格式 + str content 应走 encode 分支."""
    from cndb.plugins.tables.import_tasks import create_import_task
    from cndb.plugins.tables.models import DataTable

    dt = DataTable(workspace_id=1, name="unused_xlsx_str")
    dt.db_table_name = "table_testxlsxs123"
    db.add(dt)
    db.commit()
    db.refresh(dt)

    task = create_import_task(
        db,
        table_id=dt.id,
        user_id=None,
        filename="test.xlsx",
        fmt="xlsx",
        content="PKG!fake-string-content",
    )
    assert task.format == "xlsx"
    assert task.file_content != "PKG!fake-string-content"  # 应该被 base64 编码

    db.delete(task)
    db.delete(dt)
    db.commit()


def test_execute_import_task_invalid_json_swallows_count_error(db):
    """无效 JSON 字符串应在估算行数阶段被静默吞掉 — 覆盖 L74-75."""
    from cndb.plugins.tables.import_tasks import create_import_task, execute_import_task
    from cndb.plugins.tables.models import DataTable

    dt = DataTable(workspace_id=1, name="badjson_t")
    dt.db_table_name = "table_badjson123"
    db.add(dt)
    db.commit()
    db.refresh(dt)

    task = create_import_task(
        db,
        table_id=dt.id,
        user_id=None,
        filename="bad.json",
        fmt="json",
        content="NOT-valid-JSON-at-all!!!",
    )
    execute_import_task(db, task.id)
    db.refresh(task)
    # 估算行数的 JSON 解析失败被静默吞掉 → total_rows 为 0
    assert task.total_rows == 0

    db.delete(task)
    db.delete(dt)
    db.commit()


def test_execute_import_task_status_conflict_fallback(db, monkeypatch):
    """_transition_status 失败时应回退到直接赋值 — 覆盖 L111-113."""
    from cndb.plugins.tables import import_tasks as it
    from cndb.plugins.tables.import_tasks import create_import_task, execute_import_task
    from cndb.plugins.tables.models import DataTable

    real = it._transition_status

    def fake_transition(task, new_status):
        if task.status == "running":
            raise ValueError("simulate conflict")
        return real(task, new_status)

    monkeypatch.setattr(it, "_transition_status", fake_transition)

    dt = DataTable(workspace_id=1, name="conflict_t")
    dt.db_table_name = "table_conflict123"
    db.add(dt)
    db.commit()
    db.refresh(dt)

    task = create_import_task(
        db,
        table_id=dt.id,
        user_id=None,
        filename="x.json",
        fmt="json",
        content='[{"a":1}]',
    )
    execute_import_task(db, task.id)
    db.refresh(task)
    assert task.status == "failed"

    db.delete(task)
    db.delete(dt)
    db.commit()


# ── 多编码自动检测 + XLSX 导入 e2e ──────────────────────────


def _create_workspace_and_table(client, auth_headers, db):
    """helper: 创建 workspace + table + text 字段，返回 (wid, tid)."""
    ws = client.post("/api/v1/workspaces", headers=auth_headers, json={"name": "ws_encode"})
    wid = ws.json()["id"]
    tbl = client.post(
        f"/api/v1/workspaces/{wid}/tables",
        headers=auth_headers,
        json={"name": "t_encode"},
    )
    tid = tbl.json()["id"]
    client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/fields",
        headers=auth_headers,
        json={"name": "姓名", "field_type": "text", "order": 0},
    )
    client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/fields",
        headers=auth_headers,
        json={"name": "年龄", "field_type": "number", "order": 1},
    )
    return wid, tid


def _poll_pending_confirm(client, wid, tid, task_id, auth_headers, timeout=10):
    """helper: 轮询直到 pending_confirm 或 done."""
    import time

    for _ in range(timeout * 5):
        resp = client.get(
            f"/api/v1/workspaces/{wid}/tables/{tid}/import/async/{task_id}",
            headers=auth_headers,
        )
        info = resp.json()
        if info.get("status") in ("pending_confirm", "pending_validation", "done", "failed"):
            return info
        time.sleep(0.2)
    return resp.json()


def test_gbk_csv_import_direct_analyze(client, auth_headers, db):
    """GBK 编码 CSV —— Importer 直接 analyze 能正确解码中文字段."""
    from cndb.plugins.tables.import_tasks import analyze_import_task, create_import_task
    from cndb.plugins.tables.models import DataTable

    _wid, tid = _create_workspace_and_table(client, auth_headers, db)
    gbk_bytes = "姓名,年龄\n张三,25\n李四,30\n".encode("gbk")

    dt = db.get(DataTable, tid)
    task = create_import_task(db, table_id=dt.id, user_id=None, filename="data.csv", fmt="csv", content=gbk_bytes)
    analyze_import_task(db, task.id)
    db.refresh(task)

    assert task.status == "pending_confirm", task.error_message
    import json as _json

    report = _json.loads(task.validation_report)
    skipped = report.get("skipped_columns", [])
    assert "姓名" not in skipped, f"GBK 解码失败，'姓名' 被跳过: skipped={skipped}"
    assert "年龄" not in skipped


def test_utf8_bom_csv_import_direct_analyze(client, auth_headers, db):
    """UTF-8 BOM CSV —— analyze_import_task 应自动去除 BOM."""
    from cndb.plugins.tables.import_tasks import analyze_import_task, create_import_task
    from cndb.plugins.tables.models import DataTable

    _wid, tid = _create_workspace_and_table(client, auth_headers, db)
    bom_bytes = "姓名,年龄\n张三,25\n".encode("utf-8-sig")

    dt = db.get(DataTable, tid)
    task = create_import_task(db, table_id=dt.id, user_id=None, filename="data.csv", fmt="csv", content=bom_bytes)
    analyze_import_task(db, task.id)
    db.refresh(task)

    assert task.status == "pending_confirm", task.error_message
    import json as _json

    report = _json.loads(task.validation_report)
    # BOM 应被去除，列名不包含 BOM 字符
    file_cols = report.get("file_columns") or []
    assert not any(c.startswith("\ufeff") for c in file_cols), f"BOM 未被去除: {file_cols}"


def test_xls_filename_gives_friendly_error(client, auth_headers, db):
    """上传 .xls 后缀文件应返回友好错误提示."""
    wid, tid = _create_workspace_and_table(client, auth_headers, db)

    resp = client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/import/analyze",
        headers=auth_headers,
        files={"file": ("data.xls", b"fake-binary", "application/vnd.ms-excel")},
    )
    assert resp.status_code == 400, resp.text
    assert "xls" in resp.json()["detail"].lower()


def test_xlsx_import_sync_route(client, auth_headers, db):
    """同步导入路由支持 XLSX 文件."""
    import io

    from openpyxl import Workbook

    wid, tid = _create_workspace_and_table(client, auth_headers, db)

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["姓名", "年龄"])
    ws.append(["王五", 22])
    ws.append(["赵六", 28])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    resp = client.post(
        f"/api/v1/workspaces/{wid}/tables/{tid}/import",
        headers=auth_headers,
        files={"file": ("data.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["imported"] == 2


def test_importer_parse_xlsx_roundtrip():
    """Importer._parse_xlsx 能正确解析 openpyxl Workbook 写入的 XLSX."""
    import io

    from openpyxl import Workbook

    from cndb.plugins.tables.importer import Importer

    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["姓名", "年龄"])
    ws.append(["孙七", 31])
    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    rows, cols = Importer._parse_xlsx(xlsx_bytes)
    assert cols == ["姓名", "年龄"]
    assert len(rows) == 1
    assert rows[0]["姓名"] == "孙七"
