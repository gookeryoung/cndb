"""bulk + transfer 模块测试 —— 纯函数可独立测，路由走 API."""

import json

import pytest

from cndb.plugins.accounts.models import User
from cndb.plugins.tables import ddl, transfer
from cndb.plugins.tables.models import DataField, DataTable, DataView
from cndb.plugins.workspaces.models import Workspace, WorkspaceRole

# ── transfer 纯函数测试 ──────────────────────────────


class TestTransferPure:
    def test_export_json(self):
        rows = [{"姓名": "A", "年龄": 1}, {"姓名": "B", "年龄": 2}]
        text = transfer.export_rows_to_json(rows)
        assert json.loads(text) == rows

    def test_export_json_empty(self):
        assert transfer.export_rows_to_json([]) == "[]"

    def test_export_csv(self):
        rows = [{"姓名": "A", "年龄": 1}]
        csv_text = transfer.export_rows_to_csv(rows)
        assert "姓名,年龄" in csv_text
        assert "A,1" in csv_text

    def test_export_csv_empty(self):
        assert transfer.export_rows_to_csv([]) == ""

    def test_export_xlsx(self):
        rows = [{"姓名": "A", "年龄": 1}]
        data = transfer.export_rows_to_xlsx(rows)
        assert isinstance(data, bytes)
        assert len(data) > 0

    def test_export_xlsx_empty(self):
        data = transfer.export_rows_to_xlsx([])
        assert isinstance(data, bytes)

    def test_guess_format(self):
        assert transfer.guess_format_from_filename("a.xlsx") == "xlsx"
        assert transfer.guess_format_from_filename("a.CSV") == "csv"
        assert transfer.guess_format_from_filename("a.json") == "json"
        with pytest.raises(ValueError):
            transfer.guess_format_from_filename("a.txt")


# ── bulk router API 测试 ─────────────────────────────


@pytest.fixture
def owner(db):
    u = User(username="bulk_owner", nickname="Owner")
    u.set_password("passw0rd")
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


@pytest.fixture
def auth_owner(client, owner):
    r = client.post("/api/v1/accounts/auth/login", json={"login": owner.username, "password": "passw0rd"})
    assert r.status_code == 200
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture
def ws(db, owner):
    from cndb.plugins.workspaces.models import WorkspaceMember

    w = Workspace(name="BWS", created_by_id=owner.id)
    db.add(w)
    db.flush()
    db.add(WorkspaceMember(workspace_id=w.id, user_id=owner.id, role=WorkspaceRole.OWNER))
    db.commit()
    db.refresh(w)
    return w


@pytest.fixture
def table(db, ws):
    engine = db.get_bind()
    dt = DataTable(workspace_id=ws.id, name="BulkTable")
    dt.ensure_db_name()
    db.add(dt)
    db.flush()
    df = DataField(table_id=dt.id, name="姓名", field_type="text", order=0)
    df.ensure_db_name()
    db.add(df)
    db.flush()
    df2 = DataField(table_id=dt.id, name="年龄", field_type="number", order=1)
    df2.ensure_db_name()
    db.add(df2)
    db.commit()
    db.refresh(dt)
    ddl.create_table(engine, dt)
    # 加几条数据
    from cndb.plugins.tables import records as rec

    rec.create_row(engine, dt, {"姓名": "张三", "年龄": 20})
    rec.create_row(engine, dt, {"姓名": "李四", "年龄": 30})
    rec.create_row(engine, dt, {"姓名": "王五", "年龄": 40})
    return dt


class TestBulkAPI:
    def test_bulk_create(self, client, ws, table, auth_owner):
        """覆盖 POST /records/bulk-create —— 批量新建行."""
        # 先用空 rows 测试 400
        r0 = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/records/bulk-create",
            json={"rows": []},
            headers=auth_owner,
        )
        assert r0.status_code == 400
        # 正常批量创建（values 包装格式）
        r = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/records/bulk-create",
            json={
                "rows": [
                    {"values": {"姓名": "批量甲", "年龄": 25}},
                    {"values": {"姓名": "批量乙", "年龄": 28}},
                ],
            },
            headers=auth_owner,
        )
        assert r.status_code == 201
        assert r.json()["created"] == 2
        assert len(r.json()["ids"]) == 2

    def test_bulk_delete(self, client, ws, table, auth_owner):
        r = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/records/bulk-delete",
            json={"row_ids": [1, 2]},
            headers=auth_owner,
        )
        assert r.status_code == 200
        assert r.json()["deleted"] == 2

    def test_bulk_update(self, client, ws, table, auth_owner):
        r = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/records/bulk-update",
            json={"row_ids": [1, 2], "values": {"年龄": 99}},
            headers=auth_owner,
        )
        assert r.status_code == 200
        assert r.json()["updated"] == 2

    def test_bulk_update_empty(self, client, ws, table, auth_owner):
        r = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/records/bulk-update",
            json={"row_ids": [], "values": {"年龄": 1}},
            headers=auth_owner,
        )
        assert r.status_code == 400

    def test_export_json(self, client, ws, table, auth_owner):
        r = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=json",
            headers=auth_owner,
        )
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/json")
        data = json.loads(r.text)
        assert len(data) == 3

    def test_export_csv(self, client, ws, table, auth_owner):
        r = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=csv",
            headers=auth_owner,
        )
        assert r.status_code == 200
        assert "姓名" in r.text
        assert "张三" in r.text

    def test_export_invalid_format(self, client, ws, table, auth_owner):
        r = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=pdf",
            headers=auth_owner,
        )
        assert r.status_code == 400

    def test_import_json(self, client, ws, table, auth_owner):
        import io

        data = json.dumps([{"姓名": "导入1", "年龄": 100}, {"姓名": "导入2", "年龄": 200}])
        r = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/import",
            files={"file": ("import.json", io.BytesIO(data.encode()), "application/json")},
            headers=auth_owner,
        )
        assert r.status_code == 200
        assert r.json()["imported"] == 2

    def test_import_invalid_filename(self, client, ws, table, auth_owner):
        import io

        r = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/import",
            files={"file": ("test.txt", io.BytesIO(b"x"), "text/plain")},
            headers=auth_owner,
        )
        # guess_format 抛 ValueError 被 except 捕获 → 400
        assert r.status_code == 400

    def test_bulk_update_values_empty(self, client, ws, table, auth_owner):
        r = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/records/bulk-update",
            json={"row_ids": [1, 2], "values": {}},
            headers=auth_owner,
        )
        assert r.status_code == 400

    def test_export_xlsx(self, client, ws, table, auth_owner):
        r = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=xlsx",
            headers=auth_owner,
        )
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
        assert len(r.content) > 0

    def test_import_csv(self, client, ws, table, auth_owner):
        import io

        csv_content = "姓名,年龄\n赵六,60\n钱七,70\n"
        r = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/import",
            files={"file": ("import.csv", io.BytesIO(csv_content.encode()), "text/csv")},
            headers=auth_owner,
        )
        assert r.status_code == 200
        assert r.json()["imported"] == 2

    def test_import_xlsx(self, client, ws, table, auth_owner):
        import io

        from openpyxl import Workbook

        wb = Workbook()
        ws_wb = wb.active
        assert ws_wb is not None
        ws_wb.append(["姓名", "年龄"])
        ws_wb.append(["孙八", 80])
        ws_wb.append(["周九", 90])
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        r = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/import",
            files={"file": ("import.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=auth_owner,
        )
        assert r.status_code == 200
        assert r.json()["imported"] == 2

    def test_import_invalid_json_not_array(self, client, ws, table, auth_owner):
        import io

        r = client.post(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/import",
            files={"file": ("bad.json", io.BytesIO(b'"just a string"'), "application/json")},
            headers=auth_owner,
        )
        assert r.status_code == 400

    def test_export_with_view_id(self, client, ws, table, db, auth_owner):
        """带 view_id 参数时，应按视图 filters 筛选导出."""
        dv = DataView(
            table_id=table.id,
            owner_id=None,
            name="只看年长",
            view_type="grid",
            filter_type="AND",
            filters=[{"field_name": "年龄", "op": ">", "value": 25}],
            sortings=[],
            field_options={},
            field_order=[],
            view_options={},
            is_default=False,
            order=0,
        )
        db.add(dv)
        db.commit()
        db.refresh(dv)

        r = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=json&view_id={dv.id}",
            headers=auth_owner,
        )
        assert r.status_code == 200
        data = json.loads(r.text)
        # 原数据：张三 20, 李四 30, 王五 40 —— 筛选 > 25 后只剩 2 条
        assert len(data) == 2
        names = {row["姓名"] for row in data}
        assert names == {"李四", "王五"}

    def test_export_with_view_id_sorts(self, client, ws, table, db, auth_owner):
        """带 sortings 的视图导出应尊重排序."""
        dv = DataView(
            table_id=table.id,
            owner_id=None,
            name="按年龄降序",
            view_type="grid",
            filter_type="AND",
            filters=[],
            sortings=[{"field_name": "年龄", "direction": "desc"}],
            field_options={},
            field_order=[],
            view_options={},
            is_default=False,
            order=0,
        )
        db.add(dv)
        db.commit()
        db.refresh(dv)

        r = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=csv&view_id={dv.id}",
            headers=auth_owner,
        )
        assert r.status_code == 200
        # CSV 含 id 列，换行符可能是 \r\n
        lines = [line.strip("\r") for line in r.text.strip().split("\n") if line.strip()]
        # 第一行表头（id,姓名,年龄），然后应按王五 40 → 李四 30 → 张三 20 降序
        assert "姓名" in lines[0] and "年龄" in lines[0] and "id" in lines[0]
        # 找到姓名列和年龄列的索引
        headers = lines[0].split(",")
        name_idx = headers.index("姓名")
        age_idx = headers.index("年龄")
        # 验证顺序：40 → 30 → 20
        assert lines[1].split(",")[age_idx] == "40"
        assert lines[2].split(",")[age_idx] == "30"
        assert lines[3].split(",")[age_idx] == "20"
        assert lines[1].split(",")[name_idx] == "王五"
        assert lines[2].split(",")[name_idx] == "李四"
        assert lines[3].split(",")[name_idx] == "张三"

    def test_export_with_view_id_not_found(self, client, ws, table, db, auth_owner):
        """view_id 不存在或不属于当前表 → 404."""
        # 不存在的 view_id
        r = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=json&view_id=99999",
            headers=auth_owner,
        )
        assert r.status_code == 404

        # 属于其他表的 view_id
        other_table = DataTable(workspace_id=ws.id, name="OtherTable")
        other_table.ensure_db_name()
        db.add(other_table)
        db.commit()
        db.refresh(other_table)
        dv2 = DataView(
            table_id=other_table.id,
            owner_id=None,
            name="其他表的视图",
            view_type="grid",
            filters=[],
            sortings=[],
            field_options={},
            field_order=[],
            view_options={},
            is_default=False,
            order=0,
        )
        db.add(dv2)
        db.commit()

        r2 = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=json&view_id={dv2.id}",
            headers=auth_owner,
        )
        assert r2.status_code == 404

    def test_export_with_view_id_or_logic(self, client, ws, table, db, auth_owner):
        """filter_type=OR 时，满足任一条件的行都应被导出."""
        # 追加几条带不同年龄值的数据，用于 OR 条件筛选
        from cndb.plugins.tables import records as rec

        engine = db.get_bind()
        rec.create_row(engine, table, {"姓名": "赵六", "年龄": 50})
        rec.create_row(engine, table, {"姓名": "钱七", "年龄": 5})
        # 总数据：张三 20, 李四 30, 王五 40, 赵六 50, 钱七 5

        dv = DataView(
            table_id=table.id,
            owner_id=None,
            name="OR筛选: 年龄>40 或 年龄<10",
            view_type="grid",
            filter_type="OR",
            filters=[
                {"field_name": "年龄", "op": ">", "value": 40},
                {"field_name": "年龄", "op": "<", "value": 10},
            ],
            sortings=[{"field_name": "年龄", "direction": "asc"}],
            field_options={},
            field_order=[],
            view_options={},
            is_default=False,
            order=0,
        )
        db.add(dv)
        db.commit()
        db.refresh(dv)

        r = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=json&view_id={dv.id}",
            headers=auth_owner,
        )
        assert r.status_code == 200
        data = json.loads(r.text)
        # 应包含：赵六 50（>40）+ 钱七 5（<10），两条
        assert len(data) == 2
        names = {row["姓名"] for row in data}
        assert names == {"赵六", "钱七"}
        # 排序：升序 → 钱七在前，赵六在后
        assert data[0]["姓名"] == "钱七"
        assert data[1]["姓名"] == "赵六"

    def test_export_with_view_id_xlsx(self, client, ws, table, db, auth_owner):
        """view_id + xlsx 格式组合 —— 覆盖 format 分支."""
        dv = DataView(
            table_id=table.id,
            owner_id=None,
            name="年长XLSX",
            view_type="grid",
            filter_type="AND",
            filters=[{"field_name": "年龄", "op": ">=", "value": 30}],
            sortings=[],
            field_options={},
            field_order=[],
            view_options={},
            is_default=False,
            order=0,
        )
        db.add(dv)
        db.commit()
        db.refresh(dv)

        r = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=xlsx&view_id={dv.id}",
            headers=auth_owner,
        )
        assert r.status_code == 200
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
        # XLSX 至少有文件头（PK 签名）
        assert r.content[:2] == b"PK"

        # 用 openpyxl 回读验证数据是否正确
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(r.content))
        ws_wb = wb.active
        assert ws_wb is not None
        rows = list(ws_wb.iter_rows(values_only=True))
        # 表头 + 2 行（李四 30、王五 40）
        assert len(rows) == 3
        header = rows[0]
        assert "姓名" in header
        assert "年龄" in header
        name_col = header.index("姓名")
        names = {row[name_col] for row in rows[1:]}
        assert names == {"李四", "王五"}

    def test_export_with_view_id_no_filters(self, client, ws, table, db, auth_owner):
        """视图存在但 filters/sortings 都为空 —— 导出应等价全表."""
        dv = DataView(
            table_id=table.id,
            owner_id=None,
            name="空筛选视图",
            view_type="grid",
            filter_type="AND",
            filters=[],
            sortings=[],
            field_options={},
            field_order=[],
            view_options={},
            is_default=False,
            order=0,
        )
        db.add(dv)
        db.commit()
        db.refresh(dv)

        # 带 view_id 但视图无筛选
        r_with = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=json&view_id={dv.id}",
            headers=auth_owner,
        )
        # 不带 view_id（全表）
        r_without = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=json",
            headers=auth_owner,
        )
        assert r_with.status_code == 200
        assert r_without.status_code == 200

        data_with = json.loads(r_with.text)
        data_without = json.loads(r_without.text)
        assert len(data_with) == len(data_without) == 3
        # 内容一致
        names_with = {row["姓名"] for row in data_with}
        names_without = {row["姓名"] for row in data_without}
        assert names_with == names_without

    def test_export_excludes_trashed_rows(self, client, ws, table, db, auth_owner):
        """软删行（_trashed=True）不应出现在导出结果中."""
        import sqlalchemy

        engine = db.get_bind()
        # 用 DataTable.db_table_name 直接引用动态表
        sa_table = sqlalchemy.table(table.db_table_name, sqlalchemy.column("id"), sqlalchemy.column("_trashed"))

        with engine.begin() as conn:
            conn.execute(sa_table.update().where(sa_table.c.id == 1).values(_trashed=True))

        # 导出应只剩李四、王五两条
        r = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=json",
            headers=auth_owner,
        )
        assert r.status_code == 200
        data = json.loads(r.text)
        assert len(data) == 2
        names = {row["姓名"] for row in data}
        assert names == {"李四", "王五"}

        # 即使带一个匹配软删行的视图筛选，也不应把软删行导出来
        dv = DataView(
            table_id=table.id,
            owner_id=None,
            name="查找已软删的张三",
            view_type="grid",
            filter_type="AND",
            filters=[{"field_name": "姓名", "op": "=", "value": "张三"}],
            sortings=[],
            field_options={},
            field_order=[],
            view_options={},
            is_default=False,
            order=0,
        )
        db.add(dv)
        db.commit()
        db.refresh(dv)

        r2 = client.get(
            f"/api/v1/workspaces/{ws.id}/tables/{table.id}/export?format=json&view_id={dv.id}",
            headers=auth_owner,
        )
        # 软删条件 AND 业务条件（姓名=张三）最终不命中任何行
        data2 = json.loads(r2.text)
        assert len(data2) == 0


__all__ = []
