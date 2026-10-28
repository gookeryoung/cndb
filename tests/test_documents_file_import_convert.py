"""通用文件导入 —— 转换管线单元测试（P0）.

覆盖：空白规整 / 标题清洗 / 解码兜底 / md 标题识别 / docx 与 html 转换 /
读取器批次校验与逐条隔离。对应需求 C4-C8 与 ADR 0010 P0。
"""

from __future__ import annotations

import base64
import io
import struct
import xml.etree.ElementTree as ET
import zlib

import pytest
from docx import Document as DocxDocument
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls

from cndb.core.errors import BadRequestError
from cndb.plugins.documents.services.importers import docx as docx_converter
from cndb.plugins.documents.services.importers import html as html_converter
from cndb.plugins.documents.services.importers import markdown as markdown_converter
from cndb.plugins.documents.services.importers import mhtml as mhtml_converter
from cndb.plugins.documents.services.importers import normalize as normalize_module
from cndb.plugins.documents.services.importers import reader
from cndb.plugins.documents.services.importers.errors import UnsupportedImport


def _png_bytes() -> bytes:
    """构造一张合法的 1×1 RGBA PNG（手工拼块，避免额外图片依赖）."""

    def chunk(kind: bytes, payload: bytes) -> bytes:
        head = struct.pack(">I", len(payload)) + kind + payload
        return head + struct.pack(">I", zlib.crc32(kind + payload) & 0xFFFFFFFF)

    signature = b"\x89PNG\r\n\x1a\n"
    ihdr = struct.pack(">IIBBBBB", 1, 1, 8, 6, 0, 0, 0)
    idat = zlib.compress(b"\x00\x00\x00\x00\x00")
    return signature + chunk(b"IHDR", ihdr) + chunk(b"IDAT", idat) + chunk(b"IEND", b"")


#: 1×1 PNG（供 python-docx add_picture 使用）
PNG_BYTES = _png_bytes()


def _make_docx(build) -> bytes:
    document = DocxDocument()
    build(document)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# ── normalize ──


class TestNormalize:
    def test_trailing_whitespace_removed(self):
        assert normalize_module.normalize_markdown("hello   \nworld\t\n") == "hello\nworld"

    def test_inline_spaces_collapsed(self):
        assert normalize_module.normalize_markdown("a    b\t c") == "a b c"

    def test_blank_lines_collapsed(self):
        assert normalize_module.normalize_markdown("a\n\n\n\nb") == "a\n\nb"

    def test_leading_blank_lines_dropped(self):
        assert normalize_module.normalize_markdown("\n\n\na") == "a"

    def test_trailing_blank_lines_dropped(self):
        assert normalize_module.normalize_markdown("a\n\n\n") == "a"

    def test_invalid_indent_removed_but_list_kept(self):
        assert normalize_module.normalize_markdown("  plain\n  - item\n    code") == "plain\n  - item\n    code"

    def test_quote_indent_kept(self):
        assert normalize_module.normalize_markdown("  > 引用") == "  > 引用"

    def test_inline_code_spaces_preserved(self):
        assert normalize_module.normalize_markdown("a  `b   c`  d") == "a `b   c` d"

    def test_fence_content_preserved(self):
        text = "```\na  b\n\n\n```\n"
        assert normalize_module.normalize_markdown(text) == "```\na  b\n\n\n```"

    def test_table_row_preserved(self):
        text = "| a |  b |\n|---|---|\n| 1 |  2 |"
        assert normalize_module.normalize_markdown(text) == text

    def test_fence_after_blank_line(self):
        assert normalize_module.normalize_markdown("a\n\n```\ncode\n```") == "a\n\n```\ncode\n```"

    def test_unterminated_fence_trailing_blank_trimmed(self):
        assert normalize_module.normalize_markdown("```\na\n\n") == "```\na"

    def test_front_matter_preserved_and_body_normalized(self):
        text = "---\ntitle: x  \n---\n\n# H\n\n\nbody  \n"
        assert normalize_module.normalize_markdown(text) == "---\ntitle: x  \n---\n\n# H\n\nbody"

    def test_empty_and_whitespace_only(self):
        assert normalize_module.normalize_markdown("") == ""
        assert normalize_module.normalize_markdown("   \n\t\n") == ""

    def test_crlf_normalized(self):
        assert normalize_module.normalize_markdown("a\r\nb\r\n") == "a\nb"

    def test_unterminated_fence_kept(self):
        assert normalize_module.normalize_markdown("```\na  b") == "```\na  b"

    def test_clean_title(self):
        assert normalize_module.clean_title("  标题\x00  甲 ") == "标题 甲"
        assert normalize_module.clean_title("") == "未命名"
        assert normalize_module.clean_title("x" * 600).__len__() == 500

    def test_decode_text_bom_and_gb18030(self):
        assert normalize_module.decode_text(b"\xef\xbb\xbfhi") == "hi"
        assert normalize_module.decode_text("中文".encode("gb18030")) == "中文"

    def test_decode_text_chardet_high_confidence(self, monkeypatch):
        import chardet

        monkeypatch.setattr(chardet, "detect", lambda data: {"encoding": "cp1252", "confidence": 0.9})
        assert normalize_module.decode_text(b"\xe0\xff") == "àÿ"

    def test_decode_text_chardet_low_confidence(self, monkeypatch):
        import chardet

        monkeypatch.setattr(chardet, "detect", lambda data: {"encoding": "cp1252", "confidence": 0.1})
        assert normalize_module.decode_text(b"\xe0\xff") == "àÿ"

    def test_decode_text_ascii_detection_skipped(self, monkeypatch):
        import chardet

        monkeypatch.setattr(chardet, "detect", lambda data: {"encoding": "ascii", "confidence": 0.99})
        assert normalize_module.decode_text(b"\xe0\xff") == "àÿ"

    def test_decode_text_unknown_codec_falls_back(self, monkeypatch):
        import chardet

        monkeypatch.setattr(chardet, "detect", lambda data: {"encoding": "no-such-codec", "confidence": 0.9})
        assert normalize_module.decode_text(b"\xe0\xff") == "àÿ"


# ── markdown 转换 ──


class TestMarkdownConvert:
    def test_front_matter_title(self):
        title, body, _ = markdown_converter.convert_markdown(
            b"---\ntitle: \xe6\xa0\x87\xe9\xa2\x98\n---\n\n\xe6\xad\xa3\xe6\x96\x87\n", "a.md"
        )
        assert title == "标题"
        assert body.strip() == "正文"

    def test_h1_title_stripped(self):
        title, body, warnings = markdown_converter.convert_markdown(
            b"# \xe6\xa0\x87\xe9\xa2\x98\n\n\xe6\xad\xa3\xe6\x96\x87\n", "a.md"
        )
        assert title == "标题"
        assert body.strip() == "正文"
        assert warnings

    def test_filename_fallback(self):
        title, body, warnings = markdown_converter.convert_markdown(
            b"\xe5\x8f\xaa\xe6\x9c\x89\xe6\xad\xa3\xe6\x96\x87", "notes/笔记.md"
        )
        assert title == "笔记"
        assert body == "只有正文"
        assert warnings

    def test_leading_blank_then_h1(self):
        title, _, warnings = markdown_converter.convert_markdown(
            b"\n\n# \xe6\xa0\x87\xe9\xa2\x98\n\n\xe6\xad\xa3\xe6\x96\x87", "a.md"
        )
        assert title == "标题"
        assert warnings

    def test_blank_body_falls_back_to_filename(self):
        title, body, warnings = markdown_converter.convert_markdown(b"\n\n\n", "a.md")
        assert title == "a"
        assert body.strip() == ""
        assert warnings


# ── docx 转换 ──


class TestDocxConvert:
    def test_heading_title_stripped(self):
        data = _make_docx(lambda d: (d.add_heading("我的标题", level=1), d.add_paragraph("正文")))
        title, body, warnings = docx_converter.convert_docx(data, "x.docx")
        assert title == "我的标题"
        assert body == "正文"
        assert warnings

    def test_core_property_title_keeps_heading(self):
        def build(document):
            document.core_properties.title = "核心标题"
            document.add_heading("章节", level=1)
            document.add_paragraph("正文")

        title, body, _ = docx_converter.convert_docx(_make_docx(build), "x.docx")
        assert title == "核心标题"
        assert "# 章节" in body

    def test_heading_levels_and_formatting(self):
        def build(document):
            document.add_heading("二级", level=2)
            paragraph = document.add_paragraph()
            run = paragraph.add_run("粗体")
            run.bold = True
            paragraph.add_run("与")
            italic = paragraph.add_run("斜体")
            italic.italic = True

        _, body, _ = docx_converter.convert_docx(_make_docx(build), "x.docx")
        assert "## 二级" in body
        assert "**粗体**与*斜体*" in body

    def test_lists(self):
        def build(document):
            document.add_paragraph("项一", style="List Bullet")
            document.add_paragraph("项二", style="List Number")

        _, body, _ = docx_converter.convert_docx(_make_docx(build), "x.docx")
        assert "- 项一" in body
        assert "1. 项二" in body

    def test_table(self):
        def build(document):
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "甲"
            table.cell(0, 1).text = "乙"
            table.cell(1, 0).text = "1"
            table.cell(1, 1).text = "2"

        _, body, _ = docx_converter.convert_docx(_make_docx(build), "x.docx")
        assert "| 甲 | 乙 |" in body
        assert "| --- | --- |" in body
        assert "| 1 | 2 |" in body

    def test_image_with_sink(self):
        def build(document):
            document.add_paragraph("前")
            document.add_picture(io.BytesIO(PNG_BYTES))
            document.add_paragraph("后")

        calls: list[tuple[bytes, str]] = []

        def sink(data: bytes, name: str) -> str:
            calls.append((data, name))
            return "KEY123"

        _, body, _ = docx_converter.convert_docx(_make_docx(build), "x.docx", image_sink=sink)
        assert "file:KEY123" in body
        assert calls and calls[0][1] == "image1.png"

    def test_image_without_sink_is_placeholder(self):
        def build(document):
            document.add_paragraph("前")
            document.add_picture(io.BytesIO(PNG_BYTES))

        _, body, warnings = docx_converter.convert_docx(_make_docx(build), "x.docx")
        assert "图片未导入" in body
        assert any("未导入" in message for message in warnings)

    def test_image_sink_returning_none_is_placeholder(self):
        def build(document):
            document.add_picture(io.BytesIO(PNG_BYTES))

        _, body, warnings = docx_converter.convert_docx(_make_docx(build), "x.docx", image_sink=lambda d, n: None)
        assert "图片未导入" in body
        assert warnings

    def test_filename_fallback_when_no_heading(self):
        data = _make_docx(lambda d: d.add_paragraph("只有正文"))
        title, _, warnings = docx_converter.convert_docx(data, "报告.docx")
        assert title == "报告"
        assert warnings

    def test_corrupt_file_raises(self):
        with pytest.raises(UnsupportedImport):
            docx_converter.convert_docx(b"not-a-docx", "x.docx")

    def test_bold_italic(self):
        def build(document):
            run = document.add_paragraph().add_run("双")
            run.bold = True
            run.italic = True

        _, body, _ = docx_converter.convert_docx(_make_docx(build), "x.docx")
        assert "***双***" in body

    def test_title_style_becomes_title(self):
        title, body, _ = docx_converter.convert_docx(
            _make_docx(lambda d: d.add_paragraph("标题样式", style="Title")), "x.docx"
        )
        assert title == "标题样式"
        assert body == ""

    def test_empty_paragraph_skipped(self):
        data = _make_docx(lambda d: (d.add_paragraph(""), d.add_paragraph("正文")))
        _, body, _ = docx_converter.convert_docx(data, "x.docx")
        assert body == "正文"

    def test_hyperlink_fallback_text(self):
        def build(document):
            paragraph = document.add_paragraph()
            hyperlink = OxmlElement("w:hyperlink")
            run = OxmlElement("w:r")
            text = OxmlElement("w:t")
            text.text = "链接文字"
            run.append(text)
            hyperlink.append(run)
            paragraph._p.append(hyperlink)

        _, body, _ = docx_converter.convert_docx(_make_docx(build), "x.docx")
        assert "链接文字" in body

    def test_empty_table_block_skipped(self):
        def build(document):
            document.add_paragraph("正文")
            document.element.body.append(parse_xml(f"<w:tbl {nsdecls('w')}/>"))

        _, body, _ = docx_converter.convert_docx(_make_docx(build), "x.docx")
        assert body == "正文"

    def test_blip_without_relation_is_skipped(self):
        def build(document):
            drawing = (
                f"<w:drawing {nsdecls('w', 'wp', 'a', 'pic', 'r')}><wp:inline><a:graphic><a:graphicData>"
                "<pic:pic><pic:blipFill><a:blip/></pic:blipFill></pic:pic></a:graphicData></a:graphic>"
                "</wp:inline></w:drawing>"
            )
            document.add_paragraph().add_run()._element.append(parse_xml(drawing))

        _, body, _ = docx_converter.convert_docx(_make_docx(build), "x.docx")
        assert body.strip() == ""

    def test_blip_with_dangling_relation_is_skipped(self):
        def build(document):
            drawing = (
                f"<w:drawing {nsdecls('w', 'wp', 'a', 'pic', 'r')}><wp:inline><a:graphic><a:graphicData>"
                '<pic:pic><pic:blipFill><a:blip r:embed="rId9999"/></pic:blipFill></pic:pic>'
                "</a:graphicData></a:graphic></wp:inline></w:drawing>"
            )
            document.add_paragraph().add_run()._element.append(parse_xml(drawing))

        _, body, _ = docx_converter.convert_docx(_make_docx(build), "x.docx")
        assert body.strip() == ""


# ── html 转换 ──


class TestHtmlConvert:
    def test_title_tag_wins(self):
        text = "<html><head><title>标题甲</title></head><body><h1>忽略我</h1><p>正文</p></body></html>"
        title, body, _ = html_converter.convert_html(text.encode(), "x.html")
        assert title == "标题甲"
        assert body == "# 忽略我\n\n正文"

    def test_h1_as_title_stripped(self):
        title, body, warnings = html_converter.convert_html(
            "<body><h1>我的标题</h1><p>正文</p></body>".encode(), "x.html"
        )
        assert title == "我的标题"
        assert body == "正文"
        assert warnings

    def test_filename_fallback(self):
        title, body, warnings = html_converter.convert_html(b"<p>\xe6\xad\xa3\xe6\x96\x87</p>", "page.html")
        assert title == "page"
        assert body == "正文"
        assert warnings

    def test_headings_paragraph_inline(self):
        html = (
            "<body><h3>小节</h3><p>Hello <strong>粗</strong> 与 <em>斜</em> 用 <code>pip</code></p>"
            "<p>见 <a href='https://e.x'>链接</a></p></body>"
        )
        _, body, _ = html_converter.convert_html(html.encode("utf-8"), "x.html")
        assert "### 小节" in body
        assert "Hello **粗** 与 *斜* 用 `pip`" in body
        assert "[链接](https://e.x)" in body

    def test_lists_with_nesting(self):
        html = "<body><ul><li>甲</li><li>乙<ul><li>子</li></ul></li></ul><ol><li>一</li><li>二</li></ol></body>"
        _, body, _ = html_converter.convert_html(html.encode("utf-8"), "x.html")
        assert "- 甲" in body
        assert "  - 子" in body
        assert "2. 二" in body

    def test_blockquote_and_hr_and_pre(self):
        html = (
            "<body><blockquote><p>引用</p></blockquote><hr>"
            "<pre><code class='language-python'>print(1)</code></pre></body>"
        )
        _, body, _ = html_converter.convert_html(html.encode("utf-8"), "x.html")
        assert "> 引用" in body
        assert "---" in body
        assert "```python\nprint(1)\n```" in body

    def test_table(self):
        html = "<body><table><tr><th>甲</th><th>乙</th></tr><tr><td>1</td><td>2</td></tr></table></body>"
        _, body, _ = html_converter.convert_html(html.encode("utf-8"), "x.html")
        assert body == "| 甲 | 乙 |\n| --- | --- |\n| 1 | 2 |"

    def test_external_image_kept(self):
        _, body, _ = html_converter.convert_html(
            b"<body><p><img src='http://e.x/a.png' alt='\xe5\x9b\xbe'></p></body>", "x.html"
        )
        assert "![图](http://e.x/a.png)" in body

    def test_data_uri_image_with_sink(self):
        payload = base64.b64encode(PNG_BYTES).decode("ascii")
        html = f"<body><p><img src='data:image/png;base64,{payload}' alt='图'></p></body>"
        calls: list[tuple[bytes, str]] = []

        def sink(data: bytes, name: str) -> str:
            calls.append((data, name))
            return "KEY9"

        _, body, _ = html_converter.convert_html(html.encode("utf-8"), "x.html", image_sink=sink)
        assert "file:KEY9" in body
        assert calls[0][1] == "image1.png"

    def test_data_uri_image_without_sink_placeholder(self):
        payload = base64.b64encode(PNG_BYTES).decode("ascii")
        html = f"<p><img src='data:image/png;base64,{payload}' alt='图'></p>"
        _, body, warnings = html_converter.convert_html(html.encode("utf-8"), "x.html")
        assert "图片未导入" in body
        assert warnings

    def test_bad_data_uri_placeholder(self):
        _, body, warnings = html_converter.convert_html(b"<p><img src='data:image/png,plain' alt='x'></p>", "x.html")
        assert "图片未导入" in body
        assert warnings

    def test_script_and_style_dropped(self):
        html = "<body><script>var x=1;</script><style>.a{}</style><p>正文</p></body>"
        _, body, _ = html_converter.convert_html(html.encode("utf-8"), "x.html")
        assert body == "正文"

    def test_broken_image_without_src(self):
        _, body, _ = html_converter.convert_html("<body><p><img alt='x'>文字</p></body>".encode(), "x.html")
        assert body == "文字"

    def test_unknown_inline_tag_passthrough(self):
        _, body, _ = html_converter.convert_html(b"<body><p>a <span>b</span> c</p></body>", "x.html")
        assert body == "a b c"

    def test_malformed_html_tolerated(self):
        _, body, _ = html_converter.convert_html(b"<body><p>abc", "x.html")
        assert body == "abc"

    def test_empty_raises(self):
        with pytest.raises(UnsupportedImport):
            html_converter.convert_html(b"   ", "x.html")


class TestHtmlLocalImages:
    """同批上传的本地图片：相对 src 解析 → 站内附件（无 ADR，修复导入丢图）."""

    def test_relative_src_inlined(self):
        html = "<body><p><img src='./图片/图%201.png?x=1#y' alt='图'></p></body>"
        calls: list[tuple[bytes, str]] = []

        def sink(data: bytes, name: str) -> str:
            calls.append((data, name))
            return "LOCAL1"

        _, body, warnings = html_converter.convert_html(
            html.encode(),
            "报告/page.html",
            image_sink=sink,
            local_images={"报告/图片/图 1.png": PNG_BYTES},
        )
        assert "![图](file:LOCAL1)" in body
        assert calls == [(PNG_BYTES, "图 1.png")]
        assert [warning for warning in warnings if "图片" in warning] == []

    def test_parent_relative_src_and_cache(self):
        html = "<body><p><img src='../img/a.png'><img src='../img/a.png' alt='重复'></p></body>"
        calls: list[tuple[bytes, str]] = []

        def sink(data: bytes, name: str) -> str:
            calls.append((data, name))
            return "K"

        _, body, _ = html_converter.convert_html(
            html.encode(),
            "docs/sub/page.html",
            image_sink=sink,
            local_images={"docs/img/a.png": PNG_BYTES},
        )
        assert body.count("file:K") == 2
        assert len(calls) == 1

    def test_missing_local_image_placeholder(self):
        _, body, warnings = html_converter.convert_html(
            b"<p><img src='img/a.png' alt='x'></p>",
            "page.html",
            image_sink=lambda d, n: "K",
            local_images={},
        )
        assert "图片未导入" in body
        assert any("未随附" in warning for warning in warnings)

    def test_invalid_local_path_placeholder(self):
        _, body, warnings = html_converter.convert_html(
            b"<p><img src='/abs/a.png'><img src='../../outside.png'></p>",
            "docs/page.html",
            image_sink=lambda d, n: "K",
            local_images={"a.png": PNG_BYTES},
        )
        assert body.count("图片未导入") == 2
        assert len([warning for warning in warnings if "本地图片" in warning]) == 2

    def test_local_sink_rejected_keeps_placeholder(self):
        _, body, warnings = html_converter.convert_html(
            b"<p><img src='a.png'></p>",
            "page.html",
            image_sink=lambda d, n: None,
            local_images={"a.png": PNG_BYTES},
        )
        assert "图片未导入" in body
        assert warnings

    def test_without_batch_index_keeps_relative_url(self):
        _, body, _ = html_converter.convert_html(b"<p><img src='img/a.png' alt='x'></p>", "page.html")
        assert "![x](img/a.png)" in body

    def test_external_image_not_looked_up(self):
        _, body, _ = html_converter.convert_html(
            b"<p><img src='https://e.x/a.png'><img src='//cdn.x/b.png'></p>",
            "page.html",
            local_images={"https://e.x/a.png": PNG_BYTES},
        )
        assert "![](https://e.x/a.png)" in body
        assert "![](//cdn.x/b.png)" in body


class TestHtmlConvertEdges:
    """html 渲染的边界与容错分支（标签未闭合 / 空洞容器 / 裸文本 / 异常 data URI）."""

    def test_self_closing_void_tag(self):
        _, body, _ = html_converter.convert_html(b"<body><p>a<br/>b</p></body>", "x.html")
        assert body == "a\nb"

    def test_stray_and_void_end_tags(self):
        _, body, _ = html_converter.convert_html(b"<body><p>a</span></br>b</p></body>", "x.html")
        assert body == "ab"

    def test_blockquote_multi_block_and_empty(self):
        html = "<body><blockquote><p>一</p><p>二</p></blockquote><blockquote></blockquote></body>"
        _, body, _ = html_converter.convert_html(html.encode(), "x.html")
        assert body == "> 一\n> 二"

    def test_bare_text_container_and_block_br(self):
        html = "<body>裸文本<div>块<p>段</p></div><br><span>行内</span></body>"
        _, body, _ = html_converter.convert_html(html.encode(), "x.html")
        assert "裸文本" in body
        assert "段" in body
        assert "行内" in body

    def test_table_without_rows(self):
        _, body, _ = html_converter.convert_html(b"<body><table></table><p>x</p></body>", "x.html")
        assert body == "x"

    def test_table_row_without_cells(self):
        _, body, _ = html_converter.convert_html(b"<body><table><tr></tr></table><p>y</p></body>", "x.html")
        assert body == "y"

    def test_pre_without_code_and_plain_class(self):
        html = "<body><pre>纯文本</pre><pre><code class='hljs'>z</code></pre></body>"
        _, body, _ = html_converter.convert_html(html.encode(), "x.html")
        assert "```\n纯文本\n```" in body
        assert "```\nz\n```" in body

    def test_list_with_stray_child(self):
        _, body, _ = html_converter.convert_html("<body><ul>杂文本<li>项</li></ul>".encode(), "x.html")
        assert body == "- 项"

    def test_inline_edge_tags(self):
        html = (
            "<body><ul><li>a<script>x</script>b</li><li><b></b></li><li>x<br>y</li><li><del>删</del></li></ul></body>"
        )
        _, body, _ = html_converter.convert_html(html.encode(), "x.html")
        # <script> 内容被跳过，且相邻行内文本按源码原样拼合（不额外插空格）
        assert "- ab" in body
        assert "- ~~删~~" in body
        assert "\n-\n" in body

    def test_table_cell_pipe_escaped(self):
        html = "<body><table><tr><th>a|b</th></tr><tr><td>1</td></tr></table></body>"
        _, body, _ = html_converter.convert_html(html.encode(), "x.html")
        assert "a\\|b" in body

    def test_table_rowspan_and_colspan_expanded(self):
        html = (
            "<body><table>"
            "<tr><th>类别</th><th>指标</th><th>值</th></tr>"
            "<tr><td rowspan='2'>规模</td><td>甲</td><td>1</td></tr>"
            "<tr><td>乙</td><td>2</td></tr>"
            "<tr><td>结构</td><td colspan='2'>合并</td></tr>"
            "</table></body>"
        )
        _, body, _ = html_converter.convert_html(html.encode(), "x.html")
        assert body.splitlines()[2:] == ["| 规模 | 甲 | 1 |", "| 规模 | 乙 | 2 |", "| 结构 | 合并 |  |"]

    def test_list_items_join_inline_without_added_spaces(self):
        html = "<body><ol><li><strong>格局</strong>：解释</li><li>甲 <em>乙</em></li></ol></body>"
        _, body, _ = html_converter.convert_html(html.encode(), "x.html")
        assert body == "1. **格局**：解释\n2. 甲 *乙*"

    def test_same_document_link_rewritten_to_heading_anchor(self):
        html = (
            "<body><nav><a href='file:///C:/docs/page.html#s01'>导言</a>"
            "<a href='other.html#x'>外部</a><a href='page.html#nope'>未知</a>"
            "<a href='#s01'>纯锚点</a></nav>"
            "<section id='s01'><h2>导言：先看结论</h2></section></body>"
        )
        _, body, _ = html_converter.convert_html(html.encode(), "dir/page.html", base_url="file:///C:/docs/page.html")
        # 章节 id 映射到预览标题锚点（与前端 outline.headingAnchor 同算法），跳转才生效
        assert "[导言](#doc-h-导言先看结论)" in body
        assert "[纯锚点](#doc-h-导言先看结论)" in body
        assert "[外部](other.html#x)" in body
        # 原文没有对应标题时保留原 fragment
        assert "[未知](#nope)" in body

    def test_anchor_index_falls_back_for_non_heading_and_skipped_targets(self):
        html = (
            "<body><h1 id='top'>文档标题</h1><h2>资料</h2>"
            "<ol><li id='cite-1'>[1] 来源</li></ol>"
            "<p><a href='p.html#cite-1'>1</a><a href='p.html#top'>顶</a></p></body>"
        )
        title, body, _ = html_converter.convert_html(html.encode(), "p.html")
        assert title == "文档标题"
        # 引注等无标题锚点退化到最近标题；被剥离为标题的 h1 不产生死锚点
        assert "](#doc-h-资料)" in body
        assert "](#top)" in body

    def test_rendered_svg_in_pre_becomes_svg_image(self):
        html = (
            "<body><pre class='mermaid'><svg viewBox='0 0 10 10'><style>#mermaid-1{fill:#000}</style>"
            "<text>步骤一</text></svg></pre><p>正文</p></body>"
        )
        sunk: list[tuple[bytes, str]] = []

        def sink(data: bytes, name: str) -> str:
            sunk.append((data, name))
            return f"K{len(sunk)}"

        _, body, warnings = html_converter.convert_html(html.encode(), "x.html", image_sink=sink)
        # 渲染型图表落为 .svg 图片附件，不再把 SVG 内的 CSS 灌进正文
        assert body == "![图表](file:K1)\n\n正文"
        assert len(sunk) == 1
        assert sunk[0][1] == "image1.svg"
        assert b"<svg" in sunk[0][0]
        assert b"mermaid-1{fill:#000}" in sunk[0][0]  # <style> 保留（配色靠它）
        assert "font-family" not in body
        assert not any("SVG" in warning for warning in warnings)

    def test_rendered_svg_without_sink_is_placeholder(self):
        html = "<body><pre class='mermaid'><svg viewBox='0 0 1 1'><text>步骤</text></svg></pre><p>正文</p></body>"
        _, body, warnings = html_converter.convert_html(html.encode(), "x.html")
        assert body == "_[图表未导入：image1.svg]_\n\n正文"
        assert not any("SVG" in warning for warning in warnings)  # 与既有图片口径一致：无 sink 时占位不告警（D8）

    def test_chart_figure_svg_becomes_image_and_keeps_caption(self):
        html = (
            "<body><figure><figcaption>图1-1 占比</figcaption>"
            "<div><svg viewBox='0 0 1 1'><style>.x{}</style><text>柱状</text></svg></div></figure></body>"
        )
        _, body, warnings = html_converter.convert_html(html.encode(), "x.html", image_sink=lambda data, name: "K")
        assert body == "图1-1 占比\n\n![图表](file:K)"
        assert not any("SVG" in warning for warning in warnings)

    def test_inline_svg_foreign_object_void_tags_self_closed(self):
        # <img> 引入 SVG 走 XML 解析：foreignObject 内的 <br> 必须自闭合，否则整图失败
        html = "<body><svg viewBox='0 0 1 1'><foreignObject><div>上<br>下</div></foreignObject></svg></body>"
        captured: list[bytes] = []
        html_converter.convert_html(html.encode(), "x.html", image_sink=lambda data, name: captured.append(data) or "K")
        svg = captured[0].decode("utf-8")
        assert "<br/>" in svg
        assert "</br>" not in svg
        ET.fromstring(svg)  # 合法 XML：不抛异常

    def test_inline_svg_scripts_and_event_handlers_stripped(self):
        html = (
            "<body><svg viewBox='0 0 1 1' onload='go()'><script>alert(1)</script>"
            "<a xlink:href='javascript:alert(2)'><text>点</text></a></svg></body>"
        )
        captured: list[bytes] = []
        html_converter.convert_html(html.encode(), "x.html", image_sink=lambda data, name: captured.append(data) or "K")
        svg = captured[0].decode("utf-8")
        assert "<script" not in svg
        assert "onload" not in svg
        assert "javascript:" not in svg
        ET.fromstring(svg)

    def test_data_uri_invalid_base64(self):
        _, body, warnings = html_converter.convert_html(b"<p><img src='data:image/png;base64,A' alt='x'></p>", "x.html")
        assert "图片未导入" in body
        assert warnings

    def test_data_uri_unknown_mime(self):
        payload = base64.b64encode(PNG_BYTES).decode("ascii")
        names: list[str] = []

        def sink(data: bytes, name: str) -> str:
            del data
            names.append(name)
            return "K"

        for mime, expected in (("image/x-icon", ".x-icon"), ("application/octet-stream", ".bin")):
            html = f"<p><img src='data:{mime};base64,{payload}'></p>"
            html_converter.convert_html(html.encode(), "x.html", image_sink=sink)
            assert names[-1].endswith(expected)

    def test_sink_returning_none_keeps_placeholder(self):
        payload = base64.b64encode(PNG_BYTES).decode("ascii")
        html = f"<p><img src='data:image/png;base64,{payload}'></p>"
        _, body, warnings = html_converter.convert_html(html.encode(), "x.html", image_sink=lambda d, n: None)
        assert "图片未导入" in body
        assert warnings

    def test_empty_h1_falls_back_to_filename(self):
        title, _, _ = html_converter.convert_html(b"<body><h1></h1><p>x</p></body>", "x.html")
        assert title == "x"

    def test_feed_error_wrapped(self, monkeypatch):
        def boom(self, data):
            del self, data
            raise RuntimeError("kaboom")

        monkeypatch.setattr(html_converter._Builder, "feed", boom)
        with pytest.raises(UnsupportedImport):
            html_converter.convert_html(b"<p>x</p>", "x.html")


# ── reader ──


class TestReader:
    def test_is_supported(self):
        assert reader.is_supported("a.md")
        assert reader.is_supported("a.MARKDOWN")
        assert reader.is_supported("a.DOCX")
        assert reader.is_supported("a.htm")
        assert not reader.is_supported("a.txt")
        assert not reader.is_supported("a")

    def test_normalize_path(self):
        assert reader.normalize_path("dir\\a.md") == "dir/a.md"
        assert reader.normalize_path("/abs/a.md") is None
        assert reader.normalize_path("../a.md") is None
        assert reader.normalize_path(".hidden/a.md") is None
        assert reader.normalize_path("dir/.a.md") is None
        assert reader.normalize_path("") is None

    def test_validate_batch(self):
        with pytest.raises(BadRequestError):
            reader.validate_batch([])
        assert reader.validate_batch([reader.UploadedItem("a.md", b"xy")]) == 2

    def test_validate_batch_limits(self, monkeypatch):
        monkeypatch.setattr(reader, "MAX_FILES", 2)
        items = [reader.UploadedItem(f"{i}.md", b"x") for i in range(3)]
        with pytest.raises(BadRequestError):
            reader.validate_batch(items)
        monkeypatch.setattr(reader, "MAX_FILES", 100)
        monkeypatch.setattr(reader, "MAX_TOTAL_BYTES", 4)
        with pytest.raises(BadRequestError):
            reader.validate_batch([reader.UploadedItem("a.md", b"12345")])

    def test_convert_file_unsupported(self):
        with pytest.raises(UnsupportedImport):
            reader.convert_file(reader.UploadedItem("a.txt", b"x"))
        with pytest.raises(UnsupportedImport):
            reader.convert_file(reader.UploadedItem("a.md", b""))

    def test_convert_file_oversize(self, monkeypatch):
        monkeypatch.setattr(reader, "MAX_FILE_BYTES", 2)
        with pytest.raises(UnsupportedImport):
            reader.convert_file(reader.UploadedItem("a.md", b"12345"))

    def test_convert_file_markdown(self):
        doc = reader.convert_file(reader.UploadedItem("dir/note.md", "#  标题  \n\n\n正文  \n".encode()))
        assert doc.title == "标题"
        assert doc.body == "正文"
        assert doc.suffix == ".md"

    def test_convert_file_parse_error_wrapped(self, monkeypatch):
        def boom(data, path, *, image_sink=None):
            del data, path, image_sink
            raise RuntimeError("kaboom")

        monkeypatch.setattr(reader.html_converter, "convert_html", boom)
        with pytest.raises(UnsupportedImport):
            reader.convert_file(reader.UploadedItem("a.html", b"<p>x</p>"))

    def test_build_local_image_index(self):
        items = [
            reader.UploadedItem("图片/a.png", PNG_BYTES),
            reader.UploadedItem("图片/空.png", b""),
            reader.UploadedItem("说明.md", b"x"),
        ]
        assert reader.build_local_image_index(items) == {"图片/a.png": PNG_BYTES}

    def test_convert_batch_inlines_local_html_image(self):
        items = [
            reader.UploadedItem("站点/页面.html", "<body><p><img src='./图片/a.png' alt='a'></p></body>".encode()),
            reader.UploadedItem("站点/图片/a.png", PNG_BYTES),
        ]
        keys: list[str] = []

        def sink(data: bytes, name: str) -> str:
            assert data == PNG_BYTES
            assert name == "a.png"
            keys.append(name)
            return "BATCH1"

        docs, skipped = reader.convert_batch(items, image_sink=sink)
        assert skipped == []
        page = next(doc for doc in docs if doc.suffix == ".html")
        image = next(doc for doc in docs if doc.asset)
        assert "![a](file:BATCH1)" in page.body
        assert image.path == "站点/图片/a.png"
        assert keys == ["a.png"]

    def test_convert_batch_isolates_failures(self):
        items = [
            reader.UploadedItem("ok.md", b"# T\n"),
            reader.UploadedItem("bad.txt", b"x"),
            reader.UploadedItem("broken.docx", b"nope"),
        ]
        docs, skipped = reader.convert_batch(items)
        assert [doc.path for doc in docs] == ["ok.md"]
        assert {entry["path"] for entry in skipped} == {"bad.txt", "broken.docx"}
        assert all(entry["reason"] for entry in skipped)


# ── mhtml 转换（ADR 0049）──


def _qp(payload: bytes) -> str:
    import quopri

    return quopri.encodestring(payload).decode("ascii")


def _b64(payload: bytes) -> str:
    return base64.encodebytes(payload).decode("ascii")


def _make_mhtml(
    html: str,
    parts: list[dict] | None = None,
    *,
    boundary: str = "----MultipartBoundary--mht49----",
) -> bytes:
    """构造 Chrome 风格 MHTML 字节流。

    ``parts`` 每项：``{"location": str, "content_id": str | "", "mime": str,
    "data": bytes, "encoding": "base64" | "quoted-printable"}``。
    """
    lines = [
        "From: <Saved by Blink>",
        "Snapshot-Content-Location: https://example.com/page",
        "MIME-Version: 1.0",
        f'Content-Type: multipart/related; type="text/html"; boundary="{boundary}"',
        "",
        f"--{boundary}",
        "Content-Type: text/html; charset=UTF-8",
        "Content-Transfer-Encoding: quoted-printable",
        "Content-Location: https://example.com/page",
        "",
        _qp(html.encode("utf-8")),
        "",
    ]
    for part in parts or []:
        lines += [
            f"--{boundary}",
            f"Content-Type: {part['mime']}",
            f"Content-Transfer-Encoding: {part.get('encoding', 'base64')}",
        ]
        if part.get("content_id"):
            lines.append(f"Content-ID: <{part['content_id']}>")
        if part.get("location"):
            lines.append(f"Content-Location: {part['location']}")
        lines += ["", _b64(part["data"]) if part.get("encoding") != "quoted-printable" else _qp(part["data"]), ""]
    lines.append(f"--{boundary}--")
    return "\r\n".join(lines).encode("ascii")


class TestMhtmlConvert:
    def test_basic_convert_title_and_body(self):
        data = _make_mhtml("<html><head><title>页面标题</title></head><body><h2>小节</h2><p>正文内容</p></body></html>")
        title, body, warnings, assets = mhtml_converter.convert_mhtml(data, "page.mhtml")
        assert title == "页面标题"
        assert "## 小节" in body
        assert "正文内容" in body
        assert warnings == []
        assert assets == []

    def test_images_sunk_as_file_refs(self):
        html = (
            "<html><head><title>图文页</title></head><body>"
            '<p><img src="https://example.com/img/a.png" alt="甲"></p>'
            '<p><img src="./b.jpg" alt="乙"></p>'
            "</body></html>"
        )
        parts = [
            {"location": "https://example.com/img/a.png", "mime": "image/png", "data": PNG_BYTES},
            {"location": "https://example.com/img/b.jpg", "mime": "image/jpeg", "data": b"JPGDATA"},
        ]
        data = _make_mhtml(html, parts)
        sunk: list[tuple[bytes, str]] = []

        def sink(payload: bytes, name: str) -> str:
            sunk.append((payload, name))
            return f"K{len(sunk)}"

        _title, body, warnings, assets = mhtml_converter.convert_mhtml(data, "page.mhtml", image_sink=sink)
        assert "![甲](file:K1)" in body
        assert "![乙](file:K2)" in body
        assert sunk[0] == (PNG_BYTES, "a.png")
        assert sunk[1] == (b"JPGDATA", "b.jpg")
        assert warnings == []
        assert {rel for rel, _ in assets} == {"_assets/a.png", "_assets/b.jpg"}
        assert dict(assets)["_assets/a.png"] == PNG_BYTES

    def test_sink_none_keeps_relative_refs_without_warnings(self):
        html = '<html><head><title>无告警页</title></head><body><p><img src="https://example.com/img/a.png" alt="甲"></p></body></html>'
        parts = [{"location": "https://example.com/img/a.png", "mime": "image/png", "data": PNG_BYTES}]
        data = _make_mhtml(html, parts)
        _, body, warnings, assets = mhtml_converter.convert_mhtml(data, "page.mhtml")
        assert "![甲](_assets/a.png)" in body
        assert "图片未导入" not in body
        assert warnings == []
        assert dict(assets)["_assets/a.png"] == PNG_BYTES

    def test_inline_svg_becomes_attachment(self):
        html = (
            "<html><head><title>图表页</title></head><body><p>前</p>"
            "<svg viewBox='0 0 10 10'><style>.a{fill:red}</style><text>柱状</text></svg>"
            "<p>后</p></body></html>"
        )
        sunk: list[tuple[bytes, str]] = []

        def sink(payload: bytes, name: str) -> str:
            sunk.append((payload, name))
            return f"K{len(sunk)}"

        _title, body, warnings, _assets = mhtml_converter.convert_mhtml(
            _make_mhtml(html), "page.mhtml", image_sink=sink
        )
        assert "![图表](file:K1)" in body
        assert sunk[0][1] == "image1.svg"
        assert b"<svg" in sunk[0][0]
        assert warnings == []

    def test_cid_reference_matches(self):
        html = (
            '<html><head><title>.cid页</title></head><body><p><img src="cid:logo@example" alt="标"></p></body></html>'
        )
        parts = [
            {
                "location": "https://example.com/logo.png",
                "content_id": "logo@example",
                "mime": "image/png",
                "data": PNG_BYTES,
            }
        ]
        data = _make_mhtml(html, parts)
        sunk: list[bytes] = []
        _, body, warnings, _ = mhtml_converter.convert_mhtml(
            data, "page.mhtml", image_sink=lambda payload, name: sunk.append(payload) or "CID1"
        )
        assert "![标](file:CID1)" in body
        assert sunk == [PNG_BYTES]
        assert warnings == []

    def test_relative_src_basename_fallback(self):
        html = (
            '<html><head><title>相对图</title></head><body><p><img src="images/图片 名.png" alt="图"></p></body></html>'
        )
        parts = [
            {
                "location": "https://example.com/assets/%E5%9B%BE%E7%89%87%20%E5%90%8D.png",
                "mime": "image/png",
                "data": b"P",
            }
        ]
        data = _make_mhtml(html, parts)
        _, body, warnings, assets = mhtml_converter.convert_mhtml(data, "page.mhtml")
        assert "![图](_assets/" in body
        assert "图片未导入" not in body
        assert warnings == []
        assert len(assets) == 1

    def test_quoted_printable_html_with_cjk(self):
        # QP 编码的真实用途是文本 part（HTML）；图片素材生产者一律 base64
        html = "<html><head><title>中文页</title></head><body><h1>标题行</h1><p>中文正文与「标点」</p></body></html>"
        parts = [{"location": "https://example.com/a.png", "mime": "image/png", "data": PNG_BYTES}]
        data = _make_mhtml(html, parts)
        title, body, warnings, assets = mhtml_converter.convert_mhtml(data, "page.mhtml")
        assert title == "中文页"
        assert "中文正文与「标点」" in body
        assert warnings == []
        assert dict(assets)["_assets/a.png"] == PNG_BYTES

    def test_missing_resource_placeholder_with_warning(self):
        html = '<html><body><p><img src="cid:missing@x" alt="失"></p></body></html>'
        data = _make_mhtml(html)
        _, body, warnings, _ = mhtml_converter.convert_mhtml(data, "page.mhtml")
        assert "图片未导入" in body
        assert any("missing@x" in w for w in warnings)

    def test_ambiguous_basename_not_matched(self):
        html = '<html><body><p><img src="same.png" alt="重"></p></body></html>'
        parts = [
            {"location": "https://a.com/x/same.png", "mime": "image/png", "data": b"1"},
            {"location": "https://b.com/y/same.png", "mime": "image/png", "data": b"2"},
        ]
        data = _make_mhtml(html, parts)
        _, body, warnings, _ = mhtml_converter.convert_mhtml(data, "page.mhtml")
        assert "图片未导入" in body
        assert warnings

    def test_external_image_kept_as_url(self):
        html = '<html><head><title>外链</title></head><body><p><img src="https://cdn.example.com/away.png" alt="外"></p></body></html>'
        data = _make_mhtml(html)
        _, body, warnings, assets = mhtml_converter.convert_mhtml(data, "page.mhtml")
        assert "![外](https://cdn.example.com/away.png)" in body
        assert warnings == []
        assert assets == []

    def test_css_part_discarded(self):
        html = '<html><head><link rel="stylesheet" href="style.css"></head><body><p>正文</p></body></html>'
        parts = [{"location": "https://example.com/style.css", "mime": "text/css", "data": b"body{}"}]
        data = _make_mhtml(html, parts)
        _, _, _, assets = mhtml_converter.convert_mhtml(data, "page.mhtml")
        assert assets == []

    def test_unreferenced_image_still_sunk(self):
        html = "<html><body><p>无图正文</p></body></html>"
        parts = [{"location": "https://example.com/extra.png", "mime": "image/png", "data": b"EXTRA"}]
        data = _make_mhtml(html, parts)
        sunk: list[str] = []
        _, _, _, assets = mhtml_converter.convert_mhtml(
            data, "page.mhtml", image_sink=lambda payload, name: sunk.append(name) or "U1"
        )
        assert sunk == ["extra.png"]
        assert assets == [("_assets/extra.png", b"EXTRA")]

    def test_duplicate_names_deduped(self):
        html = (
            "<html><body>"
            '<p><img src="https://a.com/x/a.png" alt="一"></p>'
            '<p><img src="https://b.com/y/a.png" alt="二"></p>'
            "</body></html>"
        )
        parts = [
            {"location": "https://a.com/x/a.png", "mime": "image/png", "data": b"1"},
            {"location": "https://b.com/y/a.png", "mime": "image/png", "data": b"2"},
        ]
        data = _make_mhtml(html, parts)
        names: list[str] = []

        def sink(payload: bytes, name: str) -> str:
            names.append(name)
            return f"K{name}"

        _title, _body, _, assets = mhtml_converter.convert_mhtml(data, "page.mhtml", image_sink=sink)
        assert sorted(names) == ["a-2.png", "a.png"]
        assert dict(assets)["_assets/a-2.png"] == b"2"

    def test_mhtml_same_document_links_use_fragments(self):
        html = (
            "<html><head><title>锚点页</title></head><body>"
            '<nav><a href="https://example.com/page#s01">导言</a>'
            '<a href="https://example.com/other#x">外部</a></nav>'
            "<section id='s01'><h2>导言</h2></section></body></html>"
        )
        _, body, warnings, _ = mhtml_converter.convert_mhtml(_make_mhtml(html), "page.mhtml")
        assert "[导言](#doc-h-导言)" in body
        assert "[外部](https://example.com/other#x)" in body
        assert warnings == []

    def test_empty_and_broken_rejected(self):
        with pytest.raises(UnsupportedImport):
            mhtml_converter.convert_mhtml(_make_mhtml(""), "page.mhtml")
        with pytest.raises(UnsupportedImport):
            mhtml_converter.convert_mhtml(b"not a mime message", "page.mhtml")

    def test_no_html_part_rejected(self):
        with pytest.raises(UnsupportedImport):
            mhtml_converter.convert_mhtml(b"From: x\r\nMIME-Version: 1.0\r\n\r\nno multipart here", "page.mhtml")


class TestReaderMhtml:
    def test_convert_file_mhtml(self):
        data = _make_mhtml("<html><head><title>T</title></head><body><p>正文</p></body></html>")
        doc = reader.convert_file(reader.UploadedItem("page.mhtml", data), mhtml_mode="convert")
        assert doc.title == "T"
        assert doc.suffix == ".mhtml"
        assert doc.body == "正文"

    def test_convert_file_mht_suffix(self):
        data = _make_mhtml("<html><body><p>存档</p></body></html>")
        doc = reader.convert_file(reader.UploadedItem("page.mht", data), mhtml_mode="convert")
        assert doc.title == "page"
        assert doc.suffix == ".mht"

    def test_convert_file_mhtml_archive_default(self):
        """0061 C2：默认 archive 原样存档，不转 Markdown."""
        data = _make_mhtml("<html><head><title>T</title></head><body><p>正文</p></body></html>")
        doc = reader.convert_file(reader.UploadedItem("page.mhtml", data))
        assert doc.asset is True
        assert doc.content_type == "mhtml"
        assert doc.mime == "multipart/related"
        assert doc.body == ""
        assert doc.resources == []

    def test_convert_file_mhtml_archive_over_32mb_skipped(self, monkeypatch):
        """archive 模式同样受 32MB 上限约束（先于模式分流检查）。"""
        data = _make_mhtml("<html><body><p>超大</p></body></html>")
        monkeypatch.setattr(reader, "MHTML_MAX_FILE_BYTES", 4)
        with pytest.raises(UnsupportedImport):
            reader.convert_file(reader.UploadedItem("page.mhtml", data), mhtml_mode="archive")

    def test_convert_file_mhtml_unknown_mode_rejected(self):
        data = _make_mhtml("<html><body><p>正文</p></body></html>")
        with pytest.raises(UnsupportedImport):
            reader.convert_file(reader.UploadedItem("page.mhtml", data), mhtml_mode="bogus")

    def test_mhtml_over_8mb_but_under_32mb_accepted(self, monkeypatch):
        data = _make_mhtml("<html><body><p>大文件</p></body></html>")
        assert len(data) <= reader.MAX_FILE_BYTES  # 本例不必真造 8MB，用monkeypatch 缩阈值验证分派
        monkeypatch.setattr(reader, "MAX_FILE_BYTES", 4)
        doc = reader.convert_file(reader.UploadedItem("page.mhtml", data), mhtml_mode="convert")
        assert doc.body == "大文件"

    def test_mhtml_over_32mb_skipped(self, monkeypatch):
        data = _make_mhtml("<html><body><p>超大</p></body></html>")
        monkeypatch.setattr(reader, "MHTML_MAX_FILE_BYTES", 4)
        with pytest.raises(UnsupportedImport):
            reader.convert_file(reader.UploadedItem("page.mhtml", data))

    def test_batch_isolates_broken_mhtml(self):
        items = [
            reader.UploadedItem("ok.mhtml", _make_mhtml("<html><body><p>好</p></body></html>")),
            reader.UploadedItem("bad.mhtml", b"junk"),
        ]
        docs, skipped = reader.convert_batch(items, mhtml_mode="convert")
        assert [doc.path for doc in docs] == ["ok.mhtml"]
        assert {entry["path"] for entry in skipped} == {"bad.mhtml"}


class TestMhtmlPreview:
    """build_preview_html（需求 0064 C1/C2）：自包含 HTML，命中 img 内联 data URI，内联样式保留."""

    def test_inlines_matched_images_and_keeps_style(self):
        html = (
            '<html><head><meta charset="utf-8"><title>预览页</title>'
            "<style>body{color:red}</style></head><body>"
            '<p><img src="https://example.com/img/a.png" alt="甲"></p></body></html>'
        )
        parts = [{"location": "https://example.com/img/a.png", "mime": "image/png", "data": PNG_BYTES}]
        out = mhtml_converter.build_preview_html(_make_mhtml(html, parts))
        assert "<style>body{color:red}</style>" in out
        assert "data:image/png;base64," in out
        assert "https://example.com/img/a.png" not in out

    def test_cid_reference_inlined(self):
        html = '<html><head><title>cid 页</title></head><body><img src="cid:logo" alt="标"></body></html>'
        parts = [{"location": "", "content_id": "logo", "mime": "image/png", "data": PNG_BYTES}]
        out = mhtml_converter.build_preview_html(_make_mhtml(html, parts))
        assert "data:image/png;base64," in out

    def test_unmatched_external_ref_preserved(self):
        html = '<html><head><title>外链页</title></head><body><img src="https://cdn.example.com/x.png" alt="外"></body></html>'
        out = mhtml_converter.build_preview_html(_make_mhtml(html))
        assert 'src="https://cdn.example.com/x.png"' in out

    def test_script_tag_preserved_but_never_executes_in_sandbox(self):
        """预览不改写 HTML 结构；脚本隔离由端点 CSP sandbox + 前端 iframe sandbox 兜底."""
        html = "<html><head><title>脚本页</title></head><body><script>alert(1)</script><p>正文</p></body></html>"
        out = mhtml_converter.build_preview_html(_make_mhtml(html))
        assert "<script>alert(1)</script>" in out

    def test_unparsable_raises(self):
        with pytest.raises(UnsupportedImport):
            mhtml_converter.build_preview_html(b"junk")
