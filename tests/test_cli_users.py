"""cndb users CLI 集成测试.

测试策略：
- 直接调用 cli_users 函数（create/delete/list/import），不走 subprocess
- 用 tmp_path 创建临时 SQLite db，设置 settings.DATABASE_URL
- 覆盖 CRUD + 级联删除 + 批量导入（csv + xlsx + dry-run + 冲突跳过）
"""

from __future__ import annotations

import argparse
from pathlib import Path

import pytest

from cndb.core.config import settings
from cndb.models.base import Base


@pytest.fixture
def cli_db(tmp_path: Path, monkeypatch) -> Path:
    """为 CLI 函数准备独立的临时 SQLite DB，并确保所有模型导入."""
    from sqlalchemy import create_engine

    db_path = tmp_path / "cli_users_test.db"
    monkeypatch.setattr(settings, "DATABASE_URL", f"sqlite:///{db_path}")

    # import 所有插件模型确保 Base.metadata 完整
    import cndb.plugins.accounts.models
    import cndb.plugins.reports.models
    import cndb.plugins.tables.models
    import cndb.plugins.workspaces.models  # noqa: F401

    # 用临时 db_path 新建 engine，避免误用模块级 engine（绑定旧 URL）
    engine = create_engine(f"sqlite:///{db_path}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    engine.dispose()
    return db_path


def _ns(**kwargs) -> argparse.Namespace:
    """构造 argparse.Namespace 的简写."""
    ns = argparse.Namespace()
    for k, v in kwargs.items():
        setattr(ns, k, v)
    return ns


# ── cmd_create ─


class TestCmdCreate:
    def test_create_basic_user(self, cli_db):
        from cndb.cli.users import cmd_create

        args = _ns(username="alice", password="pw1234", email=None, nickname=None, role="user", is_superuser=False)
        result = cmd_create(args)
        assert result.username == "alice"
        assert result.role == "user"
        assert result.user_id == 1
        assert result.password == "pw1234"

    def test_create_auto_password(self, cli_db):
        from cndb.cli.users import cmd_create

        args = _ns(username="bob", password=None, email=None, nickname=None, role="user", is_superuser=False)
        result = cmd_create(args)
        assert len(result.password) == 12
        assert "自动生成" in result.note

    def test_create_system_admin(self, cli_db):
        from cndb.cli.users import cmd_create

        args = _ns(
            username="sa1", password="pw1234", email=None, nickname=None, role="system_admin", is_superuser=False
        )
        result = cmd_create(args)
        assert result.role == "system_admin"

    def test_create_superuser_flag(self, cli_db):
        from cndb.cli.users import cmd_create

        args = _ns(username="su1", password="pw1234", email=None, nickname=None, role="user", is_superuser=True)
        result = cmd_create(args)
        assert "[superuser]" in result.note

    def test_create_nickname_defaults_to_role_cn(self, cli_db):
        from cndb.cli.users import _get_session

        args = _ns(
            username="nicktest", password="pw1234", email=None, nickname=None, role="system_admin", is_superuser=False
        )
        from cndb.cli.users import cmd_create

        cmd_create(args)

        db = _get_session()
        try:
            from cndb.plugins.accounts.models import User

            u = db.query(User).filter(User.username == "nicktest").first()
            assert u is not None
            assert u.nickname == "系统管理员"
        finally:
            db.close()

    def test_create_duplicate_username_fails(self, cli_db):
        from cndb.cli.users import cmd_create

        args = _ns(username="dup", password="pw1234", email=None, nickname=None, role="user", is_superuser=False)
        cmd_create(args)

        with pytest.raises(SystemExit) as ei:
            cmd_create(
                _ns(username="dup", password="pw1234", email=None, nickname=None, role="user", is_superuser=False)
            )
        assert "已存在" in str(ei.value)

    def test_create_invalid_role_fails(self, cli_db):
        from cndb.cli.users import cmd_create

        with pytest.raises(SystemExit) as ei:
            cmd_create(
                _ns(
                    username="badrole",
                    password="pw1234",
                    email=None,
                    nickname=None,
                    role="super_hero",
                    is_superuser=False,
                )
            )
        assert "无效角色" in str(ei.value)

    def test_create_short_password_fails(self, cli_db):
        from cndb.cli.users import cmd_create

        with pytest.raises(SystemExit) as ei:
            cmd_create(
                _ns(username="shortpw", password="12", email=None, nickname=None, role="user", is_superuser=False)
            )
        assert "至少 6 位" in str(ei.value)


# ── cmd_delete ─


class TestCmdDelete:
    def test_delete_user_no_owned_workspace(self, cli_db):
        from cndb.cli.users import cmd_create, cmd_delete

        cmd_create(_ns(username="plain", password="pw1234", email=None, nickname=None, role="user", is_superuser=False))
        args = _ns(target="plain", cascade=False, yes=False, check_only=False)
        result = cmd_delete(args)
        assert result["deleted_username"] == "plain"
        assert result["deleted_workspace_count"] == 0

    def test_delete_owned_workspace_without_cascade_rejects(self, cli_db):
        from cndb.cli.users import _get_session, cmd_create, cmd_delete

        # 先创建用户，再让他成为某个 workspace 的 OWNER
        cmd_create(
            _ns(username="owner1", password="pw1234", email=None, nickname=None, role="user", is_superuser=False)
        )

        db = _get_session()
        try:
            from cndb.plugins.accounts.models import User
            from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

            u = db.query(User).filter(User.username == "owner1").first()
            ws = Workspace(name="ws_owned")
            db.add(ws)
            db.flush()
            db.add(
                WorkspaceMember(
                    workspace_id=ws.id,
                    user_id=u.id,
                    role=WorkspaceRole.OWNER,
                )
            )
            db.commit()
        finally:
            db.close()

        # 不传 --cascade 应被拒绝
        args = _ns(target="owner1", cascade=False, yes=False, check_only=False)
        with pytest.raises(SystemExit) as ei:
            cmd_delete(args)
        assert ei.value.code == 2

    def test_delete_owned_workspace_with_cascade_and_yes(self, cli_db):
        from cndb.cli.users import _get_session, cmd_create, cmd_delete

        cmd_create(
            _ns(username="owner2", password="pw1234", email=None, nickname=None, role="user", is_superuser=False)
        )

        db = _get_session()
        try:
            from cndb.plugins.accounts.models import User
            from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

            u = db.query(User).filter(User.username == "owner2").first()
            for name in ("ws_a", "ws_b"):
                ws = Workspace(name=name)
                db.add(ws)
                db.flush()
                db.add(
                    WorkspaceMember(
                        workspace_id=ws.id,
                        user_id=u.id,
                        role=WorkspaceRole.OWNER,
                    )
                )
            db.commit()
        finally:
            db.close()

        args = _ns(target="owner2", cascade=True, yes=True, check_only=False)
        result = cmd_delete(args)
        assert result["deleted_username"] == "owner2"
        assert result["deleted_workspace_count"] == 2

        # 确认工作区已被删
        db = _get_session()
        try:
            from cndb.plugins.workspaces.models import Workspace

            assert db.query(Workspace).count() == 0
        finally:
            db.close()

    def test_delete_check_only(self, cli_db):
        from cndb.cli.users import _get_session, cmd_create, cmd_delete_check

        cmd_create(
            _ns(username="owner3", password="pw1234", email=None, nickname=None, role="user", is_superuser=False)
        )

        db = _get_session()
        try:
            from cndb.plugins.accounts.models import User
            from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

            u = db.query(User).filter(User.username == "owner3").first()
            ws = Workspace(name="ws_check")
            db.add(ws)
            db.flush()
            db.add(WorkspaceMember(workspace_id=ws.id, user_id=u.id, role=WorkspaceRole.OWNER))
            db.commit()
        finally:
            db.close()

        check = cmd_delete_check(_ns(target="owner3"))
        assert check is not None
        assert check.can_delete_without_cascade is False
        assert len(check.owned_workspaces) == 1

    def test_delete_not_found(self, cli_db):
        from cndb.cli.users import cmd_delete

        with pytest.raises(SystemExit):
            cmd_delete(_ns(target="no_such_user", cascade=False, yes=False, check_only=False))


# ── cmd_list ─


class TestCmdList:
    def test_list_all(self, cli_db):
        from cndb.cli.users import cmd_create, cmd_list

        for i in range(3):
            cmd_create(
                _ns(
                    username=f"list_u{i}", password="pw1234", email=None, nickname=None, role="user", is_superuser=False
                )
            )
        users = cmd_list(_ns(role=None, active=False, inactive=False))
        assert len(users) == 3

    def test_list_filter_by_role(self, cli_db):
        from cndb.cli.users import cmd_create, cmd_list

        cmd_create(_ns(username="l1", password="pw1234", email=None, nickname=None, role="user", is_superuser=False))
        cmd_create(
            _ns(username="l2", password="pw1234", email=None, nickname=None, role="system_admin", is_superuser=False)
        )
        cmd_create(
            _ns(username="l3", password="pw1234", email=None, nickname=None, role="system_admin", is_superuser=False)
        )

        admins = cmd_list(_ns(role="system_admin", active=False, inactive=False))
        assert len(admins) == 2
        plain = cmd_list(_ns(role="user", active=False, inactive=False))
        assert len(plain) == 1

    def test_list_filter_invalid_role(self, cli_db):
        from cndb.cli.users import cmd_list

        with pytest.raises(SystemExit):
            cmd_list(_ns(role="bad_role", active=False, inactive=False))


# ── cmd_import ─


class TestCmdImport:
    def test_import_csv_basic(self, cli_db, tmp_path: Path):
        from cndb.cli.users import cmd_import

        csv_file = tmp_path / "users.csv"
        csv_file.write_text(
            "username,password,email,nickname,role\n"
            "csv_a,pw1234,a@x.com,A,user\n"
            "csv_b,pw1234,,B,system_admin\n"
            "csv_c,,c@x.com,C,user\n",
            encoding="utf-8",
        )
        report = cmd_import(_ns(file=str(csv_file), dry_run=False))
        assert report.total == 3
        assert len(report.ok) == 3
        # 第三行 password 空 → 应自动生成
        ok_c = next(r for r in report.ok if r.username == "csv_c")
        assert ok_c.generated_password is not None
        assert len(ok_c.generated_password) == 12

    def test_import_csv_with_chinese_headers(self, cli_db, tmp_path: Path):
        from cndb.cli.users import cmd_import

        csv_file = tmp_path / "users_cn.csv"
        csv_file.write_text(
            "用户名,密码,邮箱,昵称,角色\ncn_user,pw1234,,CN,user\n",
            encoding="utf-8",
        )
        report = cmd_import(_ns(file=str(csv_file), dry_run=False))
        assert len(report.ok) == 1

    def test_import_csv_conflict_skips(self, cli_db, tmp_path: Path):
        from cndb.cli.users import cmd_create, cmd_import

        cmd_create(
            _ns(username="dup_import", password="pw1234", email=None, nickname=None, role="user", is_superuser=False)
        )

        csv_file = tmp_path / "users_dup.csv"
        csv_file.write_text(
            "username,password\ndup_import,pw1234\nnew_user,pw1234\n",
            encoding="utf-8",
        )
        report = cmd_import(_ns(file=str(csv_file), dry_run=False))
        assert len(report.ok) == 1
        assert len(report.skipped) == 1
        assert report.skipped[0].username == "dup_import"

    def test_import_csv_dry_run(self, cli_db, tmp_path: Path):
        from cndb.cli.users import cmd_import

        csv_file = tmp_path / "users_dry.csv"
        csv_file.write_text(
            "username,password,role\ndry_a,pw1234,user\ndry_b,pw1234,bad_role\n",
            encoding="utf-8",
        )
        report = cmd_import(_ns(file=str(csv_file), dry_run=True))
        # 第一行校验通过，第二行失败
        assert len(report.ok) == 1
        assert len(report.failed) == 1
        # DB 里不应有数据
        from cndb.cli.users import _get_session
        from cndb.plugins.accounts.models import User

        db = _get_session()
        try:
            assert db.query(User).count() == 0
        finally:
            db.close()

    def test_import_xlsx_basic(self, cli_db, tmp_path: Path):
        from openpyxl import Workbook

        from cndb.cli.users import cmd_import

        wb = Workbook()
        ws = wb.active
        ws.append(["username", "password", "email", "role"])
        ws.append(["xl_a", "pw1234", "a@x.com", "user"])
        ws.append(["xl_b", "pw1234", "", "security_admin"])
        ws.append(["xl_c", "", "", "user"])  # 空密码 → 自动生成
        xlsx_path = tmp_path / "users.xlsx"
        wb.save(xlsx_path)

        report = cmd_import(_ns(file=str(xlsx_path), dry_run=False))
        assert len(report.ok) == 3

    def test_import_invalid_format(self, cli_db, tmp_path: Path):
        from cndb.cli.users import cmd_import

        bad = tmp_path / "users.txt"
        bad.write_text("nonsense\n", encoding="utf-8")
        with pytest.raises(ValueError):
            cmd_import(_ns(file=str(bad), dry_run=False))

    def test_import_file_not_found(self, cli_db, tmp_path: Path):
        from cndb.cli.users import cmd_import

        with pytest.raises(FileNotFoundError):
            cmd_import(_ns(file=str(tmp_path / "no.csv"), dry_run=False))

    def test_import_empty_username_fails(self, cli_db, tmp_path: Path):
        from cndb.cli.users import cmd_import

        csv_file = tmp_path / "users_empty.csv"
        csv_file.write_text(
            "username,password\n,pw1234\n",
            encoding="utf-8",
        )
        report = cmd_import(_ns(file=str(csv_file), dry_run=False))
        assert len(report.failed) == 1
        assert "缺少 username" in report.failed[0].message


# ── print_* 输出函数 ─


class TestPrintFunctions:
    def test_print_create_result_with_note(self, capsys, cli_db):
        from cndb.cli.users import CreateResult, print_create_result

        r = CreateResult(user_id=1, username="alice", role="user", password="pw1234", note="（自动生成） [superuser]")
        print_create_result(r)
        out = capsys.readouterr().out
        assert "alice" in out
        assert "pw1234" in out
        assert "自动生成" in out

    def test_print_create_result_no_note(self, capsys, cli_db):
        from cndb.cli.users import CreateResult, print_create_result

        r = CreateResult(user_id=2, username="bob", role="user", password="secret")
        print_create_result(r)
        out = capsys.readouterr().out
        assert "bob" in out
        assert "secret" in out

    def test_print_delete_check_none(self, capsys, cli_db):
        from cndb.cli.users import print_delete_check

        print_delete_check(None)
        out = capsys.readouterr().out
        assert "不存在" in out

    def test_print_delete_check_with_workspaces(self, capsys, cli_db):
        from cndb.cli.users import DeleteCheckResult, print_delete_check
        from cndb.plugins.workspaces.models import Workspace

        ws = Workspace(name="ws1")
        ws.id = 99
        check = DeleteCheckResult(
            user_id=1,
            username="owner",
            is_superuser=False,
            owned_workspaces=[ws],
            member_count=2,
            can_delete_without_cascade=False,
        )
        print_delete_check(check)
        out = capsys.readouterr().out
        assert "owner" in out
        assert "ws1" in out

    def test_print_delete_result_with_workspaces(self, capsys, cli_db):
        from cndb.cli.users import print_delete_result

        print_delete_result(
            {
                "deleted_username": "x",
                "deleted_user_id": 5,
                "deleted_workspace_count": 2,
                "deleted_workspace_ids": [1, 2],
            }
        )
        out = capsys.readouterr().out
        assert "x" in out
        assert "2" in out

    def test_print_delete_result_no_workspaces(self, capsys, cli_db):
        from cndb.cli.users import print_delete_result

        print_delete_result(
            {"deleted_username": "y", "deleted_user_id": 6, "deleted_workspace_count": 0, "deleted_workspace_ids": []}
        )
        out = capsys.readouterr().out
        assert "y" in out

    def test_print_list_empty(self, capsys, cli_db):
        from cndb.cli.users import print_list

        print_list([])
        out = capsys.readouterr().out
        assert "没有匹配" in out

    def test_print_list_with_users(self, capsys, cli_db):
        from cndb.cli.users import cmd_create, cmd_list, print_list

        cmd_create(_ns(username="pl1", password="pw1234", email=None, nickname=None, role="user", is_superuser=False))
        users = cmd_list(_ns(role=None, active=False, inactive=False))
        print_list(users)
        out = capsys.readouterr().out
        assert "pl1" in out
        assert "共 1 条" in out

    def test_print_import_report_full(self, capsys, cli_db):
        from cndb.cli.users import ImportReport, ImportRowResult, print_import_report

        report = ImportReport(
            total=3,
            ok=[ImportRowResult(row_index=1, status="ok", message="", username="u1", generated_password="abc123")],
            skipped=[ImportRowResult(row_index=2, status="skipped", message="已存在", username="u2")],
            failed=[ImportRowResult(row_index=3, status="failed", message="缺少 username")],
        )
        print_import_report(report, Path("users.csv"), dry_run=False)
        out = capsys.readouterr().out
        assert "u1" in out
        assert "u2" in out
        assert "缺少 username" in out
        assert "自动生成" in out  # warning about generated passwords

    def test_print_import_report_dry_run(self, capsys, cli_db):
        from cndb.cli.users import ImportReport, print_import_report

        report = ImportReport(total=1, ok=[], skipped=[], failed=[])
        print_import_report(report, Path("users.csv"), dry_run=True)
        out = capsys.readouterr().out
        assert "预演模式" in out


# ── users_command 分发 ─


class TestUsersCommand:
    def test_dispatch_create(self, capsys, cli_db, monkeypatch):
        from cndb.cli.users import users_command

        monkeypatch.setattr("cndb.cli.users.cmd_create", lambda args: None)
        monkeypatch.setattr("cndb.cli.users.print_create_result", lambda r: None)
        users_command(_ns(users_cmd="create"))

    def test_dispatch_delete_check_only(self, capsys, cli_db, monkeypatch):
        from cndb.cli.users import users_command

        monkeypatch.setattr("cndb.cli.users.cmd_delete_check", lambda args: None)
        monkeypatch.setattr("cndb.cli.users.print_delete_check", lambda c: None)
        users_command(_ns(users_cmd="delete", check_only=True))

    def test_dispatch_delete_normal(self, capsys, cli_db, monkeypatch):
        from cndb.cli.users import users_command

        monkeypatch.setattr("cndb.cli.users.cmd_delete", lambda args: {})
        monkeypatch.setattr("cndb.cli.users.print_delete_result", lambda r: None)
        users_command(_ns(users_cmd="delete", check_only=False))

    def test_dispatch_list(self, capsys, cli_db, monkeypatch):
        from cndb.cli.users import users_command

        monkeypatch.setattr("cndb.cli.users.cmd_list", lambda args: [])
        monkeypatch.setattr("cndb.cli.users.print_list", lambda u: None)
        users_command(_ns(users_cmd="list"))

    def test_dispatch_import(self, capsys, cli_db, monkeypatch):
        from cndb.cli.users import users_command

        monkeypatch.setattr("cndb.cli.users.cmd_import", lambda args: None)
        monkeypatch.setattr("cndb.cli.users.print_import_report", lambda r, f, dry_run: None)
        users_command(_ns(users_cmd="import", file="x.csv", dry_run=False))

    def test_dispatch_missing_subcommand(self, cli_db):
        from cndb.cli.users import users_command

        with pytest.raises(SystemExit) as ei:
            users_command(_ns(users_cmd=None))
        assert ei.value.code == 2


# ── 边界分支 ─


class TestEdgeCases:
    def test_create_duplicate_email_fails(self, cli_db):
        from cndb.cli.users import cmd_create

        cmd_create(
            _ns(username="e1", password="pw1234", email="dup@x.com", nickname=None, role="user", is_superuser=False)
        )
        with pytest.raises(SystemExit) as ei:
            cmd_create(
                _ns(username="e2", password="pw1234", email="dup@x.com", nickname=None, role="user", is_superuser=False)
            )
        assert "邮箱" in str(ei.value)

    def test_delete_superuser_warning(self, cli_db, capsys):
        from cndb.cli.users import cmd_create, cmd_delete

        cmd_create(_ns(username="su_del", password="pw1234", email=None, nickname=None, role="user", is_superuser=True))
        cmd_delete(_ns(target="su_del", cascade=False, yes=False, check_only=False))
        err = capsys.readouterr().err
        assert "superuser" in err

    def test_delete_check_not_found_returns_none(self, cli_db):
        from cndb.cli.users import cmd_delete_check

        assert cmd_delete_check(_ns(target="ghost")) is None

    def test_list_active_filter(self, cli_db):
        from cndb.cli.users import _get_session, cmd_create, cmd_list
        from cndb.plugins.accounts.models import User

        cmd_create(_ns(username="act1", password="pw1234", email=None, nickname=None, role="user", is_superuser=False))
        # 把第二个用户设为 inactive
        cmd_create(_ns(username="act2", password="pw1234", email=None, nickname=None, role="user", is_superuser=False))
        db = _get_session()
        try:
            u = db.query(User).filter(User.username == "act2").first()
            u.is_active = False
            db.commit()
        finally:
            db.close()

        active = cmd_list(_ns(role=None, active=True, inactive=False))
        inactive = cmd_list(_ns(role=None, active=False, inactive=True))
        assert len(active) == 1
        assert len(inactive) == 1

    def test_resolve_user_by_digit_id(self, cli_db):
        from cndb.cli.users import _get_session, _resolve_user, cmd_create

        result = cmd_create(
            _ns(username="digit_u", password="pw1234", email=None, nickname=None, role="user", is_superuser=False)
        )
        db = _get_session()
        try:
            u = _resolve_user(db, str(result.user_id))
            assert u is not None
            assert u.username == "digit_u"
        finally:
            db.close()

    def test_read_csv_gbk_fallback(self, tmp_path: Path):
        from cndb.cli.users import _read_csv

        f = tmp_path / "gbk.csv"
        f.write_bytes("用户名,密码\n张三,pw1234\n".encode("gbk"))
        rows = _read_csv(f)
        assert len(rows) == 1
        assert rows[0]["username"] == "张三"

    def test_read_csv_empty_file(self, tmp_path: Path):
        from cndb.cli.users import _read_csv

        f = tmp_path / "empty.csv"
        f.write_text("", encoding="utf-8")
        assert _read_csv(f) == []

    def test_read_xlsx_none_header(self, tmp_path: Path):
        from openpyxl import Workbook

        from cndb.cli.users import _read_xlsx

        wb = Workbook()
        ws = wb.active
        # 只写一行数据，不写表头 → header 为 None 分支
        ws.append([None, None])
        xlsx = tmp_path / "none_header.xlsx"
        wb.save(xlsx)
        rows = _read_xlsx(xlsx)
        assert isinstance(rows, list)

    def test_map_row_skips_unknown_column(self):
        """未知列头应被跳过，不进入映射结果."""
        from cndb.cli.users import _map_row

        row = {"username": "x", "unknown_col": "y"}
        mapped = _map_row(row)
        assert mapped == {"username": "x"}

    def test_validate_row_username_too_short(self):
        from cndb.cli.users import _validate_row

        err = _validate_row({"username": "a"})
        assert err is not None
        assert "2 个字符" in err

    def test_validate_row_password_too_short(self):
        from cndb.cli.users import _validate_row

        err = _validate_row({"username": "valid", "password": "12"})
        assert err is not None
        assert "6 位" in err

    def test_import_single_user_db_exception(self, cli_db, monkeypatch):
        from cndb.cli.users import ImportReport, _import_single_user

        report = ImportReport(total=1, ok=[], skipped=[], failed=[])

        def _boom(*a, **kw):
            raise RuntimeError("simulated db error")

        from cndb.plugins.accounts.models import User

        monkeypatch.setattr(User, "set_password", _boom)
        _import_single_user(report, 1, "exc_user", None, "", "user", "pw1234", None)
        assert len(report.failed) == 1
        assert "DB 错误" in report.failed[0].message
