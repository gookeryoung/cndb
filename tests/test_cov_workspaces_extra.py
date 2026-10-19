"""补充覆盖面测试：workspaces / files / accounts 的遗漏分支.

仅**新增**测试，不修改源码。集中在此前基线中未被覆盖的路径：

- ``workspaces/services/backup_io.py``：备份导入的版本校验、无字段行、
  跨表引用重映射、link 关联行回填、行级权限、文档树还原等分支；
- ``workspaces/services/storage.py``：符号链接/二进制值/配额回退；
- ``workspaces/routers/workspaces.py``：导入端点的格式校验、可见性/类型
  回退、失败清理，以及无 OWNER 成员时的详情分支；
- accounts：认证端点的 ``current_user is None`` 分支、用户管理服务、
  会话（刷新令牌）服务、通知服务、批量导入/导出边界；
- files：归属登记/清理/访问校验分支。
"""

from __future__ import annotations

import datetime as dt
from pathlib import Path
from types import SimpleNamespace
from typing import Any

import pytest

WS = "/api/v1/workspaces"
AUTH = "/api/v1/accounts/auth"


def _import(
    wid: int,
    data: dict[str, Any],
    user: Any,
    db: Any,
) -> Any:
    from cndb.plugins.workspaces.services.backup_io import _import_backup_into_workspace

    return _import_backup_into_workspace(wid, data, user, db)


def _new_workspace(client: Any, headers: dict[str, str], name: str = "cov-ws") -> int:
    r = client.post(WS, json={"name": name}, headers=headers)
    assert r.status_code == 201, r.text
    return int(r.json()["id"])


def _testuser(db: Any) -> Any:
    from cndb.plugins.accounts.models import User

    return db.query(User).filter(User.username == "testuser").one()


def _make_user(
    db: Any,
    username: str,
    *,
    role: str = "user",
    is_superuser: bool = False,
    is_active: bool = True,
    email: str | None = None,
    nickname: str = "",
) -> Any:
    from cndb.plugins.accounts.models import User

    user = User(
        username=username,
        nickname=nickname,
        role=role,
        is_superuser=is_superuser,
        is_active=is_active,
        email=email,
    )
    user.set_password("pw123456")
    db.add(user)
    db.commit()
    db.refresh(user)
    return user


# ══════════════════════════════════════════════════════════════
# workspaces/services/storage.py
# ══════════════════════════════════════════════════════════════


class TestStorageHelpers:
    def test_dir_size_skips_symlinks_and_subdirs(self, tmp_path: Path) -> None:
        from cndb.plugins.workspaces.services.storage import _dir_size

        root = tmp_path / "d"
        root.mkdir()
        (root / "f.txt").write_text("abc", encoding="utf-8")
        (root / "sub").mkdir()
        (root / "sub" / "g.txt").write_text("de", encoding="utf-8")
        target = tmp_path / "target.txt"
        target.write_text("zzzz", encoding="utf-8")
        (root / "link.txt").symlink_to(target)

        # 3 + 2 = 5：符号链接与目录条目本身不计入
        assert _dir_size(root) == 5

    def test_dir_size_missing_dir_is_zero(self, tmp_path: Path) -> None:
        from cndb.plugins.workspaces.services.storage import _dir_size

        assert _dir_size(tmp_path / "nope") == 0

    def test_value_bytes_variants(self) -> None:
        from cndb.plugins.workspaces.services.storage import _value_bytes

        assert _value_bytes(None) == 0
        assert _value_bytes(b"abc") == 3
        assert _value_bytes(bytearray(b"ab")) == 2
        assert _value_bytes(memoryview(b"abcd")) == 4
        assert _value_bytes("中") == 3
        assert _value_bytes(1234) == 4

    def test_quota_falls_back_when_invalid(self) -> None:
        from cndb.plugins.workspaces.models import DEFAULT_STORAGE_QUOTA_BYTES
        from cndb.plugins.workspaces.services.storage import workspace_quota_bytes

        assert workspace_quota_bytes(SimpleNamespace(storage_quota_bytes=-5)) == DEFAULT_STORAGE_QUOTA_BYTES
        assert workspace_quota_bytes(SimpleNamespace(storage_quota_bytes="x")) == DEFAULT_STORAGE_QUOTA_BYTES
        assert workspace_quota_bytes(SimpleNamespace(storage_quota_bytes=10)) == 10


# ══════════════════════════════════════════════════════════════
# workspaces/routers/workspaces.py 导入端点 + 详情分支
# ══════════════════════════════════════════════════════════════


class TestWorkspacesImportRouter:
    def test_invalid_backup_format_400(self, client: Any, auth_headers: dict[str, str]) -> None:
        r = client.post(f"{WS}/import", json={"json_data": {}}, headers=auth_headers)
        assert r.status_code == 400, r.text

    def test_visibility_type_fallback_and_cleanup(self, client: Any, auth_headers: dict[str, str], db: Any) -> None:
        from cndb.plugins.workspaces.models import Workspace

        before = db.query(Workspace).count()
        data = {
            "version": "9",
            "tables": [],
            "workspace": {"visibility": "bogus", "workspace_type": "bogus"},
        }
        r = client.post(f"{WS}/import", json={"json_data": data}, headers=auth_headers)
        assert r.status_code == 400, r.text
        # 导入失败应清理刚创建的工作区
        assert db.query(Workspace).count() == before

    def test_detail_without_owner_member(self, client: Any, auth_headers: dict[str, str], db: Any) -> None:
        from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

        user = _testuser(db)
        ws = Workspace(name="no-owner", created_by_id=user.id)
        db.add(ws)
        db.flush()
        db.add(WorkspaceMember(workspace_id=ws.id, user_id=user.id, role=WorkspaceRole.EDITOR))
        db.commit()

        r = client.get(f"{WS}/{ws.id}", headers=auth_headers)
        assert r.status_code == 200, r.text
        assert r.json()["owner"] is None


# ══════════════════════════════════════════════════════════════
# workspaces/services/backup_io.py 直接调用
# ══════════════════════════════════════════════════════════════


class TestBackupIoDirect:
    def _ctx(self, client: Any, auth_headers: dict[str, str], db: Any, name: str) -> tuple[int, Any]:
        wid = _new_workspace(client, auth_headers, name)
        return wid, _testuser(db)

    def test_unsupported_version_rejected(self, client: Any, auth_headers: dict[str, str], db: Any) -> None:
        from fastapi import HTTPException

        wid, user = self._ctx(client, auth_headers, db, "io-ver")
        with pytest.raises(HTTPException) as exc_info:
            _import(wid, {"version": "99", "tables": []}, user, db)
        assert exc_info.value.status_code == 400

    def test_rows_without_fields_records_error(self, client: Any, auth_headers: dict[str, str], db: Any) -> None:
        wid, user = self._ctx(client, auth_headers, db, "io-nofield")
        data = {"version": "4", "tables": [{"name": "NoField", "id": 1, "fields": [], "rows": [{"a": 1}]}]}
        resp = _import(wid, data, user, db)
        assert resp.imported_rows == 0
        assert any("无字段定义" in e for e in resp.errors)

    def test_rows_matched_and_dropped(self, client: Any, auth_headers: dict[str, str], db: Any) -> None:
        wid, user = self._ctx(client, auth_headers, db, "io-rows")
        data = {
            "version": "4",
            "tables": [
                {
                    "name": "T",
                    "id": 1,
                    "fields": [{"id": 11, "name": "名称", "field_type": "text"}],
                    "rows": [{"名称": "a"}, {"bad": "b"}],
                }
            ],
        }
        resp = _import(wid, data, user, db)
        assert resp.imported_tables == 1
        assert resp.imported_rows == 1
        assert any("无法匹配到字段" in e for e in resp.errors)

    def test_link_lookup_remap_and_backfill(self, client: Any, auth_headers: dict[str, str], db: Any) -> None:
        wid, user = self._ctx(client, auth_headers, db, "io-link")
        data = {
            "version": "4",
            "tables": [
                {
                    "name": "A",
                    "id": 1,
                    "fields": [{"id": 11, "name": "ta", "field_type": "text"}],
                    "rows": [{"ta": "a1"}],
                },
                {
                    "name": "B",
                    "id": 2,
                    "fields": [
                        {"id": 21, "name": "tb", "field_type": "text"},
                        {"id": 22, "name": "lb", "field_type": "link", "config": {"target_table_id": 1}},
                        {
                            "id": 23,
                            "name": "lu",
                            "field_type": "lookup",
                            "config": {"source_table_id": 1, "source_field_id": 11},
                        },
                    ],
                    "rows": [{"tb": "b1"}],
                    "links": [{"field": "lb", "pairs": [[0, 0]]}],
                },
            ],
        }
        resp = _import(wid, data, user, db)
        assert resp.imported_tables == 2
        assert resp.imported_rows == 2
        assert resp.errors == []

    def test_row_permissions_and_permission(self, client: Any, auth_headers: dict[str, str], db: Any) -> None:
        wid, user = self._ctx(client, auth_headers, db, "io-perm")
        data = {
            "version": "4",
            "tables": [
                {
                    "name": "RP",
                    "id": 1,
                    "fields": [{"id": 1, "name": "f", "field_type": "text"}],
                    "rows": [],
                    "permission": {"read_role": "editor", "hidden_fields": {"viewer": ["f"]}},
                    "row_permissions": [
                        123,
                        {"subject_type": "bogus", "action": "view"},
                        {"subject_type": "user", "subject_id": "1", "action": "view"},
                    ],
                }
            ],
        }
        resp = _import(wid, data, user, db)
        assert resp.imported_tables == 1
        assert resp.errors == []

    def test_documents_restored(self, client: Any, auth_headers: dict[str, str], db: Any) -> None:
        wid, user = self._ctx(client, auth_headers, db, "io-docs")
        data = {
            "version": "4",
            "tables": [],
            "document_tags": [{"name": ""}, {"name": "标签", "id": "notint"}],
            "documents": [
                {"id": 1, "title": "资料夹", "is_folder": True},
                {"id": 2, "title": "子文档", "parent_id": 1, "content": "正文"},
                {"id": 3, "title": "自引用", "parent_id": 3},
                {"id": 4, "title": "孤儿", "parent_id": 999},
            ],
        }
        resp = _import(wid, data, user, db)
        assert resp.imported_documents == 4
        assert resp.errors == []

    def test_document_import_failure_records_error(
        self, client: Any, auth_headers: dict[str, str], db: Any, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        wid, user = self._ctx(client, auth_headers, db, "io-docfail")

        def _boom(*_a: Any, **_k: Any) -> Any:
            raise RuntimeError("boom")

        monkeypatch.setattr("cndb.plugins.documents.services.content.create_document", _boom)
        data = {"version": "4", "tables": [], "documents": [{"id": 1, "title": "坏文档"}]}
        resp = _import(wid, data, user, db)
        assert resp.imported_documents == 0
        assert any("导入失败" in e for e in resp.errors)


class TestCoerceRowValueTypes:
    def test_no_date_columns_returns_early(self) -> None:
        from sqlalchemy import Column, Integer, MetaData, Table

        from cndb.plugins.workspaces.services.backup_io import _coerce_row_value_types

        metadata = MetaData()
        sa = Table("t", metadata, Column("id", Integer, primary_key=True))
        rows = [{"id": 1}]
        _coerce_row_value_types(sa, rows)
        assert rows == [{"id": 1}]


# ══════════════════════════════════════════════════════════════
# accounts/routers/auth.py —— current_user is None 分支
# ══════════════════════════════════════════════════════════════


class TestAuthNoneUserBranches:
    @staticmethod
    def _disable_auth(monkeypatch: pytest.MonkeyPatch) -> None:
        from cndb.core.config import settings

        monkeypatch.setattr(settings, "AUTH_ENABLED", False)

    def test_logout_all_requires_login(self, client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
        self._disable_auth(monkeypatch)
        r = client.post(f"{AUTH}/logout", json={"all": True})
        assert r.status_code == 401, r.text

    def test_change_password_requires_login(self, client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
        self._disable_auth(monkeypatch)
        r = client.post(f"{AUTH}/password", json={"current_password": "x", "new_password": "yyyyyy"})
        assert r.status_code == 401, r.text

    def test_list_sessions_requires_login(self, client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
        self._disable_auth(monkeypatch)
        assert client.get(f"{AUTH}/sessions").status_code == 401

    def test_revoke_session_requires_login(self, client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
        self._disable_auth(monkeypatch)
        assert client.delete(f"{AUTH}/sessions/1").status_code == 401

    def test_me_requires_login(self, client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
        self._disable_auth(monkeypatch)
        assert client.get(f"{AUTH}/me").status_code == 401

    def test_update_profile_requires_login(self, client: Any, monkeypatch: pytest.MonkeyPatch) -> None:
        self._disable_auth(monkeypatch)
        assert client.patch(f"{AUTH}/me", json={"nickname": "x"}).status_code == 401

    def test_update_profile_integrity_error(
        self, client: Any, auth_headers: dict[str, str], db: Any, monkeypatch: pytest.MonkeyPatch
    ) -> None:
        from sqlalchemy.exc import IntegrityError

        def _boom() -> None:
            raise IntegrityError("stmt", {}, Exception("dup"))

        monkeypatch.setattr(db, "commit", _boom)
        r = client.patch(f"{AUTH}/me", json={"email": "unique@example.com"}, headers=auth_headers)
        assert r.status_code == 400, r.text


# ══════════════════════════════════════════════════════════════
# accounts/services/users.py
# ══════════════════════════════════════════════════════════════


class TestUserAdminService:
    def test_update_username_and_email(self, db: Any) -> None:
        from cndb.plugins.accounts.schemas.users import UserAdminUpdateRequest
        from cndb.plugins.accounts.services.users import update_user

        actor = _make_user(db, "admin1", is_superuser=True)
        target = _make_user(db, "oldname", email="old@example.com")
        data = UserAdminUpdateRequest(username="newname", email="new@example.com")
        updated = update_user(db, actor=actor, target=target, data=data)
        assert updated.username == "newname"
        assert updated.email == "new@example.com"

    def test_update_email_conflict(self, db: Any) -> None:
        from cndb.plugins.accounts.schemas.users import UserAdminUpdateRequest
        from cndb.plugins.accounts.services.users import UserAdminError, update_user

        actor = _make_user(db, "admin2", is_superuser=True)
        _make_user(db, "taken", email="taken@example.com")
        target = _make_user(db, "t2")
        with pytest.raises(UserAdminError):
            update_user(db, actor=actor, target=target, data=UserAdminUpdateRequest(email="taken@example.com"))

    def test_update_same_role_activates(self, db: Any) -> None:
        from cndb.plugins.accounts.schemas.users import UserAdminUpdateRequest
        from cndb.plugins.accounts.services.users import update_user

        actor = _make_user(db, "admin3", is_superuser=True)
        target = _make_user(db, "t3", role="user", is_active=False)
        updated = update_user(db, actor=actor, target=target, data=UserAdminUpdateRequest(role="user", is_active=True))
        assert updated.is_active is True

    def test_update_integrity_error(self, db: Any, monkeypatch: pytest.MonkeyPatch) -> None:
        from sqlalchemy.exc import IntegrityError

        from cndb.plugins.accounts.schemas.users import UserAdminUpdateRequest
        from cndb.plugins.accounts.services.users import UserAdminError, update_user

        actor = _make_user(db, "admin4", is_superuser=True)
        target = _make_user(db, "t4")

        def _boom() -> None:
            raise IntegrityError("stmt", {}, Exception("dup"))

        monkeypatch.setattr(db, "commit", _boom)
        with pytest.raises(UserAdminError):
            update_user(db, actor=actor, target=target, data=UserAdminUpdateRequest(nickname="新昵称"))

    def test_query_audit_logs_actor_filter(self, db: Any) -> None:
        from cndb.plugins.accounts.services.users import query_audit_logs

        logs, total = query_audit_logs(db, actor_id=999)
        assert total == 0
        assert logs == []

    def test_validate_role_invalid(self) -> None:
        from cndb.plugins.accounts.services.users import UserAdminError, validate_role

        with pytest.raises(UserAdminError):
            validate_role("not-a-role")


# ══════════════════════════════════════════════════════════════
# accounts/services/sessions.py
# ══════════════════════════════════════════════════════════════


class TestSessionsService:
    def test_revoke_token_missing(self, db: Any) -> None:
        from cndb.plugins.accounts.services import sessions

        assert sessions.revoke_token(db, "no-such-refresh-token") is False

    def test_rotate_inactive_user_raises(self, db: Any) -> None:
        from cndb.core.errors import UnauthorizedError
        from cndb.core.security import hash_token
        from cndb.plugins.accounts.models import RefreshToken
        from cndb.plugins.accounts.services import sessions

        user = _make_user(db, "inactive", is_active=False)
        raw = "raw-refresh-token-value"
        db.add(
            RefreshToken(
                user_id=user.id,
                token_hash=hash_token(raw),
                expires_at=dt.datetime.now(dt.UTC).replace(tzinfo=None) + dt.timedelta(days=1),
            )
        )
        db.commit()
        with pytest.raises(UnauthorizedError):
            sessions.rotate(db, raw)


# ══════════════════════════════════════════════════════════════
# accounts/services/notifications.py
# ══════════════════════════════════════════════════════════════


class TestNotificationsService:
    def _seed(self, db: Any) -> tuple[Any, Any]:
        from cndb.plugins.accounts.models import Notification

        actor = _make_user(db, "notifactor", nickname="张三")
        recipient = _make_user(db, "notifrecipient")
        row = Notification(recipient_id=recipient.id, type="doc", actor_id=actor.id, payload={})
        db.add(row)
        db.commit()
        db.refresh(row)
        return recipient, row

    def test_serialize_falls_back_to_actor_name(self, db: Any) -> None:
        from cndb.plugins.accounts.services import notifications

        recipient, _row = self._seed(db)
        items = notifications.list_notifications(db, recipient)["items"]
        assert items[0]["actor_name"] == "张三"

    def test_unread_count(self, db: Any) -> None:
        from cndb.plugins.accounts.services import notifications

        recipient, _row = self._seed(db)
        assert notifications.unread_count(db, recipient) == 1

    def test_mark_read_twice_is_idempotent(self, db: Any) -> None:
        from cndb.plugins.accounts.services import notifications

        recipient, row = self._seed(db)
        assert notifications.mark_read(db, recipient, row.id) is True
        assert notifications.mark_read(db, recipient, row.id) is True


# ══════════════════════════════════════════════════════════════
# accounts/services/user_bulk.py
# ══════════════════════════════════════════════════════════════


class TestUserBulkExtra:
    def test_serialize_export_unsupported_format(self) -> None:
        from cndb.plugins.accounts.services.user_bulk import serialize_export

        with pytest.raises(ValueError):
            serialize_export([], "json")

    def test_read_xlsx_no_active_sheet(self, monkeypatch: pytest.MonkeyPatch) -> None:
        import openpyxl

        from cndb.plugins.accounts.services.user_bulk import _read_xlsx_rows

        monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: SimpleNamespace(active=None))
        assert _read_xlsx_rows(b"whatever") == []

    def test_read_xlsx_none_header(self, monkeypatch: pytest.MonkeyPatch) -> None:
        import openpyxl

        from cndb.plugins.accounts.services.user_bulk import _read_xlsx_rows

        fake_ws = SimpleNamespace(iter_rows=lambda **k: iter([None]))
        monkeypatch.setattr(openpyxl, "load_workbook", lambda *a, **k: SimpleNamespace(active=fake_ws))
        assert _read_xlsx_rows(b"whatever") == []


# ══════════════════════════════════════════════════════════════
# accounts/routers —— 鉴权守卫
# ══════════════════════════════════════════════════════════════


class TestAccountGuards:
    def test_ensure_authenticated_none(self) -> None:
        from fastapi import HTTPException

        from cndb.plugins.accounts.routers.preferences import _ensure_authenticated

        with pytest.raises(HTTPException) as exc_info:
            _ensure_authenticated(None)
        assert exc_info.value.status_code == 401

    def test_require_audit_reader_none(self) -> None:
        from fastapi import HTTPException

        from cndb.plugins.accounts.routers.users import _require_audit_reader

        with pytest.raises(HTTPException) as exc_info:
            _require_audit_reader(None)
        assert exc_info.value.status_code == 403


# ══════════════════════════════════════════════════════════════
# files/services.py
# ══════════════════════════════════════════════════════════════


class TestFilesServices:
    def test_record_owner_idempotent_overwrite(self, db: Any) -> None:
        from cndb.plugins.files.services import record_owner

        first = record_owner(db, file_key="covkey", workspace_id=1)
        second = record_owner(db, file_key="covkey", workspace_id=1, table_id=5)
        assert second.id == first.id
        assert second.table_id == 5

    def test_forget_owner_missing_is_noop(self, db: Any) -> None:
        from cndb.plugins.files.services import forget_owner

        forget_owner(db, "ghost-key")

    def test_assert_can_access_unknown_key_returns(self, db: Any) -> None:
        from cndb.plugins.files.services import assert_can_access

        user = _make_user(db, "fileuser")
        assert_can_access(db, workspace_id=1, file_key="ghost-key", user=user)

    def test_assert_can_attach_document_editor(self, db: Any) -> None:
        from cndb.plugins.documents.services.content import create_document
        from cndb.plugins.files.services import assert_can_attach_document
        from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

        creator = _make_user(db, "doccCreator")
        editor = _make_user(db, "doccEditor")
        ws = Workspace(name="attach-ws", created_by_id=creator.id)
        db.add(ws)
        db.flush()
        db.add(WorkspaceMember(workspace_id=ws.id, user_id=editor.id, role=WorkspaceRole.EDITOR))
        db.commit()
        doc = create_document(db, ws.id, creator, title="已发布", content="x", status="published")

        returned = assert_can_attach_document(db, document_id=doc.id, workspace_id=ws.id, user=editor)
        assert returned.id == doc.id

    def test_assert_can_attach_document_owner_admin(self, db: Any) -> None:
        from cndb.plugins.documents.services.content import create_document
        from cndb.plugins.files.services import assert_can_attach_document
        from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

        creator = _make_user(db, "doccCreator3")
        owner = _make_user(db, "doccOwner3")
        ws = Workspace(name="attach-ws3", created_by_id=creator.id)
        db.add(ws)
        db.flush()
        db.add(WorkspaceMember(workspace_id=ws.id, user_id=owner.id, role=WorkspaceRole.OWNER))
        db.commit()
        doc = create_document(db, ws.id, creator, title="归属", content="x", status="published")

        returned = assert_can_attach_document(db, document_id=doc.id, workspace_id=ws.id, user=owner)
        assert returned.id == doc.id

    def test_assert_can_attach_document_editor_granted(self, db: Any) -> None:
        from cndb.plugins.documents.models import DocumentPermission
        from cndb.plugins.documents.services.content import create_document
        from cndb.plugins.files.services import assert_can_attach_document
        from cndb.plugins.workspaces.models import Workspace

        creator = _make_user(db, "doccCreator4")
        granted = _make_user(db, "doccGranted4")
        ws = Workspace(name="attach-ws4", created_by_id=creator.id)
        db.add(ws)
        db.commit()
        doc = create_document(db, ws.id, creator, title="草稿", content="x", status="draft")
        db.add(DocumentPermission(document_id=doc.id, user_id=granted.id, role="editor"))
        db.commit()

        returned = assert_can_attach_document(db, document_id=doc.id, workspace_id=ws.id, user=granted)
        assert returned.id == doc.id

    def test_assert_can_attach_document_viewer_denied(self, db: Any) -> None:
        from cndb.core.errors import PermissionDeniedError
        from cndb.plugins.documents.services.content import create_document
        from cndb.plugins.files.services import assert_can_attach_document
        from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

        creator = _make_user(db, "doccCreator2")
        viewer = _make_user(db, "doccViewer2")
        ws = Workspace(name="attach-ws2", created_by_id=creator.id)
        db.add(ws)
        db.flush()
        db.add(WorkspaceMember(workspace_id=ws.id, user_id=viewer.id, role=WorkspaceRole.VIEWER))
        db.commit()
        doc = create_document(db, ws.id, creator, title="公开文档", content="x", status="published")

        with pytest.raises(PermissionDeniedError):
            assert_can_attach_document(db, document_id=doc.id, workspace_id=ws.id, user=viewer)
