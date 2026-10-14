"""WPS 带图片单元格 XLSX 导入测试 —— DISPLIMG 公式解析、cellimages 提取、落盘与建表导入.

WPS「插入单元格图片」在 xlsx 中的存储：
- 单元格值为 ``=DISPLIMG("ID_xxx",1)`` 公式；
- ``xl/cellimages.xml``（WPS 私有扩展）定义 ID → rId；
- ``xl/_rels/cellimages.xml.rels`` 把 rId 映射到 ``xl/media/*`` 图片二进制。

测试用合成 xlsx（openpyxl 造基础工作簿 + zipfile 注入 WPS 扩展部件）覆盖：
公式识别 / 映射提取 / 解析为元数据列表 / 类型推断 image / 建表导入全链路。
"""

from __future__ import annotations

import base64
import io
import json
import zipfile
from typing import Any

import pytest
from openpyxl import Workbook
from sqlalchemy.orm import Session

from cndb.core.config import settings
from cndb.plugins.tables.models import DataField, DataTable
from cndb.plugins.tables.services.core import ddl as _ddl
from cndb.plugins.tables.services.transfer import (
    _parse_xlsx_bytes,
    analyze_json_columns,
    match_displimg_formula,
    parse_wps_cell_images,
)

# 1x1 透明 PNG
_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)

# WPS 扩展命名空间
_NS_ETC = "http://www.wps.cn/officeDocument/2017/etcustomchars"
_NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _inject_wps_parts(xlsx_bytes: bytes, images: dict[str, tuple[str, bytes]]) -> bytes:
    """向 xlsx zip 注入 WPS 单元格图片扩展部件.

    Args:
        images: ``{displimg_id: (media 文件名, 图片二进制)}``.
    """
    pics: list[str] = []
    rels: list[str] = []
    for i, (did, (fname, _data)) in enumerate(images.items(), start=1):
        rid = f"rId{i}"
        rels.append(
            f'<Relationship Id="{rid}" '
            f'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            f'Target="media/{fname}"/>'
        )
        pics.append(
            f'<etc:cellimage><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="{i}" name="{did}"/></xdr:nvPicPr>'
            f'<xdr:blipFill><a:blip r:embed="{rid}"/></xdr:blipFill></xdr:pic></etc:cellimage>'
        )
    cellimages_xml = (
        f'<etc:cellimages xmlns:etc="{_NS_ETC}" xmlns:xdr="{_NS_XDR}" xmlns:a="{_NS_A}" xmlns:r="{_NS_R}">'
        + "".join(pics)
        + "</etc:cellimages>"
    )
    rels_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + "".join(rels)
        + "</Relationships>"
    )
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        dst.writestr("xl/cellimages.xml", cellimages_xml)
        dst.writestr("xl/_rels/cellimages.xml.rels", rels_xml)
        for fname, data in images.values():
            dst.writestr(f"xl/media/{fname}", data)
    return out.getvalue()


def _build_wps_xlsx(
    headers: list[str],
    rows: list[list[Any]],
    images: dict[str, tuple[str, bytes]],
) -> bytes:
    """构造带 WPS 单元格图片的合成 xlsx（图片单元格写 DISPLIMG 公式）."""
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return _inject_wps_parts(buf.getvalue(), images)


# ── 公式识别 ─────────────────────────────────────


class TestMatchDisplimgFormula:
    def test_standard_form(self):
        assert match_displimg_formula('=DISPLIMG("ID_1",1)') == "ID_1"

    def test_without_leading_eq(self):
        assert match_displimg_formula('DISPLIMG("ID_2",1)') == "ID_2"

    def test_xlfn_prefix(self):
        assert match_displimg_formula('_xlfn.DISPLIMG("ID_3",1)') == "ID_3"
        assert match_displimg_formula('=_xlfn.DISPLIMG("ID_3",1)') == "ID_3"

    def test_case_insensitive_and_spaces(self):
        assert match_displimg_formula('=displimg( "ID_4" , 1 )') == "ID_4"

    def test_non_formula_values(self):
        assert match_displimg_formula("hello") is None
        assert match_displimg_formula("=SUM(A1)") is None
        assert match_displimg_formula(None) is None
        assert match_displimg_formula(123) is None
        assert match_displimg_formula("DISPLIMG 在这里只是普通文本") is None


# ── cellimages 映射提取 ────────────────────────────


class TestParseWpsCellImages:
    def test_extracts_images(self):
        raw = _build_wps_xlsx(
            ["名称", "图片"],
            [["a", '=DISPLIMG("ID_A",1)']],
            {"ID_A": ("image1.png", _PNG_1PX)},
        )
        images = parse_wps_cell_images(raw)
        assert set(images) == {"ID_A"}
        assert images["ID_A"].data == _PNG_1PX
        assert images["ID_A"].filename == "image1.png"
        assert images["ID_A"].mime_type == "image/png"

    def test_multiple_images(self):
        raw = _build_wps_xlsx(
            ["名称", "图1", "图2"],
            [["a", '=DISPLIMG("ID_A",1)', '=DISPLIMG("ID_B",1)']],
            {"ID_A": ("image1.png", _PNG_1PX), "ID_B": ("image2.png", _PNG_1PX)},
        )
        images = parse_wps_cell_images(raw)
        assert set(images) == {"ID_A", "ID_B"}

    def test_non_wps_xlsx_returns_empty(self):
        wb = Workbook()
        wb.active.append(["a", "b"])
        wb.active.append([1, 2])
        buf = io.BytesIO()
        wb.save(buf)
        assert parse_wps_cell_images(buf.getvalue()) == {}

    def test_corrupt_rels_returns_empty(self):
        raw = _build_wps_xlsx(
            ["名称", "图片"],
            [["a", '=DISPLIMG("ID_A",1)']],
            {"ID_A": ("image1.png", _PNG_1PX)},
        )
        # 重写 zip：只保留 cellimages.xml，去掉 rels → 解析失败返回空
        out = io.BytesIO()
        with zipfile.ZipFile(io.BytesIO(raw)) as src, zipfile.ZipFile(out, "w") as dst:
            for item in src.infolist():
                if item.filename != "xl/_rels/cellimages.xml.rels":
                    dst.writestr(item, src.read(item.filename))
        assert parse_wps_cell_images(out.getvalue()) == {}


# ── 解析为行值 ────────────────────────────────────


class TestParseXlsxBytesWithWpsImages:
    def test_ephemeral_default_produces_meta_list(self):
        raw = _build_wps_xlsx(
            ["名称", "图片"],
            [["a", '=DISPLIMG("ID_A",1)']],
            {"ID_A": ("image1.png", _PNG_1PX)},
        )
        rows, cols = _parse_xlsx_bytes(raw)
        assert cols == ["名称", "图片"]
        assert len(rows) == 1
        val = rows[0]["图片"]
        assert isinstance(val, list) and len(val) == 1
        assert val[0]["filename"] == "image1.png"
        assert val[0]["mime_type"] == "image/png"
        assert val[0]["size"] == len(_PNG_1PX)
        assert val[0]["file_key"] == ""  # ephemeral 不落盘

    def test_unknown_displimg_id_becomes_none(self):
        raw = _build_wps_xlsx(
            ["名称", "图片"],
            [["a", '=DISPLIMG("ID_MISSING",1)']],
            {"ID_A": ("image1.png", _PNG_1PX)},
        )
        rows, _ = _parse_xlsx_bytes(raw)
        assert rows[0]["图片"] is None

    def test_custom_saver_rejection_becomes_none(self):
        raw = _build_wps_xlsx(
            ["名称", "图片"],
            [["a", '=DISPLIMG("ID_A",1)']],
            {"ID_A": ("image1.png", _PNG_1PX)},
        )
        rows, _ = _parse_xlsx_bytes(raw, image_saver=lambda _img: None)
        assert rows[0]["图片"] is None

    def test_non_image_media_skipped_by_saver(self):
        raw = _build_wps_xlsx(
            ["名称", "图片"],
            [["a", '=DISPLIMG("ID_A",1)']],
            {"ID_A": ("chart.emf", b"\x01\x02\x03")},
        )
        rows, _ = _parse_xlsx_bytes(raw)
        assert rows[0]["图片"] is None

    def test_plain_xlsx_unchanged(self):
        wb = Workbook()
        wb.active.append(["名称", "数量"])
        wb.active.append(["a", 3])
        buf = io.BytesIO()
        wb.save(buf)
        rows, cols = _parse_xlsx_bytes(buf.getvalue())
        assert cols == ["名称", "数量"]
        assert rows[0] == {"名称": "a", "数量": 3}


# ── 类型推断 ─────────────────────────────────────


class TestImageInference:
    def test_image_list_column_infers_image(self):
        rows = [
            {"名称": "a", "图片": [{"file_key": "k1.png", "filename": "i1.png", "mime_type": "image/png"}]},
            {"名称": "b", "图片": [{"file_key": "k2.png", "filename": "i2.png", "mime_type": "image/png"}]},
        ]
        cols = analyze_json_columns(rows)
        type_map = {c["name"]: c["field_type"] for c in cols}
        assert type_map["图片"] == "image"

    def test_ephemeral_values_also_infer_image(self):
        rows = [
            {"图片": [{"file_key": "", "filename": "i1.png", "mime_type": "image/png", "size": 10}]},
        ]
        cols = analyze_json_columns(rows)
        assert cols[0]["field_type"] == "image"

    def test_non_image_dict_list_stays_json(self):
        rows = [{"数据": [{"a": 1}, {"b": 2}]}]
        cols = analyze_json_columns(rows)
        assert cols[0]["field_type"] == "json"


# ── 建表导入全链路（API） ──────────────────────────


@pytest.fixture
def ws_auth(client, db):
    """工作区 + 认证 fixture."""
    from cndb.plugins.accounts.models import User
    from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

    u = User(username="wps_import_user", nickname="WPS Import User")
    u.set_password("passw0rd")
    db.add(u)
    db.flush()
    ws = Workspace(name="WpsImportWS", created_by_id=u.id)
    db.add(ws)
    db.flush()
    db.add(WorkspaceMember(workspace_id=ws.id, user_id=u.id, role=WorkspaceRole.OWNER))
    db.commit()
    r = client.post(
        "/api/v1/accounts/auth/login",
        json={"login": "wps_import_user", "password": "passw0rd"},
    )
    token = r.json()["access_token"]
    return ws.id, {"Authorization": f"Bearer {token}"}


class TestWpsImportFileApi:
    def test_analyze_reports_image_column_without_persisting(self, client, ws_auth):
        ws_id, auth = ws_auth
        raw = _build_wps_xlsx(
            ["名称", "图片"],
            [["a", '=DISPLIMG("ID_A",1)'], ["b", '=DISPLIMG("ID_B",1)']],
            {"ID_A": ("image1.png", _PNG_1PX), "ID_B": ("image2.png", _PNG_1PX)},
        )
        ws_upload_dir = settings.UPLOAD_DIR / str(ws_id)
        files_before = set(ws_upload_dir.glob("*")) if ws_upload_dir.is_dir() else set()

        r = client.post(
            f"/api/v1/workspaces/{ws_id}/import-file/analyze",
            files={"file": ("商品.xlsx", raw, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        type_map = {c["name"]: c["field_type"] for c in data["columns"]}
        assert type_map["图片"] == "image"
        assert type_map["名称"] == "text"
        assert data["total_rows"] == 2
        # 纯分析不落盘
        files_after = set(ws_upload_dir.glob("*")) if ws_upload_dir.is_dir() else set()
        assert files_after == files_before

    def test_create_table_persists_images_and_rows(self, client, ws_auth, db):
        ws_id, auth = ws_auth
        raw = _build_wps_xlsx(
            ["名称", "图片"],
            [["a", '=DISPLIMG("ID_A",1)'], ["b", '=DISPLIMG("ID_B",1)']],
            {"ID_A": ("image1.png", _PNG_1PX), "ID_B": ("image2.png", _PNG_1PX)},
        )
        r = client.post(
            f"/api/v1/workspaces/{ws_id}/import-file",
            files={"file": ("商品.xlsx", raw, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
            data={"table_name": "WPS商品表"},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert data["imported_rows"] == 2
        assert data["table_id"]

        # 字段类型：图片列推断为 image
        dt = db.get(DataTable, data["table_id"])
        assert dt is not None
        field_map = {f.name: f.field_type for f in dt.active_fields()}
        assert field_map["图片"] == "image"

        # 行值：JSON 数组，file_key 指向已落盘文件
        from cndb.plugins.tables.services.core import records as rec

        engine = db.get_bind()
        rows, _total = rec.list_rows(engine, dt, limit=10)
        assert len(rows) == 2
        for row in rows:
            val = json.loads(row["图片"])
            assert len(val) == 1
            file_key = val[0]["file_key"]
            assert file_key.endswith(".png")
            saved = settings.UPLOAD_DIR / str(ws_id) / file_key
            assert saved.is_file(), f"图片应落盘到 {saved}"
            assert saved.read_bytes() == _PNG_1PX

    def test_attachment_field_accepts_wps_images(self, client, ws_auth, db):
        """导入到已有 attachment 字段：图片元数据同构，附件字段可直接消费."""
        ws_id, _auth = ws_auth
        engine = db.get_bind()
        dt = DataTable(workspace_id=ws_id, name="附件表")
        dt.ensure_db_name()
        db.add(dt)
        db.commit()
        db.refresh(dt)
        f1 = DataField(table_id=dt.id, name="名称", field_type="text", order=0)
        f1.ensure_db_name()
        f2 = DataField(table_id=dt.id, name="照片", field_type="attachment", order=1)
        f2.ensure_db_name()
        db.add_all([f1, f2])
        db.commit()
        db.refresh(dt)
        _ddl.create_table(engine, dt)

        from cndb.plugins.tables.services.importing.importer import Importer

        raw = _build_wps_xlsx(
            ["名称", "照片"],
            [["a", '=DISPLIMG("ID_A",1)']],
            {"ID_A": ("image1.png", _PNG_1PX)},
        )
        imp = Importer(engine, db, dt)
        result = imp.execute(raw, "xlsx")
        assert result.imported_ids, "至少导入一行"

        from cndb.plugins.tables.services.core import records as rec

        rows, _total = rec.list_rows(engine, dt, limit=10)
        assert len(rows) == 1
        val = json.loads(rows[0]["照片"])
        assert val[0]["filename"] == "image1.png"
        assert (settings.UPLOAD_DIR / str(ws_id) / val[0]["file_key"]).is_file()

    def test_importer_image_field_roundtrip(self, client, ws_auth, db):
        """导入到已有 image 字段：行值为合法 JSON 数组且图片落盘."""
        ws_id, _auth = ws_auth
        engine = db.get_bind()
        dt = DataTable(workspace_id=ws_id, name="图片表")
        dt.ensure_db_name()
        db.add(dt)
        db.commit()
        db.refresh(dt)
        f1 = DataField(table_id=dt.id, name="名称", field_type="text", order=0)
        f1.ensure_db_name()
        f2 = DataField(table_id=dt.id, name="图片", field_type="image", order=1)
        f2.ensure_db_name()
        db.add_all([f1, f2])
        db.commit()
        db.refresh(dt)
        _ddl.create_table(engine, dt)

        from cndb.plugins.tables.services.core import records as rec
        from cndb.plugins.tables.services.importing.importer import Importer

        raw = _build_wps_xlsx(
            ["名称", "图片"],
            [["a", '=DISPLIMG("ID_A",1)'], ["b", '=DISPLIMG("ID_B",1)']],
            {"ID_A": ("image1.png", _PNG_1PX), "ID_B": ("image2.png", _PNG_1PX)},
        )
        imp = Importer(engine, db, dt)
        analysis = imp.analyze(raw, "xlsx")
        assert analysis.report["total"] == 2
        result = imp.execute(analysis=analysis)
        assert len(result.imported_ids) == 2

        rows, _total = rec.list_rows(engine, dt, limit=10)
        assert len(rows) == 2
        for row in rows:
            val = json.loads(row["图片"])
            assert val[0]["mime_type"] == "image/png"
            assert (settings.UPLOAD_DIR / str(ws_id) / val[0]["file_key"]).is_file()


# ── import_rows_from_xlsx 直连路径 ──────────────────


class TestImportRowsFromXlsx:
    def test_wps_images_saved_to_workspace(self, db: Session):
        engine = db.get_bind()
        dt = DataTable(workspace_id=99, name="直连导入表")
        dt.ensure_db_name()
        db.add(dt)
        db.commit()
        db.refresh(dt)
        f1 = DataField(table_id=dt.id, name="名称", field_type="text", order=0)
        f1.ensure_db_name()
        f2 = DataField(table_id=dt.id, name="图片", field_type="image", order=1)
        f2.ensure_db_name()
        db.add_all([f1, f2])
        db.commit()
        db.refresh(dt)
        _ddl.create_table(engine, dt)

        from cndb.plugins.tables.services.transfer import import_rows_from_xlsx

        raw = _build_wps_xlsx(
            ["名称", "图片"],
            [["a", '=DISPLIMG("ID_A",1)']],
            {"ID_A": ("image1.png", _PNG_1PX)},
        )
        ids = import_rows_from_xlsx(engine, dt, raw, db=db)
        assert len(ids) == 1

        from cndb.plugins.tables.services.core import records as rec

        rows, _total = rec.list_rows(engine, dt, limit=5)
        val = json.loads(rows[0]["图片"])
        assert val[0]["filename"] == "image1.png"
        assert (settings.UPLOAD_DIR / "99" / val[0]["file_key"]).is_file()
