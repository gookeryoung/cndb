"""documents 插件补充覆盖率测试（新增文件，不改动源码）.

覆盖现有测试未触达的私有分支与错误路径，主要面向：
- ``services.obsidian``：zip 边界、front-matter 边角、canvas 节点类型、正文改写、
  导入冲突/状态校验、导出命名冲突与图片改写；
- ``services.content``：错误分支、外存/压缩读取、批量软删除与恢复、出链合并、标签增删；
- ``services.search``：兜底通道、语法错误收敛、索引重建、无 workspace 分支；
- ``services.export``：PDF / DOCX 表格、分页、行内样式与字体候选解析；
- 若干路由的错误分支（公开分享、标签、vault、documents）。

写法：能直接调用服务层私有函数的一律直接调用（快速、确定），
需要走 HTTP 的用 ``client`` + ``auth_headers`` 风格。
"""

from __future__ import annotations

import base64
import datetime as dt
import gzip
import io
import json
import zipfile
from pathlib import Path
from types import SimpleNamespace

import pytest
from sqlalchemy.orm import Session

from cndb.core.config import settings
from cndb.core.errors import BadRequestError, ConflictError, NotFoundError
from cndb.plugins.accounts.models import User
from cndb.plugins.documents.models import (
    Document,
    DocumentDraft,
    DocumentPermission,
    DocumentShare,
    Tag,
)
from cndb.plugins.documents.services import content as content_service
from cndb.plugins.documents.services import document_permissions as doc_perms
from cndb.plugins.documents.services import drafts, obsidian, parse, search, visibility
from cndb.plugins.documents.services.export import docx as docx_export
from cndb.plugins.documents.services.export import pdf as pdf_export
from cndb.plugins.documents.services.export import pdf_fonts
from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

DOCS = "/api/v1/workspaces/{ws}/documents"


# ── fixtures / helpers ──


def _mk_user(db: Session, username: str) -> User:
    u = User(username=username, nickname=username)
    u.set_password("passw0rd")
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


def _login(client, username: str) -> dict[str, str]:
    r = client.post("/api/v1/accounts/auth/login", json={"login": username, "password": "passw0rd"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


@pytest.fixture
def owner(db) -> User:
    return _mk_user(db, "covdoc_owner")


@pytest.fixture
def ws(db, owner) -> Workspace:
    w = Workspace(name="CovDocWS", created_by_id=owner.id)
    db.add(w)
    db.flush()
    db.add(WorkspaceMember(workspace_id=w.id, user_id=owner.id, role=WorkspaceRole.OWNER))
    db.commit()
    db.refresh(w)
    return w


@pytest.fixture
def auth(client, owner) -> dict[str, str]:
    return _login(client, owner.username)


def _build_zip(files: dict[str, bytes | str], root: str = "") -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        for name, data in files.items():
            payload = data.encode("utf-8") if isinstance(data, str) else data
            z.writestr(f"{root}/{name}" if root else name, payload)
    return buf.getvalue()


def _post_zip(client, url: str, data: bytes, auth: dict[str, str], **form):
    return client.post(url, files={"file": ("vault.zip", data, "application/zip")}, data=form, headers=auth)


def _create_doc(client, ws: Workspace, auth: dict[str, str], *, title: str, status: str = "draft", content: str = "x"):
    r = client.post(DOCS.format(ws=ws.id), json={"title": title, "content": content, "status": status}, headers=auth)
    assert r.status_code == 201, r.text
    return r.json()


# ── obsidian：zip 读取边界 ──


class TestObsidianArchiveBoundaries:
    def test_normalize_entry_path_rejects_unsafe(self):
        assert obsidian._normalize_entry_path("") is None
        assert obsidian._normalize_entry_path("dir/") is None
        assert obsidian._normalize_entry_path("/abs.md") is None
        assert obsidian._normalize_entry_path("C:/win.md") is None
        assert obsidian._normalize_entry_path("a/../b.md") is None
        assert obsidian._normalize_entry_path(".") is None
        assert obsidian._normalize_entry_path("a\\b.md") == "a/b.md"
        assert obsidian._normalize_entry_path("./a/./b.md") == "a/b.md"

    def test_is_skipped_path_variants(self):
        assert obsidian._is_skipped_path(".obsidian/app.json") is True
        assert obsidian._is_skipped_path("node_modules/x/y.md") is True
        assert obsidian._is_skipped_path(".hidden.md") is True
        assert obsidian._is_skipped_path("note.md") is False
        assert obsidian._is_skipped_path("a/b/note.md") is False

    def test_common_root_variants(self):
        assert obsidian._common_root([]) == ""
        assert obsidian._common_root(["a/x.md", "b/y.md"]) == ""
        assert obsidian._common_root(["vault", "vault/y.md"]) == ""
        assert obsidian._common_root(["Vault/a.md", "Vault/b/c.md"]) == "Vault"

    def test_archive_rejects_empty(self):
        with pytest.raises(BadRequestError):
            obsidian.VaultArchive(b"")

    def test_archive_rejects_oversized(self, monkeypatch):
        monkeypatch.setattr(obsidian, "MAX_ARCHIVE_BYTES", 4)
        with pytest.raises(BadRequestError):
            obsidian.VaultArchive(b"12345")

    def test_archive_rejects_bad_zip(self):
        with pytest.raises(BadRequestError):
            obsidian.VaultArchive(b"not a zip at all")

    def test_archive_rejects_too_many_entries(self, monkeypatch):
        monkeypatch.setattr(obsidian, "MAX_ENTRIES", 1)
        with pytest.raises(BadRequestError):
            obsidian.VaultArchive(_build_zip({"a.md": "a", "b.md": "b"}))

    def test_archive_rejects_decompression_bomb(self, monkeypatch):
        monkeypatch.setattr(obsidian, "MAX_TOTAL_UNCOMPRESSED_BYTES", 4)
        with pytest.raises(BadRequestError):
            obsidian.VaultArchive(_build_zip({"a.md": "0123456789"}))

    def test_archive_skips_oversized_entry(self, monkeypatch):
        monkeypatch.setattr(obsidian, "MAX_ENTRY_BYTES", 4)
        archive = obsidian.VaultArchive(_build_zip({"a.md": "0123456789"}))
        assert archive.paths == []
        assert archive.skipped and "20MB" not in archive.skipped[0]["reason"]

    def test_archive_read_missing_entry(self):
        archive = obsidian.VaultArchive(_build_zip({"a.md": "a"}))
        with pytest.raises(BadRequestError):
            archive.read("missing.md")

    def test_archive_strips_common_root_and_labels_dirs(self):
        archive = obsidian.VaultArchive(_build_zip({"a.md": "x"}, root="MyVault"))
        assert archive.root_dir == "MyVault"
        assert archive.paths == ["a.md"]


class TestObsidianDecodeAndFrontMatter:
    def test_decode_markdown_gb18030_and_latin1(self):
        assert obsidian._decode_markdown("中文".encode("gb18030")) == "中文"
        assert obsidian._decode_markdown(b"\x80") == "\x80"

    def test_clean_scalar_strips_matching_quotes(self):
        assert obsidian._clean_scalar("'abc'") == "abc"
        assert obsidian._clean_scalar('"xy"') == "xy"
        assert obsidian._clean_scalar("  a b  ") == "a b"

    def test_as_text_and_as_list(self):
        assert obsidian._as_text(["a", "b"]) == ""
        assert obsidian._as_text(123) == "123"
        assert obsidian._as_list("a, b c") == ["a", "b", "c"]
        assert obsidian._as_list(["x", " y "]) == ["x", "y"]

    def test_parse_front_matter_edge_lines(self):
        fm = "\n".join(
            [
                "---",
                "# comment",
                "- orphan",
                ": no key",
                "plain text",
                "scalarfirst: hello",
                "  - listitem",
                "emptyitem:",
                "  - ",
                "quoted: 'single'",
                "inline: [x, y]",
                "---",
                "body",
            ]
        )
        meta, body = obsidian.parse_vault_front_matter(fm)
        assert body == "body"
        assert meta["scalarfirst"] == ["listitem"]
        assert meta["emptyitem"] == []
        assert meta["quoted"] == "single"
        assert meta["inline"] == ["x", "y"]
        assert "orphan" not in meta

    def test_parse_front_matter_without_block(self):
        meta, body = obsidian.parse_vault_front_matter("no front matter")
        assert meta == {}
        assert body == "no front matter"

    def test_normalize_tag_and_title_helpers(self):
        assert obsidian._normalize_tag_name("##标签 ") == "标签"
        assert obsidian._clean_title("") == "未命名"
        assert obsidian._clean_title("  a\n b ") == "a b"
        assert obsidian._safe_filename("a/b:c") == "a_b_c"


class TestObsidianIndex:
    def test_resolve_empty_variants(self):
        idx = obsidian._VaultIndex()
        assert idx.resolve_note("   ", ()) is None
        assert idx.resolve_attachment("   ", ()) is None

    def test_resolve_attachment_with_folder(self):
        idx = obsidian._VaultIndex()
        idx.attach_by_path["sub/x.png"] = "sub/x.png"
        idx.attach_by_name["x.png"] = "sub/x.png"
        assert idx.resolve_attachment("x.png", ("sub",)) == "sub/x.png"

    def test_title_of_falls_back_to_stem(self):
        idx = obsidian._VaultIndex()
        assert idx.title_of("dir/note.md") == "note"


class TestObsidianCanvas:
    def test_canvas_all_node_types(self):
        payload = {
            "nodes": [
                {"id": "n1", "type": "text", "text": "hi"},
                {"id": "n2", "type": "text"},
                "notadict",
                {"id": "n4", "type": "file"},
                {"id": "n5", "type": "file", "file": "a.png"},
                {"id": "n6", "type": "file", "file": "note.md", "subpath": "#h"},
                {"id": "n7", "type": "link", "url": "https://e.com"},
                {"id": "n8", "type": "link"},
                {"id": "n9", "type": "group", "label": "G"},
                {"id": "n10", "type": "group"},
            ],
            "edges": [
                {"fromNode": "n1", "toNode": "n5", "label": "L"},
                "notadict",
                {"fromNode": "n1", "toNode": "n7"},
            ],
        }
        md = obsidian._canvas_to_markdown(json.dumps(payload).encode(), "x.canvas")
        assert "## 卡片" in md
        assert "### 文本卡片 1" in md
        assert "_（空卡片）_" in md
        assert "![[a.png]]" in md
        assert "[[note.md#h]]" in md
        assert "- <https://e.com>" in md
        assert "### 分组：G" in md
        assert "### 分组：分组 10" in md
        assert "卡片 1 → 卡片 5（L）" in md
        assert "卡片 1 → 卡片 7" in md
        assert "## 原始 JSON" in md

    def test_canvas_empty_nodes_and_edges(self):
        md = obsidian._canvas_to_markdown(b'{"nodes": [], "edges": []}', "e.canvas")
        assert "## 卡片" not in md
        assert "## 连线" not in md

    def test_canvas_non_dict_payload(self):
        md = obsidian._canvas_to_markdown(b'"just a string"', "s.canvas")
        assert "结构无法识别" in md

    def test_base_to_markdown(self):
        md = obsidian._base_to_markdown(b"filters:\n  and: []\n", "v.base")
        assert "~~~yaml" in md


class TestObsidianRewrite:
    def test_embed_anchor_only_degrades(self):
        note = obsidian.VaultNote(path="n.md", folder=(), title="N", body="![[#anchor]]")
        idx = obsidian._VaultIndex()
        res = obsidian.rewrite_note_body(note, idx, store_attachment=lambda _t: "k")
        # 无目标嵌入先原样保留，随后 [[...]] 规则把 [[#anchor]] 退化为纯文本 "anchor"
        assert res.body == "!anchor"

    def test_embed_attachment_deduplicated(self):
        note = obsidian.VaultNote(path="n.md", folder=(), title="N", body="![[img.png]]\n\n![[img.png]]")
        idx = obsidian._VaultIndex()
        idx.attach_by_path["img.png"] = "img.png"
        res = obsidian.rewrite_note_body(note, idx, store_attachment=lambda _t: "key1")
        assert res.attachments == ["img.png"]
        assert res.body.count("file:key1") == 2

    def test_embed_unresolved_when_store_returns_none(self):
        note = obsidian.VaultNote(path="n.md", folder=(), title="N", body="![[img.png]]")
        idx = obsidian._VaultIndex()
        idx.attach_by_path["img.png"] = "img.png"
        res = obsidian.rewrite_note_body(note, idx, store_attachment=lambda _t: None)
        # 嵌入与随后的双链规则各处理一次同一路径，故缺失计数为 2
        assert res.missing_attachments == 2
        assert res.unresolved == ["![[img.png]]", "[[img.png]]"]
        assert res.body == "![[img.png]]"

    def test_md_image_unresolved_kept(self):
        note = obsidian.VaultNote(path="n.md", folder=(), title="N", body="![](missing.png)")
        idx = obsidian._VaultIndex()
        res = obsidian.rewrite_note_body(note, idx, store_attachment=lambda _t: "k")
        assert res.body == "![](missing.png)"

    def test_wikilink_path_prefix_stripped(self):
        note = obsidian.VaultNote(path="n.md", folder=(), title="N", body="[[dir/未知.md#h|a]]")
        idx = obsidian._VaultIndex()
        res = obsidian.rewrite_note_body(note, idx, store_attachment=lambda _t: "k")
        assert "[[未知#h|a]]" in res.body
        assert res.normalized_links == 1


class TestObsidianWarnings:
    def test_anchor_only_warning(self):
        notes = [obsidian.VaultNote(path="n.md", folder=(), title="N", body="[[#本文锚点]]")]
        archive = obsidian.VaultArchive(_build_zip({"n.md": "# N\n[[#x]]\n"}))
        warnings = obsidian._analyze_warnings(notes, archive)
        codes = {w["code"] for w in warnings}
        assert "anchor_only_links" in codes

    def test_referenced_attachments_branches(self):
        notes = [
            obsidian.VaultNote(
                path="s.md",
                folder=(),
                title="S",
                body="![[目标]]\n![[缺图.png]]\n![](https://e.com/a.png)\n![](缺图2.png)",
            ),
            obsidian.VaultNote(path="目标.md", folder=(), title="目标", body="x"),
        ]
        idx = obsidian._build_index(notes, [])
        found = obsidian._referenced_attachments(notes[0], idx)
        assert found == []

    def test_sibling_helpers(self):
        sibs = {0: {"a": SimpleNamespace()}}
        assert obsidian._existing_child(sibs, 0, "A") is not None
        assert obsidian._existing_child(sibs, 0, "zzz") is None
        assert obsidian._next_title(sibs, 0, "a") == "a (2)"
        assert obsidian._folder_chain(("a", "b")) == [("a",), ("a", "b")]

    def test_front_matter_value_and_render(self):
        assert obsidian._front_matter_value("") == '""'
        assert obsidian._front_matter_value("a: b") == '"a: b"'
        assert obsidian._front_matter_value("plain") == "plain"
        doc = SimpleNamespace(created_at=None, updated_at=None)
        fm = obsidian._render_front_matter("T", [], doc)
        assert "tags:" not in fm and "created:" not in fm and "updated:" not in fm
        assert obsidian._iso("not-a-datetime") == ""

    def test_batch_content_helpers_empty(self, db):
        assert obsidian._latest_contents(db, []) == {}
        assert obsidian._tags_by_document(db, []) == {}


# ── obsidian：analyze / import / export HTTP ──


class TestObsidianAnalyzeHttp:
    def test_analyze_defaults_and_attachment_resolution(self, client, ws, auth):
        data = _build_zip(
            {
                "目标.md": "# 目标\n",
                "来源.md": "![[目标]]\n![[缺图.png]]\n![](https://e.com/a.png)\n![](缺图2.png)\n",
                "a/x.md": "# x\n",
                "a/b/y.md": "# y\n",
            }
        )
        r = _post_zip(client, DOCS.format(ws=ws.id) + "/import/analyze", data, auth)
        assert r.status_code == 200, r.text
        payload = r.json()
        assert payload["total_folders"] == 2  # a 与 a/b

    def test_analyze_preview_limit(self, client, ws, auth):
        files = {f"n{i}.md": f"# n{i}\n" for i in range(305)}
        r = _post_zip(client, DOCS.format(ws=ws.id) + "/import/analyze", _build_zip(files), auth)
        assert r.status_code == 200, r.text
        payload = r.json()
        assert payload["total_notes"] == 305
        assert len(payload["items"]) == 300
        assert payload["preview_truncated"] is True


class TestObsidianImportHttp:
    def test_import_rejects_bad_status(self, client, ws, auth):
        r = _post_zip(client, DOCS.format(ws=ws.id) + "/import", _build_zip({"a.md": "# a\n"}), auth, status="weird")
        assert r.status_code == 400

    def test_import_rejects_empty_vault(self, client, ws, auth):
        r = _post_zip(client, DOCS.format(ws=ws.id) + "/import", _build_zip({"图片.png": b"PNG"}), auth)
        assert r.status_code == 400

    def test_import_nested_folders_and_duplicate_attachment(self, client, ws, auth):
        data = _build_zip(
            {
                "a/x.md": "# x\n",
                "a/b/y.md": "![[图.png]]\n\n![[图.png]]\n",
                "图.png": b"PNG",
            }
        )
        r = _post_zip(client, DOCS.format(ws=ws.id) + "/import", data, auth)
        assert r.status_code == 200, r.text
        stats = r.json()
        assert stats["created"] == 2
        assert stats["attachments"] == 1

    def test_import_note_title_collides_with_folder(self, client, ws, auth):
        client.post(DOCS.format(ws=ws.id), json={"title": "索引页", "content": "", "is_folder": True}, headers=auth)
        data = _build_zip({"索引页.md": "---\ntitle: 索引页\n---\n正文\n"})
        r = _post_zip(client, DOCS.format(ws=ws.id) + "/import", data, auth)
        assert r.status_code == 200, r.text
        assert r.json()["created"] == 1
        titles = {item["title"] for item in client.get(DOCS.format(ws=ws.id) + "/tree", headers=auth).json()["items"]}
        assert "索引页 (2)" in titles

    def test_import_overwrite_applies_tags(self, client, ws, auth):
        data = _build_zip({"索引页.md": "---\ntitle: 索引页\ntags: [项目]\n---\n正文\n"})
        assert _post_zip(client, DOCS.format(ws=ws.id) + "/import", data, auth).status_code == 200
        r = _post_zip(client, DOCS.format(ws=ws.id) + "/import", data, auth, conflict="overwrite")
        assert r.status_code == 200, r.text
        assert r.json()["updated"] == 1


class TestObsidianExportHttp:
    def test_export_name_collision(self, client, db, ws, owner, auth):
        for _ in range(2):
            db.add(
                Document(
                    workspace_id=ws.id,
                    title="重复命名文档",
                    content_type="markdown",
                    status="published",
                    created_by_id=owner.id,
                )
            )
        db.commit()
        r = client.get(DOCS.format(ws=ws.id) + "/export/vault", headers=auth)
        assert r.status_code == 200, r.text
        with zipfile.ZipFile(io.BytesIO(r.content)) as archive:
            names = archive.namelist()
        assert any(n.endswith("重复命名文档.md") for n in names)
        assert any(n.endswith("重复命名文档 (2).md") for n in names)

    def test_export_image_rewrite_and_missing(self, client, db, ws, owner, auth):
        key = "imgkey.png"
        updir = Path(settings.UPLOAD_DIR) / str(ws.id)
        updir.mkdir(parents=True, exist_ok=True)
        (updir / key).write_bytes(b"PNGDATA")
        content_service.create_document(
            db,
            ws.id,
            owner,
            title="图文",
            content=f"![a](file:{key})\n\n![b](file:{key})\n\n![c](file:missing.png)",
            status="published",
        )
        r = client.get(DOCS.format(ws=ws.id) + "/export/vault", headers=auth)
        assert r.status_code == 200, r.text
        with zipfile.ZipFile(io.BytesIO(r.content)) as archive:
            names = archive.namelist()
            note = archive.read(next(n for n in names if n.endswith("图文.md"))).decode("utf-8")
        assert f"![[attachments/{key}]]" in note
        assert "file:missing.png" in note


# ── content：错误分支与批量操作 ──


class TestContentServiceBranches:
    def test_get_any_state_or_404_missing(self, db, ws):
        with pytest.raises(NotFoundError):
            content_service.get_document_any_state_or_404(db, ws.id, 999999)

    def test_get_workspace_missing(self, db):
        with pytest.raises(NotFoundError):
            content_service._get_workspace(db, 999999)

    def test_load_content_inline_compressed(self):
        raw = "gz:" + base64.b64encode(gzip.compress(b"hello")).decode("ascii")
        version = SimpleNamespace(content=raw, content_size=5, version_number=1)
        doc = SimpleNamespace(id=1, content_path=None)
        assert content_service.load_content(version, doc) == "hello"

    def test_load_content_empty_version(self):
        assert content_service.load_content(None, SimpleNamespace(id=1, content_path=None)) == ""

    def test_create_document_parent_not_folder(self, db, ws, owner):
        parent = content_service.create_document(db, ws.id, owner, title="P")
        with pytest.raises(ConflictError):
            content_service.create_document(db, ws.id, owner, title="C", parent_id=parent.id)

    def test_update_content_skips_role_check(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T", content="v1")
        updated = content_service.update_content(db, ws.id, doc.id, owner, content="v2", enforce_role=False)
        assert updated.id == doc.id

    def test_update_content_unchanged_returns_early(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T", content="same")
        same = content_service.update_content(db, ws.id, doc.id, owner, content="same")
        assert content_service.version_count(db, doc.id) == 1
        assert same.id == doc.id

    def test_rollback_missing_target(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T", content="v1")
        with pytest.raises(NotFoundError):
            content_service.rollback_version(db, ws.id, doc.id, owner, 999)

    def test_move_target_reason_branches(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T")
        assert content_service.move_target_reason(db, ws.id, doc, None) is None
        assert content_service.move_target_reason(db, ws.id, doc, doc.id) == "不能将文档移动到自身"
        assert content_service.move_target_reason(db, ws.id, doc, 999999) == "目标文件夹不存在"

    def test_soft_delete_documents_skips_already_deleted(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T")
        content_service.soft_delete_document(db, ws.id, doc.id, owner)
        db.refresh(doc)
        assert content_service.soft_delete_documents(db, ws.id, [doc], owner) == 0

    def test_mark_deleted_subtree_on_deleted_node(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T")
        content_service.soft_delete_document(db, ws.id, doc.id, owner)
        db.refresh(doc)
        assert content_service._mark_deleted_subtree(db, doc, owner, dt.datetime.now(dt.UTC)) == 0

    def test_restore_keeps_parent_alive(self, db, ws, owner):
        folder = content_service.create_document(db, ws.id, owner, title="F", is_folder=True)
        child = content_service.create_document(db, ws.id, owner, title="C", parent_id=folder.id)
        content_service.soft_delete_document(db, ws.id, child.id, owner)
        _doc, restored, reparented = content_service.restore_document(db, ws.id, child.id, owner)
        assert restored == 1
        assert reparented is False

    def test_same_delete_batch_none_handling(self):
        now = dt.datetime.now(dt.UTC)
        assert content_service._same_delete_batch(None, None) is True
        assert content_service._same_delete_batch(None, now) is False

    def test_find_backlinks_with_viewer(self, db, ws, owner):
        content_service.create_document(db, ws.id, owner, title="目标")
        source = content_service.create_document(db, ws.id, owner, title="来源", content="见 [[目标]]")
        results = content_service.find_backlinks(db, ws.id, "目标", viewer=owner, is_admin=False)
        # 反链返回的是引用方（正文含 [[目标]]），不是被引用的目标文档本身
        assert [d.id for d in results] == [source.id]

    def test_resolve_outgoing_links_merges_duplicates(self, db, ws, owner):
        doc = content_service.create_document(
            db, ws.id, owner, title="S", content="[[#锚点]]\n[[目标]]\n[[目标#章节|别名]]"
        )
        links = content_service.resolve_outgoing_links(db, ws.id, doc)
        assert len(links) == 1
        assert links[0]["title"] == "目标"
        assert links[0]["anchor"] == "章节"
        assert links[0]["alias"] == "别名"

    def test_apply_tags_add_readd_remove(self, db, ws, owner):
        tag1 = Tag(workspace_id=ws.id, name="t1")
        tag2 = Tag(workspace_id=ws.id, name="t2")
        db.add_all([tag1, tag2])
        db.commit()
        doc = content_service.create_document(db, ws.id, owner, title="T")
        content_service._apply_tags(db, doc, [tag1.id], owner)
        content_service._apply_tags(db, doc, [tag1.id], owner)
        content_service._apply_tags(db, doc, [tag2.id], owner)
        db.commit()
        remaining = content_service.serialize_document(db, doc)["tag_ids"]
        assert remaining == [tag2.id]

    def test_normalize_ts_strips_timezone(self):
        assert content_service._normalize_ts("2024-01-01T00:00:00+08:00") == "2024-01-01 00:00:00"

    def test_collaborator_counts_empty(self, db):
        assert content_service._collaborator_counts(db, []) == {}


# ── content：documents 路由错误分支 ──


class TestDocumentsRouterBranches:
    def test_list_nonexistent_workspace(self, client, auth):
        r = client.get(DOCS.format(ws=999999), headers=auth)
        assert r.status_code == 404

    def test_list_with_parent_and_status_filters(self, client, ws, auth):
        _create_doc(client, ws, auth, title="根文档", status="draft")
        r = client.get(DOCS.format(ws=ws.id), params={"parent_id": 0, "status": "draft"}, headers=auth)
        assert r.status_code == 200, r.text

    def test_folder_content(self, client, ws, auth):
        r = client.post(DOCS.format(ws=ws.id), json={"title": "空文件夹", "is_folder": True}, headers=auth)
        assert r.status_code == 201, r.text
        folder = r.json()
        detail = client.get(f"{DOCS.format(ws=ws.id)}/{folder['id']}/content", headers=auth)
        assert detail.status_code == 200, detail.text
        assert detail.json()["content"] == ""
        assert detail.json()["draft"] is None

    def test_patch_sort_order_and_tags(self, client, db, ws, owner, auth):
        doc = _create_doc(client, ws, auth, title="可更新", content="body")
        tag = Tag(workspace_id=ws.id, name="补丁标签")
        db.add(tag)
        db.commit()
        r = client.patch(
            f"{DOCS.format(ws=ws.id)}/{doc['id']}",
            json={"sort_order": 5, "tag_ids": [tag.id]},
            headers=auth,
        )
        assert r.status_code == 200, r.text
        assert r.json()["sort_order"] == 5
        assert tag.id in r.json()["tag_ids"]

    def test_diff_missing_version(self, client, ws, auth):
        doc = _create_doc(client, ws, auth, title="差分子", content="v1")
        r = client.get(f"{DOCS.format(ws=ws.id)}/{doc['id']}/diff", params={"base": 999, "head": 1}, headers=auth)
        assert r.status_code == 404

    def test_editor_granted_direct_write(self, client, db, ws, owner, auth):
        # 工作区 editor + 关闭 allow_edit：常规保存走修改申请；被授予文档级 editor 后可直编草稿
        editor = _mk_user(db, "covdoc_weditor")
        db.add(WorkspaceMember(workspace_id=ws.id, user_id=editor.id, role=WorkspaceRole.EDITOR))
        db.commit()
        ws.allow_edit = False
        db.commit()
        editor_auth = _login(client, editor.username)
        doc = _create_doc(client, ws, auth, title="授予直编", status="draft", content="原始正文")
        db.add(DocumentPermission(document_id=doc["id"], user_id=editor.id, role="editor"))
        db.commit()
        r = client.put(
            f"{DOCS.format(ws=ws.id)}/{doc['id']}/content",
            json={"content": "被授予者直接编辑"},
            headers=editor_auth,
        )
        assert r.status_code == 200, r.text


# ── search ──


class TestSearchServiceBranches:
    def test_search_without_workspace(self, db, ws, owner):
        search.ensure_fts_schema(db)
        content_service.create_document(db, ws.id, owner, title="无关", content="正文")
        items, total = search.search_documents(db, "不存在的词")
        assert items == []
        assert total == 0

    def test_mark_empty_text(self):
        assert search._mark("x", "") == ""

    def test_mark_multiple_and_missing_terms(self):
        out = search._mark("流失 失踪", "客户流失预警，再次流失")
        assert out.count("<mark>流失</mark>") == 2
        assert "失踪" not in out

    def test_snippet_without_mark(self):
        assert search._snippet("no marks here") == "no marks here"

    def test_substring_no_terms(self, db):
        assert search._search_by_substring(db, "AND OR", [], {}, limit=5, offset=0) == ([], 0)

    def test_syntax_error_translated_to_bad_request(self, db, ws, owner, monkeypatch):
        content_service.create_document(db, ws.id, owner, title="X", content="y")
        monkeypatch.setattr(search, "escape_query", lambda _q: "(((")
        with pytest.raises(BadRequestError):
            search.search_documents(db, "x", workspace_id=ws.id)


# ── export: pdf / docx / pdf_fonts ──


class TestExportRenderers:
    def test_render_pdf_with_table_and_pagebreak(self):
        md = "| a | b |\n| --- | --- |\n| 1 | 2 |\n\ntext\n\n---PAGE---\n\n| c |\n| --- |\n| 3 |"
        out = pdf_export._render_pdf(md, title="T")
        assert out[:4] == b"%PDF"

    def test_make_pdf_table_ragged_rows(self):
        table = pdf_export._make_pdf_table([["a"], ["b", "c"]])
        assert table is not None
        empty_row = pdf_export._make_pdf_table([[], ["a"]])
        assert empty_row is not None

    def test_pdf_escape_and_rml(self):
        assert pdf_export._pdf_escape("a<b>&c") == "a&lt;b&gt;&amp;c"
        assert "<b>x</b>" in pdf_export._pdf_markdown_to_rml("**x**")

    def test_render_docx_with_table_pagebreak(self):
        out = docx_export._render_docx(
            "| a | b |\n| --- | --- |\n| 1 | 2 |\n\npara **b** `c`\n\n---PAGE---\n\n###### 六级\n\n| --- |\n",
            title="T",
        )
        assert out[:2] == b"PK"

    def test_render_docx_without_title(self):
        out = docx_export._render_docx("正文 **粗** `码`")
        assert out[:2] == b"PK"

    def test_add_rich_runs_edge_cases(self):
        from docx import Document

        para = Document().add_paragraph()
        docx_export._add_rich_runs(para, "`code` **bold** plain")
        assert para.text == "code bold plain"

    def test_docx_image_width_fallback(self):
        from docx.shared import Cm

        width = docx_export._docx_image_width(SimpleNamespace(sections=[]))
        assert width == Cm(14)


class TestPdfFonts:
    def test_is_cjk_family(self):
        assert pdf_fonts._is_cjk_family("") is False
        assert pdf_fonts._is_cjk_family("WenQuanYi Micro Hei") is True
        assert pdf_fonts._is_cjk_family("DejaVu Sans") is False

    def test_candidates_parsing(self, monkeypatch, tmp_path):
        real = tmp_path / "real.ttf"
        real.write_bytes(b"x")

        def fake_run(cmd, timeout=3.0):
            if cmd[0] == "fc-match":
                fam = cmd[1].split(":")[0]
                if fam == "WenQuanYi Micro Hei":
                    return 0, f"{real}|WenQuanYi Micro Hei|1"
                if fam == "WenQuanYi Zen Hei":
                    return 0, "no-pipes"
                if fam == "Source Han Sans CN":
                    return 0, f"{tmp_path / 'missing.ttf'}|Source Han Sans CN|0"
                if fam == "Source Han Sans SC":
                    return 0, f"{real}|DejaVu Sans|0"
                if fam == "Noto Sans SC":
                    return 0, f"{real}|Noto Sans SC|xx"
                return 1, ""
            if cmd[0] == "fc-list":
                return 0, ("\n" + f"{real}:Family:index=3\n" + "no-file.ttf:Family\n" + "no-colon-line\n")
            return -1, ""

        monkeypatch.setattr(pdf_fonts, "_run_command", fake_run)
        monkeypatch.setattr(pdf_fonts, "_LINUX_FONT_DIRS", [])
        cands = pdf_fonts._cjk_font_candidates()
        assert ("WQY-MicroHei", str(real), 1) in cands
        assert ("NotoSansSC", str(real), None) in cands

    def test_candidates_oserror_scanned_dir(self, monkeypatch, tmp_path):
        import pathlib

        monkeypatch.setattr(pdf_fonts, "_run_command", lambda cmd, timeout=3.0: (-1, ""))
        monkeypatch.setattr(pdf_fonts, "_LINUX_FONT_DIRS", [str(tmp_path)])
        target = str(tmp_path)
        orig_rglob = pathlib.Path.rglob

        def fake_rglob(self, pattern):
            if str(self) == target:
                raise OSError("boom")
            return orig_rglob(self, pattern)

        monkeypatch.setattr(pathlib.Path, "rglob", fake_rglob)
        assert isinstance(pdf_fonts._cjk_font_candidates(), list)

    def test_ensure_pdf_font_resets_on_error(self, monkeypatch):
        original = pdf_fonts._pdf_font_state[0]
        pdf_fonts._pdf_font_state[0] = None

        def boom():
            raise RuntimeError("no reportlab fonts")

        monkeypatch.setattr(pdf_fonts, "_cjk_font_candidates", boom)
        try:
            assert pdf_fonts._ensure_pdf_font() is None
        finally:
            pdf_fonts._pdf_font_state[0] = original


# ── routers：public / tags / vault ──


class TestPublicShareRouter:
    def test_expired_share(self, client, db, ws, owner, auth):
        doc = _create_doc(client, ws, auth, title="分享过期", status="published", content="X")
        db.add(
            DocumentShare(
                document_id=doc["id"],
                token="expired-token",
                is_active=True,
                expires_at=dt.datetime(2000, 1, 1),
            )
        )
        db.commit()
        r = client.get("/api/v1/documents/share/expired-token")
        assert r.status_code == 403

    def test_folder_share_without_content(self, client, db, ws, owner, auth):
        folder = client.post(
            DOCS.format(ws=ws.id), json={"title": "分享文件夹", "is_folder": True, "status": "published"}, headers=auth
        ).json()
        db.add(DocumentShare(document_id=folder["id"], token="folder-token", is_active=True))
        db.commit()
        r = client.get("/api/v1/documents/share/folder-token")
        assert r.status_code == 200, r.text
        assert "content" not in r.json()

    def test_password_required_for_get(self, client, db, ws, auth):
        doc = _create_doc(client, ws, auth, title="加密分享", status="published", content="X")
        db.add(
            DocumentShare(
                document_id=doc["id"],
                token="pwd-token",
                is_active=True,
                password_hash="$2b$04$abcdefghijklmnopqrstuv",
            )
        )
        db.commit()
        r = client.get("/api/v1/documents/share/pwd-token")
        assert r.status_code == 403

    def test_post_password_on_open_share(self, client, db, ws, auth):
        doc = _create_doc(client, ws, auth, title="无需密码", status="published", content="X")
        db.add(DocumentShare(document_id=doc["id"], token="open-token", is_active=True))
        db.commit()
        r = client.post("/api/v1/documents/share/open-token", json={"password": "anything"})
        assert r.status_code == 400

    def test_password_share_unpublished_doc(self, client, db, ws, auth):
        import bcrypt

        doc = _create_doc(client, ws, auth, title="未发布加密", status="draft", content="X")
        pwd_hash = bcrypt.hashpw(b"secret", bcrypt.gensalt(4)).decode()
        db.add(DocumentShare(document_id=doc["id"], token="draft-pwd", is_active=True, password_hash=pwd_hash))
        db.commit()
        r = client.post("/api/v1/documents/share/draft-pwd", json={"password": "secret"})
        assert r.status_code == 404

    def test_doc_workspace_missing(self, client, db, ws, auth):
        db.add(DocumentShare(document_id=999999, token="ghost-token", is_active=True))
        db.commit()
        r = client.get("/api/v1/documents/share/ghost-token")
        assert r.status_code == 404


class TestTagsRouterBranches:
    TAG_API = "/api/v1/workspaces/{ws}/document-tags"

    def test_list_nonexistent_workspace(self, client, auth):
        r = client.get(self.TAG_API.format(ws=999999), headers=auth)
        assert r.status_code == 404

    def test_list_non_member(self, client, ws, auth_headers):
        r = client.get(self.TAG_API.format(ws=ws.id), headers=auth_headers)
        assert r.status_code == 403

    def test_patch_missing_tag(self, client, ws, auth):
        r = client.patch(self.TAG_API.format(ws=ws.id) + "/999999", json={"color": "#111111"}, headers=auth)
        assert r.status_code == 404

    def test_patch_description_only(self, client, ws, auth):
        tag = client.post(self.TAG_API.format(ws=ws.id), json={"name": "desc-tag"}, headers=auth).json()
        r = client.patch(self.TAG_API.format(ws=ws.id) + f"/{tag['id']}", json={"description": "说明"}, headers=auth)
        assert r.status_code == 200, r.text
        assert r.json()["description"] == "说明"

    def test_patch_color_only(self, client, ws, auth):
        tag = client.post(self.TAG_API.format(ws=ws.id), json={"name": "color-tag"}, headers=auth).json()
        r = client.patch(self.TAG_API.format(ws=ws.id) + f"/{tag['id']}", json={"color": "#123456"}, headers=auth)
        assert r.status_code == 200, r.text
        assert r.json()["color"] == "#123456"

    def test_patch_rename_conflict(self, client, ws, auth):
        client.post(self.TAG_API.format(ws=ws.id), json={"name": "原名"}, headers=auth)
        other = client.post(self.TAG_API.format(ws=ws.id), json={"name": "另一个"}, headers=auth).json()
        r = client.patch(self.TAG_API.format(ws=ws.id) + f"/{other['id']}", json={"name": "原名"}, headers=auth)
        assert r.status_code == 409


class TestVaultRouterBranches:
    ANALYZE = DOCS + "/import/analyze"

    def test_empty_upload(self, client, ws, auth):
        r = client.post(
            self.ANALYZE.format(ws=ws.id),
            files={"file": ("vault.zip", b"", "application/zip")},
            headers=auth,
        )
        assert r.status_code == 400

    def test_nonexistent_workspace(self, client, auth):
        r = _post_zip(client, self.ANALYZE.format(ws=999999), _build_zip({"a.md": "# a\n"}), auth)
        assert r.status_code == 404

    def test_non_member(self, client, ws, auth_headers):
        r = _post_zip(client, self.ANALYZE.format(ws=ws.id), _build_zip({"a.md": "# a\n"}), auth_headers)
        assert r.status_code == 403


# ── parse 补充 ──


class TestParseExtra:
    def test_parse_front_matter_skips_blank_comment_and_no_colon(self):
        text = "---\n\n# c\nnokey value\nkey: value\n---\nbody"
        meta, body = parse.parse_front_matter(text)
        assert meta == {"key": "value"}
        assert body == "body"

    def test_rewrite_wikilink_noop_when_empty_old(self):
        assert parse.rewrite_wikilink_target("[[a]]", "", "b") == ("[[a]]", 0)

    def test_iter_wikilinks_skips_fenced(self):
        content = "```\n[[fenced]]\n```\n[[real]]\n"
        assert parse.extract_wikilinks(content) == ["real"]

    def test_unified_diff_counts(self):
        _diff, added, removed = parse.unified_diff("a\n", "a\nb\n")
        assert added == 1
        assert removed == 0


# ── obsidian：更多细分分支 ──


class TestObsidianExtraBranches:
    def test_archive_skips_illegal_and_system_paths(self):
        archive = obsidian.VaultArchive(_build_zip({"../evil.md": "x", ".obsidian/app.json": "y", "ok.md": "z"}))
        assert archive.paths == ["ok.md"]
        reasons = [entry["reason"] for entry in archive.skipped]
        assert any("非法" in r for r in reasons)
        assert any("隐藏" in r for r in reasons)

    def test_size_of_registered_and_missing(self):
        archive = obsidian.VaultArchive(_build_zip({"a.md": "abc"}))
        assert archive.size_of("a.md") == 3
        assert archive.size_of("missing.md") == 0

    def test_parse_front_matter_block_list_skips_empty_and_reuses_bucket(self):
        fm = "\n".join(["---", "tags:", "  - ''", "  - a", "  - b", "---", "body"])
        meta, body = obsidian.parse_vault_front_matter(fm)
        assert meta["tags"] == ["a", "b"]
        assert body == "body"

    def test_canvas_invalid_json_keeps_raw(self):
        md = obsidian._canvas_to_markdown(b"{not json", "b.canvas")
        assert "JSON 解析失败" in md

    def test_embed_resolves_note_and_leaves_unresolved(self):
        notes = [obsidian.VaultNote(path="目标.md", folder=(), title="目标", body="x")]
        idx = obsidian._build_index(notes, [])
        note = obsidian.VaultNote(path="n.md", folder=(), title="N", body="![[目标]]\n\n![[没有]]")
        res = obsidian.rewrite_note_body(note, idx, store_attachment=lambda _t: "k")
        assert "[[目标]]" in res.body
        assert "![[没有]]" in res.unresolved

    def test_wikilink_resolves_note_and_attachment(self):
        notes = [obsidian.VaultNote(path="dir/Note.md", folder=("dir",), title="Note", body="x")]
        idx = obsidian._build_index(notes, ["img.png"])
        note = obsidian.VaultNote(path="n.md", folder=(), title="N", body="[[dir/Note.md]]\n\n[[img.png]]")
        res = obsidian.rewrite_note_body(note, idx, store_attachment=lambda _t: "key")
        assert "[[Note]]" in res.body
        assert "![img.png](file:key)" in res.body


class TestContentCycleGuards:
    def test_restore_subtree_visit_guard(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T")
        assert content_service._restore_subtree(db, doc, None, visited={doc.id}) == 0

    def test_collect_purge_targets_visit_guard(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T")
        assert content_service._collect_purge_targets(db, doc, [], {doc.id}) == 0


class TestSearchEscapeAndCache:
    def test_reset_fts_schema_cache(self, db):
        search.ensure_fts_schema(db)
        search.reset_fts_schema_cache()
        assert search._fts_state["ready"] is False

    def test_escape_query_unclosed_phrase_prefix_and_symbols(self):
        assert search.escape_query('"未闭合短语') == '"未闭合短语"'
        assert search.escape_query("pre*") == '"pre"*'
        assert search.escape_query("ok ---") == '"ok"'

    def test_escape_query_boolean_positions(self):
        assert search.escape_query("a AND b") == '"a" AND "b"'
        assert search.escape_query("AND a") == '"AND" "a"'
        assert search.escape_query("a AND") == '"a" "AND"'

    def test_escape_query_rejects_empty_and_symbol_only(self):
        with pytest.raises(BadRequestError):
            search.escape_query("   ")
        with pytest.raises(BadRequestError):
            search.escape_query("()[]{}")


class TestVisibilityServiceBranches:
    def test_is_document_admin_superuser(self, db, ws):
        su = _mk_user(db, "covdoc_su")
        su.is_superuser = True
        db.commit()
        assert visibility.is_document_admin(db, ws.id, su) is True

    def test_visible_document_ids_empty_and_filtered(self, db, ws, owner):
        assert visibility.visible_document_ids(db, ws.id, owner, []) == set()
        doc = content_service.create_document(db, ws.id, owner, title="V", status="draft")
        assert visibility.visible_document_ids(db, ws.id, owner, [doc.id]) == {doc.id}
        outsider = _mk_user(db, "covdoc_vis_outsider")
        assert visibility.visible_document_ids(db, ws.id, outsider, [doc.id]) == set()

    def test_is_workspace_owner_missing_and_member(self, db, ws, owner):
        assert visibility.is_workspace_owner(db, 999999, owner) is False
        assert visibility.is_workspace_owner(db, ws.id, owner) is True


class TestDocumentPermissionsBranches:
    def test_workspace_missing(self, db):
        with pytest.raises(NotFoundError):
            doc_perms._workspace(db, 999999)

    def test_can_manage_superuser(self, db, ws, owner):
        su = _mk_user(db, "covdoc_perm_su2")
        su.is_superuser = True
        db.commit()
        doc = content_service.create_document(db, ws.id, owner, title="T")
        assert doc_perms.can_manage(db, doc, su) is True

    def test_granted_user_ids(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T")
        grantee = _mk_user(db, "covdoc_grantee")
        db.add(DocumentPermission(document_id=doc.id, user_id=grantee.id, role="viewer"))
        db.commit()
        assert doc_perms.granted_user_ids(db, doc.id) == {grantee.id}

    def test_upsert_permission_guard_and_update_branches(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T")
        member = _mk_user(db, "covdoc_perm_member")
        db.add(WorkspaceMember(workspace_id=ws.id, user_id=member.id, role=WorkspaceRole.VIEWER))
        outsider = _mk_user(db, "covdoc_perm_outsider")
        su = _mk_user(db, "covdoc_perm_su")
        su.is_superuser = True
        db.commit()

        with pytest.raises(BadRequestError):
            doc_perms.upsert_permission(db, doc, operator=owner, target_user_id=member.id, role="bogus")
        with pytest.raises(BadRequestError):
            doc_perms.upsert_permission(db, doc, operator=owner, target_user_id=owner.id, role="viewer")
        with pytest.raises(BadRequestError):
            doc_perms.upsert_permission(db, doc, operator=su, target_user_id=owner.id, role="viewer")
        with pytest.raises(NotFoundError):
            doc_perms.upsert_permission(db, doc, operator=owner, target_user_id=999999, role="viewer")
        with pytest.raises(BadRequestError):
            doc_perms.upsert_permission(db, doc, operator=owner, target_user_id=outsider.id, role="viewer")

        doc_perms.upsert_permission(db, doc, operator=owner, target_user_id=member.id, role="viewer")
        updated = doc_perms.upsert_permission(db, doc, operator=owner, target_user_id=member.id, role="editor")
        assert updated["role"] == "editor"

    def test_remove_permission_missing_returns_false(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T")
        member = _mk_user(db, "covdoc_remove_member")
        db.add(WorkspaceMember(workspace_id=ws.id, user_id=member.id, role=WorkspaceRole.VIEWER))
        db.commit()
        assert doc_perms.remove_permission(db, doc, operator=owner, target_user_id=member.id) is False


class TestDraftsServiceBranches:
    def test_serialize_draft_none_and_present(self, db, ws, owner):
        assert drafts.serialize_draft(None) is None
        doc = content_service.create_document(db, ws.id, owner, title="T", content="v1")
        saved = drafts.save_draft(db, ws.id, doc.id, owner, content="v2")
        assert drafts.serialize_draft(saved)["content"] == "v2"

    def test_draft_payload_unchanged_is_none(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T", content="v1")
        db.add(
            DocumentDraft(
                document_id=doc.id,
                user_id=owner.id,
                content="v1",
                content_size=2,
                content_hash=parse.content_hash("v1"),
                base_version_number=1,
            )
        )
        db.commit()
        assert drafts.draft_payload(db, doc, owner) is None

    def test_submit_draft_folder_conflict(self, db, ws, owner):
        folder = content_service.create_document(db, ws.id, owner, title="F", is_folder=True)
        with pytest.raises(ConflictError):
            drafts.submit_draft(db, ws.id, folder.id, owner, content="x")

    def test_submit_draft_base_above_latest(self, db, ws, owner):
        doc = content_service.create_document(db, ws.id, owner, title="T", content="v1")
        with pytest.raises(BadRequestError):
            drafts.submit_draft(db, ws.id, doc.id, owner, content="v2", base_version_number=99)


class TestParseFenceHandling:
    def test_rewrite_wikilink_target_skips_fenced_region(self):
        content = "```\n[[old]]\n```\n[[old]]\n"
        new_content, count = parse.rewrite_wikilink_target(content, "old", "new")
        assert count == 1
        assert "```\n[[old]]\n```" in new_content
        assert "[[new]]" in new_content


class TestDocxExtraBranches:
    def test_render_ending_with_table_flushes(self):
        out = docx_export._render_docx("| a | b |\n| --- | --- |\n| 1 | 2 |")
        assert out[:2] == b"PK"

    def test_heading_fallback_when_style_missing(self, monkeypatch):
        from docx.document import Document as DocxDocument

        def _boom(self, text, level=1):
            raise KeyError("no such style")

        monkeypatch.setattr(DocxDocument, "add_heading", _boom)
        out = docx_export._render_docx("# 标题\n\n正文")
        assert out[:2] == b"PK"
