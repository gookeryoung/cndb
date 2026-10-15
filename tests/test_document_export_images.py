"""导出时嵌入站内图片测试.

覆盖：`file:` 引用解析与安全校验（路径穿越 / 缺失 / 非法 key）、DOCX 真实嵌入图片、
PDF 嵌入图片 XObject、引用缺失时降级为占位文字、Markdown 导出保留引用不改写、
以及经导出接口的端到端验证（上传文件 → 导出 docx 内含图片）。
"""

from __future__ import annotations

import io
import zipfile

import pytest
from fastapi.testclient import TestClient
from sqlalchemy.orm import Session

from cndb.core.config import settings
from cndb.plugins.accounts.models import User
from cndb.plugins.documents.services.export import render_export
from cndb.plugins.documents.services.export.images import (
    is_safe_file_key,
    make_image_resolver,
    referenced_file_keys,
    resolve_upload_path,
)
from cndb.plugins.workspaces.models import Workspace, WorkspaceMember, WorkspaceRole

DOCS = "/api/v1/workspaces/{ws}/documents"

# 1x1 PNG（合法图片字节，供嵌入测试使用）
PNG_1PX = bytes.fromhex(
    "89504e470d0a1a0a0000000d49484452000000010000000108060000001f15c489"
    "0000000a49444154789c6300010000050001a5f645400000000049454e44ae426082"
)

IMAGE_LINE = "![示意图](file:abc123.png)"
TEXT_WITH_IMAGE = f"# 标题\n\n正文段落。\n\n{IMAGE_LINE}\n\n结尾段落。\n"


def _login(client: TestClient, username: str) -> dict[str, str]:
    r = client.post("/api/v1/accounts/auth/login", json={"login": username, "password": "passw0rd"})
    assert r.status_code == 200, r.text
    return {"Authorization": f"Bearer {r.json()['access_token']}"}


def _make_user(db: Session, username: str) -> User:
    u = User(username=username, nickname=username)
    u.set_password("passw0rd")
    db.add(u)
    db.commit()
    db.refresh(u)
    return u


@pytest.fixture
def owner(db) -> User:
    return _make_user(db, "exp_owner")


@pytest.fixture
def ws(db, owner) -> Workspace:
    w = Workspace(name="ExportWS", created_by_id=owner.id)
    db.add(w)
    db.flush()
    db.add(WorkspaceMember(workspace_id=w.id, user_id=owner.id, role=WorkspaceRole.OWNER))
    db.commit()
    db.refresh(w)
    return w


@pytest.fixture
def image_resolver(ws):
    """在工作区上传目录里放一张真实图片，并返回解析器"""
    upload_dir = settings.UPLOAD_DIR / str(ws.id)
    upload_dir.mkdir(parents=True, exist_ok=True)
    (upload_dir / "abc123.png").write_bytes(PNG_1PX)
    return make_image_resolver(settings.UPLOAD_DIR, ws.id)


class TestResolvePath:
    def test_safe_key_rules(self):
        assert is_safe_file_key("abc123.png")
        for bad in ("../secret.png", "a/b.png", "a\\b.png", ".hidden", "", "x" * 300):
            assert not is_safe_file_key(bad), bad

    def test_resolve_inside_workspace_dir(self, ws):
        upload_dir = settings.UPLOAD_DIR / str(ws.id)
        upload_dir.mkdir(parents=True, exist_ok=True)
        target = upload_dir / "ok.png"
        target.write_bytes(PNG_1PX)

        assert resolve_upload_path(settings.UPLOAD_DIR, ws.id, "ok.png") == target.resolve()

    def test_traversal_and_missing_return_none(self, ws):
        upload_dir = settings.UPLOAD_DIR / str(ws.id)
        upload_dir.mkdir(parents=True, exist_ok=True)
        outside = settings.UPLOAD_DIR / "outside.png"
        outside.write_bytes(PNG_1PX)

        # 穿越尝试一律拒绝（即便目标文件真实存在）
        assert resolve_upload_path(settings.UPLOAD_DIR, ws.id, "../outside.png") is None
        assert resolve_upload_path(settings.UPLOAD_DIR, ws.id, "not-there.png") is None

    def test_referenced_keys_dedup_in_order(self):
        md = f"{IMAGE_LINE}\n\n![b](file:def.png)\n\n{IMAGE_LINE}\n"
        assert referenced_file_keys(md) == ["abc123.png", "def.png"]


class TestDocxImages:
    def test_embed_real_image(self, image_resolver):
        data = render_export(TEXT_WITH_IMAGE, "docx", "标题", images=image_resolver)
        from docx import Document

        doc = Document(io.BytesIO(data))
        assert len(doc.inline_shapes) == 1
        texts = "\n".join(p.text for p in doc.paragraphs)
        assert "图片缺失" not in texts

    def test_missing_file_falls_back_to_placeholder(self):
        data = render_export(TEXT_WITH_IMAGE, "docx", "标题", images=lambda _k: None)
        from docx import Document

        doc = Document(io.BytesIO(data))
        assert len(doc.inline_shapes) == 0
        assert "图片缺失：示意图" in "\n".join(p.text for p in doc.paragraphs)

    def test_without_resolver_keeps_markdown_line(self):
        # 未传解析器（如历史调用方）时不抛错，仍输出占位文字
        data = render_export(TEXT_WITH_IMAGE, "docx", "标题")
        from docx import Document

        doc = Document(io.BytesIO(data))
        assert len(doc.inline_shapes) == 0

    def test_broken_image_bytes_fall_back(self, ws):
        upload_dir = settings.UPLOAD_DIR / str(ws.id)
        upload_dir.mkdir(parents=True, exist_ok=True)
        (upload_dir / "broken.png").write_bytes(b"not-an-image")

        data = render_export(
            "![坏图](file:broken.png)", "docx", "标题", images=make_image_resolver(settings.UPLOAD_DIR, ws.id)
        )
        from docx import Document

        doc = Document(io.BytesIO(data))
        assert len(doc.inline_shapes) == 0
        assert "图片无法嵌入：坏图" in "\n".join(p.text for p in doc.paragraphs)


class TestPdfImages:
    def test_embed_real_image(self, image_resolver):
        data = render_export(TEXT_WITH_IMAGE, "pdf", "标题", images=image_resolver)
        assert data.startswith(b"%PDF")
        # reportlab 把位图作为 XObject 图像写入
        assert b"/Subtype /Image" in data or b"/Subtype/Image" in data

    def test_missing_file_placeholder_only(self):
        data = render_export(TEXT_WITH_IMAGE, "pdf", "标题", images=lambda _k: None)
        assert data.startswith(b"%PDF")
        assert b"/Subtype /Image" not in data and b"/Subtype/Image" not in data

    def test_broken_image_bytes_fall_back(self, ws):
        upload_dir = settings.UPLOAD_DIR / str(ws.id)
        upload_dir.mkdir(parents=True, exist_ok=True)
        (upload_dir / "broken.png").write_bytes(b"not-an-image")

        data = render_export(
            "![坏图](file:broken.png)", "pdf", "标题", images=make_image_resolver(settings.UPLOAD_DIR, ws.id)
        )
        assert data.startswith(b"%PDF")
        assert b"/Subtype /Image" not in data and b"/Subtype/Image" not in data


class TestMarkdownImages:
    def test_markdown_keeps_internal_reference(self):
        data = render_export(TEXT_WITH_IMAGE, "md", "标题")
        assert IMAGE_LINE.encode("utf-8") in data


class TestExportEndpointEmbedsImage:
    def test_export_docx_contains_embedded_image(self, client, ws, owner, db):
        auth = _login(client, owner.username)
        # 上传一张图片（走真实上传接口）
        up = client.post(
            f"/api/v1/workspaces/{ws.id}/files",
            files={"file": ("图表.png", PNG_1PX, "image/png")},
            headers=auth,
        )
        assert up.status_code == 201, up.text
        file_key = up.json()["file_key"]

        # 创建引用该图片的文档（直接写库，避免走改变更流程）
        from cndb.plugins.documents.services import content as content_service

        doc = content_service.create_document(
            db, ws.id, owner, title="带图文档", content=f"# 带图\n\n![图表](file:{file_key})\n"
        )
        assert doc is not None

        r = client.get(f"{DOCS.format(ws=ws.id)}/{doc.id}/export", params={"format": "docx"}, headers=auth)
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("application/vnd.openxmlformats")

        # docx 是 zip：检查 media 目录里确实打包了图片
        with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
            media = [n for n in zf.namelist() if n.startswith("word/media/")]
        assert media, "导出 docx 未包含图片资源"

    def test_export_pdf_with_image(self, client, ws, owner, db):
        auth = _login(client, owner.username)
        up = client.post(
            f"/api/v1/workspaces/{ws.id}/files",
            files={"file": ("图表.png", PNG_1PX, "image/png")},
            headers=auth,
        )
        file_key = up.json()["file_key"]

        from cndb.plugins.documents.services import content as content_service

        doc = content_service.create_document(db, ws.id, owner, title="带图 PDF", content=f"![图表](file:{file_key})\n")

        r = client.get(f"{DOCS.format(ws=ws.id)}/{doc.id}/export", params={"format": "pdf"}, headers=auth)
        assert r.status_code == 200, r.text
        assert r.content.startswith(b"%PDF")
        assert b"/Subtype /Image" in r.content or b"/Subtype/Image" in r.content


class TestMarkdownArchive:
    def test_rewrite_collects_assets_and_keeps_unresolved_refs(self, image_resolver):
        from cndb.plugins.documents.services.export.images import rewrite_markdown_images

        md = "图片：\n\n![示意图](file:abc123.png)\n\n![缺失](file:missing.png)\n\n![重复](file:abc123.png)\n"
        rewritten, assets = rewrite_markdown_images(md, image_resolver)

        assert "![示意图](assets/abc123.png)" in rewritten
        assert "![重复](assets/abc123.png)" in rewritten
        # 解析不到的引用原样保留（不静默丢信息）
        assert "![缺失](file:missing.png)" in rewritten
        assert list(assets) == ["assets/abc123.png"]
        assert assets["assets/abc123.png"] == PNG_1PX

    def test_archive_contains_markdown_and_assets(self, image_resolver):
        from cndb.plugins.documents.services.export import render_markdown_archive

        data = render_markdown_archive(TEXT_WITH_IMAGE, "标题", image_resolver, "标题")
        assert data is not None
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            names = zf.namelist()
            assert "标题.md" in names
            assert "assets/abc123.png" in names
            assert zf.read("assets/abc123.png") == PNG_1PX
            inner = zf.read("标题.md").decode("utf-8")
        assert "![示意图](assets/abc123.png)" in inner
        assert "file:" not in inner

    def test_no_resolvable_images_returns_none(self):
        from cndb.plugins.documents.services.export import render_markdown_archive

        assert render_markdown_archive("# 纯文字\n", "标题", lambda _k: None, "标题") is None
        assert render_markdown_archive("![x](file:missing.png)", "标题", lambda _k: None, "标题") is None
        assert render_markdown_archive("![x](https://e.com/a.png)", "标题", lambda _k: None, "标题") is None


class TestMarkdownExportEndpoint:
    def _make_doc_with_image(self, client, ws, owner, db):
        auth = _login(client, owner.username)
        up = client.post(
            f"/api/v1/workspaces/{ws.id}/files",
            files={"file": ("图表.png", PNG_1PX, "image/png")},
            headers=auth,
        )
        assert up.status_code == 201, up.text
        file_key = up.json()["file_key"]

        from cndb.plugins.documents.services import content as content_service

        doc = content_service.create_document(
            db, ws.id, owner, title="带图 Markdown", content=f"# 带图\n\n![图表](file:{file_key})\n"
        )
        return auth, doc, file_key

    def test_markdown_with_images_returns_zip(self, client, ws, owner, db):
        auth, doc, file_key = self._make_doc_with_image(client, ws, owner, db)

        r = client.get(f"{DOCS.format(ws=ws.id)}/{doc.id}/export", params={"format": "md"}, headers=auth)
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("application/zip")
        assert ".zip" in r.headers["content-disposition"]

        with zipfile.ZipFile(io.BytesIO(r.content)) as zf:
            names = zf.namelist()
            assert f"assets/{file_key}" in names
            md_name = next(n for n in names if n.endswith(".md"))
            inner = zf.read(md_name).decode("utf-8")
        assert f"![图表](assets/{file_key})" in inner

    def test_markdown_without_images_stays_plain(self, client, ws, owner, db):
        auth = _login(client, owner.username)
        from cndb.plugins.documents.services import content as content_service

        doc = content_service.create_document(db, ws.id, owner, title="纯文字", content="# 纯文字\n")

        r = client.get(f"{DOCS.format(ws=ws.id)}/{doc.id}/export", params={"format": "md"}, headers=auth)
        assert r.status_code == 200, r.text
        assert r.headers["content-type"].startswith("text/markdown")
        assert r.content.decode("utf-8") == "# 纯文字\n"
