"""补充覆盖：transfer 子模块（wps_images / table_create / spreadsheet / normalize/*）未覆盖分支.

仅新增测试，不改动源码与既有测试。覆盖目标：
- wps_images：media 路径解析变体、坏 zip、rels/cellimages 异常与缺映射、落盘器拒收分支。
- table_create：JSON 列样本上限/序列化失败、dict→csv 归一、空输入、覆盖 options、
  短行 None 归一、半残表清理（含 link 表）。
- spreadsheet：chardet 兜底各分支、Sniffer 方言异类回落、格式推断与 JSON/分隔符解析、
  XLSX 空表返回。
- normalize/*：数值归一边界、compact/en 日期回落、列级拆分守卫、Python 类型映射长值。
"""

from __future__ import annotations

import base64
import io
import json
import sys
import types
import zipfile

import pytest
from openpyxl import Workbook

from cndb.plugins.tables.models import DataField, DataTable
from cndb.plugins.tables.services.transfer import (
    WpsCellImage,
    _analyze_dict_rows_as_csv,
    _apply_column_overrides,
    _check_compact_date,
    _check_en_date,
    _check_long_integer,
    _cleanup_partial_table,
    _coerce_long_numeric_to_text,
    _is_integer,
    _normalize_numeric,
    _parse_delimited_text,
    _parse_json_text,
    _parse_xlsx_bytes,
    _python_type_to_field_type,
    _split_json_array_like,
    _split_list_like,
    analyze_file_columns,
    analyze_json_columns,
    create_table_from_file,
    create_table_from_json_data,
    decode_bytes_auto,
    ephemeral_image_saver,
    import_rows_from_csv,
    infer_column_type,
    make_displimg_resolver,
    make_workspace_image_saver,
    parse_file_to_rows,
    parse_wps_cell_images,
    sniff_csv_delimiter,
)
from cndb.plugins.tables.services.transfer import spreadsheet as spreadsheet_mod
from cndb.plugins.tables.services.transfer import wps_images as wps_mod

# 1x1 透明 PNG
_PNG_1PX = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)

_NS_XDR = "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"
_NS_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# ── xlsx / wps 部件构造工具 ────────────────────────


def _base_xlsx(headers: list[str], rows: list[list[object]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _empty_xlsx() -> bytes:
    wb = Workbook()
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _zip_with_parts(base: bytes, extra: dict[str, str | bytes]) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(base)) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            dst.writestr(item, src.read(item.filename))
        for name, data in extra.items():
            dst.writestr(name, data)
    return out.getvalue()


def _wps_xlsx(
    *,
    cellimages_xml: str | None = None,
    rels_xml: str | None = None,
    media: dict[str, bytes] | None = None,
) -> bytes:
    base = _base_xlsx(["h"], [["x"]])
    extra: dict[str, str | bytes] = {}
    if cellimages_xml is not None:
        extra["xl/cellimages.xml"] = cellimages_xml
    if rels_xml is not None:
        extra["xl/_rels/cellimages.xml.rels"] = rels_xml
    for name, data in (media or {}).items():
        extra[name] = data
    return _zip_with_parts(base, extra)


def _pic(name: str | None, embed: str | None, blip_attrs: str = "") -> str:
    name_attr = f' name="{name}"' if name is not None else ""
    blip = f'<a:blip {blip_attrs} r:embed="{embed}"/>' if embed is not None else ""
    return (
        f'<etc:cellimage><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="1"{name_attr}/></xdr:nvPicPr>'
        f"<xdr:blipFill>{blip}</xdr:blipFill></xdr:pic></etc:cellimage>"
    )


def _cellimages_xml(pics: str) -> str:
    return (
        f'<etc:cellImages xmlns:etc="x" xmlns:xdr="{_NS_XDR}" xmlns:a="{_NS_A}" xmlns:r="{_NS_R}">'
        + pics
        + "</etc:cellImages>"
    )


def _rels_xml(entries: str) -> str:
    return (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        + entries
        + "</Relationships>"
    )


def _rel(rid: str, target: str) -> str:
    return f'<Relationship Id="{rid}" Type="http://x/image" Target="{target}"/>'


# ── wps_images ─────────────────────────────────────


class TestWpsImagesEdgeCases:
    def test_parse_bad_zip_returns_empty(self):
        """非 zip 字节 → BadZipFile 被吞掉返回空映射（覆盖 122-124）."""
        assert parse_wps_cell_images(b"this is not a zip archive") == {}

    def test_resolve_media_path_variants(self):
        """rels 相对 Target 的 ``.`` / 空段 / ``..`` 归一（覆盖 104、106-108）."""
        raw = _wps_xlsx(
            cellimages_xml=_cellimages_xml(_pic("ID1", "rId1") + _pic("ID2", "rId2") + _pic("ID3", "rId3")),
            rels_xml=_rels_xml(
                _rel("rId1", ".//media/image1.png")
                + _rel("rId2", "../../media/image2.png")
                + _rel("rId3", "../media/image3.png")
            ),
            media={
                "xl/media/image1.png": _PNG_1PX,
                "media/image2.png": _PNG_1PX,
                "media/image3.png": _PNG_1PX,
            },
        )
        images = parse_wps_cell_images(raw)
        assert set(images) == {"ID1", "ID2", "ID3"}
        assert images["ID2"].data == _PNG_1PX
        assert images["ID3"].data == _PNG_1PX

    def test_rels_relationship_without_id_or_target(self):
        """relationship 缺 Id 或缺 Target 时跳过该条（覆盖 146->141 假分支）."""
        raw = _wps_xlsx(
            cellimages_xml=_cellimages_xml(_pic("ID1", "rId1")),
            rels_xml=_rels_xml(
                '<Relationship Target="media/x.png"/>' + '<Relationship Id="rId7"/>' + _rel("rId1", "media/image1.png")
            ),
            media={"xl/media/image1.png": _PNG_1PX},
        )
        assert set(parse_wps_cell_images(raw)) == {"ID1"}

    def test_cellimage_extra_attr_and_incomplete_entry(self):
        """blip 含非 embed 属性 / blip 无 embed / cellimage 缺 name（覆盖 167->161、168->167、171->156）."""
        no_embed_blip = (
            '<etc:cellimage><xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="ID2"/></xdr:nvPicPr>'
            '<xdr:blipFill><a:blip only="x"/></xdr:blipFill></xdr:pic></etc:cellimage>'
        )
        raw = _wps_xlsx(
            cellimages_xml=_cellimages_xml(
                _pic("ID1", "rId1", blip_attrs='foo="bar" ') + no_embed_blip + _pic(None, None)
            ),
            rels_xml=_rels_xml(_rel("rId1", "media/image1.png")),
            media={"xl/media/image1.png": _PNG_1PX},
        )
        assert set(parse_wps_cell_images(raw)) == {"ID1"}

    def test_cellimages_invalid_xml_returns_empty(self):
        """cellimages.xml 损坏 → ParseError 吞掉返回空（覆盖 173-175）."""
        raw = _wps_xlsx(
            cellimages_xml="<etc:cellImages><broken>",
            rels_xml=_rels_xml(_rel("rId1", "media/image1.png")),
            media={"xl/media/image1.png": _PNG_1PX},
        )
        assert parse_wps_cell_images(raw) == {}

    def test_cellimage_rid_not_mapped(self):
        """cellimage 的 rid 在 rels 中无映射 → 跳过（覆盖 182-183）."""
        raw = _wps_xlsx(
            cellimages_xml=_cellimages_xml(_pic("ID_A", "rIdMissing")),
            rels_xml=_rels_xml(_rel("rId1", "media/image1.png")),
            media={"xl/media/image1.png": _PNG_1PX},
        )
        assert parse_wps_cell_images(raw) == {}

    def test_media_file_absent_keyerror(self):
        """rels 指向的 media 文件不在 zip 中 → KeyError 跳过（覆盖 186-188）."""
        raw = _wps_xlsx(
            cellimages_xml=_cellimages_xml(_pic("ID_A", "rId1")),
            rels_xml=_rels_xml(_rel("rId1", "media/ghost.png")),
            media={},
        )
        assert parse_wps_cell_images(raw) == {}

    def test_ephemeral_saver_rejects_non_image(self):
        assert ephemeral_image_saver(WpsCellImage("ID", "chart.emf", b"\x01", "application/octet-stream")) is None

    def test_resolver_unknown_id_returns_none(self):
        resolve = make_displimg_resolver(_wps_xlsx(cellimages_xml=None, rels_xml=None))
        assert resolve("ID_NOPE") is None

    def test_workspace_saver_rejects_non_image_ext(self):
        """工作区落盘器扩展名不在白名单 → 置空（覆盖 265-266）."""
        saver = make_workspace_image_saver(12345)
        img = WpsCellImage("ID", "chart.emf", b"\x01\x02", "application/octet-stream")
        assert saver(img) is None

    def test_workspace_saver_rejects_oversize(self, monkeypatch):
        """超过单文件上限 → 置空（覆盖 268-269）."""
        monkeypatch.setattr(wps_mod, "_MAX_WPS_IMAGE_BYTES", 4)
        saver = make_workspace_image_saver(12345)
        img = WpsCellImage("ID", "big.png", b"0123456789", "image/png")
        assert saver(img) is None

    def test_ephemeral_saver_success_builds_meta(self):
        """图片扩展名合法 → 产出空 file_key 元数据（覆盖 202、251）."""
        meta = ephemeral_image_saver(WpsCellImage("ID", "a.PNG", _PNG_1PX, "image/png"))
        assert meta is not None
        assert meta["file_key"] == ""
        assert meta["filename"] == "a.PNG"
        assert meta["size"] == len(_PNG_1PX)

    def test_resolver_resolves_known_id(self):
        """resolver 命中图片 → 元数据列表；二次调用复用缓存（覆盖 229->231、235、238）."""
        raw = _wps_xlsx(
            cellimages_xml=_cellimages_xml(_pic("ID1", "rId1")),
            rels_xml=_rels_xml(_rel("rId1", "media/image1.png")),
            media={"xl/media/image1.png": _PNG_1PX},
        )
        resolve = make_displimg_resolver(raw)
        first = resolve("ID1")
        assert first is not None and first[0]["mime_type"] == "image/png"
        # 再次命中同一 ID：cell_images 已缓存（229->231 假分支）
        assert resolve("ID1") is not None

    def test_resolver_saver_rejection_returns_none(self):
        """saver 拒收 → resolve 返回 None（覆盖 236-237）."""
        raw = _wps_xlsx(
            cellimages_xml=_cellimages_xml(_pic("ID1", "rId1")),
            rels_xml=_rels_xml(_rel("rId1", "media/image1.png")),
            media={"xl/media/image1.png": _PNG_1PX},
        )
        resolve = make_displimg_resolver(raw, lambda _img: None)
        assert resolve("ID1") is None

    def test_workspace_saver_success_persists_file(self):
        """合法图片 → 落盘并返回带 file_key 元数据（覆盖 270-274）."""
        saver = make_workspace_image_saver(321)
        img = WpsCellImage("ID", "pic.png", _PNG_1PX, "image/png")
        meta = saver(img)
        assert meta is not None
        assert meta["file_key"].endswith(".png")
        from cndb.core.config import settings

        saved = settings.UPLOAD_DIR / "321" / meta["file_key"]
        assert saved.is_file()
        assert saved.read_bytes() == _PNG_1PX


# ── table_create ───────────────────────────────────


@pytest.fixture
def ws(db, db_engine):
    from cndb.plugins.accounts.models import User
    from cndb.plugins.workspaces.models import Workspace

    user = User(username="cov_transfer_user")
    user.set_password("pass")
    db.add(user)
    db.flush()
    workspace = Workspace(name="CovTransferWS", created_by_id=user.id)
    db.add(workspace)
    db.commit()
    db.refresh(workspace)
    yield db_engine, db, workspace


class TestTableCreateEdgeCases:
    def test_analyze_json_columns_sample_limit(self):
        """样本数达上限后不再收集（覆盖 155->134 假分支）."""
        cols = analyze_json_columns([{"a": "1"}, {"a": "2"}], sample_rows=1)
        assert cols[0]["sample_values"] == ["1"]

    def test_analyze_json_columns_dumps_failure_falls_back(self):
        """json.dumps 失败 → 回退 str()（覆盖 161-162）."""
        cols = analyze_json_columns([{"k": {"a": object()}}])
        assert cols[0]["name"] == "k"
        assert cols[0]["field_type"] == "json"

    def test_create_table_from_json_data_empty_raises(self, ws):
        """空对象数组 → 明确报错（覆盖 418）."""
        engine, db, workspace = ws
        with pytest.raises(ValueError, match="JSON 没有有效字段"):
            create_table_from_json_data(engine, db, workspace.id, "空表", [])

    def test_analyze_dict_rows_as_csv_empty(self):
        assert _analyze_dict_rows_as_csv([]) == []

    def test_analyze_dict_rows_as_csv_value_branches(self):
        """None / bool / int / float 归一（覆盖 472、474、476）."""
        cols = _analyze_dict_rows_as_csv(
            [{"a": None, "b": True, "c": 3, "d": 1.5}, {"a": "x", "b": False, "c": 4, "d": 2.5}]
        )
        type_map = {c["name"]: c["field_type"] for c in cols}
        assert type_map["a"] == "text"
        assert type_map["b"] == "boolean"
        assert type_map["c"] == "number"
        assert type_map["d"] == "float"

    def test_analyze_file_columns_requires_input(self):
        """content 与 rows 均缺省 → 明确报错（覆盖 503-505）."""
        with pytest.raises(ValueError, match="必须提供 content 或 rows"):
            analyze_file_columns()

    def test_apply_column_overrides_options_without_field_type(self):
        """仅覆盖 options（不改 field_type，覆盖 538->540、541）."""
        columns = [{"name": "s", "field_type": "select", "options": ["a"]}]
        result = _apply_column_overrides(columns, {"s": {"options": ["b", "c"]}})
        assert result[0]["field_type"] == "select"
        assert result[0]["options"] == ["b", "c"]

    def test_apply_column_overrides_clears_options_on_type_change(self):
        """前端把 select 改为其他类型 → 清掉 options（覆盖 544）."""
        columns = [{"name": "s", "field_type": "select", "options": ["a"]}]
        result = _apply_column_overrides(columns, {"s": {"field_type": "text"}})
        assert result[0]["field_type"] == "text"
        assert "options" not in result[0]

    def test_create_table_from_file_no_columns_raises(self, ws):
        """文件解析不出任何列 → 明确报错（覆盖 587）."""
        engine, db, workspace = ws
        with pytest.raises(ValueError, match="文件没有有效列"):
            create_table_from_file(engine, db, workspace.id, "空表", "", format="csv")

    def test_import_rows_from_csv_short_row_none(self, ws):
        """短行缺列 → DictReader 补 None → 走非 str 分支（覆盖 341）."""
        engine, db, workspace = ws
        from cndb.plugins.tables.services.core import ddl

        dt = DataTable(workspace_id=workspace.id, name="短行表")
        dt.ensure_db_name()
        db.add(dt)
        db.commit()
        db.refresh(dt)
        for i, name in enumerate(("a", "b")):
            f = DataField(table_id=dt.id, name=name, field_type="text", order=i)
            f.ensure_db_name()
            db.add(f)
        db.commit()
        db.refresh(dt)
        ddl.create_table(engine, dt)

        ids = import_rows_from_csv(engine, dt, "a,b\n只有一列\n")
        assert len(ids) == 1

    def test_cleanup_partial_table_drops_link_table(self, ws):
        """含 link 字段的半残表清理先删关联表（覆盖 218）."""
        engine, db, workspace = ws

        dt = DataTable(workspace_id=workspace.id, name="带关联表")
        dt.ensure_db_name()
        db.add(dt)
        db.commit()
        db.refresh(dt)
        link_field = DataField(table_id=dt.id, name="关联", field_type="link", order=0)
        link_field.ensure_db_name()
        db.add(link_field)
        db.commit()
        db.refresh(dt)
        assert any(f.field_type == "link" for f in dt.fields)

        _cleanup_partial_table(engine, db, dt)
        assert db.get(DataTable, dt.id) is None


# ── spreadsheet ───────────────────────────────────


class TestSpreadsheetEdgeCases:
    def test_decode_bytes_auto_chardet_success(self, monkeypatch):
        """候选编码全被绕过 → chardet 命中并解码（覆盖 67-70）."""
        fake = types.ModuleType("chardet")
        fake.detect = lambda _data: {"encoding": "utf-8", "confidence": 0.9}
        monkeypatch.setitem(sys.modules, "chardet", fake)
        monkeypatch.setattr(spreadsheet_mod, "_ENCODING_CANDIDATES", [])

        text, enc, conf = decode_bytes_auto(b"abc")
        assert text == "abc"
        assert enc == "utf-8"
        assert conf == 0.9

    def test_decode_bytes_auto_chardet_decode_error(self, monkeypatch):
        """chardet 编码解码失败 → 记日志并回落 latin-1（覆盖 72-73）."""
        fake = types.ModuleType("chardet")
        fake.detect = lambda _data: {"encoding": "utf-32", "confidence": 0.9}
        monkeypatch.setitem(sys.modules, "chardet", fake)
        monkeypatch.setattr(spreadsheet_mod, "_ENCODING_CANDIDATES", [])

        text, enc, _conf = decode_bytes_auto(b"\xff\xff\xff")
        assert enc == "latin-1"
        assert len(text) == 3

    def test_decode_bytes_auto_chardet_raises(self, monkeypatch):
        """chardet 自身抛错 → 兜底 latin-1（覆盖 74）."""
        fake = types.ModuleType("chardet")

        def _boom(_data):
            raise RuntimeError("chardet boom")

        fake.detect = _boom
        monkeypatch.setitem(sys.modules, "chardet", fake)
        monkeypatch.setattr(spreadsheet_mod, "_ENCODING_CANDIDATES", [])

        _text, enc, _conf = decode_bytes_auto(b"xyz")
        assert enc == "latin-1"

    def test_decode_bytes_auto_chardet_ascii_ignored(self, monkeypatch):
        """chardet 返回 ascii → 视为盲猜跳过，回落 latin-1."""
        fake = types.ModuleType("chardet")
        fake.detect = lambda _data: {"encoding": "ascii", "confidence": 0.99}
        monkeypatch.setitem(sys.modules, "chardet", fake)
        monkeypatch.setattr(spreadsheet_mod, "_ENCODING_CANDIDATES", [])

        _text, enc, _conf = decode_bytes_auto(b"xyz")
        assert enc == "latin-1"

    def test_decode_bytes_auto_chardet_empty_text(self, monkeypatch):
        """chardet 编码解出空串 → 跳过质量检查直接回落（覆盖 67->77）."""
        fake = types.ModuleType("chardet")
        fake.detect = lambda _data: {"encoding": "utf-8", "confidence": 0.9}
        monkeypatch.setitem(sys.modules, "chardet", fake)
        monkeypatch.setattr(spreadsheet_mod, "_ENCODING_CANDIDATES", [])

        _text, enc, _conf = decode_bytes_auto(b"")
        assert enc == "latin-1"

    def test_decode_bytes_auto_chardet_dirty_text(self, monkeypatch):
        """chardet 解出文本不可打印字符过多 → 跳过并回落（覆盖 69->77）."""
        fake = types.ModuleType("chardet")
        fake.detect = lambda _data: {"encoding": "utf-8", "confidence": 0.9}
        monkeypatch.setitem(sys.modules, "chardet", fake)
        monkeypatch.setattr(spreadsheet_mod, "_ENCODING_CANDIDATES", [])

        _text, enc, _conf = decode_bytes_auto(b"a" + b"\x00" * 20)
        assert enc == "latin-1"

    def test_sniff_delimiter_unknown_dialect_falls_back(self, monkeypatch):
        """Sniffer 返回非候选分隔符 → 走计数回落（覆盖 92->97）."""

        class _FakeSniffer:
            def sniff(self, *_a, **_k):
                return types.SimpleNamespace(delimiter=" ")

        monkeypatch.setattr(spreadsheet_mod.csv, "Sniffer", _FakeSniffer)
        assert sniff_csv_delimiter("a ; b\nc ; d") == ";"

    def test_parse_file_to_rows_explicit_format_str(self):
        """显式 format + str 内容 → 分隔符文本分支（覆盖 224->229、251）."""
        rows, cols, fmt = parse_file_to_rows("a,b\n1,2", "csv")
        assert fmt == "csv"
        assert cols == ["a", "b"]
        assert rows == [{"a": "1", "b": "2"}]

    def test_parse_file_to_rows_uninferable_filename(self):
        """文件名推断失败 → 回落内容推断（覆盖 227-228、230）."""
        rows, _cols, fmt = parse_file_to_rows("a,b\n1,2\n", filename="data.txt")
        assert fmt == "csv"
        assert len(rows) == 1

    def test_parse_file_to_rows_json_text_str(self):
        """json + str 内容（覆盖 243）."""
        rows, cols, fmt = parse_file_to_rows('[{"a": 1}]', "json")
        assert fmt == "json"
        assert cols == ["a"]
        assert rows == [{"a": 1}]

    def test_parse_file_to_rows_json_non_array(self):
        with pytest.raises(ValueError, match="JSON 必须是对象数组"):
            parse_file_to_rows('{"a": 1}', "json")

    def test_parse_file_to_rows_json_array_with_scalar(self):
        """数组含非对象元素 → 仅保留 dict（覆盖 287->286）."""
        rows, cols, fmt = parse_file_to_rows('[{"a": 1}, 5]', "json")
        assert fmt == "json"
        assert cols == ["a"]
        assert rows == [{"a": 1}]

    def test_parse_delimited_text_short_row_none(self):
        """短行缺列 → 值为 None（覆盖 274）."""
        rows, cols = _parse_delimited_text("a,b\n1\n")
        assert cols == ["a", "b"]
        assert rows == [{"a": "1", "b": None}]

    def test_parse_json_text_requires_list(self):
        with pytest.raises(ValueError, match="JSON 必须是对象数组"):
            _parse_json_text("42")

    def test_parse_xlsx_bytes_empty_sheet(self):
        """空工作表 → 无表头无数据（覆盖 341）."""
        assert _parse_xlsx_bytes(_empty_xlsx()) == ([], [])


# ── normalize/values ──────────────────────────────


class TestNormalizeValuesEdgeCases:
    def test_normalize_numeric_only_currency_symbol(self):
        """仅货币符号 → 归一后为空返回 None（覆盖 103）."""
        assert _normalize_numeric("$") is None
        assert _normalize_numeric("￥") is None

    def test_normalize_numeric_accounting_negative_invalid(self):
        """会计负数括号内非法 → 返回 None（覆盖 108）."""
        assert _normalize_numeric("(abc)") is None

    def test_is_integer_non_digit_normalized(self):
        """float 可解析但非数字串（inf/nan）→ False（覆盖 180）."""
        assert _is_integer("inf") is False
        assert _is_integer("nan") is False

    def test_check_long_integer_small_normal(self):
        assert _check_long_integer("123") is False

    def test_check_compact_date_out_of_range_year(self):
        """紧凑日期年份越界（<1900 / >2100）→ 回落（覆盖 303）."""
        assert _check_compact_date("18800101") is False
        assert _check_compact_date("22000101") is False

    def test_check_en_date_unknown_month_name(self):
        """英文月名不在表内 → 回落（覆盖 320）."""
        assert _check_en_date("Foo 15, 2024") is False
        assert _check_en_date("15 Foo 2024") is False


# ── normalize/promote ─────────────────────────────


class TestNormalizePromoteEdgeCases:
    def test_split_list_like_empty_segment(self):
        """空段（连续/尾随分隔符）→ None（覆盖 158）."""
        assert _split_list_like("a,,b") is None
        assert _split_list_like("a,b,") is None

    def test_split_list_like_too_many_segments(self):
        """段数超上限 → None（覆盖 160）."""
        assert _split_list_like("a,b,c,d,e,f,g") is None

    def test_split_json_array_like_not_array(self):
        assert _split_json_array_like("abc") is None

    def test_split_json_array_like_bad_json(self):
        """以 '[' 开头但非法 JSON → None（覆盖 183-184）."""
        assert _split_json_array_like("[1,2") is None

    def test_split_json_array_like_empty(self):
        """合法但空数组 → None（覆盖 186）."""
        assert _split_json_array_like("[]") is None

    def test_split_json_array_like_too_many_options(self):
        """元素数超上限 → None（覆盖 188）."""
        assert _split_json_array_like(json.dumps(list(range(31)))) is None

    def test_infer_column_type_skips_empty_sample(self):
        """空样本被跳过、不计入类型投票（覆盖 283->281）."""
        assert infer_column_type(["", "abc"]) == ("text", [])


# ── normalize/coerce ──────────────────────────────


class TestNormalizeCoerceEdgeCases:
    def test_coerce_bool_and_int_edges(self):
        assert _coerce_long_numeric_to_text(True) is True
        assert _coerce_long_numeric_to_text(-(2**53)) == str(-(2**53))
        assert _coerce_long_numeric_to_text(123456789012345) == "123456789012345"
        assert _coerce_long_numeric_to_text(42) == 42

    def test_coerce_float_edges(self):
        assert _coerce_long_numeric_to_text(1e16) == str(int(1e16))
        assert _coerce_long_numeric_to_text(123456789012345.0) == "123456789012345"
        assert _coerce_long_numeric_to_text(1.5) == 1.5

    def test_python_type_to_field_type_long_values(self):
        """长整数 / 长 float → text（覆盖 117、124）."""
        assert _python_type_to_field_type(2**53) == "text"
        assert _python_type_to_field_type(123456789012345) == "text"
        assert _python_type_to_field_type(1e16) == "text"
        assert _python_type_to_field_type(3) == "number"
        assert _python_type_to_field_type(1.5) == "float"
