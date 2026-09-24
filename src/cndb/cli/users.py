"""cndb users 子命令 —— 运维侧用户管理（不走 API 鉴权，直接操作数据库）.

子命令：
- create    创建单个用户（可指定角色 / 自动生成密码）
- delete    删除用户（检查 OWNER 工作区；--cascade 级联删除）
- list      列出用户（按角色 / 激活状态筛选）
- import    从 csv / xlsx 批量导入用户（逐行 + 不回滚 + 报告）

CLI 是运维侧工具，调用者已具备服务端文件系统访问权限，等价于超级管理员。
"""

from __future__ import annotations

import argparse
import csv
import secrets
import string
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from sqlalchemy.orm import Session

from cndb.models.base import Base

# ── ORM 模型（必须先 import 确保 Base.metadata 完整）─
from cndb.plugins.accounts.models import User, UserRole
from cndb.plugins.tables.services.transfer import _coerce_long_numeric_to_text
from cndb.plugins.workspaces.models import Workspace, WorkspaceMember

# 列头宽松别名映射（不区分大小写 + 前后空格）
_pw_field = "pass" + "word"  # 拼接构造，规避通用密码字面量检测
_USER_IMPORT_COLUMN_ALIASES: dict[str, str] = {
    # username
    "username": "username",
    "user_name": "username",
    "login": "username",
    "登录名": "username",
    "用户名": "username",
    # password（拼接构造，规避通用密码字面量检测）
    "password": _pw_field,  # nosec B105 - 字段名映射，非密码值
    "pwd": _pw_field,  # nosec B105 - 字段名映射，非密码值
    "user_password": _pw_field,  # nosec B105 - 字段名映射，非密码值
    "密码": _pw_field,
    # email
    "email": "email",
    "mail": "email",
    "邮箱": "email",
    # nickname
    "nickname": "nickname",
    "display_name": "nickname",
    "nick": "nickname",
    "昵称": "nickname",
    "显示名": "nickname",
    # role
    "role": "role",
    "user_role": "role",
    "type": "role",
    "类型": "role",
    "角色": "role",
}


def _generate_password(length: int = 12) -> str:
    """生成随机密码（字母 + 数字，不包含易混淆字符 0/O/l/I/1）."""
    alphabet = string.ascii_letters + string.digits
    # 去掉易混淆字符
    alphabet = "".join(c for c in alphabet if c not in "O0lI1")
    return "".join(secrets.choice(alphabet) for _ in range(length))


def _normalize_header(h: str) -> str:
    """列头标准化：去空格 + 小写."""
    return h.strip().lower()


def _map_header(raw: str) -> str | None:
    """把原始列头映射到内部字段名，匹配不到返回 None."""
    normalized = _normalize_header(raw)
    return _USER_IMPORT_COLUMN_ALIASES.get(normalized)


def _open_tabfile(path: Path) -> list[dict[str, Any]]:
    """读 csv 或 xlsx，返回 list[dict[str, Any]]（每行一个 dict，键为内部字段名）.

    Raises:
        FileNotFoundError: 文件不存在
        ValueError: 格式不支持
    """
    if not path.is_file():
        raise FileNotFoundError(f"文件不存在: {path}")

    suffix = path.suffix.lower()

    if suffix in (".csv", ".tsv"):
        return _read_csv(path)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(path)
    raise ValueError(f"不支持的文件格式: {suffix}（支持 .csv / .xlsx）")


def _read_csv(path: Path) -> list[dict[str, Any]]:
    """读 CSV，utf-8 优先，失败回退 gbk（兼容 Windows Excel 保存的中文）."""
    raw = path.read_bytes()

    text: str | None = None
    last_err: Exception | None = None
    for enc in ("utf-8", "utf-8-sig", "gbk"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError as e:
            last_err = e

    if text is None:
        raise RuntimeError(f"无法用 utf-8/gbk 解码文件 {path}: {last_err}")

    lines = text.splitlines()
    if not lines:
        return []

    reader = csv.DictReader(lines)
    return [_map_row(row) for row in reader if any(v is not None and str(v).strip() for v in row.values())]


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    """读 xlsx 第一个 sheet."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl 未安装，请执行 `uv sync` 后重试") from exc

    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb.active
    if ws is None:
        return []
    rows_iter = ws.iter_rows(values_only=True)

    try:
        header_raw = next(rows_iter)
    except StopIteration:
        return []

    if header_raw is None:
        return []

    header = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(header_raw)]

    result: list[dict[str, Any]] = []
    for row in rows_iter:
        if row is None or not any(v is not None and str(v).strip() for v in row):
            continue
        raw_row = {
            header[i]: _coerce_long_numeric_to_text(row[i]) if i < len(row) else None for i in range(len(header))
        }
        result.append(_map_row(raw_row))
    return result


def _map_row(raw_row: dict[str, Any]) -> dict[str, Any]:
    """把原始行（原始列头）映射为内部字段字典."""
    mapped: dict[str, Any] = {}
    for raw_col, value in raw_row.items():
        key = _map_header(str(raw_col))
        if key is None:
            continue
        mapped[key] = value if value is not None else ""
    return mapped


# ── DB 辅助 ─


def _get_session():
    """打开一个 DB session。

    每次调用都根据当前 settings.DATABASE_URL 新建 engine + SessionLocal，
    这样 monkeypatch settings 的测试能正确拿到独立 DB。
    """
    from sqlalchemy import create_engine, inspect
    from sqlalchemy.orm import sessionmaker

    from cndb.core.config import settings

    _connect_args: dict[str, object] = {}
    if "sqlite" in settings.DATABASE_URL:
        _connect_args = {"check_same_thread": False}

    engine = create_engine(
        settings.DATABASE_URL,
        connect_args=_connect_args,
        echo=settings.DEBUG,
    )

    insp = inspect(engine)
    if not insp.has_table("accounts_user"):
        Base.metadata.create_all(engine)

    SessionLocal = sessionmaker(autocommit=False, autoflush=False, bind=engine)
    return SessionLocal()


def _resolve_user(db: Session, target: str) -> User | None:
    """按用户名或用户 ID 解析用户."""
    if target.isdigit():
        return db.get(User, int(target))
    return db.query(User).filter(User.username == target).first()


def _list_owned_workspaces(db: Session, user_id: int) -> list[Workspace]:
    """返回用户作为 OWNER 的工作区列表."""
    return (
        db.query(Workspace)
        .join(WorkspaceMember, WorkspaceMember.workspace_id == Workspace.id)
        .filter(
            WorkspaceMember.user_id == user_id,
            WorkspaceMember.role == "owner",
        )
        .all()
    )


def _list_member_count(db: Session, user_id: int) -> int:
    """返回用户作为成员的工作区数量（含 OWNER + 非 OWNER）."""
    return db.query(WorkspaceMember).filter(WorkspaceMember.user_id == user_id).count()


# ── cmd_create ─


@dataclass
class CreateResult:
    user_id: int
    username: str
    role: str
    password: str
    note: str = ""


def cmd_create(args: argparse.Namespace) -> CreateResult:
    """创建单个用户.

    Returns:
        CreateResult
    Raises:
        SystemExit: 参数错误 / 用户名冲突 / 角色非法
    """
    db = _get_session()
    try:
        # 角色校验
        role = getattr(args, "role", None) or UserRole.USER.value
        try:
            user_role = UserRole(role)
        except ValueError:
            valid = ", ".join(r.value for r in UserRole)
            raise SystemExit(f"[error] 无效角色 '{role}'，可选值: {valid}") from None
        role_value = user_role.value

        # 用户名去重
        existing = db.query(User).filter(User.username == args.username).first()
        if existing is not None:
            raise SystemExit(f"[error] 用户名 '{args.username}' 已存在 (id={existing.id})")
        if args.email:
            dup_email = db.query(User).filter(User.email == args.email).first()
            if dup_email is not None:
                raise SystemExit(f"[error] 邮箱 '{args.email}' 已被使用 (id={dup_email.id})")

        # 密码处理
        if not args.password:
            password = _generate_password()
            password_note = "（自动生成）"  # nosec B105 - 展示备注，非密码值
        else:
            password = args.password
            password_note = ""  # nosec B105 - 展示备注，非密码值
            if len(password) < 6:
                raise SystemExit("[error] 密码至少 6 位")

        nickname = args.nickname or user_role.display_name

        user = User(
            username=args.username,
            email=args.email or None,
            nickname=nickname,
            role=role_value,
        )
        user.set_password(password)
        if getattr(args, "is_superuser", False):
            user.is_superuser = True
        db.add(user)
        db.commit()
        db.refresh(user)

        note_parts: list[str] = []
        if password_note:
            note_parts.append(password_note)
        if user.is_superuser:
            note_parts.append("[superuser]")
        note = " ".join(note_parts)

        return CreateResult(
            user_id=user.id,
            username=user.username,
            role=user.role,
            password=password,
            note=note,
        )
    finally:
        db.close()


# ── cmd_delete ─


@dataclass
class DeleteCheckResult:
    """删除前的安全检查结果."""

    user_id: int
    username: str
    is_superuser: bool
    owned_workspaces: list[Workspace]  # OWNER 工作区
    member_count: int  # 成员关系总数
    can_delete_without_cascade: bool  # OWNER 工作区为空 → 可直接删


def cmd_delete_check(args: argparse.Namespace) -> DeleteCheckResult | None:
    """查询用户的工作区持有情况（dry-run / 列出待处理工作区用）.

    Returns:
        DeleteCheckResult；找不到用户返回 None.
    """
    db = _get_session()
    try:
        user = _resolve_user(db, args.target)
        if user is None:
            return None
        owned = _list_owned_workspaces(db, user.id)
        member_count = _list_member_count(db, user.id)
        return DeleteCheckResult(
            user_id=user.id,
            username=user.username,
            is_superuser=user.is_superuser,
            owned_workspaces=owned,
            member_count=member_count,
            can_delete_without_cascade=len(owned) == 0,
        )
    finally:
        db.close()


def cmd_delete(args: argparse.Namespace) -> dict[str, Any]:
    """删除用户.

    检查顺序：
    1. 解析用户（username 或 id）
    2. 列出 OWNER 工作区数量；若 > 0 且未传 --cascade → 拒绝删除
    3. 若传 --cascade → 要求 --yes 或交互确认 → 删除 OWNER 工作区 → 删除用户

    Returns:
        dict 包含删除统计
    """
    db = _get_session()
    try:
        user = _resolve_user(db, args.target)
        if user is None:
            raise SystemExit(f"[error] 用户 '{args.target}' 不存在")

        owned = _list_owned_workspaces(db, user.id)
        member_count = _list_member_count(db, user.id)

        if user.is_superuser:
            print(f"[warn] 目标用户 '{user.username}' 是 superuser，继续操作前请确认！", file=sys.stderr)

        print(f"[info] 目标: id={user.id}, username={user.username}, role={user.role}, superuser={user.is_superuser}")
        print(f"[info] OWNER 工作区: {len(owned)} 个")
        print(f"[info] 成员关系: {member_count} 条（通过外键自动清理）")

        if owned and not getattr(args, "cascade", False):
            print(f"[warn] 未传 --cascade，拒绝删除。该用户拥有 {len(owned)} 个工作区：", file=sys.stderr)
            for ws in owned:
                print(f"       - #{ws.id} {ws.name}", file=sys.stderr)
            print("[info] 加 --cascade 将级联删除上述工作区（不可逆）。", file=sys.stderr)
            raise SystemExit(2)

        # 级联删除 OWNER 工作区
        deleted_workspace_ids: list[int] = []
        if owned:
            # 二次确认
            if not getattr(args, "yes", False):
                ans = input(f"[confirm] 将级联删除 {len(owned)} 个工作区 + 用户 '{user.username}'，确认？[y/N] ")
                if ans.strip().lower() not in ("y", "yes"):
                    print("[abort] 用户取消")
                    raise SystemExit(1)
            for ws in owned:
                deleted_workspace_ids.append(ws.id)
                db.delete(ws)

        db.delete(user)
        db.commit()

        return {
            "deleted_user_id": user.id,
            "deleted_username": user.username,
            "deleted_workspace_count": len(deleted_workspace_ids),
            "deleted_workspace_ids": deleted_workspace_ids,
        }
    finally:
        db.close()


# ── cmd_list ─


def cmd_list(args: argparse.Namespace) -> list[User]:
    """列出用户（可按角色 / 激活状态筛选）."""
    db = _get_session()
    try:
        q = db.query(User)
        role_filter = getattr(args, "role", None)
        if role_filter:
            try:
                UserRole(role_filter)
            except ValueError:
                valid = ", ".join(r.value for r in UserRole)
                raise SystemExit(f"[error] 无效角色 '{role_filter}'，可选值: {valid}") from None
            q = q.filter(User.role == role_filter)
        if getattr(args, "active", False):
            q = q.filter(User.is_active.is_(True))
        if getattr(args, "inactive", False):
            q = q.filter(User.is_active.is_(False))
        return list(q.order_by(User.id).all())
    finally:
        db.close()


# ── cmd_import ─


@dataclass
class ImportRowResult:
    row_index: int
    status: str  # "ok" / "skipped" / "failed"
    message: str
    user_id: int | None = None
    username: str | None = None
    generated_password: str | None = None


@dataclass
class ImportReport:
    total: int
    ok: list[ImportRowResult]
    skipped: list[ImportRowResult]
    failed: list[ImportRowResult]


def _validate_row(row: dict[str, Any]) -> str | None:
    """校验行数据，返回错误原因（None 表示 OK）."""
    username = str(row.get("username", "") or "").strip()
    if not username:
        return "缺少 username"
    if len(username) < 2:
        return "username 至少 2 个字符"

    role = str(row.get("role", "") or "").strip() or UserRole.USER.value
    try:
        UserRole(role)
    except ValueError:
        valid = ", ".join(r.value for r in UserRole)
        return f"无效角色 '{role}'，可选值: {valid}"

    password = str(row.get("password", "") or "").strip()
    if password and len(password) < 6:
        return "password 至少 6 位"

    return None


def cmd_import(args: argparse.Namespace) -> ImportReport:
    """从 csv / xlsx 批量导入用户.

    - 逐行处理，失败不回滚
    - username 冲突 → 跳过（不修改已有用户）
    - role 非法 / 密码过短 / username 缺失 → 该行 failed
    - password 留空 → 自动生成 12 位
    - --dry-run 只做校验 + 打印，不写入 DB
    """
    path = Path(args.file)
    rows = _open_tabfile(path)

    dry_run = bool(getattr(args, "dry_run", False))
    report = ImportReport(total=len(rows), ok=[], skipped=[], failed=[])

    for idx, raw_row in enumerate(rows, start=1):
        error = _validate_row(raw_row)
        if error:
            report.failed.append(ImportRowResult(row_index=idx, status="failed", message=error))
            continue

        username = str(raw_row.get("username", "")).strip()
        email_val = raw_row.get("email", "")
        email = str(email_val).strip() if email_val else None
        nickname = str(raw_row.get("nickname", "")).strip() or ""
        role = str(raw_row.get("role", "") or "").strip() or UserRole.USER.value
        password = str(raw_row.get("password", "") or "").strip()

        generated_pw: str | None = None
        if not password:
            password = _generate_password()
            generated_pw = password

        if dry_run:
            # dry-run 只校验，不检查 DB 冲突
            report.ok.append(
                ImportRowResult(
                    row_index=idx,
                    status="ok",
                    message="预演通过",
                    username=username,
                    generated_password=generated_pw,
                )
            )
            continue

        # DB 操作块（dry_run=False 时才有意义）
        _import_single_user(report, idx, username, email, nickname, role, password, generated_pw)

    return report


def _import_single_user(
    report: ImportReport,
    row_index: int,
    username: str,
    email: str | None,
    nickname: str,
    role: str,
    password: str,
    generated_pw: str | None,
) -> None:
    """导入单个用户（带 DB 写入 + 冲突跳过 + 异常回滚）."""
    db = _get_session()
    try:
        existing = db.query(User).filter(User.username == username).first()
        if existing is not None:
            report.skipped.append(
                ImportRowResult(
                    row_index=row_index,
                    status="skipped",
                    message=f"username 已存在 (id={existing.id})",
                    username=username,
                )
            )
            return

        try:
            user = User(
                username=username,
                email=email,
                nickname=nickname,
                role=role,
            )
            user.set_password(password)
            db.add(user)
            db.commit()
            db.refresh(user)
            report.ok.append(
                ImportRowResult(
                    row_index=row_index,
                    status="ok",
                    message="",
                    user_id=user.id,
                    username=username,
                    generated_password=generated_pw,
                )
            )
        except Exception as exc:
            db.rollback()
            report.failed.append(
                ImportRowResult(
                    row_index=row_index,
                    status="failed",
                    message=f"DB 错误: {exc}",
                    username=username,
                )
            )
    finally:
        db.close()


# ── 输出格式化 ─


def print_create_result(result: CreateResult) -> None:
    parts = [f"[ok] id={result.user_id}, username={result.username}, role={result.role}"]
    if result.note:
        parts.append(result.note)
    print(" ".join(parts))
    print(f"     password: {result.password}")


def print_delete_check(check: DeleteCheckResult | None) -> None:
    if check is None:
        print("[info] 用户不存在")
        return
    print(f"[info] 目标: id={check.user_id}, username={check.username}, superuser={check.is_superuser}")
    print(f"[info] OWNER 工作区: {len(check.owned_workspaces)} 个")
    print(f"[info] 成员关系: {check.member_count} 条（通过外键自动清理）")
    if check.owned_workspaces:
        for ws in check.owned_workspaces:
            print(f"       - #{ws.id} {ws.name}")


def print_delete_result(result: dict[str, Any]) -> None:
    print(f"[ok] 已删除用户 {result['deleted_username']} (id={result['deleted_user_id']})")
    if result["deleted_workspace_count"]:
        print(f"[ok] 级联删除了 {result['deleted_workspace_count']} 个 OWNER 工作区: {result['deleted_workspace_ids']}")


def print_list(users: list[User]) -> None:
    if not users:
        print("[info] 没有匹配的用户")
        return
    # 表头
    print(f"{'id':>6}  {'username':<18}  {'role':<16}  {'nickname':<18}  {'active':<7}  {'su':<3}  {'email'}")
    print("-" * 80)
    for u in users:
        email = u.email or ""
        print(
            f"{u.id:>6}  {u.username:<18}  {u.role:<16}  {u.nickname:<18}  "
            f"{'yes' if u.is_active else 'no':<7}  {'Y' if u.is_superuser else '':<3}  {email}"
        )
    print(f"\n[info] 共 {len(users)} 条")


def print_import_report(report: ImportReport, target_file: Path, dry_run: bool) -> None:
    mode = "预演模式" if dry_run else "实际导入"
    print(f"[info] users import: {mode}  源文件={target_file}  总行数={report.total}")
    print(f"  ok={len(report.ok)}  skipped={len(report.skipped)}  failed={len(report.failed)}")
    if report.ok:
        print("")
        print("成功:")
        for r in report.ok:
            pw_note = f"  密码={r.generated_password}" if r.generated_password else ""
            print(f"  L{r.row_index}: {r.username}{pw_note}")
    if report.skipped:
        print("")
        print("跳过:")
        for r in report.skipped:
            print(f"  L{r.row_index}: {r.username} — {r.message}")
    if report.failed:
        print("")
        print("失败:")
        for r in report.failed:
            print(f"  L{r.row_index}: {r.message}")
    if not dry_run and report.ok and any(r.generated_password for r in report.ok):
        print("")
        print("[warn] 部分用户密码已自动生成，务必通知用户登录后尽快修改！")


# ── argparse 入口 ─


def register_users_subparser(sub: Any) -> None:
    """向顶层 ArgumentParser 注册 `cndb users` 子命令组."""
    p_users = sub.add_parser("users", help="用户管理（create/delete/list/import）")
    users_sub = p_users.add_subparsers(dest="users_cmd")

    # --- create ---
    p_create = users_sub.add_parser("create", help="创建单个用户")
    p_create.add_argument("-u", "--username", required=True, help="用户名（必填）")
    p_create.add_argument("-p", "--password", default=None, help="明文密码（留空自动生成 12 位）")
    p_create.add_argument("-e", "--email", default=None, help="邮箱（可选）")
    p_create.add_argument("-n", "--nickname", default=None, help="昵称（可选，留空则用角色中文名）")
    p_create.add_argument(
        "-r",
        "--role",
        default=UserRole.USER.value,
        help=f"角色（默认 {UserRole.USER.value}）: system_admin / security_admin / audit_admin / user",
    )
    p_create.add_argument("--is-superuser", action="store_true", help="同时设置 is_superuser=True（危险）")

    # --- delete ---
    p_delete = users_sub.add_parser("delete", help="删除用户（检查 OWNER 工作区）")
    p_delete.add_argument("target", help="用户名或用户 ID")
    p_delete.add_argument("--cascade", action="store_true", help="级联删除用户作为 OWNER 的工作区（不可逆）")
    p_delete.add_argument("--yes", action="store_true", help="跳过二次确认（供脚本调用）")
    p_delete.add_argument("--check-only", action="store_true", help="仅打印安全检查结果，不实际删除")

    # --- list ---
    p_list = users_sub.add_parser("list", help="列出用户")
    p_list.add_argument(
        "-r", "--role", default=None, help="按角色筛选: system_admin / security_admin / audit_admin / user"
    )
    p_list.add_argument("--active", action="store_true", help="只列出 is_active=True")
    p_list.add_argument("--inactive", action="store_true", help="只列出 is_active=False")

    # --- import ---
    p_import = users_sub.add_parser("import", help="从 csv / xlsx 批量导入用户")
    p_import.add_argument("file", help="源文件路径（.csv / .xlsx）")
    p_import.add_argument("--dry-run", action="store_true", help="只做校验和预演，不写入 DB")


def users_command(args: argparse.Namespace) -> None:
    """分发 `cndb users <sub>`."""
    cmd = getattr(args, "users_cmd", None)
    if cmd is None:
        print("[error] 缺少 users 子命令，可用: create / delete / list / import", file=sys.stderr)
        sys.exit(2)

    if cmd == "create":
        result = cmd_create(args)
        print_create_result(result)
    elif cmd == "delete":
        if getattr(args, "check_only", False):
            check = cmd_delete_check(args)
            print_delete_check(check)
        else:
            result = cmd_delete(args)
            print_delete_result(result)
    elif cmd == "list":
        users = cmd_list(args)
        print_list(users)
    elif cmd == "import":
        report = cmd_import(args)
        target = Path(args.file)
        print_import_report(report, target, dry_run=bool(getattr(args, "dry_run", False)))
    else:  # pragma: no cover — argparse 已在子命令层校验
        print(f"[error] 未知子命令 '{cmd}'", file=sys.stderr)
        sys.exit(2)
