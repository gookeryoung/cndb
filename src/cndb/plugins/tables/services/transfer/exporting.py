"""行导出序列化 — JSON/CSV/XLSX 导出、link 值序列化、公式注入防护."""

from __future__ import annotations

import csv
import io
import json
from typing import Any

from cndb.plugins.tables.models import DataTable
from cndb.plugins.tables.services.core.links import is_link_field


def _serialize_link_value(value: Any) -> Any:
    """导出时把 link 摘要列表压成分号分隔 id 串，其余原样."""
    if isinstance(value, list) and value and isinstance(value[0], dict):
        return ";".join(str(item.get("id", "")) for item in value)
    return value


def _exportable_rows(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """行响应统一做导出序列化（目前仅 link 摘要列表）."""
    return [{k: _serialize_link_value(v) for k, v in row.items()} for row in rows]


def _parse_link_import_value(table: DataTable, row: dict[str, Any]) -> dict[str, Any]:
    """导入时把 link 列的字符串/单值归一为 id 列表."""
    link_names = {f.name for f in table.active_fields() if is_link_field(f)}
    if not link_names:
        return row
    result = dict(row)
    for name in link_names:
        value = result.get(name)
        if value is None:
            continue
        if isinstance(value, list):
            continue
        text = str(value).strip()
        if not text:
            result[name] = []
            continue
        try:
            result[name] = [int(part) for part in text.replace(",", ";").split(";")]
        except ValueError as exc:
            raise ValueError(f"link 字段 {name} 的导入值必须是分号分隔的行 id: {value!r}") from exc
    return result


def export_rows_to_json(rows: list[dict[str, Any]]) -> str:
    """把行列表序列化为 JSON 字符串."""
    return json.dumps(rows, ensure_ascii=False, indent=2, default=str)


# ── 公式注入防护 ──────────────────────────────────────

# Excel/WPS 打开 CSV 时会把以这些前缀开头的单元格当作公式求值（OWASP 缓解：前缀 '）
_CSV_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def _sanitize_csv_cell(value: Any) -> Any:
    """CSV 公式注入防护：以危险前缀开头的字符串值前加 ``'`` 前缀.

    仅处理 str —— int/float/bool/date 等序列化后为纯字面量，不会被当作公式.
    """
    if isinstance(value, str) and value.startswith(_CSV_FORMULA_PREFIXES):
        return "'" + value
    return value


def _fix_xlsx_formula_cells(ws: Any, row_idx: int) -> None:
    """把 XLSX 某行中以 ``=`` 开头被 openpyxl 判为公式的单元格强制回字符串类型.

    openpyxl 对 ``=`` 开头字符串自动设 ``data_type='f'``，Excel 打开即执行；
    赋值后强制 ``data_type='s'`` 值原样保留。``+``/``-``/``@`` 开头在 XLSX 中
    本就存为字符串类型，无需处理.
    """
    for cell in ws[row_idx]:
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"


def export_rows_to_csv(rows: list[dict[str, Any]]) -> str:
    """把行列表转为 CSV 字符串（link 摘要序列化为分号分隔 id）.

    公式注入防护：以 ``=``/``+``/``-``/``@``/Tab/CR 开头的字符串单元格（含表头）
    前加 ``'`` 前缀，避免 Excel/WPS 打开时被当作公式求值.
    """
    rows = _exportable_rows(rows)
    if not rows:
        return ""
    buf = io.StringIO()
    # 表头一并转义；已知列名同步替换为转义后列名，未知列名保留交由 extrasaction 处理
    key_map = {k: _sanitize_csv_cell(k) for k in rows[0]}
    writer = csv.DictWriter(buf, fieldnames=list(key_map.values()))
    writer.writeheader()
    for row in rows:
        writer.writerow({key_map.get(k, k): _sanitize_csv_cell(v) for k, v in row.items()})
    return buf.getvalue()


def export_rows_to_xlsx(rows: list[dict[str, Any]]) -> bytes:
    """把行列表转为 XLSX 字节串（link 摘要序列化为分号分隔 id）.

    公式注入防护：``=`` 开头字符串单元格强制回字符串类型，Excel 打开不被求值.
    """
    from openpyxl import Workbook

    rows = _exportable_rows(rows)
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    if not rows:
        return _wb_to_bytes(wb)
    # Header
    fieldnames = list(rows[0].keys())
    ws.append(fieldnames)
    _fix_xlsx_formula_cells(ws, 1)
    for r in rows:
        ws.append([r.get(k) for k in fieldnames])
        _fix_xlsx_formula_cells(ws, ws.max_row)
    return _wb_to_bytes(wb)


def _wb_to_bytes(wb: Any) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
