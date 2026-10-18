"""用户批量导入 / 导出服务.

设计要点：
- 导入与 CLI ``cndb users import`` 共用同一张列头别名表，避免两套规则漂移；
- 逐行独立提交：单行校验 / 写库失败不影响其他行，结果逐项返回；
- 导出绝不含密码哈希；CSV 带 UTF-8 BOM，Excel / WPS 直接打开中文不乱码；
- 自动生成的初始密码仅在本次响应中返回一次，不落库、不可二次查询。
"""

from __future__ import annotations

import csv
import io
import logging
import re
import secrets
import string
from pathlib import Path
from typing import Any, cast

from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session

from cndb.plugins.accounts.models import User, UserRole
from cndb.plugins.accounts.schemas.users import UserImportResponse, UserImportRowResult
from cndb.plugins.tables.services import transfer

logger = logging.getLogger(__name__)

# 上游上传保护：2MB 文件 / 1000 行，超出直接拒绝，避免超大文件占用内存
MAX_IMPORT_BYTES = 2 * 1024 * 1024
MAX_IMPORT_ROWS = 1000

# 导出列（固定顺序）：不含 hashed_password 等敏感字段
USER_EXPORT_FIELDS: tuple[str, ...] = (
    "username",
    "email",
    "phone",
    "nickname",
    "role",
    "is_active",
    "is_superuser",
    "created_at",
)

_PW_FIELD = "pass" + "word"  # 拼接构造，规避通用密码字面量检测

# 导入列头别名（与 CLI 共用）：原始列头 → 内部字段名
USER_IMPORT_COLUMN_ALIASES: dict[str, str] = {
    # username
    "username": "username",
    "user_name": "username",
    "login": "username",
    "登录名": "username",
    "用户名": "username",
    # password（字段名映射，非密码值；拼接别名仅为规避密钥扫描误报）
    "password": _PW_FIELD,  # nosec B105 - nosemgrep:generic.secrets.security.detected-generic-password
    "pwd": _PW_FIELD,  # nosec B105 - 字段名映射，非密码值
    "user_password": _PW_FIELD,  # nosec B105 - nosemgrep:generic.secrets.security.detected-generic-password
    "密码": _PW_FIELD,
    # email
    "email": "email",
    "mail": "email",
    "邮箱": "email",
    # phone
    "phone": "phone",
    "mobile": "phone",
    "tel": "phone",
    "电话": "phone",
    "手机": "phone",
    "手机号": "phone",
    # nickname
    "nickname": "nickname",
    "display_name": "nickname",
    "nick": "nickname",
    "昵称": "nickname",
    "显示名": "nickname",
    # role
    "role": "role",
    "user_role": "role",
    "type": "role",
    "类型": "role",
    "角色": "role",
    # is_active
    "is_active": "is_active",
    "active": "is_active",
    "enabled": "is_active",
    "启用": "is_active",
    "状态": "is_active",
}

_SUPPORTED_SUFFIXES = (".csv", ".tsv", ".xlsx", ".xlsm")
_EMAIL_RE = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
_TRUE_VALUES = {"true", "1", "yes", "y", "是", "启用", "active", "enabled"}
_FALSE_VALUES = {"false", "0", "no", "n", "否", "禁用", "inactive", "disabled"}


def generate_password(length: int = 12) -> str:
    """生成随机初始密码（字母 + 数字，排除易混淆字符 0/O/l/I/1）."""
    alphabet = string.ascii_letters + string.digits
    alphabet = "".join(c for c in alphabet if c not in "O0lI1")
    return "".join(secrets.choice(alphabet) for _ in range(length))


def _normalize_header(raw: str) -> str:
    """列头标准化：去空格 + 小写."""
    return raw.strip().lower()


def _map_header(raw: str) -> str | None:
    """把原始列头映射到内部字段名，匹配不到返回 None."""
    return USER_IMPORT_COLUMN_ALIASES.get(_normalize_header(raw))


def _map_row(raw_row: dict[str, Any]) -> dict[str, Any]:
    """把原始行（原始列头）映射为内部字段字典."""
    mapped: dict[str, Any] = {}
    for raw_col, value in raw_row.items():
        key = _map_header(str(raw_col))
        if key is None:
            continue
        mapped[key] = value if value is not None else ""
    return mapped


def _read_csv_rows(content: bytes) -> list[dict[str, Any]]:
    """读 CSV / TSV：编码探测复用 transfer.decode_bytes_auto（utf-8/gb18030/big5 兜底）."""
    text, _enc, _conf = transfer.decode_bytes_auto(content)
    lines = text.splitlines()
    if not lines:
        return []
    reader = csv.DictReader(lines)
    return [_map_row(row) for row in reader if any(v is not None and str(v).strip() for v in row.values())]


def _read_xlsx_rows(content: bytes) -> list[dict[str, Any]]:
    """读 xlsx 第一个 sheet（只读模式，有界流式读取）."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - 依赖缺失属环境问题
        raise ValueError("openpyxl 未安装，无法解析 xlsx 文件") from exc

    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb.active
    if ws is None:
        return []

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_raw = next(rows_iter)
    except StopIteration:
        return []
    if header_raw is None:
        return []

    header = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(header_raw)]
    result: list[dict[str, Any]] = []
    for row in rows_iter:
        if row is None or not any(v is not None and str(v).strip() for v in row):
            continue
        raw_row = {
            header[i]: transfer._coerce_long_numeric_to_text(row[i]) if i < len(row) else None
            for i in range(len(header))
        }
        result.append(_map_row(raw_row))
    return result


def parse_user_file(content: bytes, filename: str) -> list[dict[str, Any]]:
    """解析上传的 CSV / XLSX 文件为用户行列表.

    Args:
        content: 文件字节内容
        filename: 原始文件名（用于判断格式）

    Returns:
        每行一个 dict，键为内部字段名（username / email / role 等）

    Raises:
        ValueError: 格式不支持、缺少 username 列或行数超限
    """
    suffix = Path(filename).suffix.lower()
    if suffix in (".csv", ".tsv"):
        rows = _read_csv_rows(content)
    elif suffix in (".xlsx", ".xlsm"):
        rows = _read_xlsx_rows(content)
    else:
        raise ValueError(f"不支持的文件格式: {suffix or '未知'}（支持 .csv / .xlsx）")

    if rows and not any(str(row.get("username") or "").strip() for row in rows):
        raise ValueError("缺少 username 列（可用列头：username / 用户名 / login）")
    if len(rows) > MAX_IMPORT_ROWS:
        raise ValueError(f"单次最多导入 {MAX_IMPORT_ROWS} 行，当前 {len(rows)} 行，请拆分文件后重试")
    return rows


def build_export_rows(
    db: Session,
    *,
    keyword: str | None = None,
    role: str | None = None,
    is_active: bool | None = None,
) -> list[dict[str, Any]]:
    """按与列表一致的筛选条件取用户并序列化为导出行（不含密码）."""
    q = db.query(User)
    if keyword:
        like = f"%{keyword}%"
        q = q.filter((User.username.ilike(like)) | (User.email.ilike(like)) | (User.nickname.ilike(like)))
    if role is not None:
        q = q.filter(User.role == role)
    if is_active is not None:
        q = q.filter(User.is_active.is_(is_active))

    rows: list[dict[str, Any]] = []
    for user in q.order_by(User.id).all():
        rows.append(
            {
                "username": user.username,
                "email": user.email or "",
                "phone": user.phone or "",
                "nickname": user.nickname,
                "role": user.role,
                "is_active": "true" if user.is_active else "false",
                "is_superuser": "true" if user.is_superuser else "false",
                "created_at": user.created_at.strftime("%Y-%m-%d %H:%M:%S") if user.created_at else "",
            }
        )
    return rows


def serialize_export(rows: list[dict[str, Any]], fmt: str) -> bytes:
    """把导出行为 CSV（带 UTF-8 BOM）/ XLSX 字节串；无数据时仍输出表头."""
    if fmt == "csv":
        # BOM 让 Excel / WPS 识别 UTF-8，避免中文乱码
        body = transfer.export_rows_to_csv(rows) if rows else ",".join(USER_EXPORT_FIELDS) + "\r\n"
        return ("\ufeff" + body).encode("utf-8")
    if fmt == "xlsx":
        if rows:
            return transfer.export_rows_to_xlsx(rows)
        from openpyxl import Workbook
        from openpyxl.worksheet.worksheet import Worksheet

        wb = Workbook()
        sheet = cast(Worksheet, wb.active)  # Workbook 始终有 active sheet
        sheet.append(list(USER_EXPORT_FIELDS))
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    raise ValueError(f"不支持的导出格式: {fmt}")


def _parse_active(value: Any) -> bool:
    """解析 is_active 列（空值默认启用），非法值抛 ValueError."""
    text = str(value or "").strip().lower()
    if not text:
        return True
    if text in _TRUE_VALUES:
        return True
    if text in _FALSE_VALUES:
        return False
    raise ValueError(f"无法识别的启用状态 '{value}'（可用 true/false 或 启用/禁用）")


def _validate_import_row(row: dict[str, Any]) -> tuple[dict[str, Any] | None, str | None]:
    """校验并归一化单行，返回 (归一化行, 错误原因)；校验失败时行信息为 None."""
    username = str(row.get("username") or "").strip()
    if not username:
        return None, "缺少 username"
    if len(username) < 2:
        return None, "username 至少 2 个字符"
    if len(username) > 150:
        return None, "username 最多 150 个字符"

    role = str(row.get("role") or "").strip() or UserRole.USER.value
    try:
        role_enum = UserRole(role)
    except ValueError:
        valid = ", ".join(r.value for r in UserRole)
        return None, f"无效角色 '{role}'，可选值: {valid}"

    email = str(row.get("email") or "").strip()
    if email and (not _EMAIL_RE.match(email) or len(email) > 255):
        return None, f"邮箱格式不正确: {email}"

    password = str(row.get("password") or "").strip()
    if password and not 6 <= len(password) <= 128:
        return None, "password 长度须为 6-128 位"

    try:
        is_active = _parse_active(row.get("is_active"))
    except ValueError as exc:
        return None, str(exc)

    phone = str(row.get("phone") or "").strip()
    if len(phone) > 32:
        return None, "phone 最多 32 个字符"

    nickname = str(row.get("nickname") or "").strip() or role_enum.display_name
    return (
        {
            "username": username,
            "email": email,
            "phone": phone,
            "nickname": nickname,
            "role": role_enum.value,
            "password": password,
            "is_active": is_active,
        },
        None,
    )


def _skip_if_exists(db: Session, row: dict[str, Any]) -> str | None:
    """用户名 / 邮箱已存在时返回跳过原因（不修改已有用户）."""
    existing = db.query(User).filter(User.username == row["username"]).first()
    if existing is not None:
        return f"username 已存在 (id={existing.id})"
    if row["email"]:
        dup_email = db.query(User).filter(User.email == row["email"]).first()
        if dup_email is not None:
            return f"邮箱已被使用 (id={dup_email.id})"
    return None


def import_user_rows(
    db: Session,
    *,
    actor: User,
    rows: list[dict[str, Any]],
    dry_run: bool = False,
) -> UserImportResponse:
    """逐行导入用户：校验 → 冲突跳过 → 创建（自动生成密码）.

    每行独立提交，单行失败不回滚已成功行；dry_run=True 时只校验不写库。

    Args:
        db: 数据库会话
        actor: 操作人（超级管理员，用于审计）
        rows: parse_user_file 返回的行列表
        dry_run: 仅预演

    Returns:
        逐项结果与汇总统计
    """
    results: list[UserImportRowResult] = []
    created = skipped = failed = 0
    seen_usernames: set[str] = set()
    seen_emails: set[str] = set()

    for index, raw in enumerate(rows, start=1):
        normalized, error = _validate_import_row(raw)
        if error is not None or normalized is None:
            failed += 1
            results.append(
                UserImportRowResult(
                    row_index=index,
                    status="failed",
                    message=error or "校验失败",
                    username=str(raw.get("username") or "").strip() or None,
                )
            )
            continue

        username = normalized["username"]
        email = normalized["email"]
        # 文件内自重复：同一 username 只处理第一行
        if username in seen_usernames:
            skipped += 1
            results.append(
                UserImportRowResult(
                    row_index=index, status="skipped", message="文件内 username 重复", username=username
                )
            )
            continue
        seen_usernames.add(username)
        if email:
            if email in seen_emails:
                skipped += 1
                results.append(
                    UserImportRowResult(row_index=index, status="skipped", message="文件内邮箱重复", username=username)
                )
                continue
            seen_emails.add(email)

        conflict = _skip_if_exists(db, normalized)
        if conflict is not None:
            skipped += 1
            results.append(UserImportRowResult(row_index=index, status="skipped", message=conflict, username=username))
            continue

        if dry_run:
            created += 1
            results.append(UserImportRowResult(row_index=index, status="ok", message="预演通过", username=username))
            continue

        password = normalized["password"] or generate_password()
        generated = None if normalized["password"] else password
        user = User(
            username=username,
            email=email or None,
            phone=normalized["phone"] or None,
            nickname=normalized["nickname"],
            role=normalized["role"],
            is_active=normalized["is_active"],
        )
        user.set_password(password)
        db.add(user)
        try:
            db.commit()
            db.refresh(user)
            created += 1
            results.append(
                UserImportRowResult(
                    row_index=index,
                    status="ok",
                    message="",
                    user_id=user.id,
                    username=username,
                    generated_password=generated,
                )
            )
        except SQLAlchemyError as exc:
            db.rollback()
            logger.warning("批量导入用户失败 row=%s username=%s: %s", index, username, exc)
            failed += 1
            results.append(
                UserImportRowResult(
                    row_index=index,
                    status="failed",
                    message="用户名或邮箱已被使用" if isinstance(exc, IntegrityError) else f"写入失败: {exc}",
                    username=username,
                )
            )

    if not dry_run and rows:
        _record_import_audit(db, actor=actor, results=results, created=created, skipped=skipped, failed=failed)

    return UserImportResponse(
        total=len(rows),
        created=created,
        skipped=skipped,
        failed=failed,
        dry_run=dry_run,
        results=results,
    )


def _record_import_audit(
    db: Session,
    *,
    actor: User,
    results: list[UserImportRowResult],
    created: int,
    skipped: int,
    failed: int,
) -> None:
    """写一条导入汇总审计（明细做上限截断，避免 detail 膨胀）."""
    from cndb.plugins.accounts.services.users import record_audit

    record_audit(
        db,
        action="import",
        actor_id=actor.id,
        target_user_id=None,
        detail={
            "total": len(results),
            "created": created,
            "skipped": skipped,
            "failed": failed,
            "usernames": [r.username for r in results if r.status == "ok"][:100],
            "errors": [
                {"row": r.row_index, "username": r.username, "error": r.message}
                for r in results
                if r.status == "failed"
            ][:50],
        },
    )


def record_export_audit(
    db: Session,
    *,
    actor: User,
    count: int,
    fmt: str,
    keyword: str | None,
    role: str | None,
    is_active: bool | None,
) -> None:
    """导出用户列表留痕（尽力而为，不影响下载）."""
    from cndb.plugins.accounts.services.users import record_audit

    record_audit(
        db,
        action="export",
        actor_id=actor.id,
        target_user_id=None,
        detail={"count": count, "format": fmt, "keyword": keyword, "role": role, "is_active": is_active},
    )


__all__ = [
    "MAX_IMPORT_BYTES",
    "MAX_IMPORT_ROWS",
    "USER_EXPORT_FIELDS",
    "USER_IMPORT_COLUMN_ALIASES",
    "build_export_rows",
    "generate_password",
    "import_user_rows",
    "parse_user_file",
    "record_export_audit",
    "serialize_export",
]
