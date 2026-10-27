"""DOCX 导出器：Markdown 风格标题/加粗/代码/表格 → Word 文档。"""

from __future__ import annotations

import io
import logging
import re
from typing import Any, cast

from cndb.plugins.documents.services.charts import to_png
from cndb.plugins.documents.services.export.chart_refs import (
    ChartResolver,
    ResolvedChart,
    parse_chart_ref_line,
)
from cndb.plugins.documents.services.export.images import IMAGE_LINE_RE, ImageResolver
from cndb.plugins.documents.services.export.markdown_style import (
    _MD_BOLD_RE,
    _MD_HEADING_RE,
    _MD_TABLE_SEP_RE,
)
from cndb.plugins.documents.services.export.table_refs import (
    ResolvedTable,
    TableResolver,
    parse_table_ref_line,
)
from cndb.plugins.documents.services.export.theme import DEFAULT_THEME

logger = logging.getLogger(__name__)


def _render_docx(
    markdown_text: str,
    title: str = "",
    images: ImageResolver | None = None,
    tables: TableResolver | None = None,
    charts: ChartResolver | None = None,
) -> bytes:
    # images：站内图片解析器；``![alt](file:<key>)`` 独占一行时嵌入真实图片，
    # 解析失败写占位文字，避免导出物静默丢图
    # tables：数据表引用解析器；``{{table:…}}`` 独占一行时渲染为真实表格
    # charts：图表引用解析器；``{{chart:…}}`` 独占一行时嵌入整幅图 PNG（2 倍分辨率）
    """把文档 Markdown 正文渲染为 .docx 字节流.

    支持标题 / 加粗 / 行内代码 / GFM 表格 / ``---PAGE---`` 换页，
    并按默认导出预设应用字体、字号与表头配色。
    """
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor
    from docx.styles.style import ParagraphStyle as DocxParagraphStyle

    preset = DEFAULT_THEME
    doc = Document()
    if title:
        doc.core_properties.title = title

    # 全局正文样式（Normal）应用主题正文字体/字号，含中文东亚字体映射
    normal = cast(DocxParagraphStyle, doc.styles["Normal"])
    normal.font.name = preset.body_font
    normal.font.size = Pt(preset.body_size)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), preset.body_font)

    def _apply_run_font(run: Any, font: str, size: float | None = None, color: str | None = None) -> None:
        """给 run 设置字体名（含东亚映射）、字号、颜色."""
        run.font.name = font
        run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
        if size is not None:
            run.font.size = Pt(size)
        if color is not None:
            run.font.color.rgb = RGBColor.from_string(color.lstrip("#"))

    def _add_themed_heading(text: str, level: int) -> None:
        """按主题添加标题段落；主题色/字号/字体应用于标题 run."""
        idx = min(level, 3) - 1
        try:
            p = doc.add_heading(text, level=level)
        except (KeyError, ValueError):
            # 某些 Word 版本没有 Heading 5+
            p = doc.add_paragraph()
            run = p.add_run(text)
            run.bold = True
            _apply_run_font(run, preset.heading_font, preset.heading_sizes[idx], preset.heading_colors[idx])
            return
        for run in p.runs:
            _apply_run_font(run, preset.heading_font, preset.heading_sizes[idx], preset.heading_colors[idx])

    table_buffer: list[str] = []
    in_table = False

    def _is_table_row(line: str) -> bool:
        s = line.strip()
        return s.startswith("|") and s.endswith("|")

    def _write_table(rows: list[list[str]]) -> None:
        """把二维文本写入 docx 表格对象（首行加粗 + 主题底纹）."""
        if not rows:
            return
        max_cols = max(len(r) for r in rows)
        norm_rows = [r + [""] * (max_cols - len(r)) for r in rows]
        t = doc.add_table(rows=len(norm_rows), cols=max_cols)
        t.style = "Table Grid"
        for i, row_cells in enumerate(norm_rows):
            for j, cell_text in enumerate(row_cells):
                cell = t.cell(i, j)
                p = cell.paragraphs[0]
                if i == 0:
                    run = p.add_run(cell_text)
                    run.bold = True
                    # 表头主题：底纹 + 文字色
                    shd = cell._tc.get_or_add_tcPr().makeelement(qn("w:shd"), {})
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), preset.table_header_bg.lstrip("#"))
                    cell._tc.get_or_add_tcPr().append(shd)
                    _apply_run_font(run, preset.body_font, color=preset.table_header_color)
                else:
                    _add_rich_runs(p, cell_text, preset.table_body_font, preset.code_font)

    def _flush_table() -> None:
        """把缓冲的 markdown 表格行写入 docx 表格对象（跳过分隔行）."""
        nonlocal in_table, table_buffer
        rows = [[c.strip() for c in raw.strip().strip("|").split("|")] for raw in table_buffer if raw.strip()]
        # 跳过 markdown 分隔行（| --- | :---: |）
        rows = [r for r in rows if not all(_MD_TABLE_SEP_RE.match(c or "-") for c in r)]
        _write_table(rows)
        in_table = False
        table_buffer = []

    def _add_table_ref(resolved: ResolvedTable) -> None:
        """渲染数据表引用：坏引用降级为占位段落，正常引用输出说明行 + 表格."""
        if resolved.error:
            doc.add_paragraph(resolved.error)
            return
        caption = doc.add_paragraph()
        run = caption.add_run(resolved.caption)
        run.italic = True
        caption_size = preset.body_size - 1
        _apply_run_font(run, preset.body_font, caption_size)
        _write_table([resolved.header or [], *(resolved.rows or [])])

    def _add_chart_ref(resolved: ResolvedChart, chart_id: int) -> None:
        """渲染图表引用：坏引用降级为占位段落，正常引用嵌入整幅图 PNG（D6，2 倍分辨率）."""
        if resolved.error:
            doc.add_paragraph(resolved.error)
            return
        rendered = resolved.rendered
        if rendered is None:  # pragma: no cover - 解析器契约保证二选一
            doc.add_paragraph(f"[图表无法嵌入：{resolved.title or f'图表 #{chart_id}'}]")
            return
        try:
            doc.add_picture(io.BytesIO(to_png(rendered, scale=2)), width=_docx_image_width(doc))
        except Exception:  # pragma: no cover - Pillow 绘制失败兜底
            doc.add_paragraph(f"[图表无法嵌入：{resolved.title or f'图表 #{chart_id}'}]")

    for raw_line in markdown_text.split("\n"):
        line = raw_line.rstrip()

        if _is_table_row(line):
            if not in_table:
                in_table = True
                table_buffer = []
            table_buffer.append(line)
            continue
        if in_table:
            _flush_table()

        # 数据表引用：独占一行时渲染为真实表格
        if tables is not None:
            spec = parse_table_ref_line(line)
            if spec is not None:
                _add_table_ref(tables(spec))
                continue

        # 图表引用：独占一行时渲染为整幅图（坏引用降级为占位文字）
        if charts is not None:
            chart_spec = parse_chart_ref_line(line)
            if chart_spec is not None:
                _add_chart_ref(charts(chart_spec), chart_spec.chart_id)
                continue

        if not line.strip():
            doc.add_paragraph("")
            continue

        m = _MD_HEADING_RE.match(line)
        if m:
            level = min(len(m.group(1)), 9)
            text = m.group(2).strip()
            _add_themed_heading(text, level)
            continue

        # 换页标记
        if line.strip() == "---PAGE---":
            doc.add_page_break()
            continue

        # 站内图片：独占一行时嵌入真实图片
        img_match = IMAGE_LINE_RE.match(line)
        if img_match:
            alt = img_match.group("alt") or "图片"
            resolved = images(img_match.group("key")) if images is not None else None
            if resolved is None:
                doc.add_paragraph(f"[图片缺失：{alt}]")
            else:
                data, _name = resolved
                try:
                    doc.add_picture(io.BytesIO(data), width=_docx_image_width(doc))
                except Exception:
                    doc.add_paragraph(f"[图片无法嵌入：{alt}]")
            continue

        # 普通段落，处理粗体和代码
        p = doc.add_paragraph()
        _add_rich_runs(p, line, preset.body_font, preset.code_font)

    # 收尾：提交剩余表格
    if in_table:
        _flush_table()

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _docx_image_width(doc: Any) -> Any:
    """图片最大宽度 = 正文可用宽度（页宽减左右页边距）；取不到时回退 14cm."""
    from docx.shared import Cm

    try:
        section = doc.sections[0]
        usable = section.page_width - section.left_margin - section.right_margin
        if usable and usable > 0:
            return usable
    except (IndexError, AttributeError):  # pragma: no cover - 极端文档结构
        pass
    return Cm(14)


def _add_rich_runs(paragraph: Any, text: str, body_font: str = "Calibri", code_font: str = "Courier New") -> None:
    """在 docx 段落中按 **bold** 和 `code` 分段添加 Run，正文/代码分别应用主题字体."""
    # 先用代码段切分（优先级高）
    code_parts = re.split(r"(`[^`]+`)", text)
    for part in code_parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = code_font
            continue
        # 再按粗体切分
        bold_parts = _MD_BOLD_RE.split(part)
        for i, bp in enumerate(bold_parts):
            if i % 2 == 1:
                # 奇数位是粗体内容
                run = paragraph.add_run(bp)
                run.bold = True
                run.font.name = body_font
            elif bp:
                run = paragraph.add_run(bp)
                run.font.name = body_font


__all__ = ["_add_rich_runs", "_docx_image_width", "_render_docx"]
